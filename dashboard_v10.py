"""
Dashboard Profissional v10 - Cockpit executivo com MariaDB, custos reais,
Jogos Santa Casa, Novadis, Delta e margem por produto.
"""

import streamlit as st
import pandas as pd
import plotly.express as px
import plotly.graph_objects as go
from plotly.subplots import make_subplots
import numpy as np
from datetime import datetime, timedelta
from data_loader_v10 import DataLoaderV10
from cost_manager_v2 import CostManagerV2  # NOVA VERSÃO COM DESPESIFY
from product_categorizer import ProductCategorizer
import streamlit_authenticator as stauth
import yaml
from yaml.loader import SafeLoader
from pathlib import Path
from io import BytesIO
from html import escape
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.utils.dataframe import dataframe_to_rows

# Configuração da página
st.set_page_config(
    page_title="Dashboard v10 - Café Martins",
    page_icon="📊",
    layout="wide",
    initial_sidebar_state="expanded"
)

# Estilos CSS customizados
st.markdown("""
    <style>
    /* Estilos profissionais */
    .main {
        background-color: #f8f9fa;
    }

    /* Cards de métricas */
    .metric-card {
        background: white;
        color: #172026;
        padding: 20px;
        border-radius: 10px;
        box-shadow: 0 2px 4px rgba(0,0,0,0.1);
        margin: 10px 0;
        transition: transform 0.2s;
    }

    .metric-card * {
        color: inherit;
    }

    .metric-card:hover {
        transform: translateY(-2px);
        box-shadow: 0 4px 8px rgba(0,0,0,0.15);
    }

    /* Título principal */
    .main-title {
        color: #1f77b4;
        font-size: 2.5rem;
        font-weight: 700;
        text-align: center;
        margin-bottom: 30px;
        text-shadow: 2px 2px 4px rgba(0,0,0,0.1);
    }

    .subtitle {
        color: #666;
        font-size: 1.2rem;
        text-align: center;
        margin-bottom: 40px;
    }

    /* Badges de categoria */
    .category-badge {
        display: inline-block;
        padding: 5px 15px;
        border-radius: 20px;
        font-weight: 600;
        margin: 5px;
        color: white;
    }

    /* Separadores */
    .section-divider {
        margin: 30px 0;
        border-top: 2px solid #e0e0e0;
    }

    /* Tabs customizadas - responsivas com múltiplas linhas */
    .stTabs [data-baseweb="tab-list"] {
        gap: 8px;
        flex-wrap: wrap !important;
        overflow-x: visible !important;
    }

    .stTabs [data-baseweb="tab"] {
        min-height: 50px;
        height: auto !important;
        background-color: white;
        border-radius: 8px 8px 0 0;
        padding: 8px 12px;
        font-weight: 600;
        color: #333 !important;
        font-size: 14px !important;
        white-space: nowrap;
        flex-shrink: 0;
    }

    .stTabs [aria-selected="true"] {
        background-color: #1f77b4;
        color: white !important;
    }

    /* Garantir que o texto das tabs é visível */
    .stTabs button div {
        color: inherit !important;
    }

    /* Responsividade para ecrãs pequenos */
    @media (max-width: 1200px) {
        .stTabs [data-baseweb="tab"] {
            font-size: 12px !important;
            padding: 6px 10px;
            min-height: 45px;
        }
    }

    @media (max-width: 768px) {
        .stTabs [data-baseweb="tab"] {
            font-size: 11px !important;
            padding: 5px 8px;
            min-height: 40px;
        }
    }

    /* Sidebar */
    .css-1d391kg {
        background-color: #f8f9fa;
    }

    /* Botões */
    .stButton > button {
        width: 100%;
        border-radius: 8px;
        font-weight: 600;
        transition: all 0.3s;
    }

    .stButton > button:hover {
        transform: scale(1.02);
        box-shadow: 0 4px 8px rgba(0,0,0,0.2);
    }

    /* Dataframe */
    .dataframe {
        border-radius: 10px;
        overflow: hidden;
    }

    /* Loading spinner */
    .stSpinner > div {
        border-color: #1f77b4 !important;
    }

    /* Info boxes */
    .info-box {
        background: #e3f2fd;
        color: #17324d;
        padding: 15px;
        border-radius: 8px;
        border-left: 4px solid #1f77b4;
        margin: 15px 0;
    }

    .warning-box {
        background: #fff3e0;
        color: #5f370e;
        padding: 15px;
        border-radius: 8px;
        border-left: 4px solid #ff9800;
        margin: 15px 0;
    }

    .success-box {
        background: #e8f5e9;
        color: #184d27;
        padding: 15px;
        border-radius: 8px;
        border-left: 4px solid #4caf50;
        margin: 15px 0;
    }
    </style>
""", unsafe_allow_html=True)

st.markdown("""
    <style>
    :root {
        --cm-ink: #172026;
        --cm-muted: #667085;
        --cm-line: #d7dde5;
        --cm-panel: #ffffff;
        --cm-bg: #f5f7f9;
        --cm-green: #12805c;
        --cm-red: #b42318;
        --cm-blue: #1b5f8f;
        --cm-gold: #936b10;
    }

    .main .block-container {
        max-width: 1540px;
        padding-top: 1.25rem;
        padding-bottom: 3rem;
    }

    .v10-topbar {
        border-bottom: 1px solid var(--cm-line);
        padding: 0 0 16px 0;
        margin-bottom: 16px;
    }

    .v10-title {
        color: var(--cm-ink);
        font-size: 2.1rem;
        line-height: 1.05;
        font-weight: 780;
        letter-spacing: 0;
        margin: 0;
    }

    .v10-subtitle {
        color: var(--cm-muted);
        font-size: 0.98rem;
        margin-top: 6px;
    }

    .v10-chip-row {
        display: flex;
        flex-wrap: wrap;
        gap: 8px;
        margin-top: 10px;
    }

    .v10-chip {
        border: 1px solid var(--cm-line);
        border-radius: 999px;
        padding: 4px 10px;
        color: #344054;
        background: #fff;
        font-size: 0.8rem;
        white-space: nowrap;
    }

    .v10-kpi {
        background: #ffffff;
        color: var(--cm-ink);
        border: 1px solid var(--cm-line);
        border-top: 4px solid var(--cm-blue);
        border-radius: 8px;
        padding: 14px 16px 14px 16px;
        min-height: 132px;
        box-shadow: 0 1px 2px rgba(16, 24, 40, 0.06);
    }

    .v10-kpi * {
        color: inherit;
    }

    .v10-kpi-label {
        color: var(--cm-muted) !important;
        font-size: 0.78rem;
        text-transform: uppercase;
        font-weight: 740;
        letter-spacing: 0;
        margin-bottom: 8px;
    }

    .v10-kpi-value {
        color: var(--cm-ink) !important;
        font-size: 1.85rem;
        font-weight: 820;
        line-height: 1.05;
        margin-bottom: 8px;
    }

    .v10-kpi-delta {
        font-size: 0.88rem;
        font-weight: 700;
    }

    .v10-positive { color: var(--cm-green) !important; }
    .v10-negative { color: var(--cm-red) !important; }
    .v10-neutral { color: var(--cm-muted) !important; }

    .v10-panel {
        background: #ffffff;
        color: var(--cm-ink);
        border: 1px solid var(--cm-line);
        border-radius: 8px;
        padding: 14px 16px;
        min-height: 100%;
        box-shadow: 0 1px 2px rgba(16, 24, 40, 0.06);
    }

    .v10-panel * {
        color: inherit;
    }

    .v10-panel-title {
        color: var(--cm-ink) !important;
        font-size: 1rem;
        font-weight: 760;
        margin-bottom: 8px;
    }

    .v10-action {
        border-left: 4px solid var(--cm-blue);
        padding: 9px 10px;
        margin: 8px 0;
        background: #f8fbfd;
        color: #24313a !important;
        font-size: 0.92rem;
    }

    .v10-action-high { border-left-color: var(--cm-red); background: #fff8f7; }
    .v10-action-medium { border-left-color: var(--cm-gold); background: #fffaf0; }
    .v10-action-low { border-left-color: var(--cm-green); background: #f6fbf8; }

    .v10-score {
        display: flex;
        align-items: baseline;
        gap: 8px;
    }

    .v10-score-number {
        font-size: 2.3rem;
        font-weight: 840;
        color: var(--cm-ink);
        line-height: 1;
    }

    .v10-score-label {
        color: var(--cm-muted);
        font-size: 0.88rem;
    }

    div[data-testid="stMetric"] {
        background: #ffffff !important;
        color: #172026 !important;
        border: 1px solid #d7dde5;
        border-top: 4px solid #1b5f8f;
        border-radius: 8px;
        padding: 12px 14px;
        box-shadow: 0 1px 2px rgba(16, 24, 40, 0.06);
    }

    div[data-testid="stMetric"] * {
        color: #172026 !important;
    }

    div[data-testid="stMetric"] label,
    div[data-testid="stMetricLabel"],
    div[data-testid="stMetricLabel"] * {
        color: #475467 !important;
    }

    div[data-testid="stMetricValue"],
    div[data-testid="stMetricValue"] * {
        color: #172026 !important;
    }

    div[data-testid="stMetricDelta"],
    div[data-testid="stMetricDelta"] * {
        color: #344054 !important;
    }
    </style>
""", unsafe_allow_html=True)

st.markdown("""
    <style>
    :root {
        --cm-ink: #18222d;
        --cm-muted: #667085;
        --cm-line: #d8ded8;
        --cm-panel: #ffffff;
        --cm-bg: #f3f5f1;
        --cm-slate: #22333b;
        --cm-teal: #2f5d62;
        --cm-rust: #9b5a3c;
        --cm-gold: #b88746;
        --cm-green: #2f7d68;
        --cm-red: #b54740;
    }

    html, body, [data-testid="stAppViewContainer"] {
        background: var(--cm-bg) !important;
        color: var(--cm-ink);
    }

    section[data-testid="stSidebar"] {
        background: #fbfaf7 !important;
        border-right: 1px solid var(--cm-line);
    }

    .v10-shell-brand {
        background: var(--cm-slate);
        color: #ffffff;
        border: 1px solid rgba(255, 255, 255, 0.12);
        border-radius: 8px;
        padding: 14px 14px 12px 14px;
        margin: 4px 0 16px 0;
    }

    .v10-shell-brand-title {
        color: #ffffff;
        font-size: 1.2rem;
        font-weight: 850;
        line-height: 1.1;
    }

    .v10-shell-brand-subtitle {
        color: rgba(255, 255, 255, 0.72);
        font-size: 0.82rem;
        margin-top: 5px;
    }

    .v10-topbar {
        background: #fffefa;
        color: var(--cm-ink);
        border: 1px solid var(--cm-line);
        border-left: 5px solid var(--cm-teal);
        border-radius: 8px;
        padding: 16px 18px;
        margin-bottom: 14px;
        box-shadow: 0 1px 2px rgba(16, 24, 40, 0.06);
    }

    .v10-title {
        color: var(--cm-ink) !important;
        font-size: 1.82rem;
        font-weight: 850;
    }

    .v10-subtitle {
        color: var(--cm-muted) !important;
    }

    .v10-chip {
        background: #f4f7f4;
        border-color: var(--cm-line);
        color: #344054;
        font-weight: 650;
    }

    .v10-kpi,
    div[data-testid="stMetric"] {
        border: 1px solid var(--cm-line) !important;
        border-left: 4px solid var(--cm-teal) !important;
        border-top: 0 !important;
        background: #fffefa !important;
        border-radius: 8px !important;
        box-shadow: 0 1px 2px rgba(16, 24, 40, 0.06) !important;
    }

    .v10-kpi:nth-of-type(2n) {
        border-left-color: var(--cm-rust) !important;
    }

    .v10-kpi-label,
    div[data-testid="stMetricLabel"],
    div[data-testid="stMetricLabel"] * {
        color: #667085 !important;
        letter-spacing: 0 !important;
    }

    .v10-kpi-value,
    div[data-testid="stMetricValue"],
    div[data-testid="stMetricValue"] * {
        color: var(--cm-ink) !important;
    }

    .v10-panel {
        background: #ffffff !important;
        border-color: var(--cm-line) !important;
    }

    .v10-module-map {
        display: grid;
        grid-template-columns: repeat(4, minmax(0, 1fr));
        gap: 10px;
        margin: 6px 0 18px 0;
    }

    .v10-module {
        background: #fffefa;
        border: 1px solid var(--cm-line);
        border-radius: 8px;
        padding: 12px 13px;
        min-height: 92px;
    }

    .v10-module-kicker {
        color: var(--cm-rust);
        font-size: 0.74rem;
        font-weight: 820;
        text-transform: uppercase;
        letter-spacing: 0;
        margin-bottom: 5px;
    }

    .v10-module-title {
        color: var(--cm-ink);
        font-size: 0.98rem;
        font-weight: 800;
        margin-bottom: 4px;
    }

    .v10-module-copy {
        color: var(--cm-muted);
        font-size: 0.84rem;
        line-height: 1.3;
    }

    @media (max-width: 1100px) {
        .v10-module-map {
            grid-template-columns: repeat(2, minmax(0, 1fr));
        }
    }

    @media (max-width: 680px) {
        .v10-module-map {
            grid-template-columns: 1fr;
        }
        .v10-title {
            font-size: 1.7rem;
        }
    }
    </style>
""", unsafe_allow_html=True)


# Funções auxiliares
@st.cache_resource(ttl=1800)
def carregar_dados():
    """Carrega dados com custos REAIS (Despesify) + estimados - cache de 30 minutos

    V9: Carrega direto do MariaDB - ULTRA-RÁPIDO!
    Usa cache_resource porque cost_manager tem conexão MariaDB (não serializável com pickle)
    V10: carrega direto do MariaDB e acrescenta cockpit executivo.
    """
    with st.spinner("A carregar dados operacionais..."):
        loader = DataLoaderV10(usar_despesify=True)
        df = loader.carregar_tudo_integrado_com_custos()
        # Garantir que carregar_delta está acessível e pré-carregar se necessário
        try:
            _ = loader.carregar_delta()
        except:
            pass
        cost_manager = loader.cost_manager  # CostManagerV2 com Despesify
        return df, loader, cost_manager


def formatar_moeda(valor):
    """Formata valor em euros sem casas decimais"""
    return f"{valor:,.0f}€".replace(",", "X").replace(".", ",").replace("X", ".")


def formatar_percentagem(valor):
    """Formata percentagem"""
    return f"{valor:.1f}%"


def calcular_delta(atual, anterior):
    """Calcula variação percentual - retorna 0 se anterior é 0 ou None"""
    if anterior is None or anterior == 0:
        return 0
    return ((atual - anterior) / anterior) * 100


def formatar_delta_html(valor):
    """Formata variação percentual para cartões HTML."""
    if valor is None or abs(valor) < 0.05:
        return '<span class="v10-kpi-delta v10-neutral">estável</span>'

    css_class = "v10-positive" if valor > 0 else "v10-negative"
    sinal = "+" if valor > 0 else ""
    return f'<span class="v10-kpi-delta {css_class}">{sinal}{valor:.1f}% vs período anterior</span>'


def formatar_numero(valor):
    """Formata números inteiros com separador português."""
    try:
        return f"{float(valor):,.0f}".replace(",", ".")
    except Exception:
        return "0"


def periodo_anterior_imediato(data_inicio, data_fim):
    """Retorna o período imediatamente anterior com a mesma duração."""
    dias_periodo = (data_fim - data_inicio).days + 1
    fim_anterior = data_inicio - timedelta(days=1)
    inicio_anterior = fim_anterior - timedelta(days=dias_periodo - 1)
    return inicio_anterior, fim_anterior


def resumo_vendas_periodo(df):
    """Calcula métricas base para um dataframe já filtrado."""
    if df.empty:
        return {
            "total": 0.0,
            "qtd": 0.0,
            "ticket": 0.0,
            "media_diaria": 0.0,
            "custo_total": 0.0,
            "lucro_bruto": 0.0,
            "margem_bruta": 0.0,
            "dias_com_venda": 0,
        }

    total = float(df["Valor"].sum())
    qtd = float(df["Qtd"].sum()) if "Qtd" in df.columns else float(len(df))
    vendas_diarias = df.groupby("Data")["Valor"].sum()
    custo_total = float(df["Custo_Total"].sum()) if "Custo_Total" in df.columns else 0.0
    lucro_bruto = float(df["Lucro_Bruto"].sum()) if "Lucro_Bruto" in df.columns else total - custo_total
    margem_bruta = (lucro_bruto / total * 100) if total else 0.0

    return {
        "total": total,
        "qtd": qtd,
        "ticket": total / qtd if qtd else 0.0,
        "media_diaria": float(vendas_diarias.mean()) if not vendas_diarias.empty else 0.0,
        "custo_total": custo_total,
        "lucro_bruto": lucro_bruto,
        "margem_bruta": margem_bruta,
        "dias_com_venda": int(vendas_diarias.count()),
    }


def calcular_cockpit_v10(df_filtrado, df_anterior, cost_manager, data_inicio, data_fim):
    """Constrói a síntese executiva da V10."""
    atual = resumo_vendas_periodo(df_filtrado)
    anterior = resumo_vendas_periodo(df_anterior)

    custos_ops = 0.0
    fonte_custos = "N/D"
    try:
        custos_ops, fonte_custos, _ = cost_manager.get_custos_operacionais_periodo(
            pd.Timestamp(data_inicio),
            pd.Timestamp(data_fim),
        )
        custos_ops = float(custos_ops or 0)
    except Exception:
        pass

    atual["lucro_liquido"] = atual["lucro_bruto"] - custos_ops
    atual["margem_liquida"] = (atual["lucro_liquido"] / atual["total"] * 100) if atual["total"] else 0.0
    atual["custos_ops"] = custos_ops
    atual["fonte_custos"] = fonte_custos

    if not df_filtrado.empty:
        fonte_mix = df_filtrado.groupby("Fonte")["Valor"].sum().sort_values(ascending=False)
        categoria_mix = df_filtrado.groupby("Categoria")["Valor"].sum().sort_values(ascending=False)
        produto_mix = df_filtrado.groupby("Produto")["Valor"].sum().sort_values(ascending=False)
        atual["fonte_lider"] = fonte_mix.index[0] if not fonte_mix.empty else "N/D"
        atual["categoria_lider"] = categoria_mix.index[0] if not categoria_mix.empty else "N/D"
        atual["concentracao_top5"] = (produto_mix.head(5).sum() / atual["total"] * 100) if atual["total"] else 0.0
        atual["data_max"] = df_filtrado["Data"].max()
    else:
        atual["fonte_lider"] = "N/D"
        atual["categoria_lider"] = "N/D"
        atual["concentracao_top5"] = 0.0
        atual["data_max"] = None

    deltas = {
        "total": calcular_delta(atual["total"], anterior["total"]),
        "media_diaria": calcular_delta(atual["media_diaria"], anterior["media_diaria"]),
        "ticket": calcular_delta(atual["ticket"], anterior["ticket"]),
        "lucro_bruto": calcular_delta(atual["lucro_bruto"], anterior["lucro_bruto"]),
    }

    score = 50
    score += max(min(deltas["total"], 20), -20) * 0.8
    score += max(min(atual["margem_bruta"] - 45, 20), -20) * 0.7
    score += max(min(atual["margem_liquida"], 15), -15) * 0.8
    score -= max(atual["concentracao_top5"] - 35, 0) * 0.35
    score = int(max(0, min(100, round(score))))

    if score >= 78:
        estado = "Forte"
    elif score >= 58:
        estado = "Estável"
    elif score >= 40:
        estado = "Atenção"
    else:
        estado = "Crítico"

    return {
        "atual": atual,
        "anterior": anterior,
        "deltas": deltas,
        "score": score,
        "estado": estado,
    }


def calcular_metricas_operacionais_v10(df_filtrado, loader, atual, data_inicio, data_fim):
    """Calcula sinais operacionais novos usando apenas dados já existentes no V9."""
    if df_filtrado.empty:
        return {
            "projecao_mensal": 0.0,
            "media_dia_operacao": 0.0,
            "volatilidade_diaria": 0.0,
            "dias_sem_venda": 0,
            "cobertura_break_even": 0.0,
            "break_even_diario": 0.0,
            "linhas_abaixo_objetivo": 0,
        }

    calendario = pd.date_range(pd.Timestamp(data_inicio), pd.Timestamp(data_fim), freq="D")
    dias_operacao = max(1, sum(1 for dia in calendario if dia.weekday() < 6))
    dias_com_venda = df_filtrado["Data"].dt.normalize().nunique()

    vendas_diarias = df_filtrado.groupby("Data")["Valor"].sum().sort_index()
    media_diaria = float(vendas_diarias.mean()) if not vendas_diarias.empty else 0.0
    desvio_diario = float(vendas_diarias.std()) if len(vendas_diarias) > 1 else 0.0
    volatilidade_diaria = (desvio_diario / media_diaria * 100) if media_diaria else 0.0

    data_ref = pd.Timestamp(data_fim)
    projecao_mensal = float(atual["media_diaria"] * data_ref.days_in_month)

    break_even_diario = 0.0
    try:
        break_even = loader.get_analise_break_even(df_filtrado)
        break_even_diario = float(break_even.get("vendas_break_even_diaria") or 0)
    except Exception:
        break_even_diario = 0.0

    if not np.isfinite(break_even_diario):
        break_even_diario = 0.0

    cobertura_break_even = (atual["media_diaria"] / break_even_diario * 100) if break_even_diario else 0.0

    if "Abaixo_Objetivo" in df_filtrado.columns:
        linhas_abaixo_objetivo = int(df_filtrado["Abaixo_Objetivo"].fillna(False).sum())
    elif {"Margem_Bruta_Pct", "Margem_Objetivo"}.issubset(df_filtrado.columns):
        linhas_abaixo_objetivo = int((df_filtrado["Margem_Bruta_Pct"] < df_filtrado["Margem_Objetivo"]).fillna(False).sum())
    else:
        linhas_abaixo_objetivo = 0

    return {
        "projecao_mensal": projecao_mensal,
        "media_dia_operacao": float(atual["total"] / dias_operacao) if dias_operacao else 0.0,
        "volatilidade_diaria": volatilidade_diaria,
        "dias_sem_venda": max(0, len(calendario) - dias_com_venda),
        "cobertura_break_even": cobertura_break_even,
        "break_even_diario": break_even_diario,
        "linhas_abaixo_objetivo": linhas_abaixo_objetivo,
    }


def carregar_fornecedores_v10(loader, data_inicio, data_fim):
    """Carrega síntese de fornecedores para o cockpit."""
    resultado = {
        "novadis_total": 0.0,
        "novadis_pedidos": 0,
        "delta_total": 0.0,
        "delta_faturas": 0,
        "df_novadis": pd.DataFrame(),
        "df_delta": pd.DataFrame(),
    }

    try:
        df_novadis = loader.carregar_novadis_processado(data_inicio, data_fim)
        if not df_novadis.empty:
            resultado["df_novadis"] = df_novadis
            resultado["novadis_total"] = float(df_novadis["custo_total"].sum())
            resultado["novadis_pedidos"] = int(df_novadis["numero_pedido"].nunique())
    except Exception:
        pass

    try:
        df_delta = loader.carregar_delta(data_inicio, data_fim)
        if not df_delta.empty:
            resultado["df_delta"] = df_delta
            resultado["delta_total"] = float(df_delta["Total_EUR"].sum())
            resultado["delta_faturas"] = int(df_delta["Numero_Fatura"].nunique())
    except Exception:
        pass

    return resultado


def metric_card_v10(label, value, delta=None, help_text=None):
    """Renderiza um cartão KPI compacto."""
    delta_html = formatar_delta_html(delta) if delta is not None else ""
    help_html = f'<div class="v10-subtitle">{escape(str(help_text))}</div>' if help_text else ""
    html = (
        '<div class="v10-kpi">'
        f'<div class="v10-kpi-label">{escape(str(label))}</div>'
        f'<div class="v10-kpi-value">{escape(str(value))}</div>'
        f"{delta_html}"
        f"{help_html}"
        "</div>"
    )
    st.markdown(html, unsafe_allow_html=True)


def render_mapa_modulos_v10():
    """Mostra a organização funcional do V10 sem esconder os módulos do V9."""
    modulos = [
        ("Painel", "Visão executiva", "Cockpit, visão geral e performance para perceber o dia num minuto."),
        ("Comercial", "Vendas e produtos", "Categorias, rankings, benchmark, ticket médio e detalhe exportável."),
        ("Rentabilidade", "Custos e margem", "Custos reais, margem por categoria/produto, break-even e lucro real."),
        ("Operação", "Fornecedores e produção", "Santa Casa, Novadis, Delta, forecast, encomendas e fichas técnicas."),
    ]
    html = ['<div class="v10-module-map">']
    for kicker, title, copy in modulos:
        html.append(
            '<div class="v10-module">'
            f'<div class="v10-module-kicker">{escape(kicker)}</div>'
            f'<div class="v10-module-title">{escape(title)}</div>'
            f'<div class="v10-module-copy">{escape(copy)}</div>'
            "</div>"
        )
    html.append("</div>")
    st.markdown("".join(html), unsafe_allow_html=True)


def render_sidebar_brand_v10():
    st.sidebar.markdown(
        (
            '<div class="v10-shell-brand">'
            '<div class="v10-shell-brand-title">Café Martins v10</div>'
            '<div class="v10-shell-brand-subtitle">V9 completo, reorganizado por decisão operacional</div>'
            "</div>"
        ),
        unsafe_allow_html=True,
    )


def paginas_por_area_v10():
    return {
        "Painel": [
            "🏠 Cockpit Executivo",
            "📊 Visão Geral",
            "🎯 Performance & KPIs",
        ],
        "Comercial": [
            "☕ Análise por Categoria",
            "💰 Análise Financeira",
            "📊 Comparação & Benchmarks",
            "📑 Dados Detalhados",
        ],
        "Rentabilidade": [
            "💸 Análise de Custos",
            "📊 Rentabilidade & Margens",
            "🎯 Break-Even Analysis",
            "💰 Margens de Lucro",
        ],
        "Operação": [
            "🎰 Jogos Santa Casa",
            "📦 Forecast & Encomendas",
            "📦 Encomendas Novadis",
            "📦 Encomendas Delta",
            "📋 Fichas Técnicas",
        ],
    }


def titulo_pagina_v10(label):
    nomes = {
        "🏠 Cockpit Executivo": "Cockpit Executivo",
        "📊 Visão Geral": "Visão Geral",
        "🎯 Performance & KPIs": "Performance & KPIs",
        "☕ Análise por Categoria": "Análise por Categoria",
        "💰 Análise Financeira": "Análise Financeira",
        "📊 Comparação & Benchmarks": "Comparação & Benchmarks",
        "📑 Dados Detalhados": "Dados Detalhados",
        "💸 Análise de Custos": "Análise de Custos",
        "📊 Rentabilidade & Margens": "Rentabilidade & Margens",
        "🎯 Break-Even Analysis": "Break-Even",
        "💰 Margens de Lucro": "Margens de Lucro",
        "🎰 Jogos Santa Casa": "Jogos Santa Casa",
        "📦 Forecast & Encomendas": "Forecast & Encomendas",
        "📦 Encomendas Novadis": "Encomendas Novadis",
        "📦 Encomendas Delta": "Encomendas Delta",
        "📋 Fichas Técnicas": "Fichas Técnicas",
    }
    return nomes.get(label, label)


def acoes_cockpit_v10(cockpit, df_filtrado, fornecedores):
    """Gera prioridades operacionais para o cockpit."""
    atual = cockpit["atual"]
    deltas = cockpit["deltas"]
    acoes = []

    if deltas["total"] < -5:
        acoes.append(("high", f"Vendas abaixo do período anterior ({deltas['total']:.1f}%). Rever mix diário e categorias em queda."))
    elif deltas["total"] > 5:
        acoes.append(("low", f"Vendas em crescimento ({deltas['total']:.1f}%). Proteger stock dos produtos que estão a puxar a subida."))

    if atual["margem_bruta"] < 35:
        acoes.append(("high", f"Margem bruta baixa ({atual['margem_bruta']:.1f}%). Priorizar revisão de preço/custo nos produtos com maior volume."))
    elif atual["margem_bruta"] < 50:
        acoes.append(("medium", f"Margem bruta moderada ({atual['margem_bruta']:.1f}%). Há espaço para afinar custos e fichas técnicas."))

    if atual["concentracao_top5"] > 45:
        acoes.append(("medium", f"Top 5 produtos concentram {atual['concentracao_top5']:.1f}% das vendas. Monitorizar ruturas e dependência."))

    total_fornecedores = fornecedores["novadis_total"] + fornecedores["delta_total"]
    if total_fornecedores > 0 and atual["total"] > 0 and (total_fornecedores / atual["total"]) > 0.35:
        acoes.append(("medium", f"Compras Novadis/Delta equivalem a {(total_fornecedores / atual['total'] * 100):.1f}% das vendas do período. Validar margens reais."))

    if "Abaixo_Objetivo" in df_filtrado.columns:
        abaixo = int(df_filtrado["Abaixo_Objetivo"].fillna(False).sum())
        if abaixo:
            acoes.append(("medium", f"{abaixo} linhas estão abaixo da margem objetivo. Abrir Rentabilidade & Margens para detalhe."))

    if not acoes:
        acoes.append(("low", "Indicadores principais equilibrados. Manter acompanhamento diário de vendas, margem e fornecedores."))

    return acoes[:5]


def pagina_cockpit_executivo_v10(df_filtrado, df_anterior, df_completo, loader, cost_manager, data_inicio, data_fim):
    """Página inicial V10: cockpit decisional."""
    cockpit = calcular_cockpit_v10(df_filtrado, df_anterior, cost_manager, data_inicio, data_fim)
    fornecedores = carregar_fornecedores_v10(loader, data_inicio, data_fim)
    atual = cockpit["atual"]
    deltas = cockpit["deltas"]
    operacionais = calcular_metricas_operacionais_v10(df_filtrado, loader, atual, data_inicio, data_fim)

    data_max = atual["data_max"].strftime("%d/%m/%Y") if atual["data_max"] is not None else "N/D"
    topbar_html = (
        '<div class="v10-topbar">'
        '<h1 class="v10-title">Cockpit Executivo</h1>'
        f'<div class="v10-subtitle">Café Martins · {data_inicio.strftime("%d/%m/%Y")} a {data_fim.strftime("%d/%m/%Y")} · dados até {data_max}</div>'
        '<div class="v10-chip-row">'
        f'<span class="v10-chip">Estado: {escape(cockpit["estado"])}</span>'
        f'<span class="v10-chip">Fonte líder: {escape(str(atual["fonte_lider"]))}</span>'
        f'<span class="v10-chip">Categoria líder: {escape(str(atual["categoria_lider"]))}</span>'
        f'<span class="v10-chip">Custos: {escape(str(atual["fonte_custos"]))}</span>'
        "</div>"
        "</div>"
    )
    st.markdown(topbar_html, unsafe_allow_html=True)

    render_mapa_modulos_v10()

    k1, k2, k3, k4, k5 = st.columns(5)
    with k1:
        metric_card_v10("Vendas", formatar_moeda(atual["total"]), deltas["total"], f"{formatar_numero(atual['qtd'])} unidades")
    with k2:
        metric_card_v10("Média diária", formatar_moeda(atual["media_diaria"]), deltas["media_diaria"], f"{atual['dias_com_venda']} dias com venda")
    with k3:
        metric_card_v10("Ticket médio", formatar_moeda(atual["ticket"]), deltas["ticket"])
    with k4:
        metric_card_v10("Lucro bruto", formatar_moeda(atual["lucro_bruto"]), deltas["lucro_bruto"], f"margem {atual['margem_bruta']:.1f}%")
    with k5:
        metric_card_v10("Lucro líquido", formatar_moeda(atual["lucro_liquido"]), None, f"custos ops {formatar_moeda(atual['custos_ops'])}")

    s1, s2, s3, s4 = st.columns(4)
    with s1:
        metric_card_v10(
            "Run-rate mensal",
            formatar_moeda(operacionais["projecao_mensal"]),
            None,
            f"média atual {formatar_moeda(atual['media_diaria'])}/dia",
        )
    with s2:
        metric_card_v10(
            "Cobertura break-even",
            f"{operacionais['cobertura_break_even']:.0f}%",
            None,
            f"necessário {formatar_moeda(operacionais['break_even_diario'])}/dia",
        )
    with s3:
        metric_card_v10(
            "Volatilidade diária",
            f"{operacionais['volatilidade_diaria']:.1f}%",
            None,
            "desvio das vendas diárias",
        )
    with s4:
        metric_card_v10(
            "Alertas de margem",
            formatar_numero(operacionais["linhas_abaixo_objetivo"]),
            None,
            f"{operacionais['dias_sem_venda']} dias sem venda no período",
        )

    st.markdown("")

    c1, c2, c3 = st.columns([1.1, 1.7, 1.2])
    with c1:
        st.markdown('<div class="v10-panel">', unsafe_allow_html=True)
        st.markdown('<div class="v10-panel-title">Saúde Operacional</div>', unsafe_allow_html=True)
        score_html = (
            '<div class="v10-score">'
            f'<div class="v10-score-number">{cockpit["score"]}</div>'
            f'<div class="v10-score-label">/ 100 · {escape(cockpit["estado"])}</div>'
            "</div>"
        )
        st.markdown(score_html, unsafe_allow_html=True)
        st.progress(cockpit["score"] / 100)
        st.caption(f"Top 5 produtos: {atual['concentracao_top5']:.1f}% das vendas")
        st.caption(f"Margem líquida estimada: {atual['margem_liquida']:.1f}%")
        st.markdown("</div>", unsafe_allow_html=True)

    with c2:
        vendas_diarias = df_filtrado.groupby("Data")["Valor"].sum().reset_index()
        vendas_diarias["Media_7d"] = vendas_diarias["Valor"].rolling(7, min_periods=1).mean()
        fig = go.Figure()
        fig.add_trace(go.Bar(
            x=vendas_diarias["Data"],
            y=vendas_diarias["Valor"],
            name="Vendas",
            marker_color="#b8c7d6",
            hovertemplate="%{x|%d/%m/%Y}<br>%{y:,.0f}€<extra></extra>",
        ))
        fig.add_trace(go.Scatter(
            x=vendas_diarias["Data"],
            y=vendas_diarias["Media_7d"],
            name="Média 7d",
            mode="lines",
            line=dict(color="#1b5f8f", width=3),
            hovertemplate="%{x|%d/%m/%Y}<br>%{y:,.0f}€<extra></extra>",
        ))
        fig.update_layout(
            title="Ritmo diário",
            height=360,
            margin=dict(l=10, r=10, t=45, b=10),
            hovermode="x unified",
            legend=dict(orientation="h", y=1.08, x=0.01),
        )
        st.plotly_chart(fig, use_container_width=True, key="v10_cockpit_ritmo_diario")

    with c3:
        st.markdown('<div class="v10-panel">', unsafe_allow_html=True)
        st.markdown('<div class="v10-panel-title">Prioridades</div>', unsafe_allow_html=True)
        for nivel, texto in acoes_cockpit_v10(cockpit, df_filtrado, fornecedores):
            st.markdown(
                f'<div class="v10-action v10-action-{nivel}">{escape(texto)}</div>',
                unsafe_allow_html=True,
            )
        st.markdown("</div>", unsafe_allow_html=True)

    col_mix, col_rank = st.columns([1.25, 1])
    with col_mix:
        mix = df_filtrado.groupby(["Fonte", "Categoria"], as_index=False)["Valor"].sum()
        fig_mix = px.treemap(
            mix,
            path=["Fonte", "Categoria"],
            values="Valor",
            color="Valor",
            color_continuous_scale="Blues",
            title="Mix de vendas por fonte e categoria",
        )
        fig_mix.update_layout(height=430, margin=dict(l=5, r=5, t=45, b=5))
        st.plotly_chart(fig_mix, use_container_width=True, key="v10_cockpit_mix")

    with col_rank:
        top_produtos = df_filtrado.groupby("Produto").agg(
            Vendas=("Valor", "sum"),
            Qtd=("Qtd", "sum"),
            Lucro=("Lucro_Bruto", "sum") if "Lucro_Bruto" in df_filtrado.columns else ("Valor", "sum"),
        ).reset_index().sort_values("Vendas", ascending=False).head(12)
        top_produtos["Margem"] = np.where(top_produtos["Vendas"] > 0, top_produtos["Lucro"] / top_produtos["Vendas"] * 100, 0)
        st.markdown("#### Top produtos")
        st.dataframe(
            top_produtos,
            hide_index=True,
            use_container_width=True,
            height=430,
            column_config={
                "Produto": st.column_config.TextColumn("Produto", width="medium"),
                "Vendas": st.column_config.NumberColumn("Vendas", format="%.0f€"),
                "Qtd": st.column_config.NumberColumn("Qtd", format="%.0f"),
                "Lucro": st.column_config.NumberColumn("Lucro", format="%.0f€"),
                "Margem": st.column_config.ProgressColumn("Margem", format="%.1f%%", min_value=0, max_value=100),
            },
        )

    f1, f2, f3 = st.columns(3)
    with f1:
        st.metric("Novadis", formatar_moeda(fornecedores["novadis_total"]), f"{fornecedores['novadis_pedidos']} pedidos")
    with f2:
        st.metric("Delta", formatar_moeda(fornecedores["delta_total"]), f"{fornecedores['delta_faturas']} faturas")
    with f3:
        total_fornecedores = fornecedores["novadis_total"] + fornecedores["delta_total"]
        ratio = (total_fornecedores / atual["total"] * 100) if atual["total"] else 0
        st.metric("Compras / Vendas", f"{ratio:.1f}%", formatar_moeda(total_fornecedores))

    if not df_anterior.empty:
        atual_prod = df_filtrado.groupby("Produto")["Valor"].sum()
        anterior_prod = df_anterior.groupby("Produto")["Valor"].sum()
        movimentos = pd.concat([atual_prod, anterior_prod], axis=1).fillna(0)
        movimentos.columns = ["Atual", "Anterior"]
        movimentos["Diferença"] = movimentos["Atual"] - movimentos["Anterior"]
        movimentos = movimentos.reset_index()
        mov_up = movimentos.sort_values("Diferença", ascending=False).head(8)
        mov_down = movimentos.sort_values("Diferença", ascending=True).head(8)

        m1, m2 = st.columns(2)
        with m1:
            fig_up = px.bar(
                mov_up,
                x="Diferença",
                y="Produto",
                orientation="h",
                title="Maiores contributos positivos",
                text=[formatar_moeda(v) for v in mov_up["Diferença"]],
                color_discrete_sequence=["#12805c"],
            )
            fig_up.update_layout(height=360, yaxis={"categoryorder": "total ascending"}, margin=dict(l=5, r=5, t=45, b=5))
            st.plotly_chart(fig_up, use_container_width=True, key="v10_cockpit_movers_up")

        with m2:
            fig_down = px.bar(
                mov_down,
                x="Diferença",
                y="Produto",
                orientation="h",
                title="Maiores quedas",
                text=[formatar_moeda(v) for v in mov_down["Diferença"]],
                color_discrete_sequence=["#b42318"],
            )
            fig_down.update_layout(height=360, yaxis={"categoryorder": "total descending"}, margin=dict(l=5, r=5, t=45, b=5))
            st.plotly_chart(fig_down, use_container_width=True, key="v10_cockpit_movers_down")


def filtrar_dados(df, categorias, subcategorias, fontes, data_inicio, data_fim, excluir_domingos):
    """Aplica filtros aos dados"""
    df_filtrado = df.copy()

    # Filtro de data
    df_filtrado = df_filtrado[
        (df_filtrado['Data'] >= pd.Timestamp(data_inicio)) &
        (df_filtrado['Data'] <= pd.Timestamp(data_fim))
    ]

    # Filtro de categoria
    if categorias:
        df_filtrado = df_filtrado[df_filtrado['Categoria'].isin(categorias)]

    # Filtro de subcategoria
    if subcategorias:
        df_filtrado = df_filtrado[df_filtrado['Subcategoria'].isin(subcategorias)]

    # Filtro de fonte
    if fontes:
        df_filtrado = df_filtrado[df_filtrado['Fonte'].isin(fontes)]

    # Filtro de Segunda a Sábado (Exclui Domingos)
    if excluir_domingos:
        df_filtrado = df_filtrado[df_filtrado['Dia_Semana'] < 6]

    return df_filtrado


def calcular_metricas_periodo(df, periodo_dias):
    """Calcula métricas para um período específico"""
    data_fim = df['Data'].max()
    data_inicio = data_fim - timedelta(days=periodo_dias)

    df_periodo = df[df['Data'] >= data_inicio]

    return {
        'total': df_periodo['Valor'].sum(),
        'media_diaria': df_periodo.groupby('Data')['Valor'].sum().mean(),
        'quantidade': df_periodo['Qtd'].sum() if 'Qtd' in df_periodo.columns else 0,
        'vendas': len(df_periodo)
    }


def calcular_metricas_performance(df_filtrado, df_completo=None):
    """
    Calcula métricas de performance semanal, mensal e anual com comparações
    Funciona corretamente mesmo com períodos filtrados

    Args:
        df_filtrado: DataFrame com dados filtrados (período selecionado)
        df_completo: DataFrame com todos os dados (para comparações históricas)

    Returns:
        dict com métricas de performance
    """
    if df_filtrado.empty:
        return {}

    # Se não houver dataframe completo, usar o filtrado
    df_comparacao = df_completo if df_completo is not None else df_filtrado

    data_hoje = df_filtrado['Data'].max()

    # ========== PERFORMANCE SEMANAL ==========
    # Última semana (últimos 7 dias) - usar dados filtrados
    data_inicio_semana_atual = data_hoje - timedelta(days=6)
    df_semana_atual = df_filtrado[df_filtrado['Data'] >= data_inicio_semana_atual]
    vendas_semana_atual = df_semana_atual['Valor'].sum()

    # Semana anterior (7 dias antes) - usar dados completos
    data_fim_semana_anterior = data_inicio_semana_atual - timedelta(days=1)
    data_inicio_semana_anterior = data_fim_semana_anterior - timedelta(days=6)
    df_semana_anterior = df_comparacao[(df_comparacao['Data'] >= data_inicio_semana_anterior) & (df_comparacao['Data'] <= data_fim_semana_anterior)]
    vendas_semana_anterior = df_semana_anterior['Valor'].sum() if not df_semana_anterior.empty else None
    var_semana_vs_anterior = calcular_delta(vendas_semana_atual, vendas_semana_anterior) if vendas_semana_anterior is not None else 0

    # Mesma semana ano passado - usar dados completos
    data_inicio_semana_ano_passado = data_inicio_semana_atual - timedelta(days=365)
    data_fim_semana_ano_passado = data_hoje - timedelta(days=365)
    df_semana_ano_passado = df_comparacao[(df_comparacao['Data'] >= data_inicio_semana_ano_passado) & (df_comparacao['Data'] <= data_fim_semana_ano_passado)]
    vendas_semana_ano_passado = df_semana_ano_passado['Valor'].sum() if not df_semana_ano_passado.empty else None
    var_semana_vs_ano_passado = calcular_delta(vendas_semana_atual, vendas_semana_ano_passado) if vendas_semana_ano_passado is not None else 0

    # ========== PERFORMANCE MENSAL (MTD - Month To Date) ==========
    # Mês em curso (até hoje) - usar dados filtrados
    primeiro_dia_mes = data_hoje.replace(day=1)
    df_mes_atual = df_filtrado[df_filtrado['Data'] >= primeiro_dia_mes]
    vendas_mes_atual = df_mes_atual['Valor'].sum()
    dias_mes_atual = (data_hoje - primeiro_dia_mes).days + 1

    # Mês anterior (mesmos dias) - usar dados completos
    ultimo_dia_mes_anterior = primeiro_dia_mes - timedelta(days=1)
    data_inicio_mes_anterior = ultimo_dia_mes_anterior.replace(day=1)
    data_fim_mes_anterior = data_inicio_mes_anterior + timedelta(days=dias_mes_atual - 1)
    df_mes_anterior = df_comparacao[(df_comparacao['Data'] >= data_inicio_mes_anterior) & (df_comparacao['Data'] <= data_fim_mes_anterior)]
    vendas_mes_anterior = df_mes_anterior['Valor'].sum() if not df_mes_anterior.empty else None
    var_mes_vs_anterior = calcular_delta(vendas_mes_atual, vendas_mes_anterior) if vendas_mes_anterior is not None else 0

    # Mesmo mês ano passado (até o mesmo dia) - usar dados completos
    data_inicio_smly = primeiro_dia_mes - timedelta(days=365)
    data_fim_smly = primeiro_dia_mes + timedelta(days=dias_mes_atual - 1) - timedelta(days=365)
    df_smly = df_comparacao[(df_comparacao['Data'] >= data_inicio_smly) & (df_comparacao['Data'] <= data_fim_smly)]
    vendas_smly = df_smly['Valor'].sum() if not df_smly.empty else None
    var_mes_vs_ano_passado = calcular_delta(vendas_mes_atual, vendas_smly) if vendas_smly is not None else 0

    # ========== PERFORMANCE ANUAL (YTD - Year To Date) ==========
    # Ano em curso (até hoje) - usar dados COMPLETOS para YTD correto
    primeiro_dia_ano = data_hoje.replace(month=1, day=1)
    df_ano_atual = df_comparacao[df_comparacao['Data'] >= primeiro_dia_ano]
    vendas_ano_atual = df_ano_atual['Valor'].sum()

    # Ano anterior (até mesma data) - usar dados completos
    data_ytd_ano_anterior = data_hoje - timedelta(days=365)
    primeiro_dia_ano_anterior = data_ytd_ano_anterior.replace(month=1, day=1)
    df_ano_anterior_ytd = df_comparacao[(df_comparacao['Data'] >= primeiro_dia_ano_anterior) & (df_comparacao['Data'] <= data_ytd_ano_anterior)]
    vendas_ano_anterior_ytd = df_ano_anterior_ytd['Valor'].sum() if not df_ano_anterior_ytd.empty else None
    var_ano_vs_ano_anterior = calcular_delta(vendas_ano_atual, vendas_ano_anterior_ytd) if vendas_ano_anterior_ytd is not None else 0

    # Total ano anterior (completo) - usar dados completos
    data_inicio_ano_anterior_total = primeiro_dia_ano_anterior
    data_fim_ano_anterior_total = data_ytd_ano_anterior.replace(month=12, day=31)
    df_ano_anterior_total = df_comparacao[(df_comparacao['Data'] >= data_inicio_ano_anterior_total) & (df_comparacao['Data'] <= data_fim_ano_anterior_total)]
    vendas_ano_anterior_total = df_ano_anterior_total['Valor'].sum() if not df_ano_anterior_total.empty else None

    return {
        'data_atualizacao': data_hoje.strftime('%d/%m/%Y'),
        # Semanal
        'vendas_semana_atual': vendas_semana_atual,
        'vendas_semana_anterior': vendas_semana_anterior if vendas_semana_anterior is not None else 0,
        'vendas_semana_ano_passado': vendas_semana_ano_passado if vendas_semana_ano_passado is not None else 0,
        'var_semana_vs_anterior': var_semana_vs_anterior,
        'var_semana_vs_ano_passado': var_semana_vs_ano_passado,
        # Mensal
        'vendas_mes_atual': vendas_mes_atual,
        'vendas_mes_anterior': vendas_mes_anterior if vendas_mes_anterior is not None else 0,
        'vendas_smly': vendas_smly if vendas_smly is not None else 0,
        'var_mes_vs_anterior': var_mes_vs_anterior,
        'var_mes_vs_ano_passado': var_mes_vs_ano_passado,
        'dias_mes_atual': dias_mes_atual,
        # Anual
        'vendas_ano_atual': vendas_ano_atual,
        'vendas_ano_anterior_ytd': vendas_ano_anterior_ytd if vendas_ano_anterior_ytd is not None else 0,
        'vendas_ano_anterior_total': vendas_ano_anterior_total if vendas_ano_anterior_total is not None else 0,
        'var_ano_vs_ano_anterior': var_ano_vs_ano_anterior,
    }


def criar_grafico_pizza(df, coluna, titulo, mostrar_percentagem=True):
    """Cria gráfico de pizza profissional"""
    dados = df.groupby(coluna)['Valor'].sum().sort_values(ascending=False)

    fig = go.Figure(data=[go.Pie(
        labels=dados.index,
        values=dados.values,
        hole=0.4,
        textposition='auto',
        textinfo='label+percent' if mostrar_percentagem else 'label+value',
        marker=dict(line=dict(color='white', width=2))
    )])

    fig.update_layout(
        title=titulo,
        showlegend=True,
        height=400,
        margin=dict(t=50, b=20, l=20, r=20)
    )

    return fig


def criar_grafico_barras(df, x, y, titulo, orientacao='v', cor=None):
    """Cria gráfico de barras profissional"""
    fig = px.bar(
        df,
        x=x,
        y=y,
        title=titulo,
        orientation=orientacao,
        color=cor,
        text_auto='.2f'
    )

    fig.update_layout(
        height=400,
        showlegend=True if cor else False,
        margin=dict(t=50, b=50, l=50, r=50),
        xaxis_title=x,
        yaxis_title=y
    )

    fig.update_traces(texttemplate='%{text:.2f}€', textposition='outside')

    return fig


def criar_grafico_linha(df, x, y, titulo, cor=None):
    """Cria gráfico de linha profissional"""
    fig = px.line(
        df,
        x=x,
        y=y,
        title=titulo,
        color=cor,
        markers=True
    )

    fig.update_layout(
        height=400,
        showlegend=True if cor else False,
        margin=dict(t=50, b=50, l=50, r=50),
        hovermode='x unified'
    )

    return fig


def criar_heatmap(df, titulo):
    """Cria heatmap profissional"""
    fig = go.Figure(data=go.Heatmap(
        z=df.values,
        x=df.columns,
        y=df.index,
        colorscale='Blues',
        text=df.values,
        texttemplate='%{text:.0f}€',
        textfont={"size": 10},
        colorbar=dict(title="Vendas (€)")
    ))

    fig.update_layout(
        title=titulo,
        height=500,
        margin=dict(t=50, b=50, l=100, r=50)
    )

    return fig


def criar_grafico_waterfall(valores, labels, titulo):
    """Cria gráfico waterfall"""
    fig = go.Figure(go.Waterfall(
        name="Contribuição",
        orientation="v",
        measure=["relative"] * (len(valores) - 1) + ["total"],
        x=labels,
        textposition="outside",
        text=[f"{v:,.0f}€" for v in valores],
        y=valores,
        connector={"line": {"color": "rgb(63, 63, 63)"}},
    ))

    fig.update_layout(
        title=titulo,
        showlegend=False,
        height=500,
        margin=dict(t=50, b=100, l=50, r=50)
    )

    return fig


def gerar_relatorio_excel_avancado(df_filtrado, data_inicio, data_fim, categorias_selecionadas, metricas_financeiras, cost_manager):
    """
    Gera relatório Excel avançado com múltiplas abas, gráficos e formatação profissional
    Adaptado para o contexto do Café Martins com análise de vendas, custos e rentabilidade
    """
    output = BytesIO()
    wb = openpyxl.Workbook()

    # Estilos de formatação
    header_fill = PatternFill(start_color="1f77b4", end_color="1f77b4", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=12)
    border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )

    # ===== ABA 1: RESUMO EXECUTIVO =====
    ws_resumo = wb.active
    ws_resumo.title = "Resumo Executivo"

    # Título
    ws_resumo['A1'] = "RELATÓRIO DE ANÁLISE - CAFÉ MARTINS"
    ws_resumo['A1'].font = Font(bold=True, size=16, color="1f77b4")
    ws_resumo.merge_cells('A1:E1')

    # Informações do relatório
    ws_resumo['A3'] = "Período de Análise:"
    ws_resumo['B3'] = f"{pd.to_datetime(data_inicio).strftime('%d/%m/%Y')} até {pd.to_datetime(data_fim).strftime('%d/%m/%Y')}"
    ws_resumo['A4'] = "Data de Geração:"
    ws_resumo['B4'] = datetime.now().strftime('%d/%m/%Y %H:%M:%S')
    ws_resumo['A5'] = "Categorias Analisadas:"
    ws_resumo['B5'] = ", ".join(categorias_selecionadas) if categorias_selecionadas else "Todas"

    # Métricas principais
    ws_resumo['A7'] = "MÉTRICAS PRINCIPAIS"
    ws_resumo['A7'].font = Font(bold=True, size=14, color="1f77b4")
    ws_resumo.merge_cells('A7:E7')

    ws_resumo['A9'] = "Total de Vendas (€)"
    ws_resumo['B9'] = df_filtrado['Valor'].sum()
    ws_resumo['B9'].number_format = '#,##0 €'

    ws_resumo['A10'] = "Total de Transações"
    ws_resumo['B10'] = len(df_filtrado)

    ws_resumo['A11'] = "Média Diária (€)"
    ws_resumo['B11'] = df_filtrado.groupby('Data')['Valor'].sum().mean()
    ws_resumo['B11'].number_format = '#,##0 €'

    ws_resumo['A12'] = "Ticket Médio (€)"
    total_qtd = df_filtrado['Qtd'].sum() if 'Qtd' in df_filtrado.columns else 1
    ws_resumo['B12'] = df_filtrado['Valor'].sum() / total_qtd if total_qtd > 0 else 0
    ws_resumo['B12'].number_format = '#,##0.00 €'

    # Métricas Financeiras
    ws_resumo['A14'] = "ANÁLISE FINANCEIRA"
    ws_resumo['A14'].font = Font(bold=True, size=14, color="1f77b4")
    ws_resumo.merge_cells('A14:E14')

    ws_resumo['A16'] = "Receita Total (€)"
    ws_resumo['B16'] = metricas_financeiras.get('receita_total', 0)
    ws_resumo['B16'].number_format = '#,##0 €'

    ws_resumo['A17'] = "Custos Totais (€)"
    ws_resumo['B17'] = metricas_financeiras.get('custo_total', 0)
    ws_resumo['B17'].number_format = '#,##0 €'

    ws_resumo['A18'] = "Lucro Líquido (€)"
    ws_resumo['B18'] = metricas_financeiras.get('lucro_liquido', 0)
    ws_resumo['B18'].number_format = '#,##0 €'

    ws_resumo['A19'] = "Margem Líquida (%)"
    ws_resumo['B19'] = metricas_financeiras.get('margem_liquida_pct', 0) / 100
    ws_resumo['B19'].number_format = '0.00%'

    # Ajustar largura das colunas
    ws_resumo.column_dimensions['A'].width = 30
    ws_resumo.column_dimensions['B'].width = 25

    # ===== ABA 2: DADOS DETALHADOS =====
    ws_dados = wb.create_sheet("Dados Detalhados")

    # Preparar dados para exportação
    df_export = df_filtrado.copy()
    colunas_exportar = ['Data', 'Produto', 'Categoria', 'Subcategoria', 'Valor', 'Qtd', 'Fonte']
    colunas_exportar = [col for col in colunas_exportar if col in df_export.columns]
    df_export = df_export[colunas_exportar]

    if 'Data' in df_export.columns:
        df_export['Data'] = df_export['Data'].dt.strftime('%d/%m/%Y')

    # Adicionar cabeçalhos
    headers = list(df_export.columns)
    for col_num, header in enumerate(headers, 1):
        cell = ws_dados.cell(row=1, column=col_num, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.border = border
        cell.alignment = Alignment(horizontal='center', vertical='center')

    # Adicionar dados
    for r_idx, row in enumerate(dataframe_to_rows(df_export, index=False, header=False), 2):
        for c_idx, value in enumerate(row, 1):
            cell = ws_dados.cell(row=r_idx, column=c_idx, value=value)
            cell.border = border
            if 'Valor' in headers and c_idx == headers.index('Valor') + 1:
                cell.number_format = '#,##0.00 €'

    # Ajustar largura das colunas
    for column in ws_dados.columns:
        max_length = 0
        column_letter = column[0].column_letter
        for cell in column:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = min(max_length + 2, 50)
        ws_dados.column_dimensions[column_letter].width = adjusted_width

    # ===== ABA 3: ANÁLISE POR CATEGORIA =====
    ws_categorias = wb.create_sheet("Análise por Categoria")

    # Criar tabela de resumo por categoria
    resumo_cat = df_filtrado.groupby('Categoria').agg({
        'Valor': ['sum', 'mean', 'count']
    }).reset_index()
    resumo_cat.columns = ['Categoria', 'Total (€)', 'Média (€)', 'Nº Vendas']
    resumo_cat = resumo_cat.sort_values('Total (€)', ascending=False)

    # Cabeçalhos
    ws_categorias['A1'] = "ANÁLISE POR CATEGORIA"
    ws_categorias['A1'].font = Font(bold=True, size=14, color="1f77b4")
    ws_categorias.merge_cells('A1:D1')

    for col_num, header in enumerate(resumo_cat.columns, 1):
        cell = ws_categorias.cell(row=3, column=col_num, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.border = border
        cell.alignment = Alignment(horizontal='center', vertical='center')

    # Dados
    for r_idx, row in enumerate(dataframe_to_rows(resumo_cat, index=False, header=False), 4):
        for c_idx, value in enumerate(row, 1):
            cell = ws_categorias.cell(row=r_idx, column=c_idx, value=value)
            cell.border = border
            if c_idx in [2, 3]:
                cell.number_format = '#,##0.00 €'

    # Gráfico de barras
    if len(resumo_cat) > 0:
        chart = BarChart()
        chart.title = "Vendas por Categoria"
        chart.style = 10
        chart.y_axis.title = 'Valor (€)'
        chart.x_axis.title = 'Categoria'

        data = Reference(ws_categorias, min_col=2, min_row=3, max_row=3+len(resumo_cat))
        cats = Reference(ws_categorias, min_col=1, min_row=4, max_row=3+len(resumo_cat))
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(cats)
        chart.height = 12
        chart.width = 20

        ws_categorias.add_chart(chart, "F3")

    # Ajustar largura
    for col in ['A', 'B', 'C', 'D']:
        ws_categorias.column_dimensions[col].width = 20

    # ===== ABA 4: RENTABILIDADE =====
    ws_rent = wb.create_sheet("Rentabilidade")

    ws_rent['A1'] = "ANÁLISE DE RENTABILIDADE"
    ws_rent['A1'].font = Font(bold=True, size=14, color="1f77b4")
    ws_rent.merge_cells('A1:E1')

    # Análise de rentabilidade por categoria
    analise_rent = cost_manager.analisar_rentabilidade_categorias(df_filtrado)

    if not analise_rent.empty:
        # Cabeçalhos
        ws_rent['A3'] = "Categoria"
        ws_rent['B3'] = "Vendas (€)"
        ws_rent['C3'] = "Custos (€)"
        ws_rent['D3'] = "Lucro Bruto (€)"
        ws_rent['E3'] = "Margem (%)"

        for col in ['A', 'B', 'C', 'D', 'E']:
            cell = ws_rent[f'{col}3']
            cell.fill = header_fill
            cell.font = header_font
            cell.border = border
            cell.alignment = Alignment(horizontal='center', vertical='center')

        # Dados
        row_num = 4
        for categoria, row in analise_rent.iterrows():
            ws_rent[f'A{row_num}'] = categoria
            ws_rent[f'B{row_num}'] = row['Valor']
            ws_rent[f'C{row_num}'] = row['Custo_Total']
            ws_rent[f'D{row_num}'] = row['Lucro_Bruto']
            ws_rent[f'E{row_num}'] = row['Margem_Bruta_Pct'] / 100

            for col in ['A', 'B', 'C', 'D', 'E']:
                ws_rent[f'{col}{row_num}'].border = border

            ws_rent[f'B{row_num}'].number_format = '#,##0 €'
            ws_rent[f'C{row_num}'].number_format = '#,##0 €'
            ws_rent[f'D{row_num}'].number_format = '#,##0 €'
            ws_rent[f'E{row_num}'].number_format = '0.00%'

            row_num += 1

        # Gráfico de margens
        if len(analise_rent) > 0:
            chart = BarChart()
            chart.title = "Margem Bruta por Categoria"
            chart.style = 11
            chart.y_axis.title = 'Margem (%)'
            chart.x_axis.title = 'Categoria'

            data = Reference(ws_rent, min_col=5, min_row=3, max_row=3+len(analise_rent))
            cats = Reference(ws_rent, min_col=1, min_row=4, max_row=3+len(analise_rent))
            chart.add_data(data, titles_from_data=True)
            chart.set_categories(cats)
            chart.height = 12
            chart.width = 20

            ws_rent.add_chart(chart, "G3")

    # Ajustar largura
    for col in ['A', 'B', 'C', 'D', 'E']:
        ws_rent.column_dimensions[col].width = 18

    # ===== ABA 5: ANÁLISE TEMPORAL =====
    ws_temporal = wb.create_sheet("Análise Temporal")

    ws_temporal['A1'] = "ANÁLISE TEMPORAL"
    ws_temporal['A1'].font = Font(bold=True, size=14, color="1f77b4")
    ws_temporal.merge_cells('A1:C1')

    # Vendas diárias
    vendas_diarias = df_filtrado.groupby('Data')['Valor'].sum().reset_index()
    vendas_diarias['Data'] = vendas_diarias['Data'].dt.strftime('%d/%m/%Y')
    vendas_diarias.columns = ['Data', 'Vendas (€)']

    # Cabeçalhos
    for col_num, header in enumerate(vendas_diarias.columns, 1):
        cell = ws_temporal.cell(row=3, column=col_num, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.border = border
        cell.alignment = Alignment(horizontal='center', vertical='center')

    # Dados
    for r_idx, row in enumerate(dataframe_to_rows(vendas_diarias, index=False, header=False), 4):
        for c_idx, value in enumerate(row, 1):
            cell = ws_temporal.cell(row=r_idx, column=c_idx, value=value)
            cell.border = border
            if c_idx == 2:
                cell.number_format = '#,##0 €'

    # Gráfico de linha
    if len(vendas_diarias) > 1:
        line_chart = LineChart()
        line_chart.title = "Evolução Temporal das Vendas"
        line_chart.style = 12
        line_chart.y_axis.title = 'Vendas (€)'
        line_chart.x_axis.title = 'Data'

        data = Reference(ws_temporal, min_col=2, min_row=3, max_row=3+len(vendas_diarias))
        cats = Reference(ws_temporal, min_col=1, min_row=4, max_row=3+len(vendas_diarias))
        line_chart.add_data(data, titles_from_data=True)
        line_chart.set_categories(cats)
        line_chart.height = 12
        line_chart.width = 20

        ws_temporal.add_chart(line_chart, "E3")

    # Ajustar largura
    ws_temporal.column_dimensions['A'].width = 20
    ws_temporal.column_dimensions['B'].width = 20

    # Salvar o workbook
    wb.save(output)
    output.seek(0)
    return output


# ============= FUNÇÕES PARA JOGOS SANTA CASA =============

@st.cache_data(ttl=1800)
def carregar_dados_santa_casa():
    """Carrega dados dos Jogos Santa Casa - cache de 30 minutos"""
    data_file = Path('/home/jorge/Documentos/Santa casa/dados/dados.csv')

    if not data_file.exists():
        st.error(f"Arquivo não encontrado: {data_file}")
        return None

    try:
        df = pd.read_csv(data_file, sep=';', encoding='utf-8')
    except Exception as e:
        st.error(f"Erro ao carregar dados: {e}")
        return None

    # Converter Data para datetime
    df['Data'] = pd.to_datetime(df['Data'], format='%d-%m-%Y', errors='coerce')

    # Remover linhas de PRESTAÇÃO DE CONTAS dos dados principais
    df = df[df['Categoria'] != 'PRESTAÇÃO DE CONTAS'].copy()

    # Converter colunas numéricas (OTIMIZADO: vectorizado)
    numeric_cols = ['Qt Maços', 'Vendas ilíquidas (€)', 'Remunerações (€)', 'Prémios (€)', 'Valor (€)']
    for col in numeric_cols:
        if col in df.columns:
            # Vectorizado: converter formato europeu para float
            df[col] = df[col].astype(str).str.strip()
            # Formato europeu: 1.234,56 -> remove pontos de milhar, troca vírgula por ponto
            mask_europeu = df[col].str.contains(',', na=False)
            df.loc[mask_europeu, col] = df.loc[mask_europeu, col].str.replace('.', '', regex=False).str.replace(',', '.', regex=False)
            df[col] = pd.to_numeric(df[col], errors='coerce').fillna(0.0)

    # Adicionar colunas temporais
    df['Ano'] = df['Data'].dt.year
    df['Mes'] = df['Data'].dt.month
    df['Ano_Mes'] = df['Data'].dt.to_period('M').astype(str)
    df['Semana'] = df['Data'].dt.isocalendar().week
    df['Trimestre'] = df['Data'].dt.quarter
    df['Dia_Semana'] = df['Data'].dt.dayofweek

    # Ordenar por data
    df = df.sort_values('Data')

    return df


@st.cache_data(ttl=3600)
def carregar_objetivos():
    """Carrega objetivos semanais do CSV (cacheado por 1 hora)."""
    objetivos_file = Path('/home/jorge/Documentos/Streamlit/objetivos_semanais.csv')

    if not objetivos_file.exists():
        # Criar arquivo padrão se não existir
        df_objetivos = pd.DataFrame({
            'Jogo': [],
            'Objetivo': []
        })
        df_objetivos.to_csv(objetivos_file, sep=';', index=False)
        return {}

    try:
        df_objetivos = pd.read_csv(objetivos_file, sep=';')
        return dict(zip(df_objetivos['Jogo'], df_objetivos['Objetivo']))
    except Exception as e:
        st.warning(f"Erro ao carregar objetivos: {e}")
        return {}


def guardar_objetivos(objetivos_dict):
    """Guarda objetivos semanais em CSV."""
    objetivos_file = Path('/home/jorge/Documentos/Streamlit/objetivos_semanais.csv')

    df_objetivos = pd.DataFrame({
        'Jogo': list(objetivos_dict.keys()),
        'Objetivo': list(objetivos_dict.values())
    })

    try:
        df_objetivos.to_csv(objetivos_file, sep=';', index=False)
        return True
    except Exception as e:
        st.error(f"Erro ao guardar objetivos: {e}")
        return False


def pagina_jogos_santa_casa(df, data_inicio=None, data_fim=None):
    """Página de análise dos Jogos Santa Casa"""
    if df is None or len(df) == 0:
        st.error("Dados não carregados")
        return

    # Filtrar dados pela data se fornecida
    df_filtrado = df.copy()
    if data_inicio is not None and data_fim is not None:
        df_filtrado = df_filtrado[
            (pd.to_datetime(df_filtrado['Data']).dt.date >= data_inicio) &
            (pd.to_datetime(df_filtrado['Data']).dt.date <= data_fim)
        ]

    st.markdown('<h2 style="text-align: center; color: #1f77b4;">🎰 Análise de Jogos - Santa Casa</h2>',
                unsafe_allow_html=True)
    st.caption("📊 Análise de vendas ilíquidas semanais dos Jogos Santa Casa")

    # KPIs principais
    col1, col2, col3, col4 = st.columns(4)

    total_vendas = df_filtrado['Vendas ilíquidas (€)'].sum()
    total_remuneracoes = df_filtrado['Remunerações (€)'].sum()
    total_premios = df_filtrado['Prémios (€)'].sum()
    num_jogos = df_filtrado['Jogo'].nunique()

    with col1:
        st.metric("Total Vendas Ilíquidas", f"{total_vendas:,.0f}€")
    with col2:
        st.metric("Total Remunerações", f"{total_remuneracoes:,.0f}€")
    with col3:
        st.metric("Total Prémios", f"{total_premios:,.0f}€")
    with col4:
        st.metric("Nº de Jogos", num_jogos)

    st.divider()

    # Tabs para análises
    tab1, tab2, tab3, tab4, tab5, tab6, tab7, tab8 = st.tabs([
        "📈 Evolução Temporal",
        "🎮 Por Jogo",
        "🎰 Por Categoria",
        "📊 Comparação Semanal",
        "📈 MoM (Mês a Mês)",
        "📅 YoY (Ano a Ano)",
        "📆 Semana a Semana",
        "🔮 Previsão"
    ])

    with tab1:
        st.subheader("Evolução de Vendas ao Longo do Tempo")

        # Dados semanais
        vendas_semanal = df_filtrado.groupby('Data').agg({
            'Vendas ilíquidas (€)': 'sum',
            'Remunerações (€)': 'sum',
            'Prémios (€)': 'sum'
        }).reset_index()

        fig = go.Figure()

        fig.add_trace(go.Scatter(
            x=vendas_semanal['Data'],
            y=vendas_semanal['Vendas ilíquidas (€)'],
            name='Vendas Ilíquidas',
            mode='lines+markers',
            line=dict(color='#1f77b4', width=2),
            fill='tozeroy',
            fillcolor='rgba(31, 119, 180, 0.2)'
        ))

        fig.update_layout(
            title="Evolução de Vendas Ilíquidas por Semana",
            xaxis_title="Data (Semana)",
            yaxis_title="Vendas (€)",
            hovermode='x unified',
            height=500
        )

        st.plotly_chart(fig, use_container_width=True)

        # Estatísticas
        st.markdown("#### 📊 Estatísticas Semanais")
        stats = {
            'Semana com Maior Vendas': f"{vendas_semanal['Vendas ilíquidas (€)'].max():,.0f}€",
            'Semana com Menor Vendas': f"{vendas_semanal['Vendas ilíquidas (€)'].min():,.0f}€",
            'Média Semanal': f"{vendas_semanal['Vendas ilíquidas (€)'].mean():,.0f}€",
            'Desvio Padrão': f"{vendas_semanal['Vendas ilíquidas (€)'].std():,.0f}€"
        }

        col1, col2, col3, col4 = st.columns(4)
        with col1:
            st.metric("Máximo", stats['Semana com Maior Vendas'])
        with col2:
            st.metric("Mínimo", stats['Semana com Menor Vendas'])
        with col3:
            st.metric("Média", stats['Média Semanal'])
        with col4:
            st.metric("Desvio", stats['Desvio Padrão'])

    with tab2:
        st.subheader("Análise por Jogo")

        # Top Jogos
        vendas_jogo = df_filtrado.groupby('Jogo').agg({
            'Vendas ilíquidas (€)': ['sum', 'mean', 'count'],
            'Data': 'nunique'
        }).reset_index()

        vendas_jogo.columns = ['Jogo', 'Total Vendas', 'Média Semanal', 'Nº Transações', 'Semanas Ativas']
        vendas_jogo = vendas_jogo.sort_values('Total Vendas', ascending=False)

        # Gráfico Top 15 Jogos
        top_jogos = vendas_jogo.head(15)

        fig = px.bar(
            top_jogos,
            x='Jogo',
            y='Total Vendas',
            title='Top 15 Jogos - Vendas Ilíquidas Totais',
            labels={'Total Vendas': 'Vendas (€)', 'Jogo': 'Jogo'},
            height=500
        )

        fig.update_traces(
            text=top_jogos['Total Vendas'].apply(lambda x: f'{x:,.0f}€'),
            textposition='outside'
        )

        st.plotly_chart(fig, use_container_width=True)

        # Tabela detalhada
        st.markdown("#### 📋 Resumo Completo por Jogo")
        vendas_jogo_display = vendas_jogo.copy()
        vendas_jogo_display['Total Vendas'] = vendas_jogo_display['Total Vendas'].apply(lambda x: f'{x:,.0f}€')
        vendas_jogo_display['Média Semanal'] = vendas_jogo_display['Média Semanal'].apply(lambda x: f'{x:,.0f}€')
        st.dataframe(vendas_jogo_display, use_container_width=True, hide_index=True)

    with tab3:
        st.subheader("Análise por Categoria")

        # Vendas por categoria
        vendas_cat = df_filtrado.groupby('Categoria').agg({
            'Vendas ilíquidas (€)': 'sum',
            'Jogo': 'nunique'
        }).reset_index()

        vendas_cat.columns = ['Categoria', 'Total Vendas', 'Nº Jogos']
        vendas_cat = vendas_cat.sort_values('Total Vendas', ascending=False)

        # Gráfico de pizza
        fig_pizza = px.pie(
            vendas_cat,
            values='Total Vendas',
            names='Categoria',
            title='Distribuição de Vendas por Categoria',
            height=500
        )

        st.plotly_chart(fig_pizza, use_container_width=True)

        # Gráfico de barras
        fig_barras = px.bar(
            vendas_cat,
            x='Categoria',
            y='Total Vendas',
            title='Vendas por Categoria',
            labels={'Total Vendas': 'Vendas (€)', 'Categoria': 'Categoria'},
            height=400
        )

        fig_barras.update_traces(
            text=vendas_cat['Total Vendas'].apply(lambda x: f'{x:,.0f}€'),
            textposition='outside'
        )

        st.plotly_chart(fig_barras, use_container_width=True)

        # Tabela
        st.markdown("#### 📊 Detalhes por Categoria")
        vendas_cat_display = vendas_cat.copy()
        vendas_cat_display['Total Vendas'] = vendas_cat_display['Total Vendas'].apply(lambda x: f'{x:,.0f}€')
        st.dataframe(vendas_cat_display, use_container_width=True, hide_index=True)

    with tab4:
        st.subheader("Comparação Semanal de Top Jogos")

        # Selecionar top 5 jogos
        top_5_jogos = df_filtrado.groupby('Jogo')['Vendas ilíquidas (€)'].sum().nlargest(5).index.tolist()

        # Dados semanais por jogo
        df_top = df_filtrado[df_filtrado['Jogo'].isin(top_5_jogos)]
        vendas_semana_jogo = df_top.groupby(['Data', 'Jogo'])['Vendas ilíquidas (€)'].sum().reset_index()

        fig = px.line(
            vendas_semana_jogo,
            x='Data',
            y='Vendas ilíquidas (€)',
            color='Jogo',
            title='Evolução Semanal - Top 5 Jogos',
            markers=True,
            height=500
        )

        fig.update_layout(
            xaxis_title="Semana",
            yaxis_title="Vendas (€)",
            hovermode='x unified'
        )

        st.plotly_chart(fig, use_container_width=True)

        # Tabela de comparação
        st.markdown("#### 📊 Tabela Comparativa Semanal")
        pivot_vendas = vendas_semana_jogo.pivot(index='Data', columns='Jogo', values='Vendas ilíquidas (€)').fillna(0)
        pivot_vendas = pivot_vendas.applymap(lambda x: f'{x:,.0f}€' if x > 0 else '-')
        st.dataframe(pivot_vendas, use_container_width=True)

    with tab5:
        st.subheader("📈 Análise Mês a Mês (MoM)")

        # Preparar dados mensais
        df_mes = df_filtrado.copy()
        df_mes['Ano_Mes'] = df_mes['Data'].dt.to_period('M').astype(str)
        vendas_mes = df_mes.groupby(['Ano_Mes', 'Jogo'])['Vendas ilíquidas (€)'].sum().reset_index()

        # Seletor de jogos para análise
        jogos_disponiveis = sorted(vendas_mes['Jogo'].unique())
        jogos_selecionados = st.multiselect(
            "🎮 Selecione os jogos para visualizar (deixe em branco para ver todos):",
            options=jogos_disponiveis,
            default=jogos_disponiveis[:5] if len(jogos_disponiveis) > 5 else jogos_disponiveis
        )

        if not jogos_selecionados:
            jogos_selecionados = jogos_disponiveis

        vendas_mes_filtrado = vendas_mes[vendas_mes['Jogo'].isin(jogos_selecionados)]

        col1, col2 = st.columns(2)

        with col1:
            fig = px.line(
                vendas_mes_filtrado,
                x='Ano_Mes',
                y='Vendas ilíquidas (€)',
                color='Jogo',
                title=f'Vendas Mensais por Jogo ({len(jogos_selecionados)} selecionados)',
                markers=True,
                height=450
            )
            fig.update_layout(hovermode='x unified', xaxis_tickangle=-45)
            st.plotly_chart(fig, use_container_width=True)

        with col2:
            # Crescimento MoM %
            vendas_mes_filtrado = vendas_mes_filtrado.copy()
            vendas_mes_filtrado['Vendas_Anterior'] = vendas_mes_filtrado.groupby('Jogo')['Vendas ilíquidas (€)'].shift(1)
            vendas_mes_filtrado['Crescimento_MoM'] = ((vendas_mes_filtrado['Vendas ilíquidas (€)'] - vendas_mes_filtrado['Vendas_Anterior']) /
                                              vendas_mes_filtrado['Vendas_Anterior'] * 100).fillna(0)

            fig = px.bar(
                vendas_mes_filtrado,
                x='Ano_Mes',
                y='Crescimento_MoM',
                color='Jogo',
                title='Crescimento MoM (%)',
                barmode='group',
                height=450
            )
            fig.add_hline(y=0, line_dash="dash", line_color="gray")
            fig.update_layout(xaxis_tickangle=-45)
            st.plotly_chart(fig, use_container_width=True)

        # Tabela resumida de MoM
        st.markdown("#### 📊 Tabela Resumida - Crescimento Mês a Mês")
        tabela_mom = vendas_mes_filtrado[['Jogo', 'Ano_Mes', 'Vendas ilíquidas (€)', 'Crescimento_MoM']].copy()
        tabela_mom = tabela_mom.sort_values(['Jogo', 'Ano_Mes'])
        tabela_mom['Vendas ilíquidas (€)'] = tabela_mom['Vendas ilíquidas (€)'].apply(lambda x: f'{x:,.0f}€')
        tabela_mom['Crescimento_MoM'] = tabela_mom['Crescimento_MoM'].apply(lambda x: f'{x:+.1f}%')
        tabela_mom.columns = ['Jogo', 'Mês', 'Vendas', 'Crescimento %']
        st.dataframe(tabela_mom, use_container_width=True, hide_index=True)

    with tab6:
        st.subheader("📅 Análise Ano a Ano (YoY)")

        # Preparar dados por mês do ano
        df_yoy = df_filtrado.copy()
        df_yoy['Mes'] = df_yoy['Data'].dt.month
        df_yoy['Ano'] = df_yoy['Data'].dt.year
        vendas_yoy = df_yoy.groupby(['Ano', 'Mes', 'Jogo'])['Vendas ilíquidas (€)'].sum().reset_index()
        vendas_yoy['Mes_Nome'] = pd.to_datetime(vendas_yoy['Mes'].astype(str), format='%m').dt.strftime('%b')

        # Seletor de jogo para análise YoY
        jogos_yoy = sorted(vendas_yoy['Jogo'].unique())
        jogo_selecionado_yoy = st.selectbox(
            "🎮 Selecione um jogo para análise YoY:",
            options=jogos_yoy,
            key="yoy_jogo"
        )

        vendas_yoy_jogo = vendas_yoy[vendas_yoy['Jogo'] == jogo_selecionado_yoy]

        # Gráfico principal
        fig_yoy = px.line(
            vendas_yoy_jogo,
            x='Mes',
            y='Vendas ilíquidas (€)',
            color='Ano',
            title=f'Evolução Mensal - {jogo_selecionado_yoy} (Comparação Anual)',
            markers=True,
            height=450,
            labels={'Mes': 'Mês', 'Vendas ilíquidas (€)': 'Vendas (€)'}
        )
        fig_yoy.update_layout(hovermode='x unified', xaxis_title="Mês do Ano")
        st.plotly_chart(fig_yoy, use_container_width=True)

        # Tabela comparativa
        st.markdown(f"#### 📊 Tabela Comparativa - {jogo_selecionado_yoy}")
        tabela_yoy = vendas_yoy_jogo.pivot_table(
            index='Mes_Nome',
            columns='Ano',
            values='Vendas ilíquidas (€)',
            aggfunc='sum'
        ).fillna(0)

        # Adicionar coluna de variação
        anos_unicos = sorted(vendas_yoy_jogo['Ano'].unique())
        if len(anos_unicos) >= 2:
            ano_atual = anos_unicos[-1]
            ano_anterior = anos_unicos[-2]
            if ano_atual in tabela_yoy.columns and ano_anterior in tabela_yoy.columns:
                tabela_yoy['Variação %'] = ((tabela_yoy[ano_atual] - tabela_yoy[ano_anterior]) /
                                            tabela_yoy[ano_anterior] * 100).fillna(0)

        # Formatar valores
        for col in tabela_yoy.columns:
            if col != 'Variação %':
                tabela_yoy[col] = tabela_yoy[col].apply(lambda x: f'{x:,.0f}€')
        tabela_yoy['Variação %'] = tabela_yoy.get('Variação %', pd.Series()).apply(lambda x: f'{x:+.1f}%' if isinstance(x, (int, float)) else x)

        st.dataframe(tabela_yoy, use_container_width=True)

    with tab7:
        st.subheader("📆 Análise Semana a Semana")

        # Preparar dados semanais
        df_semana = df_filtrado.copy()
        vendas_semana = df_semana.groupby(['Data', 'Jogo'])['Vendas ilíquidas (€)'].sum().reset_index()
        vendas_semana = vendas_semana.sort_values('Data')

        # Opções de visualização
        col_opts1, col_opts2 = st.columns(2)
        with col_opts1:
            tipo_viz = st.radio(
                "📊 Tipo de Visualização:",
                options=["Todos os Jogos", "Jogo Específico", "Top 5 Jogos"],
                horizontal=True
            )

        # Filtrar dados conforme a seleção
        if tipo_viz == "Jogo Específico":
            with col_opts2:
                jogo_selecionado_semana = st.selectbox(
                    "🎮 Selecione um jogo:",
                    options=sorted(vendas_semana['Jogo'].unique()),
                    key="semana_jogo"
                )
            vendas_semana_viz = vendas_semana[vendas_semana['Jogo'] == jogo_selecionado_semana]
            titulo_grafico = f"Vendas Semanais - {jogo_selecionado_semana}"

        elif tipo_viz == "Top 5 Jogos":
            top_5 = vendas_semana.groupby('Jogo')['Vendas ilíquidas (€)'].sum().nlargest(5).index.tolist()
            vendas_semana_viz = vendas_semana[vendas_semana['Jogo'].isin(top_5)]
            titulo_grafico = "Vendas Semanais - Top 5 Jogos"

        else:  # Todos os Jogos
            vendas_semana_viz = vendas_semana
            titulo_grafico = f"Vendas Semanais - Todos os Jogos ({vendas_semana['Jogo'].nunique()} jogos)"

        # Gráfico principal
        fig = px.line(
            vendas_semana_viz,
            x='Data',
            y='Vendas ilíquidas (€)',
            color='Jogo',
            title=titulo_grafico,
            markers=True,
            height=500
        )
        fig.update_layout(hovermode='x unified', xaxis_tickangle=-45)
        st.plotly_chart(fig, use_container_width=True)

        # Estatísticas por jogo
        st.markdown("#### 📊 Estatísticas Semanais")
        stats_jogo = df_filtrado.groupby('Jogo')['Vendas ilíquidas (€)'].agg([
            ('Nº Semanas', 'count'),
            ('Total', 'sum'),
            ('Média', 'mean'),
            ('Máximo', 'max'),
            ('Mínimo', 'min'),
            ('Desvio Padrão', 'std')
        ]).reset_index().sort_values('Total', ascending=False)

        stats_jogo['Total'] = stats_jogo['Total'].apply(lambda x: f'{x:,.0f}€')
        stats_jogo['Média'] = stats_jogo['Média'].apply(lambda x: f'{x:,.0f}€')
        stats_jogo['Máximo'] = stats_jogo['Máximo'].apply(lambda x: f'{x:,.0f}€')
        stats_jogo['Mínimo'] = stats_jogo['Mínimo'].apply(lambda x: f'{x:,.0f}€')
        stats_jogo['Desvio Padrão'] = stats_jogo['Desvio Padrão'].apply(lambda x: f'{x:,.0f}€')
        stats_jogo.columns = ['Jogo', 'Nº Semanas', 'Total', 'Média', 'Máximo', 'Mínimo', 'Desvio']

        st.dataframe(stats_jogo, use_container_width=True, hide_index=True)

    with tab8:
        st.subheader("🔮 Previsão vs Objetivos")

        # Carregar objetivos do CSV novo
        objetivos_guardados = carregar_objetivos()

        # Calcular tendências e fazer previsão
        df_trend = df_filtrado.copy()
        vendas_por_jogo = df_trend.groupby('Jogo')['Vendas ilíquidas (€)'].sum().reset_index()
        vendas_por_jogo = vendas_por_jogo.sort_values('Vendas ilíquidas (€)', ascending=False)

        st.markdown("#### 🎯 Comparação: Atual vs Objetivo")

        for idx, row in vendas_por_jogo.iterrows():
            jogo = row['Jogo']
            vendas_total = row['Vendas ilíquidas (€)']

            # Calcular média semanal
            df_jogo = df_trend[df_trend['Jogo'] == jogo].sort_values('Data')
            vendas_semana_jogo = df_jogo.groupby('Data')['Vendas ilíquidas (€)'].sum()
            media_semana = vendas_semana_jogo.mean() if len(vendas_semana_jogo) > 0 else 0

            # Obter objetivo do CSV guardado
            objetivo_semanal = objetivos_guardados.get(jogo, 0)

            # Projeção para 4 semanas e anual
            projecao_4sem = media_semana * 4
            projecao_anual = media_semana * 52

            # Calcular % de cumprimento vs objetivo
            pct_cumprimento = (media_semana / objetivo_semanal * 100) if objetivo_semanal > 0 else 0
            diferenca = media_semana - objetivo_semanal

            # Determinar status
            if objetivo_semanal == 0:
                status = "⚫ Sem Objetivo"
                status_tipo = "secondary"
            elif pct_cumprimento >= 100:
                status = "✅ Acima do Objetivo"
                status_tipo = "success"
            elif pct_cumprimento >= 90:
                status = "📈 Próximo do Objetivo"
                status_tipo = "info"
            else:
                status = "⚠️ Abaixo do Objetivo"
                status_tipo = "warning"

            # Mostrar métricas em colunas
            col1, col2, col3, col4, col5, col6 = st.columns(6)

            with col1:
                st.metric(f"🎮 {jogo}", f"{media_semana:,.0f}€", "Média/Sem")

            with col2:
                st.metric("🎯 Objetivo", f"{objetivo_semanal:,.0f}€", "Meta/Sem")

            with col3:
                st.metric("📊 % Cumprimento", f"{pct_cumprimento:.1f}%", f"{diferenca:+.0f}€")

            with col4:
                st.metric("📅 Proj. 4 Sem", f"{projecao_4sem:,.0f}€", "")

            with col5:
                st.metric("📈 Proj. Anual", f"{projecao_anual:,.0f}€", "")

            with col6:
                st.metric("Status", status, "")

            # Info box com análise detalhada + gráfico
            col_txt, col_graf = st.columns([2.5, 1])

            with col_txt:
                if objetivo_semanal == 0:
                    st.info(f"ℹ️ **{jogo}** - Sem objetivo definido para este ano")
                elif pct_cumprimento >= 100:
                    st.success(f"""
                    ✅ **{jogo}** está **acima do objetivo**!
                    - Objetivo semanal: {objetivo_semanal:,.0f}€
                    - Performance atual: {media_semana:,.0f}€ (+{pct_cumprimento-100:.1f}%)
                    - Projeção anual: {projecao_anual:,.0f}€
                    - Tendência: Excelente 🚀
                    """)
                elif pct_cumprimento >= 90:
                    st.info(f"""
                    📈 **{jogo}** está **próximo do objetivo**!
                    - Objetivo semanal: {objetivo_semanal:,.0f}€
                    - Performance atual: {media_semana:,.0f}€ ({pct_cumprimento:.1f}%)
                    - Diferença: {diferenca:,.0f}€
                    - Projeção anual: {projecao_anual:,.0f}€
                    - Tendência: Bom desempenho 💪
                    """)
                else:
                    st.warning(f"""
                    ⚠️ **{jogo}** está **abaixo do objetivo**.
                    - Objetivo semanal: {objetivo_semanal:,.0f}€
                    - Performance atual: {media_semana:,.0f}€ ({pct_cumprimento:.1f}%)
                    - Diferença: {diferenca:,.0f}€
                    - Projeção anual: {projecao_anual:,.0f}€
                    - Tendência: Necessário esforço adicional 💡
                    """)

            # Pequeno gráfico ao lado
            with col_graf:
                if objetivo_semanal > 0:
                    fig_mini = go.Figure(data=[
                        go.Bar(
                            x=['Atual', 'Objetivo'],
                            y=[media_semana, objetivo_semanal],
                            text=[f'{media_semana:,.0f}€', f'{objetivo_semanal:,.0f}€'],
                            textposition='outside',
                            marker=dict(
                                color=['#1f77b4' if media_semana >= objetivo_semanal else '#ff7f0e', '#2ca02c'],
                                opacity=0.8
                            )
                        )
                    ])
                    fig_mini.update_layout(
                        showlegend=False,
                        height=200,
                        margin=dict(l=20, r=20, t=20, b=20),
                        yaxis_title='€',
                        xaxis_tickfont=dict(size=10),
                        font=dict(size=9)
                    )
                    st.plotly_chart(fig_mini, use_container_width=True)

            # Opção para editar objetivo do jogo
            with st.expander(f"✏️ Editar Objetivo - {jogo}"):
                novo_objetivo = st.number_input(
                    f"Novo objetivo semanal para {jogo}",
                    value=float(objetivo_semanal) if objetivo_semanal > 0 else 0.0,
                    min_value=0.0,
                    step=100.0,
                    key=f"edit_objetivo_v6_{jogo}"
                )

                if st.button(f"💾 Guardar Objetivo - {jogo}", key=f"btn_guardar_v6_{jogo}"):
                    objetivos_guardados[jogo] = novo_objetivo
                    if guardar_objetivos(objetivos_guardados):
                        st.success(f"✅ Objetivo de €{novo_objetivo:.0f} guardado para {jogo}!")
                        st.rerun()
                    else:
                        st.error("Erro ao guardar objetivo")

            st.divider()


# ========== FUNÇÕES DE FORECAST & ANÁLISE DE ENCOMENDAS ==========

def filtrar_produtos_regulares(df, dias_minimos=30, vendas_minimas=5):
    """
    Filtra produtos com vendas regulares (exclui produtos esporádicos)

    Args:
        df: DataFrame com vendas
        dias_minimos: Período mínimo de análise (default 30 dias)
        vendas_minimas: Número mínimo de transações no período

    Returns:
        Lista de produtos regulares
    """
    if df.empty:
        return []

    # Calcular data limite
    data_max = df['Data'].max()
    data_limite = data_max - timedelta(days=dias_minimos)

    # Filtrar período
    df_periodo = df[df['Data'] >= data_limite]

    # Agrupar por produto e contar transações
    produtos_count = df_periodo.groupby('Produto').size().reset_index(name='count')

    # Filtrar produtos com vendas >= vendas_minimas
    produtos_regulares = produtos_count[produtos_count['count'] >= vendas_minimas]['Produto'].tolist()

    return produtos_regulares


def calcular_sazonalidade_produto(df_produto):
    """
    Analisa padrões de sazonalidade de um produto

    Args:
        df_produto: DataFrame filtrado para um produto específico

    Returns:
        dict com métricas de sazonalidade
    """
    if df_produto.empty:
        return {
            'media_diaria': 0,
            'media_por_dia_semana': {},
            'desvio_padrao': 0,
            'coeficiente_variacao': 0,
            'dias_com_venda': 0,
            'total_dias_periodo': 0
        }

    # Adicionar coluna Dia_Semana se não existir
    if 'Dia_Semana' not in df_produto.columns:
        df_produto = df_produto.copy()
        df_produto['Dia_Semana'] = df_produto['Data'].dt.dayofweek

    # Agrupar por Data e somar quantidade
    vendas_diarias = df_produto.groupby('Data')['Qtd'].sum()

    # Calcular métricas gerais
    media_diaria = vendas_diarias.mean()
    desvio_padrao = vendas_diarias.std()
    dias_com_venda = len(vendas_diarias)

    # Calcular período total
    data_min = df_produto['Data'].min()
    data_max = df_produto['Data'].max()
    total_dias_periodo = (data_max - data_min).days + 1

    # Calcular média por dia da semana
    media_por_dia_semana = {}
    for dia in range(7):  # 0=Segunda, 6=Domingo
        vendas_dia = df_produto[df_produto['Dia_Semana'] == dia]['Qtd'].sum()
        dias_com_venda_dia = df_produto[df_produto['Dia_Semana'] == dia]['Data'].nunique()
        media_por_dia_semana[dia] = vendas_dia / max(dias_com_venda_dia, 1)

    # Calcular coeficiente de variação
    cv = (desvio_padrao / max(media_diaria, 0.001)) * 100

    return {
        'media_diaria': media_diaria,
        'media_por_dia_semana': media_por_dia_semana,
        'desvio_padrao': desvio_padrao,
        'coeficiente_variacao': cv,
        'dias_com_venda': dias_com_venda,
        'total_dias_periodo': total_dias_periodo
    }


def calcular_tendencia_produto(df_produto, janela_dias=14):
    """
    Calcula tendência de crescimento/decrescimento usando regressão linear

    Args:
        df_produto: DataFrame do produto
        janela_dias: Janela de tempo para análise (últimos N dias)

    Returns:
        dict com tendencia_percent, direcao, confianca
    """
    if df_produto.empty or len(df_produto) < 3:
        return {
            'tendencia_percent': 0,
            'direcao': 'estavel',
            'confianca': 0
        }

    # Filtrar últimos janela_dias
    data_max = df_produto['Data'].max()
    data_inicio = data_max - timedelta(days=janela_dias)
    df_janela = df_produto[df_produto['Data'] >= data_inicio]

    if df_janela.empty or len(df_janela) < 3:
        return {
            'tendencia_percent': 0,
            'direcao': 'estavel',
            'confianca': 0
        }

    # Agrupar por data e somar quantidade
    serie = df_janela.groupby('Data')['Qtd'].sum().reset_index()
    serie = serie.sort_values('Data')

    if len(serie) < 3:
        return {
            'tendencia_percent': 0,
            'direcao': 'estavel',
            'confianca': 0
        }

    # Preparar dados para regressão
    x = np.arange(len(serie))
    y = serie['Qtd'].values

    # Regressão linear
    coefficients = np.polyfit(x, y, 1)
    slope, intercept = coefficients

    # Calcular R² (qualidade do ajuste)
    y_pred = slope * x + intercept
    ss_res = np.sum((y - y_pred) ** 2)
    ss_tot = np.sum((y - np.mean(y)) ** 2)
    r_squared = 1 - (ss_res / ss_tot) if ss_tot > 0 else 0
    r_squared = max(0, min(1, r_squared))  # Limitar entre 0 e 1

    # Converter slope em % de mudança
    valor_inicial = intercept
    valor_final = slope * len(x) + intercept

    if valor_inicial > 0:
        tendencia_percent = ((valor_final - valor_inicial) / valor_inicial) * 100
    else:
        tendencia_percent = 0

    # Classificar tendência
    if tendencia_percent > 5:
        direcao = 'crescimento'
    elif tendencia_percent < -5:
        direcao = 'queda'
    else:
        direcao = 'estavel'

    return {
        'tendencia_percent': tendencia_percent,
        'direcao': direcao,
        'confianca': r_squared
    }


def calcular_stock_seguranca(media_diaria, desvio_padrao, fator_seguranca=1.5):
    """
    Calcula stock de segurança baseado em variabilidade

    Args:
        media_diaria: Demanda média diária
        desvio_padrao: Variabilidade da demanda
        fator_seguranca: Multiplicador de segurança (default 1.5)

    Returns:
        float: Quantidade de stock de segurança
    """
    # Fórmula: stock_seguranca = fator_seguranca * desvio_padrao
    stock = fator_seguranca * desvio_padrao

    # Garantir mínimo de 20% da média diária
    stock_minimo = media_diaria * 0.2

    return max(stock, stock_minimo)


def gerar_forecast_produtos(df, cost_manager, dias_analise=60, produtos_regulares_only=True):
    """
    Gera previsão de encomenda para produtos

    Args:
        df: DataFrame completo de vendas
        cost_manager: CostManagerV2 para obter custos
        dias_analise: Período de análise histórica (default 60 dias)
        produtos_regulares_only: Se True, filtra apenas produtos regulares

    Returns:
        DataFrame com forecast completo
    """
    if df.empty:
        return pd.DataFrame()

    # Filtrar últimos dias_analise
    data_max = df['Data'].max()
    data_inicio = data_max - timedelta(days=dias_analise)
    df_analise = df[df['Data'] >= data_inicio].copy()

    # Filtrar produtos regulares se solicitado
    if produtos_regulares_only:
        produtos_regulares = filtrar_produtos_regulares(df_analise, dias_minimos=30, vendas_minimas=5)
        df_analise = df_analise[df_analise['Produto'].isin(produtos_regulares)]

    if df_analise.empty:
        return pd.DataFrame()

    # Lista para armazenar resultados
    resultados = []

    # Para cada produto
    produtos = df_analise['Produto'].unique()

    for produto in produtos:
        # Filtrar dados do produto
        df_produto = df_analise[df_analise['Produto'] == produto]

        # Pular se dados insuficientes
        if len(df_produto) < 3:
            continue

        # Obter categoria
        categoria = df_produto['Categoria'].iloc[0] if 'Categoria' in df_produto.columns else 'Outros'

        # Calcular sazonalidade
        sazonalidade = calcular_sazonalidade_produto(df_produto)
        media_diaria = sazonalidade['media_diaria']
        desvio_padrao = sazonalidade['desvio_padrao']
        cv = sazonalidade['coeficiente_variacao']
        dias_com_venda = sazonalidade['dias_com_venda']

        # Calcular tendência
        tendencia = calcular_tendencia_produto(df_produto, janela_dias=14)
        tendencia_percent = tendencia['tendencia_percent']
        tendencia_direcao = tendencia['direcao']

        # Calcular stock de segurança
        stock_seguranca = calcular_stock_seguranca(media_diaria, desvio_padrao, fator_seguranca=1.5)

        # Calcular quantidades
        qtd_semanal = media_diaria * 7
        qtd_mensal = media_diaria * 30

        # Obter custo unitário
        custo_unitario = cost_manager.get_custo_produto(produto, categoria)

        # Calcular valores
        valor_semanal = qtd_semanal * custo_unitario
        valor_mensal = qtd_mensal * custo_unitario

        # Adicionar ao resultado
        resultados.append({
            'Produto': produto,
            'Categoria': categoria,
            'Media_Diaria': media_diaria,
            'Qtd_Semanal': qtd_semanal,
            'Qtd_Mensal': qtd_mensal,
            'Stock_Seguranca': stock_seguranca,
            'Custo_Unitario': custo_unitario,
            'Valor_Encomenda_Semanal': valor_semanal,
            'Valor_Encomenda_Mensal': valor_mensal,
            'Tendencia_Percent': tendencia_percent,
            'Tendencia_Direcao': tendencia_direcao,
            'CV': cv,
            'Dias_Com_Venda': dias_com_venda
        })

    # Criar DataFrame
    df_forecast = pd.DataFrame(resultados)

    # Ordenar por Valor_Encomenda_Mensal DESC
    if not df_forecast.empty:
        df_forecast = df_forecast.sort_values('Valor_Encomenda_Mensal', ascending=False)

    return df_forecast


# Interface principal
def main():
    # Carregar dados
    try:
        df, loader, cost_manager = carregar_dados()
        categorizer = ProductCategorizer()
    except Exception as e:
        st.error(f"Erro ao carregar dados: {e}")
        return

    if df.empty:
        st.warning("Nenhum dado disponível.")
        return

    # Sidebar - Navegação
    render_sidebar_brand_v10()

    # Inicializar session_state para tab ativa
    if 'active_tab' not in st.session_state:
        st.session_state.active_tab = "🏠 Cockpit Executivo"

    paginas_por_area = paginas_por_area_v10()
    tabs_list = [pagina for paginas in paginas_por_area.values() for pagina in paginas]

    if st.session_state.active_tab not in tabs_list:
        st.session_state.active_tab = "🏠 Cockpit Executivo"

    area_atual = next(
        (area for area, paginas in paginas_por_area.items() if st.session_state.active_tab in paginas),
        "Painel",
    )

    area_selecionada = st.sidebar.selectbox(
        "Área de trabalho",
        options=list(paginas_por_area.keys()),
        index=list(paginas_por_area.keys()).index(area_atual),
        key="v10_area_selector",
    )

    paginas_disponiveis = paginas_por_area[area_selecionada]
    if st.session_state.active_tab not in paginas_disponiveis:
        st.session_state.active_tab = paginas_disponiveis[0]
        if "v10_page_selector" in st.session_state:
            del st.session_state["v10_page_selector"]

    st.session_state.active_tab = st.sidebar.selectbox(
        "Página",
        options=paginas_disponiveis,
        index=paginas_disponiveis.index(st.session_state.active_tab),
        key="v10_page_selector",
        format_func=titulo_pagina_v10,
    )

    # Sidebar - Filtros
    st.sidebar.markdown("---")
    st.sidebar.header("Filtros")

    # Filtro de data
    st.sidebar.subheader("Período")
    data_min = df['Data'].min().date()
    data_max = df['Data'].max().date()

    # Botões rápidos para anos com Pills modernas
    st.sidebar.markdown("**Atalhos de ano**")

    anos_unicos = sorted(df['Data'].dt.year.unique())
    opcoes_anos = [str(ano) for ano in anos_unicos] + ["Todos"]

    # Determinar seleção padrão
    if "selected_year" not in st.session_state:
        st.session_state.selected_year = data_max.year  # Abre por padrão no ano mais recente disponível

    if st.session_state.selected_year and str(st.session_state.selected_year) not in opcoes_anos:
        st.session_state.selected_year = data_max.year

    default_selection = str(st.session_state.selected_year) if st.session_state.selected_year else "Todos"

    ano_selecionado = st.sidebar.pills(
        "Selecionar período",
        options=opcoes_anos,
        default=default_selection,
        label_visibility="collapsed"
    )

    # Atualizar session_state baseado na seleção
    if ano_selecionado == "Todos":
        st.session_state.selected_year = None
    else:
        st.session_state.selected_year = int(ano_selecionado)

    # Aplicar filtro de ano rápido se selecionado
    if st.session_state.selected_year:
        ano = st.session_state.selected_year
        # Limitar data_fim ao máximo de dados disponíveis nesse ano
        data_max_ano = df[df['Data'].dt.year == ano]['Data'].max().date()
        date_range = (pd.Timestamp(f"{ano}-01-01").date(), data_max_ano)
    else:
        date_range = st.sidebar.date_input(
            "Período",
            value=(data_min, data_max),
            min_value=data_min,
            max_value=data_max
        )

    # Extrair data_inicio e data_fim do date_range
    if isinstance(date_range, tuple) and len(date_range) == 2:
        data_inicio, data_fim = date_range
    elif isinstance(date_range, tuple) and len(date_range) == 1:
        data_inicio = data_fim = date_range[0]
    else:
        data_inicio = data_fim = date_range

    # Filtro de categoria
    st.sidebar.subheader("Categorias")
    categorias_disponiveis = sorted(df['Categoria'].unique())
    categorias_selecionadas = st.sidebar.multiselect(
        "Selecionar Categorias",
        options=categorias_disponiveis,
        default=categorias_disponiveis
    )

    # Filtro de subcategoria (dependente da categoria)
    if categorias_selecionadas:
        subcategorias_disponiveis = sorted(
            df[df['Categoria'].isin(categorias_selecionadas)]['Subcategoria'].dropna().astype(str).unique()
        )
        subcategorias_selecionadas = st.sidebar.multiselect(
            "Selecionar Subcategorias",
            options=subcategorias_disponiveis,
            default=[]
        )
    else:
        subcategorias_selecionadas = []

    # Filtro de fonte
    st.sidebar.subheader("Fonte de dados")
    fontes_disponiveis = sorted(df['Fonte'].unique())
    fontes_selecionadas = st.sidebar.multiselect(
        "Selecionar Fontes",
        options=fontes_disponiveis,
        default=fontes_disponiveis
    )

    # Opção de Segunda a Sábado
    st.sidebar.subheader("Opções")
    excluir_domingos = st.sidebar.checkbox("Segunda a Sábado (Excluir Domingos)", value=False)

    # Botão de reset total
    if st.sidebar.button("🔄 Reset Total", use_container_width=True):
        st.cache_data.clear()
        st.cache_resource.clear()
        st.rerun()

    # Aplicar filtros
    df_filtrado = filtrar_dados(
        df,
        categorias_selecionadas,
        subcategorias_selecionadas,
        fontes_selecionadas,
        data_inicio,
        data_fim,
        excluir_domingos
    )

    if df_filtrado.empty:
        st.warning("Nenhum dado disponível com os filtros selecionados.")
        return

    data_inicio_periodo_anterior, data_fim_periodo_anterior = periodo_anterior_imediato(data_inicio, data_fim)
    df_periodo_anterior = filtrar_dados(
        df,
        categorias_selecionadas,
        subcategorias_selecionadas,
        fontes_selecionadas,
        data_inicio_periodo_anterior,
        data_fim_periodo_anterior,
        excluir_domingos
    )

    # Informações do filtro
    st.sidebar.markdown("---")
    st.sidebar.markdown(f"**Registos filtrados:** {len(df_filtrado):,}")
    st.sidebar.markdown(f"**Total original:** {len(df):,}")
    st.sidebar.markdown(f"**Período:** {(data_fim - data_inicio).days + 1} dias")

    if st.session_state.active_tab == "🏠 Cockpit Executivo":
        pagina_cockpit_executivo_v10(
            df_filtrado,
            df_periodo_anterior,
            df,
            loader,
            cost_manager,
            data_inicio,
            data_fim,
        )
        return

    # Métricas principais
    st.markdown('<div class="section-divider"></div>', unsafe_allow_html=True)

    # Calcular métricas atuais
    total_vendas = df_filtrado['Valor'].sum()
    total_quantidade = df_filtrado['Qtd'].sum() if 'Qtd' in df_filtrado.columns else 0
    media_diaria = df_filtrado.groupby('Data')['Valor'].sum().mean()
    ticket_medio = total_vendas / total_quantidade if total_quantidade > 0 else 0

    # Calcular período anterior para comparação (mesmo período do ano anterior)
    dias_periodo = (data_fim - data_inicio).days + 1
    data_inicio_anterior = data_inicio - timedelta(days=365)
    data_fim_anterior = data_fim - timedelta(days=365)

    df_anterior = filtrar_dados(
        df,
        categorias_selecionadas,
        subcategorias_selecionadas,
        fontes_selecionadas,
        data_inicio_anterior,
        data_fim_anterior,
        excluir_domingos
    )

    if not df_anterior.empty:
        total_anterior = df_anterior['Valor'].sum()
        media_anterior = df_anterior.groupby('Data')['Valor'].sum().mean()
        qtd_anterior = df_anterior['Qtd'].sum() if 'Qtd' in df_anterior.columns else 0
        ticket_anterior = total_anterior / qtd_anterior if qtd_anterior > 0 else 0

        delta_total = calcular_delta(total_vendas, total_anterior)
        delta_media = calcular_delta(media_diaria, media_anterior)
        delta_ticket = calcular_delta(ticket_medio, ticket_anterior)
    else:
        delta_total = 0
        delta_media = 0
        delta_ticket = 0

    # Calcular taxa de crescimento mensal
    if 'Mes' in df_filtrado.columns:
        vendas_por_mes = df_filtrado.groupby('Mes')['Valor'].sum()
        if len(vendas_por_mes) >= 2:
            crescimento_mensal = calcular_delta(
                vendas_por_mes.iloc[-1],
                vendas_por_mes.iloc[-2]
            )
        else:
            crescimento_mensal = 0
    else:
        crescimento_mensal = 0

    # Exibir métricas
    col1, col2, col3, col4 = st.columns(4)

    with col1:
        st.metric(
            label="💰 Total de Vendas",
            value=formatar_moeda(total_vendas),
            delta=f"{delta_total:+.1f}%" if delta_total != 0 else None
        )

    with col2:
        st.metric(
            label="📊 Média Diária",
            value=formatar_moeda(media_diaria),
            delta=f"{delta_media:+.1f}%" if delta_media != 0 else None
        )

    with col3:
        st.metric(
            label="🎯 Ticket Médio",
            value=formatar_moeda(ticket_medio),
            delta=f"{delta_ticket:+.1f}%" if delta_ticket != 0 else None
        )

    with col4:
        st.metric(
            label="📈 Crescimento Mensal",
            value=formatar_percentagem(crescimento_mensal),
            delta=f"{crescimento_mensal:+.1f}%"
        )

    st.markdown('<div class="section-divider"></div>', unsafe_allow_html=True)

    # ============================================================
    # SISTEMA DE NAVEGAÇÃO POR TABS (mantém tab ativa após rerun)
    # ============================================================

    # TAB 1 - Visão Geral
    if st.session_state.active_tab == "📊 Visão Geral":
        st.header("📊 Visão Geral das Vendas")

        # ========== MÉTRICAS DE PERFORMANCE ==========
        st.subheader("📊 Métricas de Performance")

        # Calcular métricas de performance (passou df completo para comparações históricas)
        metricas_perf = calcular_metricas_performance(df_filtrado, df)

        if metricas_perf:
            data_atualizacao = metricas_perf.get('data_atualizacao', 'N/A')
            st.caption(f"📅 Dados atualizados até: {data_atualizacao}")

            # Performance Semanal
            st.subheader("📅 Performance Semanal")
            col1, col2, col3 = st.columns(3)

            with col1:
                st.markdown("**Última Semana**")
                st.metric(
                    label="Métrica",
                    label_visibility="collapsed",
                    value=formatar_moeda(metricas_perf['vendas_semana_atual']),
                    delta=f"{metricas_perf['var_semana_vs_anterior']:+.1f}% vs Semana Anterior" if metricas_perf['vendas_semana_anterior'] > 0 else "Sem dados"
                )

            with col2:
                st.markdown("**Semana Anterior**")
                st.metric(
                    label="Métrica",
                    label_visibility="collapsed",
                    value=formatar_moeda(metricas_perf['vendas_semana_anterior'])
                )

            with col3:
                st.markdown("**Mesma Semana Ano Passado**")
                st.metric(
                    label="Métrica",
                    label_visibility="collapsed",
                    value=formatar_moeda(metricas_perf['vendas_semana_ano_passado']),
                    delta=f"{metricas_perf['var_semana_vs_ano_passado']:+.1f}% vs SWLY" if metricas_perf['vendas_semana_ano_passado'] > 0 else "Sem dados"
                )

            st.markdown("---")

            # Performance Mensal (MTD)
            st.subheader("📆 Performance Mensal (MTD)")
            col1, col2, col3 = st.columns(3)

            with col1:
                st.markdown(f"**Mês em Curso (até dia {metricas_perf['dias_mes_atual']})**")
                st.metric(
                    label="Métrica",
                    label_visibility="collapsed",
                    value=formatar_moeda(metricas_perf['vendas_mes_atual']),
                    delta=f"{metricas_perf['var_mes_vs_anterior']:+.1f}% vs Mês Anterior MTD" if metricas_perf['vendas_mes_anterior'] > 0 else "Sem dados"
                )

            with col2:
                st.markdown(f"**Mês Anterior (até dia {metricas_perf['dias_mes_atual']})**")
                st.metric(
                    label="Métrica",
                    label_visibility="collapsed",
                    value=formatar_moeda(metricas_perf['vendas_mes_anterior'])
                )

            with col3:
                st.markdown(f"**SMLY - Mesmo Mês Ano Passado (até dia {metricas_perf['dias_mes_atual']})**")
                st.metric(
                    label="Métrica",
                    label_visibility="collapsed",
                    value=formatar_moeda(metricas_perf['vendas_smly']),
                    delta=f"{metricas_perf['var_mes_vs_ano_passado']:+.1f}% vs SMLY MTD" if metricas_perf['vendas_smly'] > 0 else "Sem dados"
                )

            st.markdown("---")

            # Performance Anual
            st.subheader("📈 Performance Anual")
            col1, col2, col3 = st.columns(3)

            with col1:
                st.markdown("**Ano em Curso (YTD)**")
                st.metric(
                    label="Métrica",
                    label_visibility="collapsed",
                    value=formatar_moeda(metricas_perf['vendas_ano_atual']),
                    delta=f"{metricas_perf['var_ano_vs_ano_anterior']:+.1f}% vs Ano Anterior YTD" if metricas_perf['vendas_ano_anterior_ytd'] > 0 else "Sem dados"
                )

            with col2:
                st.markdown("**Ano Anterior (mesma altura)**")
                st.metric(
                    label="Métrica",
                    label_visibility="collapsed",
                    value=formatar_moeda(metricas_perf['vendas_ano_anterior_ytd'])
                )

            with col3:
                st.markdown("**Ano Anterior (Total)**")
                st.metric(
                    label="Métrica",
                    label_visibility="collapsed",
                    value=formatar_moeda(metricas_perf['vendas_ano_anterior_total'])
                )

            st.markdown("---")

        # KPIs por categoria
        st.subheader("🏆 Performance por Categoria")

        resumo_categorias = df_filtrado.groupby('Categoria').agg({
            'Valor': 'sum',
            'Qtd': 'sum',
            'Icone_Categoria': 'first'
        }).sort_values('Valor', ascending=False)

        cols = st.columns(min(4, len(resumo_categorias)))
        for idx, (categoria, row) in enumerate(resumo_categorias.head(8).iterrows()):
            col = cols[idx % 4]
            with col:
                percentual = (row['Valor'] / total_vendas) * 100
                # Tratar icone None ou vazio
                icone = row['Icone_Categoria'] if row['Icone_Categoria'] not in [None, 'None', ''] else '🎰'
                label_categoria = f"{icone} {categoria}"
                st.metric(
                    label=label_categoria,
                    value=formatar_moeda(row['Valor']),
                    delta=f"{formatar_percentagem(percentual)} do total"
                )

        st.markdown("---")

        # Gráficos
        col1, col2 = st.columns(2)

        with col1:
            # Gráfico de barras horizontal - Distribuição por categoria (mais legível)
            # Reutilizar dados já calculados acima para evitar novo groupby
            vendas_categoria = resumo_categorias['Valor'].sort_values(ascending=True)

            fig_categoria = go.Figure(go.Bar(
                x=vendas_categoria.values,
                y=vendas_categoria.index,
                orientation='h',
                text=[formatar_moeda(v) for v in vendas_categoria.values],
                textposition='auto',
                marker=dict(
                    color=vendas_categoria.values,
                    colorscale='Blues',
                    showscale=False
                )
            ))

            fig_categoria.update_layout(
                title="📊 Distribuição de Vendas por Categoria",
                height=400,
                margin=dict(t=50, b=50, l=150, r=50),
                xaxis_title="Vendas (€)",
                yaxis_title="",
                showlegend=False
            )

            st.plotly_chart(fig_categoria, use_container_width=True)

        with col2:
            # Top 10 produtos
            top_produtos = df_filtrado.groupby('Produto')['Valor'].sum().sort_values(ascending=True).tail(10)

            fig_top = go.Figure(go.Bar(
                x=top_produtos.values,
                y=top_produtos.index,
                orientation='h',
                text=[formatar_moeda(v) for v in top_produtos.values],
                textposition='auto',
                marker=dict(color='#1f77b4')
            ))

            fig_top.update_layout(
                title="🏆 Top 10 Produtos Mais Vendidos",
                height=400,
                margin=dict(t=50, b=50, l=200, r=50),
                xaxis_title="Vendas (€)",
                yaxis_title=""
            )

            st.plotly_chart(fig_top, use_container_width=True)

        # Evolução temporal
        st.subheader("📈 Evolução Temporal de Vendas")

        vendas_diarias = df_filtrado.groupby('Data')['Valor'].sum().reset_index()

        fig_evolucao = px.line(
            vendas_diarias,
            x='Data',
            y='Valor',
            title='Evolução Diária de Vendas',
            markers=True
        )

        fig_evolucao.update_layout(
            height=400,
            xaxis_title="Data",
            yaxis_title="Vendas (€)",
            hovermode='x unified'
        )

        fig_evolucao.add_hline(
            y=media_diaria,
            line_dash="dash",
            line_color="red",
            annotation_text=f"Média: {formatar_moeda(media_diaria)}"
        )

        st.plotly_chart(fig_evolucao, use_container_width=True)

    # TAB 2 - Análise por Categoria
    elif st.session_state.active_tab == "☕ Análise por Categoria":
        st.header("☕ Análise Detalhada por Categoria")

        # Seletor de categoria
        categoria_selecionada = st.selectbox(
            "Selecione uma categoria para análise detalhada:",
            options=sorted(df_filtrado['Categoria'].unique())
        )

        df_categoria = df_filtrado[df_filtrado['Categoria'] == categoria_selecionada]

        # Métricas da categoria
        col1, col2, col3, col4 = st.columns(4)

        total_cat = df_categoria['Valor'].sum()
        vendas_cat = len(df_categoria)
        media_cat = df_categoria['Valor'].mean()
        participacao_cat = (total_cat / total_vendas) * 100

        with col1:
            st.metric("💰 Total da Categoria", formatar_moeda(total_cat))

        with col2:
            st.metric("📊 Número de Vendas", f"{vendas_cat:,}")

        with col3:
            st.metric("📈 Média por Venda", formatar_moeda(media_cat))

        with col4:
            st.metric("🎯 Participação", formatar_percentagem(participacao_cat))

        st.markdown("---")

        # Análises da categoria
        col1, col2 = st.columns(2)

        with col1:
            # Top 10 subcategorias
            st.subheader(f"🏅 Top 10 Subcategorias")
            top_subcat = df_categoria.groupby('Subcategoria')['Valor'].sum().sort_values(ascending=True).tail(10)

            fig_subcat = go.Figure(go.Bar(
                x=top_subcat.values,
                y=top_subcat.index,
                orientation='h',
                text=[formatar_moeda(v) for v in top_subcat.values],
                textposition='auto',
                marker=dict(color='#2ca02c')
            ))

            fig_subcat.update_layout(
                height=400,
                xaxis_title="Vendas (€)",
                yaxis_title=""
            )

            st.plotly_chart(fig_subcat, use_container_width=True)

        with col2:
            # Distribuição por subcategoria
            fig_pizza_subcat = criar_grafico_pizza(
                df_categoria,
                'Subcategoria',
                '📊 Distribuição por Subcategoria'
            )
            st.plotly_chart(fig_pizza_subcat, use_container_width=True)

        # Top 20 produtos da categoria
        st.subheader(f"🌟 Top 20 Produtos - {categoria_selecionada}")

        top_produtos_cat = df_categoria.groupby('Produto').agg({
            'Valor': 'sum',
            'Qtd': 'sum',
            'Subcategoria': 'first'
        }).sort_values('Valor', ascending=False).head(20)

        top_produtos_cat['Valor'] = top_produtos_cat['Valor'].apply(formatar_moeda)
        top_produtos_cat['Qtd'] = top_produtos_cat['Qtd'].apply(lambda x: f"{x:,.0f}")

        st.dataframe(top_produtos_cat, use_container_width=True)

        # Evolução da categoria
        st.subheader(f"📈 Evolução Temporal - {categoria_selecionada}")

        evolucao_cat = df_categoria.groupby('Data')['Valor'].sum().reset_index()

        fig_evolucao_cat = px.area(
            evolucao_cat,
            x='Data',
            y='Valor',
            title=f'Evolução de Vendas - {categoria_selecionada}',
            color_discrete_sequence=['#1f77b4']
        )

        fig_evolucao_cat.update_layout(
            height=400,
            xaxis_title="Data",
            yaxis_title="Vendas (€)"
        )

        st.plotly_chart(fig_evolucao_cat, use_container_width=True)

    # TAB 3 - Performance & KPIs
    elif st.session_state.active_tab == "🎯 Performance & KPIs":
        st.header("🎯 Performance & KPIs Avançados")

        # Taxa de crescimento por categoria
        st.subheader("📊 Taxa de Crescimento vs Período homólogo")

        # Mostrar períodos de comparação
        st.caption(f"📅 **Período Atual:** {data_inicio.strftime('%d/%m/%Y')} a {data_fim.strftime('%d/%m/%Y')}")
        if not df_anterior.empty:
            st.caption(f"📅 **Período Comparação:** {data_inicio_anterior.strftime('%d/%m/%Y')} a {data_fim_anterior.strftime('%d/%m/%Y')} (mesmo período ano anterior)")

        if not df_anterior.empty:
            crescimento_cat = pd.DataFrame({
                'Atual': df_filtrado.groupby('Categoria')['Valor'].sum(),
                'Anterior': df_anterior.groupby('Categoria')['Valor'].sum()
            }).fillna(0)

            crescimento_cat['Crescimento (%)'] = crescimento_cat.apply(
                lambda row: calcular_delta(row['Atual'], row['Anterior']),
                axis=1
            )

            crescimento_cat = crescimento_cat.sort_values('Crescimento (%)', ascending=False)

            fig_crescimento = px.bar(
                x=crescimento_cat.index,
                y=crescimento_cat['Crescimento (%)'],
                #title='Taxa de Crescimento por Categoria (vs Período homólogo)',
                labels={'x': 'Categoria', 'y': 'Crescimento (%)'},
                color=crescimento_cat['Crescimento (%)'],
                color_continuous_scale='RdYlGn',
                text=[f"{v:+.1f}%" for v in crescimento_cat['Crescimento (%)']]
            )

            fig_crescimento.update_layout(height=400, showlegend=False)
            fig_crescimento.add_hline(y=0, line_dash="dash", line_color="black")

            st.plotly_chart(fig_crescimento, use_container_width=True, key="chart_crescimento_homologo")
        else:
            st.info("Dados do período anterior não disponíveis para comparação.")

        # Taxa de crescimento vs período consecutivo anterior
        st.subheader("📊 Taxa de Crescimento vs Período Consecutivo")

        # Calcular período consecutivo anterior (imediatamente antes do período atual)
        dias_periodo_atual = (data_fim - data_inicio).days + 1
        data_inicio_consecutivo = data_inicio - timedelta(days=dias_periodo_atual)
        data_fim_consecutivo = data_inicio - timedelta(days=1)

        st.caption(f"📅 **Período Atual:** {data_inicio.strftime('%d/%m/%Y')} a {data_fim.strftime('%d/%m/%Y')}")
        st.caption(f"📅 **Período Consecutivo:** {data_inicio_consecutivo.strftime('%d/%m/%Y')} a {data_fim_consecutivo.strftime('%d/%m/%Y')} ({dias_periodo_atual} dias anteriores)")

        # Filtrar dados do período consecutivo
        df_consecutivo = filtrar_dados(
            df,
            categorias_selecionadas,
            subcategorias_selecionadas,
            fontes_selecionadas,
            data_inicio_consecutivo,
            data_fim_consecutivo,
            excluir_domingos
        )

        if not df_consecutivo.empty:
            crescimento_consecutivo = pd.DataFrame({
                'Atual': df_filtrado.groupby('Categoria')['Valor'].sum(),
                'Consecutivo': df_consecutivo.groupby('Categoria')['Valor'].sum()
            }).fillna(0)

            crescimento_consecutivo['Crescimento (%)'] = crescimento_consecutivo.apply(
                lambda row: calcular_delta(row['Atual'], row['Consecutivo']),
                axis=1
            )

            crescimento_consecutivo = crescimento_consecutivo.sort_values('Crescimento (%)', ascending=False)

            fig_consecutivo = px.bar(
                x=crescimento_consecutivo.index,
                y=crescimento_consecutivo['Crescimento (%)'],
                #title='Taxa de Crescimento por Categoria (vs Período Consecutivo)',
                labels={'x': 'Categoria', 'y': 'Crescimento (%)'},
                color=crescimento_consecutivo['Crescimento (%)'],
                color_continuous_scale='RdYlGn',
                text=[f"{v:+.1f}%" for v in crescimento_consecutivo['Crescimento (%)']]
            )

            fig_consecutivo.update_layout(height=400, showlegend=False)
            fig_consecutivo.add_hline(y=0, line_dash="dash", line_color="black")

            st.plotly_chart(fig_consecutivo, use_container_width=True, key="chart_crescimento_consecutivo")
        else:
            st.info("Dados do período consecutivo não disponíveis para comparação.")

        st.markdown("---")

        # Análise de margem (Top vs Bottom performers)
        st.subheader("🏆 Top & Bottom Performers")

        col1, col2 = st.columns(2)

        with col1:
            st.markdown("**🥇 Top 10 Produtos**")
            top_performers = df_filtrado.groupby('Produto')['Valor'].sum().sort_values(ascending=False).head(10)

            for idx, (produto, valor) in enumerate(top_performers.items(), 1):
                percentual = (valor / total_vendas) * 100
                st.markdown(f"{idx}. **{produto}**: {formatar_moeda(valor)} ({percentual:.1f}%)")

        with col2:
            st.markdown("**📉 Bottom 10 Produtos**")
            bottom_performers = df_filtrado.groupby('Produto')['Valor'].sum().sort_values(ascending=True).head(10)

            for idx, (produto, valor) in enumerate(bottom_performers.items(), 1):
                percentual = (valor / total_vendas) * 100
                st.markdown(f"{idx}. **{produto}**: {formatar_moeda(valor)} ({percentual:.1f}%)")

        st.markdown("---")

        # Análise de Pareto (80/20)
        st.subheader("📈 Análise de Pareto (Regra 80/20)")

        vendas_produtos = df_filtrado.groupby('Produto')['Valor'].sum().sort_values(ascending=False)
        vendas_produtos_cum = vendas_produtos.cumsum() / vendas_produtos.sum() * 100

        # Encontrar onde atinge 80%
        produtos_80 = (vendas_produtos_cum <= 80).sum()
        percentual_produtos_80 = (produtos_80 / len(vendas_produtos)) * 100

        col1, col2, col3 = st.columns(3)

        with col1:
            st.metric(
                "🎯 Produtos que representam 80% das vendas",
                f"{produtos_80} produtos"
            )

        with col2:
            st.metric(
                "📊 Percentual do catálogo",
                formatar_percentagem(percentual_produtos_80)
            )

        with col3:
            total_produtos = len(vendas_produtos)
            st.metric(
                "📦 Total de produtos",
                f"{total_produtos} produtos"
            )

        # Gráfico de Pareto
        fig_pareto = go.Figure()

        fig_pareto.add_trace(go.Bar(
            x=list(range(1, min(51, len(vendas_produtos) + 1))),
            y=vendas_produtos.head(50).values,
            name='Vendas',
            yaxis='y',
            marker=dict(color='#1f77b4')
        ))

        fig_pareto.add_trace(go.Scatter(
            x=list(range(1, min(51, len(vendas_produtos_cum) + 1))),
            y=vendas_produtos_cum.head(50).values,
            name='Acumulado (%)',
            yaxis='y2',
            mode='lines+markers',
            marker=dict(color='#ff7f0e'),
            line=dict(width=2)
        ))

        # Linha horizontal em 80% no eixo secundário
        fig_pareto.add_shape(
            type="line",
            xref="paper", x0=0, x1=1,
            yref="y2", y0=80, y1=80,
            line=dict(color="red", width=2, dash="dash")
        )
        fig_pareto.add_annotation(
            x=0.5, xref="paper",
            y=80, yref="y2",
            text="80%",
            showarrow=False,
            bgcolor="white"
        )

        fig_pareto.update_layout(
            title='Gráfico de Pareto - Top 50 Produtos',
            xaxis=dict(title='Ranking de Produtos'),
            yaxis=dict(title='Vendas (€)', side='left'),
            yaxis2=dict(title='Acumulado (%)', side='right', overlaying='y', range=[0, 100]),
            height=500,
            hovermode='x unified'
        )

        st.plotly_chart(fig_pareto, use_container_width=True)

        st.markdown("---")

        # Correlação entre categorias
        st.subheader("🔗 Correlação entre Categorias")

        # Criar matriz de correlação
        vendas_diarias_cat = df_filtrado.pivot_table(
            values='Valor',
            index='Data',
            columns='Categoria',
            aggfunc='sum',
            fill_value=0
        )

        if len(vendas_diarias_cat.columns) > 1:
            correlacao = vendas_diarias_cat.corr()

            fig_corr = go.Figure(data=go.Heatmap(
                z=correlacao.values,
                x=correlacao.columns,
                y=correlacao.columns,
                colorscale='RdBu',
                zmid=0,
                text=np.round(correlacao.values, 2),
                texttemplate='%{text}',
                textfont={"size": 10},
                colorbar=dict(title="Correlação")
            ))

            fig_corr.update_layout(
                title='Matriz de Correlação entre Categorias',
                height=500
            )

            st.plotly_chart(fig_corr, use_container_width=True)

            st.info("A correlação varia de -1 (correlação negativa perfeita) a +1 (correlação positiva perfeita). Valores próximos de 0 indicam pouca ou nenhuma correlação.")
        else:
            st.info("Dados insuficientes para análise de correlação.")

        st.markdown("---")

        # Índice de diversificação
        st.subheader("🌈 Índice de Diversificação de Produtos")

        # Calcular índice Herfindahl-Hirschman (HHI)
        market_shares = (vendas_produtos / vendas_produtos.sum()) ** 2
        hhi = market_shares.sum() * 10000  # Multiplicar por 10000 para escala tradicional

        # Interpretar HHI
        if hhi < 1500:
            interpretacao = "📊 Mercado Altamente Competitivo/Diversificado"
            cor_hhi = "green"
        elif hhi < 2500:
            interpretacao = "⚠️ Mercado Moderadamente Concentrado"
            cor_hhi = "orange"
        else:
            interpretacao = "🎯 Mercado Altamente Concentrado"
            cor_hhi = "red"

        col1, col2, col3 = st.columns(3)

        with col1:
            st.metric(
                "Índice HHI",
                f"{hhi:.0f}"
            )

        with col2:
            st.metric(
                "Total de Produtos",
                f"{len(vendas_produtos)}"
            )

        with col3:
            st.markdown(f"**Interpretação:**")
            st.markdown(f"<p style='color: {cor_hhi}; font-weight: bold;'>{interpretacao}</p>", unsafe_allow_html=True)

        st.markdown("""
        **Sobre o Índice HHI:**
        - HHI < 1500: Mercado diversificado
        - 1500 ≤ HHI < 2500: Concentração moderada
        - HHI ≥ 2500: Alta concentração
        """)

        st.markdown("---")

        # Análise Temporal integrada
        st.subheader("📈 Análise Temporal")

        # Vendas por dia da semana
        st.markdown("**📅 Vendas por Dia da Semana**")

        dias_semana = ['Segunda', 'Terça', 'Quarta', 'Quinta', 'Sexta', 'Sábado', 'Domingo']
        df_filtrado['Dia_Semana_Nome'] = df_filtrado['Dia_Semana'].apply(lambda x: dias_semana[x])

        vendas_por_dia_semana = df_filtrado.groupby('Dia_Semana_Nome')['Valor'].sum()
        vendas_por_dia_semana = vendas_por_dia_semana.reindex(dias_semana)

        fig_dia_semana = px.bar(
            x=vendas_por_dia_semana.index,
            y=vendas_por_dia_semana.values,
            title='Vendas por Dia da Semana',
            labels={'x': 'Dia da Semana', 'y': 'Vendas (€)'},
            text=[formatar_moeda(v) for v in vendas_por_dia_semana.values],
            color=vendas_por_dia_semana.values,
            color_continuous_scale='Blues'
        )

        fig_dia_semana.update_layout(
            height=400,
            showlegend=False
        )

        st.plotly_chart(fig_dia_semana, use_container_width=True, key="chart_dia_semana")

        # Vendas por mês
        st.markdown("**📆 Vendas por Mês**")

        if 'Mes' in df_filtrado.columns and 'Ano' in df_filtrado.columns:
            df_filtrado['Mes_Ano'] = df_filtrado['Data'].dt.to_period('M').astype(str)
            vendas_por_mes = df_filtrado.groupby('Mes_Ano')['Valor'].sum().reset_index()

            fig_mes = px.bar(
                vendas_por_mes,
                x='Mes_Ano',
                y='Valor',
                title='Vendas Mensais',
                text=[formatar_moeda(v) for v in vendas_por_mes['Valor']],
                color='Valor',
                color_continuous_scale='Viridis'
            )

            fig_mes.update_layout(
                height=400,
                xaxis_title="Mês/Ano",
                yaxis_title="Vendas (€)",
                showlegend=False
            )

            st.plotly_chart(fig_mes, use_container_width=True, key="chart_vendas_mes")

        # Sazonalidade
        st.markdown("**🌊 Análise de Sazonalidade**")

        # Vendas por trimestre
        if 'Trimestre' in df_filtrado.columns:
            vendas_trimestre = df_filtrado.groupby('Trimestre')['Valor'].sum()

            fig_trimestre = px.pie(
                values=vendas_trimestre.values,
                names=[f'Q{i}' for i in vendas_trimestre.index],
                title='Distribuição por Trimestre',
                hole=0.4
            )

            fig_trimestre.update_layout(height=400)
            st.plotly_chart(fig_trimestre, use_container_width=True, key="chart_trimestre")

        # Análise Semanal
        if 'Semana' in df_filtrado.columns:
            st.markdown("**📅 Análise Semanal**")

            col1, col2 = st.columns(2)

            with col1:
                # Todas as semanas (evolução temporal)
                vendas_semana_todas = df_filtrado.groupby('Semana')['Valor'].sum().sort_index()

                fig_semana_todas = px.line(
                    x=vendas_semana_todas.index,
                    y=vendas_semana_todas.values,
                    title='Evolução Semanal - Todas as Semanas',
                    labels={'x': 'Semana do Ano', 'y': 'Vendas (€)'},
                    markers=True
                )

                fig_semana_todas.update_layout(height=400)
                st.plotly_chart(fig_semana_todas, use_container_width=True, key="chart_semana_todas")
                st.caption(f"📊 Total de {len(vendas_semana_todas)} semanas no período")

            with col2:
                # Top 20 semanas com mais vendas
                vendas_semana_top = df_filtrado.groupby('Semana')['Valor'].sum().sort_values(ascending=False).head(20)

                fig_semana_top = px.bar(
                    x=[f"S{s}" for s in vendas_semana_top.index],
                    y=vendas_semana_top.values,
                    title='Top 20 Semanas com Mais Vendas',
                    labels={'x': 'Semana do Ano', 'y': 'Vendas (€)'},
                    text=[formatar_moeda(v) for v in vendas_semana_top.values],
                    color=vendas_semana_top.values,
                    color_continuous_scale='Viridis'
                )

                fig_semana_top.update_layout(height=400, showlegend=False)
                fig_semana_top.update_traces(textposition='outside')
                st.plotly_chart(fig_semana_top, use_container_width=True, key="chart_semana_top20")
                st.caption(f"🏆 Melhores semanas: S{vendas_semana_top.index[0]} lidera com {formatar_moeda(vendas_semana_top.values[0])}")

    # TAB 4 - Análise Financeira
    elif st.session_state.active_tab == "💰 Análise Financeira":
        st.header("💰 Análise Financeira Detalhada")

        # Faturamento por categoria
        st.subheader("💵 Faturamento Total por Categoria")

        faturamento_cat = df_filtrado.groupby('Categoria')['Valor'].sum().sort_values(ascending=False)

        fig_faturamento = px.bar(
            x=faturamento_cat.index,
            y=faturamento_cat.values,
            title='Faturamento por Categoria',
            labels={'x': 'Categoria', 'y': 'Faturamento (€)'},
            text=[formatar_moeda(v) for v in faturamento_cat.values],
            color=faturamento_cat.values,
            color_continuous_scale='Greens'
        )

        fig_faturamento.update_layout(
            height=400,
            showlegend=False
        )

        st.plotly_chart(fig_faturamento, use_container_width=True)

        st.markdown("---")

        # Métricas financeiras
        col1, col2, col3 = st.columns(3)

        with col1:
            st.subheader("📊 Faturamento Médio")
            num_transacoes = len(df_filtrado)
            faturamento_medio_transacao = total_vendas / num_transacoes if num_transacoes > 0 else 0
            st.metric("Por Transação", formatar_moeda(faturamento_medio_transacao))

        with col2:
            st.subheader("🎯 Ticket Médio")
            ticket_medio_cat = df_filtrado.groupby('Categoria').apply(
                lambda x: x['Valor'].sum() / x['Qtd'].sum() if x['Qtd'].sum() > 0 else 0
            ).mean()
            st.metric("Por Categoria", formatar_moeda(ticket_medio_cat))

        with col3:
            st.subheader("📈 Projeção Mensal")
            if dias_periodo > 0:
                projecao_mensal = (total_vendas / dias_periodo) * 30
                st.metric("Estimativa", formatar_moeda(projecao_mensal))

        st.markdown("---")

        # Ticket médio por categoria
        st.subheader("🎫 Ticket Médio por Categoria")

        ticket_medio_categoria = df_filtrado.groupby('Categoria').apply(
            lambda x: x['Valor'].sum() / x['Qtd'].sum() if x['Qtd'].sum() > 0 else 0
        ).sort_values(ascending=False)

        fig_ticket = px.bar(
            x=ticket_medio_categoria.index,
            y=ticket_medio_categoria.values,
            title='Ticket Médio por Categoria',
            labels={'x': 'Categoria', 'y': 'Ticket Médio (€)'},
            text=[formatar_moeda(v) for v in ticket_medio_categoria.values],
            color=ticket_medio_categoria.values,
            color_continuous_scale='Blues'
        )

        fig_ticket.update_layout(height=400, showlegend=False)
        st.plotly_chart(fig_ticket, use_container_width=True)

        st.markdown("---")

        # Gráfico Waterfall de contribuições
        st.subheader("💧 Contribuição por Categoria (Waterfall)")

        contribuicoes = df_filtrado.groupby('Categoria')['Valor'].sum().sort_values(ascending=False)

        if len(contribuicoes) > 0:
            labels = list(contribuicoes.index) + ['Total']
            valores = list(contribuicoes.values) + [0]  # 0 para o total

            fig_waterfall = criar_grafico_waterfall(valores, labels, 'Contribuição de Cada Categoria')
            st.plotly_chart(fig_waterfall, use_container_width=True)

        st.markdown("---")

        # Projeção de faturamento
        st.subheader("🔮 Projeção de Faturamento (Regressão Linear)")

        # Preparar dados para regressão
        vendas_diarias_reg = df_filtrado.groupby('Data')['Valor'].sum().reset_index()
        vendas_diarias_reg['Dias'] = (vendas_diarias_reg['Data'] - vendas_diarias_reg['Data'].min()).dt.days

        if len(vendas_diarias_reg) >= 5:
            # Calcular regressão linear
            from numpy.polynomial import Polynomial

            x = vendas_diarias_reg['Dias'].values
            y = vendas_diarias_reg['Valor'].values

            # Fit linear
            p = Polynomial.fit(x, y, 1)

            # Projetar próximos 30 dias
            dias_futuros = np.arange(x[-1] + 1, x[-1] + 31)
            projecao = p(dias_futuros)

            # Criar gráfico
            fig_projecao = go.Figure()

            # Dados históricos
            fig_projecao.add_trace(go.Scatter(
                x=vendas_diarias_reg['Data'],
                y=vendas_diarias_reg['Valor'],
                mode='markers',
                name='Vendas Reais',
                marker=dict(size=6, color='#1f77b4')
            ))

            # Linha de tendência
            fig_projecao.add_trace(go.Scatter(
                x=vendas_diarias_reg['Data'],
                y=p(x),
                mode='lines',
                name='Tendência',
                line=dict(color='red', dash='dash')
            ))

            # Projeção
            datas_futuras = pd.date_range(
                start=vendas_diarias_reg['Data'].max() + timedelta(days=1),
                periods=30
            )

            fig_projecao.add_trace(go.Scatter(
                x=datas_futuras,
                y=projecao,
                mode='lines+markers',
                name='Projeção (30 dias)',
                line=dict(color='green', dash='dot'),
                marker=dict(size=4)
            ))

            fig_projecao.update_layout(
                title='Projeção de Vendas - Próximos 30 Dias',
                xaxis_title='Data',
                yaxis_title='Vendas (€)',
                height=500,
                hovermode='x unified'
            )

            st.plotly_chart(fig_projecao, use_container_width=True)

            # Métricas da projeção
            col1, col2, col3 = st.columns(3)

            with col1:
                st.metric("Projeção Total (30 dias)", formatar_moeda(projecao.sum()))

            with col2:
                st.metric("Média Diária Projetada", formatar_moeda(projecao.mean()))

            with col3:
                tendencia = "📈 Crescente" if p.coef[1] > 0 else "📉 Decrescente"
                st.metric("Tendência", tendencia)

            st.info("⚠️ Esta é uma projeção simples baseada em regressão linear. Considere fatores sazonais e externos para decisões estratégicas.")
        else:
            st.info("Dados insuficientes para projeção (mínimo 5 dias).")

    # TAB 5 - Comparação & Benchmarks
    elif st.session_state.active_tab == "📊 Comparação & Benchmarks":
        st.header("📊 Comparação & Benchmarks")

        # Comparação entre categorias (Scatter Plot)
        st.subheader("🎯 Comparação: Volume vs Valor Médio")

        comparacao_cat = df_filtrado.groupby('Categoria').agg({
            'Valor': ['sum', 'mean', 'count']
        })

        comparacao_cat.columns = ['Total', 'Media', 'Vendas']
        comparacao_cat = comparacao_cat.reset_index()

        # Adicionar ícones
        comparacao_cat['Icone'] = comparacao_cat['Categoria'].apply(
            lambda x: categorizer.CATEGORIAS.get(x, {}).get('icone', '📦')
        )

        fig_scatter = px.scatter(
            comparacao_cat,
            x='Vendas',
            y='Media',
            size='Total',
            color='Categoria',
            hover_name='Categoria',
            title='Análise: Número de Vendas vs Valor Médio (Tamanho = Faturamento Total)',
            labels={
                'Vendas': 'Número de Vendas',
                'Media': 'Valor Médio por Venda (€)',
                'Total': 'Faturamento Total (€)'
            },
            size_max=60
        )

        fig_scatter.update_layout(
            height=500,
            showlegend=True
        )

        st.plotly_chart(fig_scatter, use_container_width=True)

        st.markdown("---")

        # Benchmarks por fonte
        st.subheader("📍 Performance por Fonte de Dados")

        col1, col2 = st.columns(2)

        with col1:
            # Faturamento por fonte
            faturamento_fonte = df_filtrado.groupby('Fonte')['Valor'].sum().sort_values(ascending=False)

            fig_fonte = px.pie(
                values=faturamento_fonte.values,
                names=faturamento_fonte.index,
                title='Distribuição de Faturamento por Fonte',
                hole=0.4
            )

            fig_fonte.update_layout(height=400)
            st.plotly_chart(fig_fonte, use_container_width=True)

        with col2:
            # Ticket médio por fonte
            ticket_fonte = df_filtrado.groupby('Fonte').apply(
                lambda x: x['Valor'].sum() / x['Qtd'].sum() if x['Qtd'].sum() > 0 else 0
            ).sort_values(ascending=False)

            fig_ticket_fonte = px.bar(
                x=ticket_fonte.index,
                y=ticket_fonte.values,
                title='Ticket Médio por Fonte',
                labels={'x': 'Fonte', 'y': 'Ticket Médio (€)'},
                text=[formatar_moeda(v) for v in ticket_fonte.values],
                color=ticket_fonte.values,
                color_continuous_scale='Viridis'
            )

            fig_ticket_fonte.update_layout(height=400, showlegend=False)
            st.plotly_chart(fig_ticket_fonte, use_container_width=True)

        st.markdown("---")

        # Tabela de benchmarks
        st.subheader("📊 Tabela de Benchmarks Detalhados")

        benchmark = df_filtrado.groupby('Fonte').agg({
            'Valor': ['sum', 'mean', 'std'],
            'Qtd': 'sum',
            'Data': 'count'
        }).round(2)

        benchmark.columns = ['Faturamento Total', 'Média por Venda', 'Desvio Padrão', 'Qtd Total', 'Nº Vendas']
        benchmark['Ticket Médio'] = (benchmark['Faturamento Total'] / benchmark['Qtd Total']).round(2)
        benchmark['% do Total'] = ((benchmark['Faturamento Total'] / benchmark['Faturamento Total'].sum()) * 100).round(1)

        # Formatar para exibição
        benchmark_display = benchmark.copy()
        benchmark_display['Faturamento Total'] = benchmark_display['Faturamento Total'].apply(formatar_moeda)
        benchmark_display['Média por Venda'] = benchmark_display['Média por Venda'].apply(formatar_moeda)
        benchmark_display['Desvio Padrão'] = benchmark_display['Desvio Padrão'].apply(formatar_moeda)
        benchmark_display['Ticket Médio'] = benchmark_display['Ticket Médio'].apply(formatar_moeda)
        benchmark_display['% do Total'] = benchmark_display['% do Total'].apply(lambda x: f"{x}%")

        st.dataframe(benchmark_display, use_container_width=True)

        st.markdown("---")

        # Análise de sazonalidade comparativa
        st.subheader("🌊 Sazonalidade: Comparação entre Categorias")

        # Vendas por mês e categoria
        if 'Mes' in df_filtrado.columns:
            vendas_mes_cat = df_filtrado.pivot_table(
                values='Valor',
                index='Mes',
                columns='Categoria',
                aggfunc='sum',
                fill_value=0
            )

            fig_sazonalidade = go.Figure()

            for categoria in vendas_mes_cat.columns:
                fig_sazonalidade.add_trace(go.Scatter(
                    x=vendas_mes_cat.index,
                    y=vendas_mes_cat[categoria],
                    mode='lines+markers',
                    name=categoria,
                    line=dict(width=2),
                    marker=dict(size=8)
                ))

            fig_sazonalidade.update_layout(
                title='Sazonalidade por Categoria (Mensal)',
                xaxis_title='Mês',
                yaxis_title='Vendas (€)',
                height=500,
                hovermode='x unified'
            )

            st.plotly_chart(fig_sazonalidade, use_container_width=True)

        st.markdown("---")

        # Performance relativa
        st.subheader("🏅 Performance Relativa por Categoria")

        # Normalizar valores para comparação
        performance_rel = df_filtrado.groupby('Categoria').agg({
            'Valor': 'sum',
            'Qtd': 'sum'
        })

        # Normalizar (0-100)
        performance_rel['Score_Valor'] = (performance_rel['Valor'] / performance_rel['Valor'].max() * 100).round(1)
        performance_rel['Score_Qtd'] = (performance_rel['Qtd'] / performance_rel['Qtd'].max() * 100).round(1)
        performance_rel['Score_Final'] = ((performance_rel['Score_Valor'] + performance_rel['Score_Qtd']) / 2).round(1)

        performance_rel = performance_rel.sort_values('Score_Final', ascending=False)

        fig_performance = go.Figure()

        fig_performance.add_trace(go.Bar(
            name='Score Valor',
            x=performance_rel.index,
            y=performance_rel['Score_Valor'],
            marker_color='#1f77b4'
        ))

        fig_performance.add_trace(go.Bar(
            name='Score Quantidade',
            x=performance_rel.index,
            y=performance_rel['Score_Qtd'],
            marker_color='#ff7f0e'
        ))

        fig_performance.add_trace(go.Scatter(
            name='Score Final',
            x=performance_rel.index,
            y=performance_rel['Score_Final'],
            mode='lines+markers',
            marker=dict(size=10, color='red'),
            line=dict(width=3, color='red'),
            yaxis='y2'
        ))

        fig_performance.update_layout(
            title='Performance Relativa (Normalizada 0-100)',
            barmode='group',
            height=500,
            yaxis=dict(title='Score Individual'),
            yaxis2=dict(title='Score Final', overlaying='y', side='right'),
            hovermode='x unified'
        )

        st.plotly_chart(fig_performance, use_container_width=True)

    # TAB 6 - Dados Detalhados
    elif st.session_state.active_tab == "📑 Dados Detalhados":
        st.header("📑 Dados Detalhados & Exportação")

        # Opções de agrupamento
        st.subheader("⚙️ Opções de Visualização")

        col1, col2, col3 = st.columns(3)

        with col1:
            agrupar_por = st.selectbox(
                "Agrupar dados por:",
                options=['Sem agrupamento', 'Categoria', 'Subcategoria', 'Produto', 'Fonte', 'Data'],
                index=0
            )

        with col2:
            ordenar_por = st.selectbox(
                "Ordenar por:",
                options=['Valor', 'Qtd', 'Data', 'Categoria'],
                index=0
            )

        with col3:
            ordem = st.radio(
                "Ordem:",
                options=['Decrescente', 'Crescente'],
                horizontal=True
            )

        # Preparar dados para exibição
        if agrupar_por == 'Sem agrupamento':
            df_exibir = df_filtrado.copy()

            # Selecionar colunas relevantes
            colunas_exibir = ['Data', 'Produto', 'Categoria', 'Subcategoria', 'Valor', 'Qtd', 'Fonte']
            colunas_exibir = [col for col in colunas_exibir if col in df_exibir.columns]

            df_exibir = df_exibir[colunas_exibir]

            # Ordenar
            ascending = (ordem == 'Crescente')
            df_exibir = df_exibir.sort_values(by=ordenar_por, ascending=ascending)

        else:
            # Agrupar dados
            df_exibir = df_filtrado.groupby(agrupar_por).agg({
                'Valor': 'sum',
                'Qtd': 'sum',
                'Data': 'count'
            }).round(2)

            df_exibir = df_exibir.rename(columns={'Data': 'Nº Vendas'})
            df_exibir['Valor Médio'] = (df_exibir['Valor'] / df_exibir['Nº Vendas']).round(2)

            # Ordenar
            ascending = (ordem == 'Crescente')
            df_exibir = df_exibir.sort_values(by=ordenar_por if ordenar_por in df_exibir.columns else 'Valor', ascending=ascending)

            df_exibir = df_exibir.reset_index()

        # Estatísticas resumidas
        st.subheader("📊 Estatísticas Resumidas")

        col1, col2, col3, col4, col5 = st.columns(5)

        with col1:
            st.metric("Total de Registos", f"{len(df_exibir):,}")

        with col2:
            soma_total = df_filtrado['Valor'].sum()
            st.metric("Soma Total", formatar_moeda(soma_total))

        with col3:
            media_total = df_filtrado['Valor'].mean()
            st.metric("Média", formatar_moeda(media_total))

        with col4:
            mediana_total = df_filtrado['Valor'].median()
            st.metric("Mediana", formatar_moeda(mediana_total))

        with col5:
            desvio_total = df_filtrado['Valor'].std()
            st.metric("Desvio Padrão", formatar_moeda(desvio_total))

        st.markdown("---")

        # Exibir tabela
        st.subheader("📋 Tabela de Dados")

        # Formatar valores monetários para exibição
        df_display = df_exibir.copy()

        if 'Valor' in df_display.columns:
            df_display['Valor'] = df_display['Valor'].apply(formatar_moeda)

        if 'Valor Médio' in df_display.columns:
            df_display['Valor Médio'] = df_display['Valor Médio'].apply(formatar_moeda)

        if 'Qtd' in df_display.columns:
            df_display['Qtd'] = df_display['Qtd'].apply(lambda x: f"{x:,.0f}")

        # Exibir com altura ajustável
        st.dataframe(df_display, use_container_width=True, height=500)

        st.markdown("---")

        # Botões de exportação
        st.subheader("💾 Exportar Dados")

        col1, col2, col3 = st.columns(3)

        with col1:
            # Exportar para CSV
            csv = df_exibir.to_csv(index=False).encode('utf-8')
            st.download_button(
                label="📥 Descarregar CSV",
                data=csv,
                file_name=f"vendas_{datetime.now().strftime('%Y%m%d_%H%M%S')}.csv",
                mime="text/csv",
                use_container_width=True
            )
            st.caption("Exportação simples em CSV")

        with col2:
            # Exportar para Excel Simples
            buffer = BytesIO()
            with pd.ExcelWriter(buffer, engine='openpyxl') as writer:
                df_exibir.to_excel(writer, sheet_name='Vendas', index=False)

                # Adicionar resumo
                df_resumo = pd.DataFrame({
                    'Métrica': ['Total de Vendas', 'Média', 'Mediana', 'Desvio Padrão', 'Total de Registos'],
                    'Valor': [
                        soma_total,
                        media_total,
                        mediana_total,
                        desvio_total,
                        len(df_filtrado)
                    ]
                })
                df_resumo.to_excel(writer, sheet_name='Resumo', index=False)

            st.download_button(
                label="📥 Descarregar Excel",
                data=buffer.getvalue(),
                file_name=f"vendas_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                use_container_width=True
            )
            st.caption("Excel simples (2 abas)")

        with col3:
            # Exportar para Excel Avançado
            if st.button("📊 Gerar Relatório Excel Avançado", use_container_width=True):
                with st.spinner("Gerando relatório Excel profissional..."):
                    try:
                        # Calcular métricas financeiras para o relatório
                        # Converter para datetime para garantir compatibilidade com CostManagerV2
                        dt_inicio = pd.to_datetime(data_inicio).to_pydatetime()
                        dt_fim = pd.to_datetime(data_fim).to_pydatetime()
                        
                        metricas_financeiras_temp = cost_manager.calcular_metricas_financeiras(
                            df_filtrado, dt_inicio, dt_fim
                        )

                        excel_avancado = gerar_relatorio_excel_avancado(
                            df_filtrado,
                            data_inicio,
                            data_fim,
                            categorias_selecionadas,
                            metricas_financeiras_temp,
                            cost_manager
                        )

                        nome_arquivo = f"Relatorio_CafeMartins_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"

                        st.download_button(
                            label="📥 Download Relatório Completo",
                            data=excel_avancado,
                            file_name=nome_arquivo,
                            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                            use_container_width=True,
                            key="download_avancado"
                        )

                        st.success("✅ Relatório avançado gerado!")
                        st.info(
                            "**Conteúdo do relatório:**\n"
                            "- 📋 Resumo Executivo\n"
                            "- 📊 Dados Detalhados\n"
                            "- 🏆 Análise por Categoria (com gráfico)\n"
                            "- 💰 Rentabilidade (com gráfico)\n"
                            "- 📈 Análise Temporal (com gráfico)"
                        )
                    except Exception as e:
                        st.error(f"Erro ao gerar relatório: {str(e)}")
            st.caption("Relatório completo com 5 abas e gráficos")

        st.markdown("---")

        # Informações adicionais
        st.subheader("ℹ️ Informações do Dataset")

        info_cols = st.columns(3)

        with info_cols[0]:
            st.markdown("**📅 Período dos Dados:**")
            st.write(f"Início: {df_filtrado['Data'].min().strftime('%d/%m/%Y')}")
            st.write(f"Fim: {df_filtrado['Data'].max().strftime('%d/%m/%Y')}")

        with info_cols[1]:
            st.markdown("**📊 Categorias:**")
            st.write(f"Total: {df_filtrado['Categoria'].nunique()}")
            st.write(f"Subcategorias: {df_filtrado['Subcategoria'].nunique()}")

        with info_cols[2]:
            st.markdown("**📦 Produtos:**")
            st.write(f"Total: {df_filtrado['Produto'].nunique()}")
            st.write(f"Fontes: {df_filtrado['Fonte'].nunique()}")

    # TAB 7 - Análise de Custos
    elif st.session_state.active_tab == "💸 Análise de Custos":
        st.header("💸 Análise de Custos Operacionais")

        # Calcular métricas de custo (com datas para Despesify)
        data_min_filtrado = df_filtrado['Data'].min().to_pydatetime()
        data_max_filtrado = df_filtrado['Data'].max().to_pydatetime()
        metricas_financeiras = cost_manager.calcular_metricas_financeiras(
            df_filtrado, data_min_filtrado, data_max_filtrado
        )

        # Mostrar fonte dos custos operacionais
        fonte_custos = metricas_financeiras.get('fonte_custos_operacionais', 'ESTIMADO')
        if fonte_custos == 'REAL':
            st.success("✅ Custos Operacionais REAIS do Despesify (desde 1/12/2025)")
        else:
            st.warning("⚠️ Custos Operacionais ESTIMADOS (antes de 1/12/2025 ou sem dados no Despesify)")

        # KPIs de Custos
        col1, col2, col3, col4 = st.columns(4)

        with col1:
            st.metric(
                "💰 Receita Total",
                formatar_moeda(metricas_financeiras.get('receita_total', 0))
            )

        with col2:
            st.metric(
                "📦 Custos de Produtos (COGS)",
                formatar_moeda(metricas_financeiras.get('cogs_total', 0))
            )

        with col3:
            custos_ops_periodo = metricas_financeiras.get('custos_operacionais', 0)
            fonte_custos = metricas_financeiras.get('fonte_custos_operacionais', 'ESTIMADO')

            # Calcular custo mensal para referência
            dias_periodo = (data_max_filtrado - data_min_filtrado).days + 1
            if fonte_custos == 'ESTIMADO' and dias_periodo > 0 and dias_periodo < 30:
                # Mostrar custo mensal estimado para contexto
                custo_mensal_estimado = cost_manager.get_custos_operacionais_mensais_estimados()
                st.metric(
                    "🏢 Custos Operacionais",
                    formatar_moeda(custos_ops_periodo),
                    help=f"💡 Período: {dias_periodo} dia(s) | Custo Mensal Estimado: {formatar_moeda(custo_mensal_estimado)}"
                )
            else:
                st.metric(
                    "🏢 Custos Operacionais",
                    formatar_moeda(custos_ops_periodo),
                    help=f"Fonte: {fonte_custos}"
                )

        with col4:
            custo_total = metricas_financeiras.get('custo_total', 0)
            st.metric(
                "💸 Custo Total",
                formatar_moeda(custo_total)
            )

        st.markdown("---")

        # Distribuição de Custos
        col1, col2 = st.columns(2)

        with col1:
            st.subheader("📊 Distribuição de Custos")

            # Separar comissões Santa Casa de custos de produtos
            comissoes_sc = metricas_financeiras.get('comissoes_santa_casa', 0)
            custos_prods = metricas_financeiras.get('custos_produtos', 0)
            custos_ops = metricas_financeiras.get('custos_operacionais', 0)

            custos_breakdown = {
                'Comissões Santa Casa': comissoes_sc,
                'Custos de Produtos': custos_prods,
                'Custos Operacionais': custos_ops
            }

            fig_custos = go.Figure(data=[go.Pie(
                labels=list(custos_breakdown.keys()),
                values=list(custos_breakdown.values()),
                hole=0.4,
                marker=dict(colors=['#4ecdc4', '#ff6b6b', '#ffd700'])
            )])

            fig_custos.update_layout(
                title="Composição dos Custos",
                height=400
            )

            st.plotly_chart(fig_custos, use_container_width=True)

            # Info sobre custos Santa Casa
            if comissoes_sc > 0:
                # Calcular vendas de Santa Casa para percentagem correta
                vendas_sc = df_filtrado[df_filtrado['Is_Santa_Casa'] == True]['Valor'].sum() if 'Is_Santa_Casa' in df_filtrado.columns else 0
                if vendas_sc > 0:
                    percentagem_custo_sc = (comissoes_sc / vendas_sc) * 100
                    margem_sc = ((vendas_sc - comissoes_sc) / vendas_sc) * 100
                    st.info(f"💡 Custos Jogos Santa Casa: {formatar_moeda(comissoes_sc)} ({percentagem_custo_sc:.1f}% das vendas de jogos | Margem: {margem_sc:.1f}%)")

        with col2:
            st.subheader("💹 Custos Operacionais Detalhados")

            # Obter resumo de custos operacionais do período (REAL ou ESTIMADO)
            resumo_custos_op = cost_manager.get_resumo_custos_operacionais(
                data_min_filtrado, data_max_filtrado
            )

            if not resumo_custos_op.empty:
                # Usar coluna correta dependendo da fonte
                coluna_valor = 'Total_Periodo' if 'Total_Periodo' in resumo_custos_op.columns else 'Total_Mensal'

                fig_custos_op = px.bar(
                    x=resumo_custos_op.index,
                    y=resumo_custos_op[coluna_valor],
                    title=f"Custos Operacionais por Categoria ({fonte_custos})",
                    labels={'x': 'Categoria', 'y': f'Custo do Período (€)'},
                    text=[formatar_moeda(v) for v in resumo_custos_op[coluna_valor]],
                    color=resumo_custos_op[coluna_valor],
                    color_continuous_scale='Reds'
                )

                fig_custos_op.update_layout(height=400, showlegend=False)
                st.plotly_chart(fig_custos_op, use_container_width=True)

        st.markdown("---")

        # Tabela de custos operacionais
        st.subheader("📋 Detalhamento de Custos Operacionais")

        # Toggle para modo de edição
        col_edit, col_info = st.columns([1, 3])
        with col_edit:
            modo_edicao = st.toggle("✏️ Modo Edição", key="modo_edicao_custos_op", help="Ativar para editar custos operacionais ESTIMADOS (CSV)")

        with col_info:
            if modo_edicao:
                st.info("ℹ️ Modo Edição ativado. Edite os valores abaixo e clique em 'Guardar Alterações'")

        if not resumo_custos_op.empty:
            if modo_edicao and fonte_custos == 'ESTIMADO':
                # MODO EDIÇÃO - Editor de custos operacionais
                st.markdown("**✏️ Editar Custos Operacionais Mensais (CSV):**")

                # Instruções de uso
                st.info("""
                📝 **Como usar o editor:**
                - ✏️ **Editar valores:** Clique numa célula e altere o valor
                - ➕ **Adicionar novo item:** Clique no botão ➕ no final da tabela
                - 🗑️ **Remover item:** Selecione a linha (checkbox à esquerda) e clique no ícone 🗑️
                - 💾 Clique em "Guardar Alterações" para salvar
                """)

                # Carregar CSV de custos operacionais
                import os
                csv_path = os.path.join('dados_custos', 'custos_operacionais.csv')

                if os.path.exists(csv_path):
                    df_custos_csv = pd.read_csv(csv_path)

                    # Filtrar custos não calculados
                    df_editavel = df_custos_csv[df_custos_csv['Tipo'] != 'Calculado'].copy()

                    # Editor de dados
                    df_editado = st.data_editor(
                        df_editavel,
                        use_container_width=True,
                        num_rows="dynamic",
                        column_config={
                            "Categoria": st.column_config.TextColumn(
                                "Categoria",
                                required=True,
                                help="Categoria principal do custo"
                            ),
                            "Subcategoria": st.column_config.TextColumn(
                                "Subcategoria",
                                required=True,
                                help="Descrição específica do custo"
                            ),
                            "Valor_Mensal": st.column_config.NumberColumn(
                                "Valor Mensal (€)",
                                min_value=0,
                                format="€%.2f",
                                required=True,
                                help="Valor mensal do custo"
                            ),
                            "Tipo": st.column_config.SelectboxColumn(
                                "Tipo",
                                options=["Fixo", "Variável"],
                                required=True,
                                help="Fixo = valor constante, Variável = pode mudar"
                            ),
                            "Notas": st.column_config.TextColumn(
                                "Notas",
                                help="Observações adicionais"
                            )
                        },
                        hide_index=True,
                        key="editor_custos_op"
                    )

                    # Botões de ação
                    col1, col2, col3 = st.columns([2, 2, 6])

                    with col1:
                        if st.button("💾 Guardar Alterações", type="primary", key="guardar_custos_op"):
                            try:
                                # Guardar de volta ao CSV
                                df_editado.to_csv(csv_path, index=False)
                                st.success("✅ Custos operacionais guardados com sucesso!")
                                st.info("🔄 Por favor, recarregue a página (F5) para ver as alterações aplicadas.")
                            except Exception as e:
                                st.error(f"❌ Erro ao guardar: {e}")

                    with col2:
                        if st.button("🔄 Resetar", key="resetar_custos_op"):
                            st.rerun()

                    # Mostrar total
                    total_editado = df_editado['Valor_Mensal'].sum()
                    st.metric("💰 Total Mensal (Editado)", f"{total_editado:,.2f}€")

                else:
                    st.error(f"❌ Ficheiro não encontrado: {csv_path}")

            elif modo_edicao and fonte_custos == 'REAL':
                # Permitir edição de custos estimados mesmo com Despesify ativo
                st.info("💡 **Custos REAIS do Despesify estão ativos neste período.** Você pode editar os custos operacionais abaixo que ainda não estão no Despesify (ex: rendas, ordenados, segurança social). Para editar custos que já estão no Despesify, edite diretamente no sistema Despesify.")

                # Carregar CSV de custos operacionais
                import os
                csv_path = os.path.join('dados_custos', 'custos_operacionais.csv')

                if os.path.exists(csv_path):
                    df_custos_csv = pd.read_csv(csv_path)

                    # Filtrar custos não calculados
                    df_editavel = df_custos_csv[df_custos_csv['Tipo'] != 'Calculado'].copy()

                    if not df_editavel.empty:
                        st.markdown("**✏️ Editar Custos Operacionais Mensais (CSV):**")

                        # Instruções de uso
                        st.info("""
                        📝 **Como usar o editor:**
                        - ✏️ **Editar valores:** Clique numa célula e altere o valor
                        - ➕ **Adicionar novo item:** Clique no botão ➕ no final da tabela
                        - 🗑️ **Remover item:** Selecione a linha (checkbox à esquerda) e clique no ícone 🗑️
                        - 💾 Clique em "Guardar Alterações" para salvar
                        """)

                        # Editor de dados
                        df_editado = st.data_editor(
                            df_editavel,
                            use_container_width=True,
                            num_rows="dynamic",
                            column_config={
                                "Categoria": st.column_config.TextColumn(
                                    "Categoria",
                                    required=True,
                                    help="Categoria principal do custo"
                                ),
                                "Subcategoria": st.column_config.TextColumn(
                                    "Subcategoria",
                                    required=True,
                                    help="Descrição específica do custo"
                                ),
                                "Valor_Mensal": st.column_config.NumberColumn(
                                    "Valor Mensal (€)",
                                    min_value=0,
                                    format="€%.2f",
                                    required=True,
                                    help="Valor mensal do custo"
                                ),
                                "Tipo": st.column_config.SelectboxColumn(
                                    "Tipo",
                                    options=["Fixo", "Variável"],
                                    required=True,
                                    help="Fixo = valor constante, Variável = pode mudar"
                                ),
                                "Notas": st.column_config.TextColumn(
                                    "Notas",
                                    help="Observações adicionais"
                                )
                            },
                            hide_index=True,
                            key="editor_custos_op_com_despesify"
                        )

                        # Botões de ação
                        col1, col2, col3 = st.columns([2, 2, 6])

                        with col1:
                            if st.button("💾 Guardar Alterações", type="primary", key="guardar_custos_op_despesify"):
                                try:
                                    # Guardar de volta ao CSV
                                    df_editado.to_csv(csv_path, index=False)
                                    st.success("✅ Custos operacionais guardados com sucesso!")
                                    st.info("🔄 Por favor, recarregue a página (F5) para ver as alterações aplicadas.")
                                except Exception as e:
                                    st.error(f"❌ Erro ao guardar: {e}")

                        with col2:
                            if st.button("🔄 Resetar", key="resetar_custos_op_despesify"):
                                st.rerun()

                        # Mostrar total
                        total_editado = df_editado['Valor_Mensal'].sum()
                        st.metric("💰 Total Mensal (Editado - CSV)", f"{total_editado:,.2f}€")
                    else:
                        st.info("ℹ️ Não há custos estimados no CSV para editar. Todos os custos operacionais vêm do Despesify.")
                else:
                    st.warning(f"⚠️ Ficheiro não encontrado: {csv_path}")

                # Mostrar resumo dos custos do Despesify
                st.markdown("---")
                st.markdown("**📊 Custos do Despesify (não editáveis aqui):**")
                resumo_display = resumo_custos_op.copy()

                if 'Total_Periodo' in resumo_display.columns:
                    resumo_display['Total_Periodo'] = resumo_display['Total_Periodo'].apply(formatar_moeda)
                if 'Total_Mensal' in resumo_display.columns:
                    resumo_display['Total_Mensal'] = resumo_display['Total_Mensal'].apply(formatar_moeda)

                resumo_display['Percentual'] = resumo_display['Percentual'].apply(lambda x: f"{x}%")
                st.dataframe(resumo_display, use_container_width=True)

            else:
                # MODO VISUALIZAÇÃO - Mostrar apenas
                resumo_display = resumo_custos_op.copy()

                # Formatar coluna de valor (pode ser Total_Periodo ou Total_Mensal)
                if 'Total_Periodo' in resumo_display.columns:
                    resumo_display['Total_Periodo'] = resumo_display['Total_Periodo'].apply(formatar_moeda)
                if 'Total_Mensal' in resumo_display.columns:
                    resumo_display['Total_Mensal'] = resumo_display['Total_Mensal'].apply(formatar_moeda)

                resumo_display['Percentual'] = resumo_display['Percentual'].apply(lambda x: f"{x}%")

                st.dataframe(resumo_display, use_container_width=True)

        # Se tiver despesas REAIS, mostrar tabela detalhada
        if fonte_custos == 'REAL' and 'df_custos_operacionais' in metricas_financeiras:
            df_despesas_reais = metricas_financeiras['df_custos_operacionais']

            if not df_despesas_reais.empty:
                st.subheader("📄 Despesas Reais do Despesify (Detalhado)")

                # Preparar tabela para exibição
                df_display = df_despesas_reais[
                    ['Data', 'Descrição', 'Categoria_Dashboard', 'Valor_Total', 'IVA', 'Valor_Sem_IVA', 'NIF_Fornecedor']
                ].copy()

                df_display['Data'] = pd.to_datetime(df_display['Data']).dt.strftime('%d/%m/%Y')
                df_display['Valor_Total'] = df_display['Valor_Total'].apply(formatar_moeda)
                df_display['IVA'] = df_display['IVA'].apply(formatar_moeda)
                df_display['Valor_Sem_IVA'] = df_display['Valor_Sem_IVA'].apply(formatar_moeda)

                st.dataframe(df_display, use_container_width=True, height=400)

        # NOVA SEÇÃO: Todas as despesas do Despesify (incluindo compras de mercadorias)
        if cost_manager.despesify_loader:
            st.markdown("---")
            st.subheader("🏪 Todas as Despesas do Despesify (Fornecedores)")

            try:
                # Carregar TODAS as despesas (operacionais + produtos)
                resumo_completo = cost_manager.despesify_loader.get_resumo_despesas(
                    data_min_filtrado, data_max_filtrado
                )

                if resumo_completo and 'df_completo' in resumo_completo:
                    df_todas_despesas = resumo_completo['df_completo']

                    if not df_todas_despesas.empty:
                        # KPIs resumo
                        col1, col2, col3, col4 = st.columns(4)

                        with col1:
                            st.metric(
                                "💰 Total Despesas",
                                formatar_moeda(resumo_completo['total_despesas'])
                            )

                        with col2:
                            st.metric(
                                "🏢 Custos Operacionais",
                                formatar_moeda(resumo_completo['total_operacionais']),
                                help="Despesas que vão para custos operacionais"
                            )

                        with col3:
                            st.metric(
                                "📦 Compras Mercadorias",
                                formatar_moeda(resumo_completo['total_produtos']),
                                help="Compras de produtos (já nos custos de produtos CSV)"
                            )

                        with col4:
                            st.metric(
                                "📄 Nº Despesas",
                                resumo_completo['num_despesas']
                            )

                        # Tabela completa de fornecedores
                        st.markdown("**📋 Detalhes por Fornecedor:**")

                        df_fornecedores = df_todas_despesas[[
                            'Data', 'Descrição', 'Tipo_Custo', 'Categoria_Dashboard',
                            'Valor_Total', 'IVA', 'Valor_Sem_IVA', 'NIF_Fornecedor'
                        ]].copy()

                        df_fornecedores['Data'] = pd.to_datetime(df_fornecedores['Data']).dt.strftime('%d/%m/%Y')
                        df_fornecedores['Valor_Total'] = df_fornecedores['Valor_Total'].apply(formatar_moeda)
                        df_fornecedores['IVA'] = df_fornecedores['IVA'].apply(formatar_moeda)
                        df_fornecedores['Valor_Sem_IVA'] = df_fornecedores['Valor_Sem_IVA'].apply(formatar_moeda)

                        # Renomear colunas para melhor visualização
                        df_fornecedores.rename(columns={
                            'Tipo_Custo': 'Tipo',
                            'Categoria_Dashboard': 'Categoria'
                        }, inplace=True)

                        st.dataframe(df_fornecedores, use_container_width=True, height=400)

                        # Gráfico de despesas por fornecedor
                        st.markdown("**📊 Fornecedores:**")

                        top_fornecedores = df_todas_despesas.groupby('Descrição')['Valor_Total'].sum().sort_values(ascending=False)

                        fig_fornecedores = px.bar(
                            x=top_fornecedores.index,
                            y=top_fornecedores.values,
                            title=f"Todos os Fornecedores ({len(top_fornecedores)}) - Despesify",
                            labels={'x': 'Fornecedor', 'y': 'Total (€)'},
                            text=[formatar_moeda(v) for v in top_fornecedores.values],
                            color=top_fornecedores.values,
                            color_continuous_scale='Blues'
                        )

                        fig_fornecedores.update_layout(height=400, showlegend=False, xaxis_tickangle=-45)
                        st.plotly_chart(fig_fornecedores, use_container_width=True)

            except Exception as e:
                st.warning(f"⚠️ Erro ao carregar despesas completas: {e}")

        # Custos por categoria de produto
        st.markdown("---")
        st.subheader("📦 Custos de Produtos por Categoria")

        if 'Custo_Total' in df_filtrado.columns:
            custos_por_cat = df_filtrado.groupby('Categoria')['Custo_Total'].sum().sort_values(ascending=False)

            fig_custos_cat = px.bar(
                x=custos_por_cat.index,
                y=custos_por_cat.values,
                title="COGS por Categoria",
                labels={'x': 'Categoria', 'y': 'Custo Total (€)'},
                text=[formatar_moeda(v) for v in custos_por_cat.values],
                color=custos_por_cat.values,
                color_continuous_scale='OrRd'
            )

            fig_custos_cat.update_layout(height=400, showlegend=False)
            st.plotly_chart(fig_custos_cat, use_container_width=True)

        # ========== CONSULTA DE FATURAS (DESPESIFY) ==========
        st.markdown("---")
        st.subheader("🧾 Consulta de Faturas (Despesify)")

        if not cost_manager.despesify_loader:
            st.warning("⚠️ Despesify não está disponível. Esta funcionalidade requer conexão com a base de dados Despesify.")
        else:
            st.markdown("""
            Consulte faturas registadas no Despesify com informação detalhada do QR Code AT.

            **ℹ️ Nota:** O QR Code da Autoridade Tributária contém apenas **totais** e **breakdown de IVA**.
            **Não contém** descrição linha a linha dos produtos comprados.
            """)

            # Barra de pesquisa
            st.subheader("🔍 Pesquisar Fatura")

            col1, col2, col3, col4 = st.columns([3, 2, 2, 1])

            with col1:
                pesquisa_doc = st.text_input(
                    "Número do Documento",
                    placeholder="Ex: FS 2025021701A/80387",
                    help="Pesquise por número completo ou parcial"
                )

            with col2:
                pesquisa_nif = st.text_input(
                    "NIF Fornecedor",
                    placeholder="Ex: 517077582",
                    help="NIF do emitente da fatura"
                )

            with col3:
                pesquisa_nome = st.text_input(
                    "Nome Fornecedor",
                    placeholder="Ex: Sonyve",
                    help="Nome ou parte do nome"
                )

            with col4:
                st.markdown("<br>", unsafe_allow_html=True)
                pesquisar = st.button("🔍 Pesquisar", type="primary", use_container_width=True)

            # Resultados da pesquisa
            if pesquisar or pesquisa_doc or pesquisa_nif or pesquisa_nome:
                with st.spinner("Pesquisando faturas..."):
                    resultados = cost_manager.despesify_loader.pesquisar_fatura(
                        numero_documento=pesquisa_doc if pesquisa_doc else None,
                        nif_fornecedor=pesquisa_nif if pesquisa_nif else None,
                        descricao=pesquisa_nome if pesquisa_nome else None
                    )

                    if resultados.empty:
                        st.info("📭 Nenhuma fatura encontrada com os critérios especificados.")
                    else:
                        st.success(f"✅ Encontradas {len(resultados)} fatura(s)")

                        # Tabela de resultados
                        st.subheader("📋 Resultados")

                        df_display = resultados[[
                            'expense_date', 'numero_documento', 'description',
                            'nif_emitente', 'amount', 'vat_amount', 'atcud'
                        ]].copy()

                        df_display.columns = [
                            'Data', 'Nº Documento', 'Fornecedor',
                            'NIF', 'Total (€)', 'IVA (€)', 'ATCUD'
                        ]

                        df_display['Data'] = pd.to_datetime(df_display['Data']).dt.strftime('%d/%m/%Y')
                        df_display['Total (€)'] = df_display['Total (€)'].apply(lambda x: f"{x:,.2f}€")
                        df_display['IVA (€)'] = df_display['IVA (€)'].apply(lambda x: f"{x:,.2f}€" if pd.notna(x) else "N/A")

                        st.dataframe(df_display, use_container_width=True, height=300)

                        # Seleção de fatura para detalhes
                        st.markdown("---")
                        st.subheader("📄 Detalhes da Fatura")

                        numero_selecionado = st.selectbox(
                            "Selecione uma fatura para ver detalhes:",
                            options=resultados['numero_documento'].tolist(),
                            format_func=lambda x: f"{x} - {resultados[resultados['numero_documento']==x]['description'].iloc[0]}"
                        )

                        if numero_selecionado:
                            detalhes = cost_manager.despesify_loader.get_detalhes_fatura(numero_selecionado)

                            if detalhes:
                                # Informações principais
                                col1, col2 = st.columns(2)

                                with col1:
                                    st.markdown("### 📌 Informações Principais")
                                    st.markdown(f"**Fornecedor:** {detalhes.get('description', 'N/A')}")
                                    st.markdown(f"**NIF Emitente:** {detalhes.get('nif_emitente', 'N/A')}")
                                    st.markdown(f"**NIF Adquirente:** {detalhes.get('nif_adquirente', 'N/A')}")
                                    st.markdown(f"**Data:** {pd.to_datetime(detalhes.get('expense_date')).strftime('%d/%m/%Y') if detalhes.get('expense_date') else 'N/A'}")
                                    st.markdown(f"**Nº Documento:** {detalhes.get('numero_documento', 'N/A')}")
                                    st.markdown(f"**ATCUD:** {detalhes.get('atcud', 'N/A')}")
                                    st.markdown(f"**Método Pagamento:** {detalhes.get('payment_method', 'N/A')}")

                                with col2:
                                    st.markdown("### 💰 Valores")
                                    st.metric("Total", f"{detalhes.get('amount', 0):,.2f}€")
                                    st.metric("Base Tributável", f"{detalhes.get('base_tributavel', 0):,.2f}€" if detalhes.get('base_tributavel') else "N/A")
                                    st.metric("IVA Total", f"{detalhes.get('vat_amount', 0):,.2f}€" if detalhes.get('vat_amount') else "N/A")

                                # Breakdown de IVA
                                if detalhes.get('linhas_iva'):
                                    st.markdown("---")
                                    st.markdown("### 📊 Breakdown de IVA (Dados do QR Code)")

                                    linhas_iva = detalhes['linhas_iva']

                                    # Criar DataFrame para linhas de IVA
                                    df_iva = pd.DataFrame(linhas_iva)

                                    if not df_iva.empty:
                                        # Renomear colunas
                                        if 'base_tributavel' in df_iva.columns:
                                            df_iva_display = df_iva[[
                                                'base_tributavel', 'taxa_iva_percentagem', 'valor_iva'
                                            ]].copy()

                                            df_iva_display.columns = [
                                                'Base Tributável (€)', 'Taxa IVA (%)', 'Valor IVA (€)'
                                            ]

                                            df_iva_display['Base Tributável (€)'] = df_iva_display['Base Tributável (€)'].apply(lambda x: f"{x:,.2f}€")
                                            df_iva_display['Taxa IVA (%)'] = df_iva_display['Taxa IVA (%)'].apply(lambda x: f"{x:.2f}%")
                                            df_iva_display['Valor IVA (€)'] = df_iva_display['Valor IVA (€)'].apply(lambda x: f"{x:,.2f}€")

                                            st.dataframe(df_iva_display, use_container_width=True, hide_index=True)

                                            # Gráfico de IVA
                                            fig_iva = go.Figure(data=[go.Pie(
                                                labels=[f"{row['Taxa IVA (%)']} ({row['Base Tributável (€)']})" for _, row in df_iva_display.iterrows()],
                                                values=[float(row['base_tributavel']) for row in linhas_iva],
                                                hole=0.4,
                                                marker=dict(colors=['#4ecdc4', '#ff6b6b', '#ffd700', '#95e1d3'])
                                            )])

                                            fig_iva.update_layout(
                                                title="Distribuição de Base Tributável por Taxa IVA",
                                                showlegend=True,
                                                height=400
                                            )

                                            st.plotly_chart(fig_iva, use_container_width=True)

                                        st.markdown("---")
                                        st.info("""
                                        ℹ️ **Sobre os dados do QR Code AT:**
                                        - O QR Code apenas contém **totais** e **breakdown de taxas de IVA**
                                        - **Não inclui** detalhes linha a linha dos produtos/serviços
                                        - Para informação detalhada, consulte a fatura original em papel/PDF
                                        """)
                            else:
                                st.error("❌ Não foi possível obter os detalhes da fatura.")
    # TAB 8 - Rentabilidade & Margens
    elif st.session_state.active_tab == "📊 Rentabilidade & Margens":
        st.header("📊 Rentabilidade & Análise de Margens")

        # Calcular métricas de custo (com datas para Despesify)
        data_min_filtrado = df_filtrado['Data'].min().to_pydatetime()
        data_max_filtrado = df_filtrado['Data'].max().to_pydatetime()
        metricas_financeiras = cost_manager.calcular_metricas_financeiras(
            df_filtrado, data_min_filtrado, data_max_filtrado
        )

        # Mostrar fonte dos custos operacionais
        fonte_custos = metricas_financeiras.get('fonte_custos_operacionais', 'ESTIMADO')
        if fonte_custos == 'REAL':
            st.success("✅ Custos Operacionais REAIS do Despesify (desde 1/12/2025)")
        else:
            st.warning("⚠️ Custos Operacionais ESTIMADOS (antes de 1/12/2025 ou sem dados no Despesify)")

        st.markdown("---")

        # KPIs de Rentabilidade
        col1, col2, col3, col4 = st.columns(4)

        with col1:
            lucro_bruto = metricas_financeiras.get('lucro_bruto', 0)
            st.metric(
                "💚 Lucro Bruto",
                formatar_moeda(lucro_bruto)
            )

        with col2:
            lucro_liquido = metricas_financeiras.get('lucro_liquido', 0)
            delta_color = "normal" if lucro_liquido >= 0 else "inverse"
            st.metric(
                "💰 Lucro Líquido",
                formatar_moeda(lucro_liquido),
                delta="Positivo" if lucro_liquido >= 0 else "Negativo"
            )

        with col3:
            margem_bruta = metricas_financeiras.get('margem_bruta_pct', 0)
            st.metric(
                "📈 Margem Bruta %",
                formatar_percentagem(margem_bruta)
            )

        with col4:
            margem_liquida = metricas_financeiras.get('margem_liquida_pct', 0)
            st.metric(
                "📉 Margem Líquida %",
                formatar_percentagem(margem_liquida)
            )

        st.markdown("---")

        # Análise por Categoria
        st.subheader("🏆 Rentabilidade por Categoria")

        analise_cat = cost_manager.analisar_rentabilidade_categorias(df_filtrado)

        if not analise_cat.empty:
            col1, col2 = st.columns(2)

            with col1:
                # Gráfico de margens por categoria
                fig_margens = px.bar(
                    x=analise_cat.index,
                    y=analise_cat['Margem_Bruta_Pct'],
                    title="Margem Bruta por Categoria",
                    labels={'x': 'Categoria', 'y': 'Margem Bruta (%)'},
                    text=[f"{v:.1f}%" for v in analise_cat['Margem_Bruta_Pct']],
                    color=analise_cat['Margem_Bruta_Pct'],
                    color_continuous_scale='RdYlGn'
                )

                fig_margens.update_layout(height=400, showlegend=False)
                st.plotly_chart(fig_margens, use_container_width=True)

            with col2:
                # Scatter: Vendas vs Margem (usar abs para tamanho)
                # O tamanho não pode ser negativo, então usamos valores absolutos
                tamanhos = analise_cat['Lucro_Bruto'].abs()

                fig_scatter_rent = px.scatter(
                    x=analise_cat['Valor'],
                    y=analise_cat['Margem_Bruta_Pct'],
                    size=tamanhos,
                    color=analise_cat['Lucro_Bruto'],  # Cor mostra se é positivo/negativo
                    hover_name=analise_cat.index,
                    title="Vendas vs Margem (tamanho = |lucro|)",
                    labels={'x': 'Vendas Totais (€)', 'y': 'Margem Bruta (%)'},
                    color_continuous_scale='RdYlGn'
                )

                fig_scatter_rent.update_layout(height=400)
                st.plotly_chart(fig_scatter_rent, use_container_width=True)

        st.markdown("---")

        # Tabela de análise detalhada
        st.subheader("📋 Análise Detalhada por Categoria")

        if not analise_cat.empty:
            analise_display = analise_cat.copy()
            analise_display['Valor'] = analise_display['Valor'].apply(formatar_moeda)
            analise_display['Custo_Total'] = analise_display['Custo_Total'].apply(formatar_moeda)
            analise_display['Lucro_Bruto'] = analise_display['Lucro_Bruto'].apply(formatar_moeda)
            analise_display['Margem_Bruta_Pct'] = analise_display['Margem_Bruta_Pct'].apply(lambda x: f"{x:.1f}%")
            analise_display['Margem_Objetivo'] = analise_display['Margem_Objetivo'].apply(lambda x: f"{x:.1f}%")
            analise_display['Diferenca_Objetivo'] = analise_display['Diferenca_Objetivo'].apply(lambda x: f"{x:+.1f}%")

            st.dataframe(analise_display, use_container_width=True)

        st.markdown("---")

        # Produtos mais e menos rentáveis
        st.subheader("🌟 Produtos Estrela vs Produtos Problemáticos")

        col1, col2 = st.columns(2)

        with col1:
            st.markdown("**⭐ Top 10 Produtos Mais Rentáveis**")

            if 'Lucro_Bruto' in df_filtrado.columns:
                top_rent = df_filtrado.groupby('Produto').agg({
                    'Valor': 'sum',
                    'Lucro_Bruto': 'sum',
                    'Qtd': 'sum'
                }).sort_values('Lucro_Bruto', ascending=False).head(10)

                # Calcular margem corretamente: (Lucro Total / Valor Total) × 100
                top_rent['Margem_Bruta_Pct'] = (top_rent['Lucro_Bruto'] / top_rent['Valor'] * 100).round(1)

                for idx, (produto, row) in enumerate(top_rent.iterrows(), 1):
                    st.markdown(
                        f"{idx}. **{produto}**: {formatar_moeda(row['Lucro_Bruto'])} "
                        f"({row['Margem_Bruta_Pct']:.1f}% margem)"
                    )

        with col2:
            st.markdown("**⚠️ Top 10 Produtos Menos Rentáveis**")

            if 'Lucro_Bruto' in df_filtrado.columns:
                bottom_rent = df_filtrado.groupby('Produto').agg({
                    'Valor': 'sum',
                    'Lucro_Bruto': 'sum',
                    'Qtd': 'sum'
                }).sort_values('Lucro_Bruto', ascending=True).head(10)

                # Calcular margem corretamente: (Lucro Total / Valor Total) × 100
                bottom_rent['Margem_Bruta_Pct'] = (bottom_rent['Lucro_Bruto'] / bottom_rent['Valor'] * 100).round(1)

                for idx, (produto, row) in enumerate(bottom_rent.iterrows(), 1):
                    st.markdown(
                        f"{idx}. **{produto}**: {formatar_moeda(row['Lucro_Bruto'])} "
                        f"({row['Margem_Bruta_Pct']:.1f}% margem)"
                    )


    # TAB 9 - Break-Even Analysis
    elif st.session_state.active_tab == "🎯 Break-Even Analysis":
        st.header("🎯 Análise de Ponto de Equilíbrio (Break-Even)")

        # Calcular break-even
        break_even = loader.get_analise_break_even(df_filtrado)

        # KPIs de Break-Even
        col1, col2, col3 = st.columns(3)

        with col1:
            custos_fixos = break_even.get('custos_fixos_mensais', 0)
            st.metric(
                "🏢 Custos Fixos Mensais",
                formatar_moeda(custos_fixos)
            )

        with col2:
            vendas_be = break_even.get('vendas_break_even_mensal', 0)
            st.metric(
                "🎯 Vendas Break-Even (Mensal)",
                formatar_moeda(vendas_be)
            )

        with col3:
            margem_contrib = break_even.get('margem_contribuicao_pct', 0)
            st.metric(
                "📊 Margem de Contribuição",
                formatar_percentagem(margem_contrib)
            )

        st.markdown("---")

        # Análise diária
        st.subheader("📅 Análise Diária de Break-Even")

        col1, col2, col3 = st.columns(3)

        with col1:
            vendas_be_diaria = break_even.get('vendas_break_even_diaria', 0)
            st.metric(
                "🎯 Vendas Necessárias (Diária)",
                formatar_moeda(vendas_be_diaria)
            )

        with col2:
            vendas_atuais = break_even.get('vendas_atuais_diarias', 0)
            st.metric(
                "💰 Vendas Atuais (Diária)",
                formatar_moeda(vendas_atuais)
            )

        with col3:
            margem_seg = break_even.get('margem_seguranca_pct', 0)
            delta_text = "Acima" if margem_seg > 0 else "Abaixo"
            st.metric(
                "🛡️ Margem de Segurança",
                formatar_percentagem(abs(margem_seg)),
                delta=delta_text
            )

        st.markdown("---")

        # Detalhes do cálculo realista
        st.subheader("📊 Análise Detalhada do Break-Even Realista")

        col1, col2, col3 = st.columns(3)

        with col1:
            iva_medio = break_even.get('iva_medio_pct', 18)
            iva_mensal = break_even.get('iva_mensal_estimado', 0)
            st.metric(
                "🧾 IVA Médio",
                formatar_percentagem(iva_medio)
            )
            st.caption(f"Estimado/mês: {formatar_moeda(iva_mensal)}")

        with col2:
            margem_seg_pct = break_even.get('margem_seguranca_pct', 15)
            st.metric(
                "🛡️ Margem de Segurança Incluída",
                formatar_percentagem(margem_seg_pct)
            )

        with col3:
            margem_efetiva = break_even.get('margem_efetiva_apos_iva', 0)
            st.metric(
                "📉 Margem Efetiva (após IVA)",
                formatar_percentagem(margem_efetiva)
            )

        st.info("""
        **Break-Even Realista** considera:
        - ✅ Todos os custos operacionais (incluindo o seu ordenado)
        - ✅ IVA a pagar sobre as vendas (reduz a margem real)
        - ✅ Margem de segurança para períodos de menor movimento
        - ✅ Lucro mínimo esperado acima do break-even
        """)

        # Mostrar lucro mínimo esperado
        lucro_minimo = break_even.get('lucro_minimo_esperado', 0)
        if lucro_minimo > 0:
            st.success(f"💰 **Lucro Mínimo Esperado:** {formatar_moeda(lucro_minimo)}/mês acima do break-even")
        else:
            st.warning(f"⚠️ **Lucro Estimado:** {formatar_moeda(lucro_minimo)}/mês")

        st.markdown("---")

        # Gráfico de Break-Even
        st.subheader("📈 Gráfico de Break-Even")

        # Simular vendas de 0 até 150% das vendas break-even
        vendas_max = vendas_be * 1.5
        vendas_range = np.linspace(0, vendas_max, 100)

        # Calcular custos e receitas
        custos_totais = custos_fixos + (vendas_range * (1 - margem_contrib/100))
        receitas = vendas_range

        fig_be = go.Figure()

        # Linha de custos
        fig_be.add_trace(go.Scatter(
            x=vendas_range,
            y=custos_totais,
            mode='lines',
            name='Custos Totais',
            line=dict(color='red', width=2)
        ))

        # Linha de receitas
        fig_be.add_trace(go.Scatter(
            x=vendas_range,
            y=receitas,
            mode='lines',
            name='Receitas',
            line=dict(color='green', width=2)
        ))

        # Ponto de break-even
        fig_be.add_trace(go.Scatter(
            x=[vendas_be],
            y=[vendas_be],
            mode='markers',
            name='Break-Even Point',
            marker=dict(size=15, color='orange', symbol='star')
        ))

        # Vendas atuais
        if vendas_atuais > 0:
            lucro_atual = vendas_atuais - (custos_fixos + vendas_atuais * (1 - margem_contrib/100))
            fig_be.add_trace(go.Scatter(
                x=[vendas_atuais * 30],  # Mensal
                y=[vendas_atuais * 30],
                mode='markers',
                name='Vendas Atuais (Mensal)',
                marker=dict(size=12, color='blue', symbol='diamond')
            ))

        fig_be.update_layout(
            title="Análise de Break-Even: Custos vs Receitas",
            xaxis_title="Vendas Mensais (€)",
            yaxis_title="Valor (€)",
            height=500,
            hovermode='x unified'
        )

        st.plotly_chart(fig_be, use_container_width=True)

        st.markdown("---")

        # Dias para atingir break-even
        st.subheader("📆 Projeção de Break-Even")

        st.info("""
        **💡 Como interpretar "Dias para Break-Even":**

        Este cálculo responde: *"Quantos dias de operação (com as vendas e margens atuais) são necessários
        para acumular margem suficiente para pagar os custos fixos mensais?"*

        - As despesas fixas (rendas, ordenados) caem tipicamente no **dia 28** de cada mês
        - Durante o mês, o negócio acumula margem de contribuição das vendas diárias
        - O break-even mostra quantos dias de vendas são necessários para cobrir essas despesas
        - Usa **custos ESTIMADOS** completos (não parciais) do ficheiro de configuração
        """)

        dias_be = break_even.get('dias_para_break_even', 0)

        # Verificar se os custos fixos estão configurados
        if custos_fixos <= 0:
            st.error("""
            ⚠️ **Custos Fixos não configurados!**

            Para calcular o Break-Even, é necessário configurar os custos fixos mensais.

            **Opções:**
            1. Criar ficheiro `dados_custos/custos_operacionais.csv` com os seus custos
            2. Configurar a integração Despesify para custos REAIS
            """)
        elif vendas_atuais <= 0:
            st.warning("⚠️ Sem vendas no período selecionado para calcular projeção")
        elif dias_be > 0:
            col1, col2 = st.columns(2)

            with col1:
                st.metric(
                    "🗓️ Dias para Break-Even",
                    f"{dias_be:.1f} dias" if dias_be < 999 else "Mais de 999 dias"
                )

                # Mostrar detalhes do cálculo
                margem_diaria_atual = vendas_atuais * (margem_contrib/100)
                st.caption(f"""
                📊 **Detalhes do Cálculo:**
                - Custos Fixos: {formatar_moeda(custos_fixos)}/mês
                - Margem Diária: {formatar_moeda(margem_diaria_atual)}/dia
                - Dias = {formatar_moeda(custos_fixos)} ÷ {formatar_moeda(margem_diaria_atual)} = {dias_be:.1f} dias
                """)

                if dias_be <= 7:
                    st.success(f"✅ Excelente! Break-even em apenas {dias_be:.1f} dias (1 semana)")
                elif dias_be <= 15:
                    st.success(f"✅ Muito bom! Break-even em {dias_be:.1f} dias (~2 semanas)")
                elif dias_be <= 30:
                    st.success("✅ Break-even atingível no mês atual!")
                elif dias_be <= 60:
                    st.warning("⚠️ Break-even em aproximadamente 2 meses")
                elif dias_be <= 180:
                    st.error("❌ Break-even requer mais de 2 meses")
                else:
                    st.error(f"❌ Break-even requer {dias_be/30:.1f} meses ({dias_be:.0f} dias)")

            with col2:
                # Projeção de lucro por volume
                volumes = [0.5, 0.75, 1.0, 1.25, 1.5]
                lucros = []

                for v in volumes:
                    vendas_sim = vendas_be * v
                    lucro_sim = vendas_sim * (margem_contrib/100) - custos_fixos
                    lucros.append(lucro_sim)

                fig_sensibilidade = go.Figure()

                fig_sensibilidade.add_trace(go.Bar(
                    x=[f"{int(v*100)}%" for v in volumes],
                    y=lucros,
                    text=[formatar_moeda(l) for l in lucros],
                    textposition='outside',
                    marker=dict(color=lucros, colorscale='RdYlGn', cmin=-max(abs(min(lucros)), max(lucros)), cmax=max(abs(min(lucros)), max(lucros)))
                ))

                fig_sensibilidade.update_layout(
                    title="Análise de Sensibilidade (% do Break-Even)",
                    xaxis_title="Volume de Vendas",
                    yaxis_title="Lucro Líquido (€)",
                    height=400,
                    showlegend=False
                )

                st.plotly_chart(fig_sensibilidade, use_container_width=True)
        else:
            st.warning("⚠️ Dados insuficientes para calcular projeção de break-even")

        st.markdown("---")

        # Recomendações
        st.subheader("💡 Recomendações Estratégicas")

        # Calcular lucro líquido estimado (mensal)
        vendas_mensais_atuais = vendas_atuais * 30 if vendas_atuais > 0 else 0
        lucro_liquido = break_even.get('lucro_minimo_esperado', 0)

        # Se não houver lucro_minimo_esperado, calcular manualmente
        if lucro_liquido == 0 and vendas_mensais_atuais > 0:
            lucro_liquido = (vendas_mensais_atuais * (margem_contrib / 100)) - custos_fixos

        if margem_seg < 0:
            st.error(f"""
            **🚨 ALERTA: Vendas abaixo do break-even!**
            - Déficit atual: {formatar_moeda(abs(lucro_liquido))}
            - Aumento necessário nas vendas: {formatar_percentagem(abs(margem_seg))}
            - Considere reduzir custos ou aumentar preços
            """)
        elif margem_seg < 20:
            st.warning(f"""
            **⚠️ ATENÇÃO: Margem de segurança baixa**
            - Margem atual: {formatar_percentagem(margem_seg)}
            - Recomenda-se margem > 20%
            - Trabalhe para aumentar vendas ou reduzir custos
            """)
        else:
            st.success(f"""
            **✅ SITUAÇÃO SAUDÁVEL**
            - Margem de segurança: {formatar_percentagem(margem_seg)}
            - Negócio operando confortavelmente acima do break-even
            - Continue monitorando custos e otimizando margens
            """)

    # TAB 10 - Jogos Santa Casa
    elif st.session_state.active_tab == "🎰 Jogos Santa Casa":
        df_jogos = carregar_dados_santa_casa()
        pagina_jogos_santa_casa(df_jogos, data_inicio, data_fim)

    # TAB 11 - Forecast & Encomendas
    elif st.session_state.active_tab == "📦 Forecast & Encomendas":
        st.header("📦 Forecast & Previsão de Encomendas")
        st.caption("Análise de sazonalidade e recomendações de encomenda por produto")

        # ========== PARÂMETROS DE CONFIGURAÇÃO ==========
        st.subheader("⚙️ Parâmetros de Análise")

        col1, col2, col3 = st.columns(3)

        with col1:
            dias_analise_forecast = st.slider(
                "Período de análise (dias)",
                min_value=30,
                max_value=180,
                value=60,
                step=30,
                help="Período histórico para calcular médias"
            )

        with col2:
            vendas_minimas_forecast = st.number_input(
                "Vendas mínimas (regularidade)",
                min_value=3,
                max_value=20,
                value=5,
                step=1,
                help="Mínimo de vendas para considerar produto regular"
            )

        with col3:
            fator_seguranca_forecast = st.slider(
                "Fator de stock de segurança",
                min_value=1.0,
                max_value=2.5,
                value=1.5,
                step=0.1,
                help="Multiplicador para stock de segurança"
            )

        # Checkbox para incluir produtos esporádicos
        incluir_esporadicos = st.checkbox(
            "Incluir produtos com vendas esporádicas",
            value=False,
            help="Se desmarcado, mostra apenas produtos com vendas regulares"
        )

        st.markdown("---")

        # ========== GERAR FORECAST ==========
        with st.spinner("Calculando previsões..."):
            df_forecast = gerar_forecast_produtos(
                df_filtrado,
                cost_manager,
                dias_analise=dias_analise_forecast,
                produtos_regulares_only=not incluir_esporadicos
            )

        if df_forecast.empty:
            st.warning("Nenhum produto encontrado com os critérios selecionados.")
            st.info("💡 Dica: Tente aumentar o período de análise ou reduzir o número de vendas mínimas.")
        else:
            # ========== MÉTRICAS RESUMIDAS ==========
            st.subheader("📊 Resumo de Encomendas")

            col1, col2, col3, col4 = st.columns(4)

            with col1:
                total_produtos = len(df_forecast)
                st.metric(
                    "Total de Produtos",
                    total_produtos,
                    help="Produtos com vendas regulares"
                )

            with col2:
                total_semanal = df_forecast['Valor_Encomenda_Semanal'].sum()
                st.metric(
                    "Encomenda Semanal",
                    formatar_moeda(total_semanal),
                    help="Valor total estimado para encomenda semanal"
                )

            with col3:
                total_mensal = df_forecast['Valor_Encomenda_Mensal'].sum()
                st.metric(
                    "Encomenda Mensal",
                    formatar_moeda(total_mensal),
                    help="Valor total estimado para encomenda mensal"
                )

            with col4:
                produtos_crescimento = len(df_forecast[df_forecast['Tendencia_Direcao'] == 'crescimento'])
                st.metric(
                    "Produtos em Crescimento",
                    produtos_crescimento,
                    delta=f"{(produtos_crescimento/total_produtos*100):.1f}%",
                    help="Produtos com tendência de crescimento"
                )

            st.markdown("---")

            # ========== FILTROS DE CATEGORIA ==========
            st.subheader("🔍 Filtros")

            col1, col2 = st.columns(2)

            with col1:
                categorias_forecast = sorted(df_forecast['Categoria'].unique())
                categoria_selecionada_forecast = st.multiselect(
                    "Filtrar por Categoria",
                    options=categorias_forecast,
                    default=categorias_forecast,
                    help="Selecione categorias para exibir"
                )

            with col2:
                tendencia_filtro = st.multiselect(
                    "Filtrar por Tendência",
                    options=['crescimento', 'estavel', 'queda'],
                    default=['crescimento', 'estavel', 'queda'],
                    help="Filtrar produtos por tendência"
                )

            # Aplicar filtros
            df_forecast_filtrado = df_forecast[
                (df_forecast['Categoria'].isin(categoria_selecionada_forecast)) &
                (df_forecast['Tendencia_Direcao'].isin(tendencia_filtro))
            ]

            if df_forecast_filtrado.empty:
                st.warning("Nenhum produto encontrado com os filtros selecionados.")
            else:
                st.markdown("---")

                # ========== TABELA DE RECOMENDAÇÕES ==========
                st.subheader("📋 Recomendações de Encomenda")

                # Formatar DataFrame para exibição
                df_display = df_forecast_filtrado.copy()

                # Criar coluna de tendência visual
                def formatar_tendencia(row):
                    if row['Tendencia_Direcao'] == 'crescimento':
                        return f"📈 +{row['Tendencia_Percent']:.1f}%"
                    elif row['Tendencia_Direcao'] == 'queda':
                        return f"📉 {row['Tendencia_Percent']:.1f}%"
                    else:
                        return f"➡️ {row['Tendencia_Percent']:.1f}%"

                df_display['Tendencia'] = df_display.apply(formatar_tendencia, axis=1)

                # Selecionar e renomear colunas
                df_display = df_display[[
                    'Produto', 'Categoria', 'Media_Diaria', 'Qtd_Semanal',
                    'Qtd_Mensal', 'Stock_Seguranca', 'Custo_Unitario',
                    'Valor_Encomenda_Semanal', 'Valor_Encomenda_Mensal',
                    'Tendencia', 'CV', 'Dias_Com_Venda'
                ]]

                df_display.columns = [
                    'Produto', 'Categoria', 'Média Diária', 'Qtd Semanal',
                    'Qtd Mensal', 'Stock Segurança', 'Custo Unit.',
                    'Valor Semanal', 'Valor Mensal', 'Tendência',
                    'Variabilidade', 'Dias c/ Venda'
                ]

                # Formatar valores numéricos
                df_display['Média Diária'] = df_display['Média Diária'].apply(lambda x: f"{x:.1f}")
                df_display['Qtd Semanal'] = df_display['Qtd Semanal'].apply(lambda x: f"{x:.0f}")
                df_display['Qtd Mensal'] = df_display['Qtd Mensal'].apply(lambda x: f"{x:.0f}")
                df_display['Stock Segurança'] = df_display['Stock Segurança'].apply(lambda x: f"{x:.0f}")
                df_display['Custo Unit.'] = df_display['Custo Unit.'].apply(formatar_moeda)
                df_display['Valor Semanal'] = df_display['Valor Semanal'].apply(formatar_moeda)
                df_display['Valor Mensal'] = df_display['Valor Mensal'].apply(formatar_moeda)
                df_display['Variabilidade'] = df_display['Variabilidade'].apply(lambda x: f"{x:.1f}%")

                # Exibir tabela
                st.dataframe(
                    df_display,
                    use_container_width=True,
                    height=500,
                    hide_index=True
                )

                st.markdown("---")

                # ========== GRÁFICOS ==========
                st.subheader("📊 Visualizações")

                # Tabs para diferentes gráficos
                tab_g1, tab_g2, tab_g3 = st.tabs([
                    "Top Produtos (Valor)",
                    "Análise de Tendências",
                    "Comparação Ano Anterior"
                ])

                with tab_g1:
                    st.markdown("#### 💰 Top 20 Produtos por Valor de Encomenda Mensal")

                    top20 = df_forecast_filtrado.nlargest(20, 'Valor_Encomenda_Mensal')

                    fig = px.bar(
                        top20,
                        x='Valor_Encomenda_Mensal',
                        y='Produto',
                        orientation='h',
                        color='Categoria',
                        title='Top 20 Produtos - Valor Encomenda Mensal',
                        labels={'Valor_Encomenda_Mensal': 'Valor (€)', 'Produto': 'Produto'},
                        height=600
                    )

                    fig.update_layout(
                        yaxis={'categoryorder': 'total ascending'},
                        showlegend=True
                    )

                    st.plotly_chart(fig, use_container_width=True)

                with tab_g2:
                    st.markdown("#### 📈 Distribuição de Tendências")

                    # Gráfico de pizza: distribuição de tendências
                    tendencias_count = df_forecast_filtrado['Tendencia_Direcao'].value_counts().reset_index()
                    tendencias_count.columns = ['Tendencia', 'Count']

                    fig = px.pie(
                        tendencias_count,
                        values='Count',
                        names='Tendencia',
                        title='Distribuição de Tendências',
                        color='Tendencia',
                        color_discrete_map={
                            'crescimento': '#28a745',
                            'estavel': '#ffc107',
                            'queda': '#dc3545'
                        },
                        height=500
                    )

                    st.plotly_chart(fig, use_container_width=True)

                    # Scatter plot: Tendência vs Valor Encomenda
                    st.markdown("#### 🎯 Tendência vs Valor de Encomenda")

                    fig = px.scatter(
                        df_forecast_filtrado,
                        x='Tendencia_Percent',
                        y='Valor_Encomenda_Mensal',
                        size='Media_Diaria',
                        color='Categoria',
                        hover_data=['Produto', 'Qtd_Mensal'],
                        title='Relação entre Tendência e Valor de Encomenda',
                        labels={
                            'Tendencia_Percent': 'Tendência (%)',
                            'Valor_Encomenda_Mensal': 'Valor Encomenda Mensal (€)'
                        },
                        height=500
                    )

                    fig.add_hline(y=0, line_dash="dash", line_color="gray", opacity=0.5)

                    st.plotly_chart(fig, use_container_width=True)

                with tab_g3:
                    st.markdown("#### 📅 Comparação com Ano Anterior (YoY)")

                    # Selecionar produto para análise detalhada
                    produto_selecionado = st.selectbox(
                        "Selecionar Produto",
                        options=sorted(df_forecast_filtrado['Produto'].unique()),
                        help="Escolha um produto para comparar com ano anterior"
                    )

                    if produto_selecionado:
                        # Filtrar dados do produto (usar df completo, não apenas df_filtrado)
                        df_produto_viz = df[df['Produto'] == produto_selecionado].copy()

                        if not df_produto_viz.empty:
                            # Adicionar coluna de Ano e Mês
                            df_produto_viz['Ano'] = df_produto_viz['Data'].dt.year
                            df_produto_viz['Mes'] = df_produto_viz['Data'].dt.month
                            df_produto_viz['Ano_Mes'] = df_produto_viz['Data'].dt.to_period('M').astype(str)

                            # Obter ano atual e ano anterior
                            ano_atual = df_produto_viz['Ano'].max()
                            ano_anterior = ano_atual - 1

                            # Filtrar dados dos dois anos
                            df_ano_atual = df_produto_viz[df_produto_viz['Ano'] == ano_atual].copy()
                            df_ano_anterior = df_produto_viz[df_produto_viz['Ano'] == ano_anterior].copy()

                            if not df_ano_atual.empty and not df_ano_anterior.empty:
                                # Agrupar por mês
                                vendas_atual = df_ano_atual.groupby('Mes')['Qtd'].sum().reset_index()
                                vendas_atual['Ano'] = ano_atual
                                vendas_atual.columns = ['Mes', 'Qtd', 'Ano']

                                vendas_anterior = df_ano_anterior.groupby('Mes')['Qtd'].sum().reset_index()
                                vendas_anterior['Ano'] = ano_anterior
                                vendas_anterior.columns = ['Mes', 'Qtd', 'Ano']

                                # Combinar dados
                                vendas_comparacao = pd.concat([vendas_anterior, vendas_atual])

                                # Mapear nomes dos meses
                                meses_nomes = {
                                    1: 'Jan', 2: 'Fev', 3: 'Mar', 4: 'Abr',
                                    5: 'Mai', 6: 'Jun', 7: 'Jul', 8: 'Ago',
                                    9: 'Set', 10: 'Out', 11: 'Nov', 12: 'Dez'
                                }
                                vendas_comparacao['Mes_Nome'] = vendas_comparacao['Mes'].map(meses_nomes)

                                # Gráfico de barras agrupadas
                                fig = px.bar(
                                    vendas_comparacao,
                                    x='Mes_Nome',
                                    y='Qtd',
                                    color='Ano',
                                    barmode='group',
                                    title=f'Comparação Anual - {produto_selecionado}',
                                    labels={'Qtd': 'Quantidade', 'Mes_Nome': 'Mês', 'Ano': 'Ano'},
                                    color_discrete_map={ano_anterior: '#FF6B6B', ano_atual: '#4ECDC4'},
                                    height=450
                                )

                                fig.update_layout(
                                    xaxis={'categoryorder': 'array', 'categoryarray': list(meses_nomes.values())},
                                    legend=dict(orientation="h", yanchor="bottom", y=1.02, xanchor="right", x=1)
                                )

                                st.plotly_chart(fig, use_container_width=True)

                                # Calcular variações
                                st.markdown(f"##### Análise Comparativa: {produto_selecionado}")

                                # Total do ano
                                total_ano_atual = vendas_atual['Qtd'].sum()
                                total_ano_anterior = vendas_anterior['Qtd'].sum()
                                variacao_total = ((total_ano_atual - total_ano_anterior) / max(total_ano_anterior, 1)) * 100

                                # Média mensal
                                media_ano_atual = vendas_atual['Qtd'].mean()
                                media_ano_anterior = vendas_anterior['Qtd'].mean()
                                variacao_media = ((media_ano_atual - media_ano_anterior) / max(media_ano_anterior, 1)) * 100

                                col1, col2, col3, col4 = st.columns(4)

                                with col1:
                                    st.metric(
                                        f"Total {ano_atual}",
                                        f"{total_ano_atual:.0f}",
                                        delta=f"{variacao_total:+.1f}% vs {ano_anterior}"
                                    )

                                with col2:
                                    st.metric(
                                        f"Total {ano_anterior}",
                                        f"{total_ano_anterior:.0f}"
                                    )

                                with col3:
                                    st.metric(
                                        f"Média Mensal {ano_atual}",
                                        f"{media_ano_atual:.1f}",
                                        delta=f"{variacao_media:+.1f}% vs {ano_anterior}"
                                    )

                                with col4:
                                    # Obter info do forecast
                                    produto_info = df_forecast_filtrado[df_forecast_filtrado['Produto'] == produto_selecionado].iloc[0]
                                    st.metric(
                                        "Encomenda Mensal",
                                        formatar_moeda(produto_info['Valor_Encomenda_Mensal'])
                                    )

                                # Gráfico de linha de tendência
                                st.markdown("##### Evolução Temporal")

                                df_produto_viz_sorted = df_produto_viz.sort_values('Data')
                                vendas_mensais = df_produto_viz_sorted.groupby('Ano_Mes')['Qtd'].sum().reset_index()

                                fig2 = px.line(
                                    vendas_mensais,
                                    x='Ano_Mes',
                                    y='Qtd',
                                    title=f'Evolução Mensal - {produto_selecionado}',
                                    labels={'Qtd': 'Quantidade', 'Ano_Mes': 'Mês'},
                                    markers=True,
                                    height=350
                                )

                                fig2.update_layout(
                                    xaxis={'tickangle': -45}
                                )

                                st.plotly_chart(fig2, use_container_width=True)

                            elif not df_ano_atual.empty:
                                st.warning(f"Apenas dados de {ano_atual} disponíveis. Sem dados de {ano_anterior} para comparação.")

                                # Mostrar apenas dados do ano atual
                                vendas_atual = df_ano_atual.groupby('Mes')['Qtd'].sum().reset_index()
                                meses_nomes = {
                                    1: 'Jan', 2: 'Fev', 3: 'Mar', 4: 'Abr',
                                    5: 'Mai', 6: 'Jun', 7: 'Jul', 8: 'Ago',
                                    9: 'Set', 10: 'Out', 11: 'Nov', 12: 'Dez'
                                }
                                vendas_atual['Mes_Nome'] = vendas_atual['Mes'].map(meses_nomes)

                                fig = px.bar(
                                    vendas_atual,
                                    x='Mes_Nome',
                                    y='Qtd',
                                    title=f'Vendas {ano_atual} - {produto_selecionado}',
                                    labels={'Qtd': 'Quantidade', 'Mes_Nome': 'Mês'},
                                    color='Qtd',
                                    color_continuous_scale='Blues',
                                    height=400
                                )

                                st.plotly_chart(fig, use_container_width=True)

                            else:
                                st.warning("Dados insuficientes para comparação anual.")
                        else:
                            st.warning("Nenhum dado disponível para este produto.")

                st.markdown("---")

                # ========== EXPORTAÇÃO ==========
                st.subheader("💾 Exportar Dados")

                col1, col2 = st.columns(2)

                with col1:
                    # CSV
                    csv = df_forecast_filtrado.to_csv(index=False).encode('utf-8')
                    st.download_button(
                        label="📥 Descarregar CSV",
                        data=csv,
                        file_name=f"forecast_encomendas_{datetime.now().strftime('%Y%m%d_%H%M%S')}.csv",
                        mime="text/csv",
                        use_container_width=True
                    )

                with col2:
                    # Excel
                    buffer = BytesIO()
                    with pd.ExcelWriter(buffer, engine='openpyxl') as writer:
                        df_forecast_filtrado.to_excel(writer, index=False, sheet_name='Forecast')

                    st.download_button(
                        label="📥 Descarregar Excel",
                        data=buffer.getvalue(),
                        file_name=f"forecast_encomendas_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx",
                        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        use_container_width=True
                    )

        # Notas explicativas
        st.markdown("---")
        st.markdown("""
        #### 📌 Notas Explicativas

        - **Média Diária**: Quantidade média vendida por dia (baseada no período de análise)
        - **Qtd Semanal**: Média Diária × 7 dias
        - **Qtd Mensal**: Média Diária × 30 dias
        - **Stock Segurança**: Buffer adicional baseado na variabilidade das vendas
        - **Tendência**:
            - 📈 Crescimento: >5% de aumento nas vendas
            - ➡️ Estável: Entre -5% e +5%
            - 📉 Queda: <-5% de redução nas vendas
        - **Variabilidade (CV)**: Coeficiente de variação - quanto maior, mais irregular a demanda
        - **Dias c/ Venda**: Número de dias com pelo menos uma venda no período

        **Recomendação**: Produtos com alta variabilidade devem ter stock de segurança maior.
        """)

    # TAB 12 - Encomendas Novadis
    elif st.session_state.active_tab == "📦 Encomendas Novadis":
        st.header("📦 Encomendas Novadis")
        st.caption("Histórico e análise de encomendas do fornecedor Novadis")

        # Carregar dados Novadis (cached via loader)
        with st.spinner("📦 Carregando dados da Novadis..."):
            # Usar o loader da instância principal (que é cached)
            # data_inicio e data_fim já vêm dos filtros da sidebar
            df_novadis = loader.carregar_novadis_processado(data_inicio, data_fim)

        if df_novadis.empty:
            st.warning("⚠️ Nenhuma encomenda encontrada para o período selecionado.")
            st.info("Tente selecionar um período mais abrangente na barra lateral.")
        else:
            # KPIs Novadis
            col1, col2, col3, col4 = st.columns(4)

            total_gasto = df_novadis['custo_total'].sum()
            num_encomendas = df_novadis['numero_pedido'].nunique()
            itens_comprados = df_novadis['unidades_vendaveis'].sum()
            custo_medio_pedido = total_gasto / num_encomendas if num_encomendas > 0 else 0

            with col1:
                st.metric("💰 Total Gasto (c/ IVA)", formatar_moeda(total_gasto))
            
            with col2:
                st.metric("📦 Nº Encomendas", num_encomendas)
            
            with col3:
                st.metric("🍺 Unidades Compradas", f"{itens_comprados:,.0f}")
                
            with col4:
                st.metric("🧾 Custo Médio / Pedido", formatar_moeda(custo_medio_pedido))

            st.markdown("---")

            # Sub-tabs para organização
            tab_detalhes, tab_produtos, tab_evolucao = st.tabs([
                "📋 Detalhes das Encomendas", 
                "🍺 Análise por Produto",
                "📈 Evolução de Custos"
            ])

            with tab_detalhes:
                st.subheader("📋 Lista de Encomendas Detalhada")
                st.caption("💡 Se a conversão de unidades estiver errada, ajuste a coluna 'Unidades POS' e clique em Guardar.")
                
                # Preparar tabela para exibição
                df_display = df_novadis[[
                    'data', 'numero_pedido', 'produto_novadis', 
                    'quantidade_encomendada', 'preco_unitario_com_iva', 'custo_total',
                    'produto_pos', 'unidades_vendaveis'
                ]].copy()
                
                # Manter cópia original para comparar mudanças
                df_original = df_display.copy()
                
                df_display['data'] = df_display['data'].dt.strftime('%d/%m/%Y')
                # NÃO formatar como moeda aqui se quisermos que seja editável de forma simples
                # ou usar column_config para formatar na exibição
                
                df_display.columns = [
                    'Data', 'Nº Pedido', 'Produto (Novadis)', 
                    'Qtd Emb.', 'Preço Unit.', 'Total (€)',
                    'Produto (POS)', 'Unidades POS'
                ]
                
                # Editor de dados
                df_editado = st.data_editor(
                    df_display, 
                    use_container_width=True, 
                    hide_index=True,
                    column_config={
                        "Data": st.column_config.TextColumn("Data", disabled=True),
                        "Nº Pedido": st.column_config.TextColumn("Nº Pedido", disabled=True),
                        "Produto (Novadis)": st.column_config.TextColumn("Produto (Novadis)", disabled=True),
                        "Qtd Emb.": st.column_config.NumberColumn("Qtd Emb.", disabled=True),
                        "Preço Unit.": st.column_config.NumberColumn("Preço Unit.", format="%.2f€", disabled=True),
                        "Total (€)": st.column_config.NumberColumn("Total (€)", format="%.2f€", disabled=True),
                        "Produto (POS)": st.column_config.TextColumn("Produto (POS)", disabled=True),
                        "Unidades POS": st.column_config.NumberColumn(
                            "Unidades POS",
                            help="Ajuste o número total de unidades (ex: se 1 caixa tem 24 unidades, e comprou 1 caixa, coloque 24)",
                            min_value=1,
                            step=1
                        )
                    },
                    key="editor_novadis_detalhes"
                )

                if st.button("💾 Guardar Ajustes de Quantidade", type="primary"):
                    # Identificar linhas alteradas
                    cambios = []
                    for idx in range(len(df_editado)):
                        val_original = df_original.iloc[idx]['unidades_vendaveis']
                        val_novo = df_editado.iloc[idx]['Unidades POS']
                        
                        if val_original != val_novo:
                            prod_novadis = df_original.iloc[idx]['produto_novadis']
                            qtd_emb = df_original.iloc[idx]['quantidade_encomendada']
                            
                            # Calcular novas unidades por caixa (multiplicador)
                            unid_por_caixa = val_novo / qtd_emb if qtd_emb > 0 else val_novo
                            cambios.append({
                                'produto_novadis': prod_novadis,
                                'unidades_por_caixa': int(unid_por_caixa)
                            })
                    
                    if cambios:
                        df_novos_maps = pd.DataFrame(cambios).drop_duplicates(subset=['produto_novadis'])
                        arquivo_map_unid = Path('mapeamento_unidades_novadis.csv')
                        
                        if arquivo_map_unid.exists():
                            df_map_existente = pd.read_csv(arquivo_map_unid)
                            # Atualizar existentes e adicionar novos
                            df_map_existente = df_map_existente[~df_map_existente['produto_novadis'].isin(df_novos_maps['produto_novadis'])]
                            df_final_map = pd.concat([df_map_existente, df_novos_maps], ignore_index=True)
                        else:
                            df_final_map = df_novos_maps
                            
                        df_final_map.to_csv(arquivo_map_unid, index=False)
                        st.success(f"✅ {len(df_novos_maps)} mapeamentos de unidades atualizados!")
                        st.rerun()
                    else:
                        st.info("ℹ️ Nenhuma alteração detectada nas quantidades.")

            with tab_produtos:
                st.subheader("🍺 Top Produtos Comprados")
                
                # Agrupar por produto POS
                top_produtos = df_novadis.groupby('produto_pos').agg({
                    'custo_total': 'sum',
                    'unidades_vendaveis': 'sum',
                    'numero_pedido': 'nunique'
                }).reset_index().sort_values('custo_total', ascending=False)
                
                # Gráfico
                fig_prod = px.bar(
                    top_produtos.head(15),
                    x='custo_total',
                    y='produto_pos',
                    orientation='h',
                    title='Top 15 Produtos por Custo Total',
                    labels={'custo_total': 'Custo Total (€)', 'produto_pos': 'Produto'},
                    text=[formatar_moeda(v) for v in top_produtos.head(15)['custo_total']]
                )
                fig_prod.update_layout(yaxis={'categoryorder': 'total ascending'})
                st.plotly_chart(fig_prod, use_container_width=True)
                
                # Tabela resumida
                st.dataframe(
                    top_produtos.style.format({
                        'custo_total': '{:.2f}€',
                        'unidades_vendaveis': '{:,.0f}'
                    }),
                    use_container_width=True
                )

            with tab_evolucao:
                st.subheader("📈 Evolução de Custos Mensais")
                
                # Agrupar por mês
                df_novadis['mes_ano'] = df_novadis['data'].dt.to_period('M').astype(str)
                custos_mensais = df_novadis.groupby('mes_ano')['custo_total'].sum().reset_index()
                
                fig_ev = px.bar(
                    custos_mensais,
                    x='mes_ano',
                    y='custo_total',
                    title='Evolução de Gastos com Novadis',
                    labels={'custo_total': 'Total Gasto (€)', 'mes_ano': 'Mês'},
                    text=[formatar_moeda(v) for v in custos_mensais['custo_total']]
                )
                st.plotly_chart(fig_ev, use_container_width=True)

    # ============================================================
    # NOVA TAB 12.5 - Encomendas Delta
    # ============================================================
    elif st.session_state.active_tab == "📦 Encomendas Delta":
        st.header("📦 Encomendas Delta (Nabeiro)")
        st.caption("Histórico e análise de faturas do fornecedor Delta / Grupo Nabeiro")

        # Carregar dados Delta
        with st.spinner("📦 Carregando dados da Delta..."):
            df_delta = loader.carregar_delta(data_inicio, data_fim)

        if df_delta.empty:
            st.warning("⚠️ Nenhuma fatura Delta encontrada para o período selecionado.")
            st.info("Verifique se o arquivo CSV em /home/jorge/Documentos/delta/dados/faturas_nabeiro.csv existe e contém dados.")
        else:
            # KPIs Delta
            col1, col2, col3 = st.columns(3)

            total_gasto = df_delta['Total_EUR'].sum()
            num_faturas = df_delta['Numero_Fatura'].nunique()
            custo_medio_fatura = total_gasto / num_faturas if num_faturas > 0 else 0

            with col1:
                st.metric("💰 Total Gasto Delta", formatar_moeda(total_gasto))
            
            with col2:
                st.metric("📦 Nº Faturas", num_faturas)
                
            with col3:
                st.metric("🧾 Custo Médio / Fatura", formatar_moeda(custo_medio_fatura))

            st.markdown("---")

            # Sub-tabs para organização
            tab_lista, tab_produtos, tab_grafico = st.tabs([
                "📋 Lista de Faturas", 
                "🍺 Detalhes por Produto",
                "📈 Evolução de Gastos"
            ])

            with tab_lista:
                st.subheader("📋 Histórico de Faturas Delta")
                
                # Preparar tabela para exibição
                df_disp_delta = df_delta[['Data_Fatura', 'Numero_Fatura', 'Total_EUR', 'Arquivo']].copy()
                df_disp_delta['Data_Fatura'] = df_disp_delta['Data_Fatura'].dt.strftime('%d/%m/%Y')
                
                df_disp_delta.columns = ['Data', 'Nº Fatura', 'Total (€)', 'Arquivo PDF']
                
                st.dataframe(
                    df_disp_delta,
                    use_container_width=True,
                    hide_index=True,
                    column_config={
                        "Total (€)": st.column_config.NumberColumn("Total (€)", format="%.2f€")
                    }
                )

            with tab_produtos:
                st.subheader("🍺 Detalhe de Produtos (Itens)")
                # Carregar todos os itens detalhados
                df_itens_delta = loader.carregar_delta_itens()
                
                if df_itens_delta.empty:
                    st.info("Nenhum detalhe de produto disponível.")
                else:
                    # Filtrar apenas itens das faturas que estão no período
                    faturas_no_periodo = df_delta['Numero_Fatura'].astype(str).tolist()
                    df_itens_filtrado = df_itens_delta[df_itens_delta['Numero_Fatura'].astype(str).isin(faturas_no_periodo)]
                    
                    if df_itens_filtrado.empty:
                        st.info("Nenhum item encontrado para as faturas deste período.")
                    else:
                        # KPI de itens
                        total_itens = df_itens_filtrado['Quantidade'].sum()
                        top_item = df_itens_filtrado.groupby('Descricao')['Total_Liquido'].sum().idxmax()
                        
                        col1, col2 = st.columns(2)
                        with col1:
                            st.metric("Total Itens Comprados", f"{total_itens:,.0f}")
                        with col2:
                            st.metric("Produto com Maior Investimento", top_item)

                        # Tabela Detalhada
                        st.markdown("---")
                        df_disp_itens = df_itens_filtrado[[
                            'Descricao', 'Quantidade', 'Unidade', 'Preco_Unit_Liquido', 
                            'Total_Liquido', 'IVA_Pct', 'Total_Com_IVA'
                        ]].copy()
                        
                        df_disp_itens.columns = [
                            'Produto', 'Qtd', 'Un.', 'Preço Unit. (Líq)', 
                            'Total (Líq)', 'IVA %', 'Total (c/ IVA)'
                        ]
                        
                        st.dataframe(
                            df_disp_itens.sort_values('Total (Líq)', ascending=False),
                            use_container_width=True,
                            hide_index=True,
                            column_config={
                                "Preço Unit. (Líq)": st.column_config.NumberColumn(format="%.3f€"),
                                "Total (Líq)": st.column_config.NumberColumn(format="%.2f€"),
                                "Total (c/ IVA)": st.column_config.NumberColumn(format="%.2f€")
                            }
                        )

            with tab_grafico:
                st.subheader("📈 Evolução de Gastos Delta")
                
                # Agrupar por mês
                df_delta['mes_ano'] = df_delta['Data_Fatura'].dt.to_period('M').astype(str)
                gastos_mensais_delta = df_delta.groupby('mes_ano')['Total_EUR'].sum().reset_index()
                
                fig_delta = px.bar(
                    gastos_mensais_delta,
                    x='mes_ano',
                    y='Total_EUR',
                    title='Evolução de Gastos com Delta (Mensal)',
                    labels={'Total_EUR': 'Total Gasto (€)', 'mes_ano': 'Mês'},
                    text=[formatar_moeda(v) for v in gastos_mensais_delta['Total_EUR']]
                )
                st.plotly_chart(fig_delta, use_container_width=True)

    # ============================================================
    # NOVA TAB 13 - FICHAS TÉCNICAS (CUSTOS MANUAIS)
    # ============================================================
    elif st.session_state.active_tab == "📋 Fichas Técnicas":
        st.header("📋 Fichas Técnicas e Custos de Produção")
        st.caption("Gestão de custos para produtos comprados a granel (Café, Bebidas espirituosas, etc.)")

        arquivo_fichas = Path('fichas_tecnicas.csv')
        
        # Carregar dados existentes
        if arquivo_fichas.exists():
            df_fichas = pd.read_csv(arquivo_fichas)
        else:
            df_fichas = pd.DataFrame(columns=[
                'Produto_POS', 'Preco_Compra', 'Qtd_Compra', 'Unidade_Compra', 
                'Qtd_Dose', 'Unidade_Dose', 'Custo_Dose'
            ])

        # Form para adicionar/editar
        with st.expander("➕ Adicionar/Editar Ficha Técnica (Produto Simples)", expanded=df_fichas.empty):
            col1, col2, col3 = st.columns(3)
            
            with col1:
                # Obter lista de produtos do POS para sugestão
                lista_produtos_pos = sorted(df['Produto'].unique())
                nome_prod = st.selectbox("Produto no POS", options=lista_produtos_pos)
            
            with col2:
                preco_compra = st.number_input("Preço de Compra (€)", min_value=0.01, value=40.0, step=0.50)
                unid_compra = st.selectbox("Unidade de Compra", options=['Kg', 'Litro', 'Unidade', 'Gramas', 'cl'])
            
            with col3:
                # Definir quantidade base de compra (ex: 1kg = 1000g)
                if unid_compra == 'Kg': qtd_base = 1000; unid_dose_sug = 'Gramas'
                elif unid_compra == 'Litro': qtd_base = 100; unid_dose_sug = 'cl'
                else: qtd_base = 1; unid_dose_sug = unid_compra
                
                qtd_compra = st.number_input(f"Quantidade na Compra (em {unid_dose_sug})", min_value=1.0, value=float(qtd_base))
            
            col4, col5, col6 = st.columns(3)
            with col4:
                qtd_dose = st.number_input(f"Quantidade por Dose/Serviço ({unid_dose_sug})", min_value=0.01, value=8.0 if unid_dose_sug == 'Gramas' else 3.0)
            
            # Cálculo em tempo real
            custo_calculado = (preco_compra / qtd_compra) * qtd_dose
            
            with col5:
                st.metric("Custo por Dose", f"{custo_calculado:.3f}€")
            
            with col6:
                st.write("<br>", unsafe_allow_html=True)
                if st.button("💾 Guardar Ficha Técnica", use_container_width=True, type="primary"):
                    nova_ficha = pd.DataFrame([{
                        'Produto_POS': nome_prod,
                        'Preco_Compra': preco_compra,
                        'Qtd_Compra': qtd_compra,
                        'Unidade_Compra': unid_compra,
                        'Qtd_Dose': qtd_dose,
                        'Unidade_Dose': unid_dose_sug,
                        'Custo_Dose': round(custo_calculado, 4)
                    }])
                    
                    if not df_fichas.empty:
                        df_fichas = df_fichas[df_fichas['Produto_POS'] != nome_prod]
                        df_final = pd.concat([df_fichas, nova_ficha], ignore_index=True)
                    else:
                        df_final = nova_ficha
                        
                    df_final.to_csv(arquivo_fichas, index=False)
                    st.success(f"✅ Ficha técnica de '{nome_prod}' guardada!")
                    st.rerun()

        # ============================================================
        # NOVO: PRODUTOS COMPOSTOS (COCKTAILS)
        # ============================================================
        with st.expander("🧪 Criar Produto Composto / Cocktail (Ex: Martini c/ Cerveja)", expanded=False):
            st.caption("Use esta secção para produtos feitos de vários ingredientes (ex: Kir = Vinho + Cassis).")
            
            c_sel1, c_sel2 = st.columns([2, 1])
            
            with c_sel1:
                # Produto final (o que é vendido)
                prod_cocktail = st.selectbox("Produto Final (POS)", options=lista_produtos_pos, key="sel_cocktail")
            
            # Área de Ingredientes
            if 'temp_ingredientes' not in st.session_state:
                st.session_state.temp_ingredientes = []
                
            col_ing1, col_ing2, col_ing3 = st.columns([3, 2, 1])
            
            with col_ing1:
                # Fonte de ingredientes: Produtos Novadis + Fichas Técnicas
                # Precisamos carregar custos Novadis para mostrar aqui
                if 'custos_ref' not in st.session_state:
                    # Carregar Novadis (simplificado)
                    try:
                        # Tentar usar loader existente ou cache
                        with st.spinner("Carregando lista de ingredientes..."):
                            data_inicio_custos = datetime.now() - timedelta(days=180)
                            df_novadis_raw = loader.carregar_novadis_processado(data_inicio_custos, datetime.now())
                            if not df_novadis_raw.empty:
                                ref_novadis = df_novadis_raw.groupby('produto_pos')['custo_unitario'].mean().reset_index()
                                ref_novadis.columns = ['Produto', 'Custo']
                                ref_novadis['Tipo'] = 'Novadis'
                            else:
                                ref_novadis = pd.DataFrame(columns=['Produto', 'Custo', 'Tipo'])
                                
                            # Juntar com Fichas
                            if not df_fichas.empty:
                                ref_fichas = df_fichas[['Produto_POS', 'Custo_Dose']].copy()
                                ref_fichas.columns = ['Produto', 'Custo']
                                ref_fichas['Tipo'] = 'Ficha'
                                df_ref = pd.concat([ref_novadis, ref_fichas], ignore_index=True)
                            else:
                                df_ref = ref_novadis
                                
                            st.session_state.custos_ref = df_ref.sort_values('Produto')
                    except Exception as e:
                        st.session_state.custos_ref = pd.DataFrame(columns=['Produto', 'Custo', 'Tipo'])
                
                lista_ingredientes = st.session_state.custos_ref['Produto'].tolist()
                sel_ingrediente = st.selectbox("Adicionar Ingrediente", options=lista_ingredientes, key="sel_ing")
                
                # Buscar custo unitário
                custo_unit_ing = 0.0
                tipo_ing = ""
                if not st.session_state.custos_ref.empty:
                    row_ing = st.session_state.custos_ref[st.session_state.custos_ref['Produto'] == sel_ingrediente]
                    if not row_ing.empty:
                        custo_unit_ing = row_ing.iloc[0]['Custo']
                        tipo_ing = row_ing.iloc[0]['Tipo']

            with col_ing2:
                label_qtd = "Quantidade (Doses/Unid)"
                if tipo_ing == 'Novadis':
                    label_qtd = "Qtd (baseada na unidade de venda)"
                elif tipo_ing == 'Ficha':
                    label_qtd = "Qtd (doses da ficha técnica)"
                    
                qtd_ing = st.number_input(label_qtd, min_value=0.1, value=1.0, step=0.1, key="qtd_ing")
                st.caption(f"Custo ref: {custo_unit_ing:.3f}€ ({tipo_ing})")

            with col_ing3:
                st.write("<br>", unsafe_allow_html=True)
                if st.button("➕ Add", use_container_width=True):
                    custo_total_ing = custo_unit_ing * qtd_ing
                    st.session_state.temp_ingredientes.append({
                        'Ingrediente': sel_ingrediente,
                        'Qtd': qtd_ing,
                        'Custo_Unit': custo_unit_ing,
                        'Custo_Total': custo_total_ing,
                        'Tipo': tipo_ing
                    })

            # Tabela de ingredientes atuais
            if st.session_state.temp_ingredientes:
                df_ing_temp = pd.DataFrame(st.session_state.temp_ingredientes)
                st.dataframe(df_ing_temp, use_container_width=True, hide_index=True)
                
                custo_final_cocktail = df_ing_temp['Custo_Total'].sum()
                
                col_save1, col_save2 = st.columns([3, 1])
                with col_save1:
                    st.markdown(f"#### Custo Final: **{custo_final_cocktail:.3f}€**")
                
                with col_save2:
                    if st.button("💾 Guardar Composto", type="primary", use_container_width=True):
                        # Salvar como uma ficha técnica normal (com o valor calculado)
                        # Opcional: Salvar a receita detalhada em outro CSV se quiser editar depois
                        
                        nova_ficha_composta = pd.DataFrame([{
                            'Produto_POS': prod_cocktail,
                            'Preco_Compra': custo_final_cocktail, # Hack: Preço Compra = Custo Final
                            'Qtd_Compra': 1,
                            'Unidade_Compra': 'Unid',
                            'Qtd_Dose': 1,
                            'Unidade_Dose': 'Unid',
                            'Custo_Dose': round(custo_final_cocktail, 4)
                        }])
                        
                        if not df_fichas.empty:
                            df_fichas = df_fichas[df_fichas['Produto_POS'] != prod_cocktail]
                            df_final = pd.concat([df_fichas, nova_ficha_composta], ignore_index=True)
                        else:
                            df_final = nova_ficha_composta
                            
                        df_final.to_csv(arquivo_fichas, index=False)
                        st.session_state.temp_ingredientes = [] # Limpar
                        st.success(f"✅ Cocktail '{prod_cocktail}' guardado com custo {custo_final_cocktail:.3f}€")
                        st.rerun()
                
                if st.button("Limpar Lista"):
                    st.session_state.temp_ingredientes = []
                    st.rerun()

        # Mostrar Tabela de Fichas Técnicas
        if not df_fichas.empty:
            st.markdown("### 📋 Fichas Ativas")
            
            # Formatação para exibição
            df_show = df_fichas.copy()
            df_show['Custo_Dose'] = df_show['Custo_Dose'].apply(lambda x: f"{x:.3f}€")
            df_show['Preco_Compra'] = df_show['Preco_Compra'].apply(lambda x: f"{x:.2f}€")
            
            # Adicionar coluna de ação (remover)
            for idx, row in df_fichas.iterrows():
                c1, c2, c3, c4, c5 = st.columns([3, 2, 2, 2, 1])
                c1.write(f"**{row['Produto_POS']}**")
                c2.write(f"{row['Preco_Compra']}€ / {row['Qtd_Compra']}{row['Unidade_Dose']}")
                c3.write(f"Dose: {row['Qtd_Dose']}{row['Unidade_Dose']}")
                c4.write(f"Custo: **{row['Custo_Dose']:.3f}€**")
                if c5.button("🗑️", key=f"del_ficha_{idx}"):
                    df_final = df_fichas.drop(idx)
                    df_final.to_csv(arquivo_fichas, index=False)
                    st.rerun()
        else:
            st.info("ℹ️ Nenhuma ficha técnica configurada. Adicione uma acima (ex: Café).")

    # ============================================================
    # NOVA TAB 14 - MARGENS DE LUCRO (POS vs NOVADIS)
    # ============================================================
    elif st.session_state.active_tab == "💰 Margens de Lucro":
        st.header("💰 Análise de Margens Reais (POS vs Novadis)")
        st.caption("Cruzamento de vendas do POS-Café com custos reais das encomendas Novadis")

        # 1. Preparar Dados de Vendas (POS-Café)
        df_pos = df_filtrado[df_filtrado['Fonte'] == 'POS-Café'].copy()
        
        if df_pos.empty:
            st.warning("⚠️ Sem dados de vendas do POS-Café para o período selecionado.")
        else:
            # Agrupar vendas por produto
            vendas_resumo = df_pos.groupby(['Produto', 'Categoria']).agg({
                'Qtd': 'sum',
                'Valor': 'sum'
            }).reset_index()
            
            # Calcular Preço Médio de Venda (PVP Médio)
            vendas_resumo['PVP_Medio'] = vendas_resumo['Valor'] / vendas_resumo['Qtd']

            # 2. Preparar Dados de Custos (Novadis)
            # Carregar dados Novadis (usar todo o histórico disponível para ter melhor referência de preço)
            # ou usar apenas o período selecionado se preferir custos "da época"
            with st.spinner("Calculando custos médios da Novadis..."):
                # Carregar histórico expandido (últimos 6 meses) para garantir que temos preços
                # mesmo que não tenhamos comprado neste mês específico
                data_inicio_custos = datetime.now() - timedelta(days=180)
                df_novadis_custos = loader.carregar_novadis_processado(data_inicio_custos, datetime.now())
                
            if not df_novadis_custos.empty:
                custos_novadis = df_novadis_custos.groupby('produto_pos')['custo_unitario'].mean().reset_index()
                custos_novadis.columns = ['Produto', 'Custo_Novadis']
            else:
                custos_novadis = pd.DataFrame(columns=['Produto', 'Custo_Novadis'])

            # 3. Carregar Custos Manuais (CSV)
            arquivo_custos_manuais = Path('custos_margens_reais.csv')
            
            if arquivo_custos_manuais.exists():
                custos_manuais = pd.read_csv(arquivo_custos_manuais)
            else:
                custos_manuais = pd.DataFrame(columns=['Produto', 'Custo_Manual'])

            # 3.b Carregar Fichas Técnicas (Custos por Dose)
            arquivo_fichas = Path('fichas_tecnicas.csv')
            if arquivo_fichas.exists():
                df_fichas_custos = pd.read_csv(arquivo_fichas)
                custos_fichas = df_fichas_custos[['Produto_POS', 'Custo_Dose']].copy()
                custos_fichas.columns = ['Produto', 'Custo_Ficha']
            else:
                custos_fichas = pd.DataFrame(columns=['Produto', 'Custo_Ficha'])

            # 4. GESTÃO DE MAPEAMENTO (POS <-> NOVADIS)
            # Permite ao utilizador ligar produtos com nomes diferentes E quantidades diferentes
            arquivo_map = Path('mapeamento_produtos.csv')
            if arquivo_map.exists():
                df_map = pd.read_csv(arquivo_map)
                # Garantir compatibilidade com versões anteriores (adicionar coluna Multiplicador se não existir)
                if 'Multiplicador' not in df_map.columns:
                    df_map['Multiplicador'] = 1.0
            else:
                df_map = pd.DataFrame(columns=['Produto_POS', 'Produto_Custo', 'Multiplicador'])

            # Aplicar mapeamento aos custos da Novadis
            if not df_map.empty and not custos_novadis.empty:
                # Criar dicionário de custos: Nome_Custo -> Valor
                dict_custos = dict(zip(custos_novadis['Produto'], custos_novadis['Custo_Novadis']))
                
                novos_custos_map = []
                for _, row in df_map.iterrows():
                    prod_pos = row['Produto_POS']
                    prod_custo = row['Produto_Custo']
                    multiplicador = row.get('Multiplicador', 1.0)
                    
                    if prod_custo in dict_custos:
                        valor_base = dict_custos[prod_custo]
                        valor_final = valor_base * multiplicador
                        novos_custos_map.append({'Produto': prod_pos, 'Custo_Novadis': valor_final})
                
                # Adicionar os custos mapeados ao dataframe de custos
                if novos_custos_map:
                    df_novos = pd.DataFrame(novos_custos_map)
                    custos_novadis = pd.concat([custos_novadis, df_novos], ignore_index=True)
                    # Remover duplicados mantendo o último (caso haja conflito)
                    custos_novadis = custos_novadis.drop_duplicates(subset=['Produto'], keep='last')

            # 5. Cruzar Dados (Merge)
            df_margens = pd.merge(vendas_resumo, custos_novadis, on='Produto', how='left')
            df_margens = pd.merge(df_margens, custos_manuais, on='Produto', how='left')
            df_margens = pd.merge(df_margens, custos_fichas, on='Produto', how='left')

            # Determinar Custo Final (Prioridade: Ficha Técnica > Manual > Novadis)
            df_margens['Custo_Final'] = df_margens['Custo_Ficha'].combine_first(df_margens['Custo_Manual']).combine_first(df_margens['Custo_Novadis'])
            
            # Identificar produtos sem custo e produtos de custo sem venda
            produtos_sem_custo = df_margens[df_margens['Custo_Final'].isna()].copy()
            
            # --- ÁREA DE ASSOCIAÇÃO DE PRODUTOS ---
            with st.expander("🔗 Associar Produtos (Nomes Diferentes ou Quantidades)", expanded=False):
                st.info("Use esta secção para ligar produtos com nomes diferentes ou capacidades diferentes (ex: Caneca = 3x Imperial).")
                
                col_assoc1, col_assoc2, col_assoc3, col_assoc4 = st.columns([2, 2, 1, 1])
                
                # Lista de produtos POS sem custo (candidatos a mapeamento)
                lista_pos = sorted(df_pos['Produto'].unique())
                # Lista de produtos Novadis (fonte de custo)
                lista_custos = sorted(custos_novadis['Produto'].unique())
                
                with col_assoc1:
                    sel_pos = st.selectbox("Produto no POS (Venda)", options=lista_pos, key="sel_map_pos")
                
                with col_assoc2:
                    sel_custo = st.selectbox("Produto Base (Custo)", options=lista_custos, key="sel_map_custo")
                
                with col_assoc3:
                    multiplicador_input = st.number_input("Multiplicador", min_value=0.1, value=1.0, step=0.1, help="Ex: Se Caneca leva 3 imperiais, coloque 3.0")
                
                with col_assoc4:
                    st.write("<br>", unsafe_allow_html=True)
                    if st.button("➕ Associar", use_container_width=True):
                        # Adicionar ao CSV
                        novo_map = pd.DataFrame([{
                            'Produto_POS': sel_pos, 
                            'Produto_Custo': sel_custo,
                            'Multiplicador': multiplicador_input
                        }])
                        
                        if arquivo_map.exists():
                            df_existente = pd.read_csv(arquivo_map)
                            if 'Multiplicador' not in df_existente.columns:
                                df_existente['Multiplicador'] = 1.0
                                
                            # Remover se já existir mapeamento para este produto POS
                            df_existente = df_existente[df_existente['Produto_POS'] != sel_pos]
                            df_final = pd.concat([df_existente, novo_map], ignore_index=True)
                        else:
                            df_final = novo_map
                        
                        df_final.to_csv(arquivo_map, index=False)
                        st.success(f"✅ Associado: {sel_pos} = {multiplicador_input}x {sel_custo}")
                        st.rerun()

                # Mostrar associações existentes
                if arquivo_map.exists():
                    df_map_show = pd.read_csv(arquivo_map)
                    if not df_map_show.empty:
                        # Garantir coluna multiplicador
                        if 'Multiplicador' not in df_map_show.columns:
                            df_map_show['Multiplicador'] = 1.0

                        st.markdown("##### Associações Ativas:")
                        
                        # Adicionar botão de remover para cada linha
                        for idx, row in df_map_show.iterrows():
                            c1, c2, c3, c4 = st.columns([3, 3, 1, 1])
                            c1.text(f"POS: {row['Produto_POS']}")
                            c2.text(f"Base: {row['Produto_Custo']}")
                            c3.text(f"x {row['Multiplicador']}")
                            if c4.button("🗑️", key=f"del_map_{idx}"):
                                df_map_show = df_map_show.drop(idx)
                                df_map_show.to_csv(arquivo_map, index=False)
                                st.rerun()

            # --- ÁREA DE INPUT MANUAL ---
            if not produtos_sem_custo.empty:
                st.markdown("### ⚠️ Produtos sem Custo Definido")
                st.warning(f"Existem **{len(produtos_sem_custo)}** produtos vendidos sem custo associado (não encontrados na Novadis).")
                
                with st.expander("✏️ Inserir Custos Manuais", expanded=True):
                    st.caption("Edite a coluna 'Novo Custo' e clique em Guardar.")
                    
                    # Preparar dataframe para edição
                    df_editor = produtos_sem_custo[['Produto', 'Categoria', 'Qtd', 'PVP_Medio']].copy()
                    df_editor['Novo_Custo'] = 0.0
                    
                    # Editor
                    df_editado = st.data_editor(
                        df_editor,
                        column_config={
                            "Produto": st.column_config.TextColumn("Produto", disabled=True),
                            "Categoria": st.column_config.TextColumn("Categoria", disabled=True),
                            "Qtd": st.column_config.NumberColumn("Qtd Vendida", disabled=True),
                            "PVP_Medio": st.column_config.NumberColumn("PVP Médio", format="%.2f€", disabled=True),
                            "Novo_Custo": st.column_config.NumberColumn("Novo Custo Unit. (€)", min_value=0.0, format="%.2f€", required=True)
                        },
                        hide_index=True,
                        use_container_width=True,
                        key="editor_custos_faltantes"
                    )
                    
                    if st.button("💾 Guardar Custos Manuais", type="primary"):
                        # Filtrar apenas os que foram editados (Custo > 0)
                        novos_custos = df_editado[df_editado['Novo_Custo'] > 0][['Produto', 'Novo_Custo']]
                        
                        if not novos_custos.empty:
                            novos_custos.columns = ['Produto', 'Custo_Manual']
                            
                            # Carregar existente, atualizar/adicionar e salvar
                            if arquivo_custos_manuais.exists():
                                df_existente = pd.read_csv(arquivo_custos_manuais)
                                # Remover duplicados (manter novo)
                                df_existente = df_existente[~df_existente['Produto'].isin(novos_custos['Produto'])]
                                df_final = pd.concat([df_existente, novos_custos], ignore_index=True)
                            else:
                                df_final = novos_custos
                                
                            df_final.to_csv(arquivo_custos_manuais, index=False)
                            st.success(f"✅ {len(novos_custos)} custos atualizados! A página será recarregada.")
                            st.rerun()
                        else:
                            st.warning("⚠️ Nenhum custo inserido (valores > 0).")

            # --- ANÁLISE DE MARGENS ---
            st.markdown("---")
            st.subheader("📊 Relatório de Margens e Lucro")
            
            # Calcular Margens para produtos com custo
            df_analise = df_margens[df_margens['Custo_Final'].notna()].copy()
            
            if df_analise.empty:
                st.info("Aguardando inserção de custos para gerar análise.")
            else:
                df_analise['Custo_Total'] = df_analise['Custo_Final'] * df_analise['Qtd']
                df_analise['Lucro_Bruto'] = df_analise['Valor'] - df_analise['Custo_Total']
                df_analise['Margem_Pct'] = (df_analise['Lucro_Bruto'] / df_analise['Valor']) * 100
                # Garantir que Margem_Pct é numérico para evitar erros de arredondamento
                df_analise['Margem_Pct'] = pd.to_numeric(df_analise['Margem_Pct'], errors='coerce').fillna(0)
                
                # KPIs Gerais
                col1, col2, col3, col4 = st.columns(4)
                
                total_vendas_analisado = df_analise['Valor'].sum()
                total_custo_analisado = df_analise['Custo_Total'].sum()
                lucro_total = total_vendas_analisado - total_custo_analisado
                margem_media = (lucro_total / total_vendas_analisado) * 100
                
                with col1:
                    st.metric("Vendas (Analisadas)", formatar_moeda(total_vendas_analisado))
                with col2:
                    st.metric("Custos (Produtos)", formatar_moeda(total_custo_analisado))
                with col3:
                    st.metric("Lucro Bruto Real", formatar_moeda(lucro_total))
                with col4:
                    st.metric("Margem Global", f"{margem_media:.1f}%")
                
                # Visualizações
                col_g1, col_g2 = st.columns(2)
                
                with col_g1:
                    # Top Produtos por Lucro
                    top_lucro = df_analise.sort_values('Lucro_Bruto', ascending=False).head(10)
                    # Arredondar para evitar 13 casas decimais
                    top_lucro['Margem_Pct'] = top_lucro['Margem_Pct'].round(0)

                    fig_lucro = px.bar(
                        top_lucro,
                        x='Lucro_Bruto',
                        y='Produto',
                        orientation='h',
                        title='Top 10 Produtos - Maior Lucro (€)',
                        text=[formatar_moeda(v) for v in top_lucro['Lucro_Bruto']],
                        color='Margem_Pct',
                        color_continuous_scale='Viridis'
                    )
                    fig_lucro.update_layout(yaxis={'categoryorder': 'total ascending'})
                    st.plotly_chart(fig_lucro, use_container_width=True)
                    
                with col_g2:
                    # Margens Baixas (Alerta)
                    margens_baixas = df_analise[df_analise['Margem_Pct'] < 50].sort_values('Margem_Pct')
                    if not margens_baixas.empty:
                        # Preparar view arredondada
                        mb_view = margens_baixas.head(10).copy()
                        mb_view['Margem_Pct'] = mb_view['Margem_Pct'].round(0)
                        
                        fig_alert = px.bar(
                            mb_view,
                            x='Margem_Pct',
                            y='Produto',
                            orientation='h',
                            title='⚠️ Produtos com Margem < 50%',
                            text=[f"{v:.0f}%" for v in mb_view['Margem_Pct']],
                            color='Margem_Pct',
                            color_continuous_scale='RdYlGn',
                            range_color=[0, 80]
                        )
                        fig_alert.update_layout(yaxis={'categoryorder': 'total descending'})
                        st.plotly_chart(fig_alert, use_container_width=True)
                    else:
                        st.success("✅ Todos os produtos analisados têm margem superior a 50%!")

                # Tabela Detalhada
                st.subheader("📋 Detalhe por Produto")
                
                # Filtros
                col_f1, col_f2 = st.columns(2)
                with col_f1:
                    cat_filter = st.multiselect("Filtrar Categoria", options=sorted(df_analise['Categoria'].unique()))
                with col_f2:
                    search_prod = st.text_input("Pesquisar Produto")
                    
                df_table = df_analise.copy()
                if cat_filter:
                    df_table = df_table[df_table['Categoria'].isin(cat_filter)]
                if search_prod:
                    df_table = df_table[df_table['Produto'].str.contains(search_prod, case=False)]
                
                # Formatar e Mostrar
                df_display = df_table[[
                    'Produto', 'Categoria', 'Qtd', 'PVP_Medio', 
                    'Custo_Final', 'Valor', 'Lucro_Bruto', 'Margem_Pct'
                ]].sort_values('Lucro_Bruto', ascending=False)
                
                df_display.columns = [
                    'Produto', 'Categoria', 'Qtd', 'PVP Médio', 
                    'Custo Unit.', 'Vendas Total', 'Lucro', 'Margem %'
                ]
                
                # Configuração da tabela
                st.dataframe(
                    df_display,
                    column_config={
                        "PVP Médio": st.column_config.NumberColumn(format="%.2f€"),
                        "Custo Unit.": st.column_config.NumberColumn(format="%.2f€"),
                        "Vendas Total": st.column_config.NumberColumn(format="%.0f€"),
                        "Lucro": st.column_config.NumberColumn(format="%.0f€"),
                        "Margem %": st.column_config.ProgressColumn(
                            format="%.1f%%",
                            min_value=0,
                            max_value=100,
                        ),
                    },
                    hide_index=True,
                    use_container_width=True
                )
                
                # Botão para Editar Custos Existentes (Override)
                with st.expander("🛠️ Ajustar Custos Existentes (Override)"):
                    st.write("Se o custo automático da Novadis estiver incorreto, você pode forçar um valor manual aqui.")
                    
                    df_override = df_analise[['Produto', 'Custo_Final']].copy()
                    df_override.columns = ['Produto', 'Custo_Atual']
                    df_override['Novo_Custo_Manual'] = 0.0
                    
                    df_override_edit = st.data_editor(
                        df_override,
                        column_config={
                            "Produto": st.column_config.TextColumn(disabled=True),
                            "Custo_Atual": st.column_config.NumberColumn(format="%.2f€", disabled=True),
                            "Novo_Custo_Manual": st.column_config.NumberColumn("Novo Custo (€)", min_value=0.0, format="%.2f€")
                        },
                        hide_index=True,
                        key="override_custos"
                    )
                    
                    if st.button("💾 Atualizar Custos Manuais"):
                        novos_overrides = df_override_edit[df_override_edit['Novo_Custo_Manual'] > 0][['Produto', 'Novo_Custo_Manual']]
                        if not novos_overrides.empty:
                            novos_overrides.columns = ['Produto', 'Custo_Manual']
                            
                            if arquivo_custos_manuais.exists():
                                df_existente = pd.read_csv(arquivo_custos_manuais)
                                df_existente = df_existente[~df_existente['Produto'].isin(novos_overrides['Produto'])]
                                df_final = pd.concat([df_existente, novos_overrides], ignore_index=True)
                            else:
                                df_final = novos_overrides
                                
                            df_final.to_csv(arquivo_custos_manuais, index=False)
                            st.success("✅ Custos atualizados!")
                            st.rerun()


if __name__ == '__main__':
    # Configuração de autenticação
    config_file = Path(__file__).parent / 'config.yaml'

    # Verificar se o arquivo de configuração existe
    if not config_file.exists():
        st.error(f'Arquivo de configuração não encontrado: {config_file}')
        st.stop()

    with open(config_file) as file:
        config = yaml.load(file, Loader=SafeLoader)

    # Criar objeto de autenticação
    authenticator = stauth.Authenticate(
        config['credentials'],
        config['cookie']['name'],
        config['cookie']['key'],
        config['cookie']['expiry_days']
    )

    # Renderizar formulário de login
    try:
        authenticator.login()
    except Exception as e:
        st.error(f'Erro ao processar login: {e}')
        st.stop()

    # Verificar status de autenticação
    if st.session_state.get("authentication_status"):
        # Usuário autenticado - adicionar logout na sidebar
        authenticator.logout(location="sidebar", use_container_width=True)
        st.sidebar.write(f'Bem-vindo, **{st.session_state["name"]}**!')

        # Executar aplicação principal
        main()

    elif st.session_state.get("authentication_status") is False:
        st.error('Utilizador ou password incorretos')
    elif st.session_state.get("authentication_status") is None:
        st.warning('Por favor, introduza as suas credenciais')
