"""
Dashboard Interativo de Análise de Vendas - Jogos Santa Casa
Utilize: streamlit run dashboard.py
"""

import streamlit as st
import pandas as pd
import plotly.express as px
import plotly.graph_objects as go
from plotly.subplots import make_subplots
import numpy as np
from pathlib import Path
import re
from datetime import datetime, timedelta
from sklearn.linear_model import LinearRegression
from sklearn.preprocessing import PolynomialFeatures
import streamlit_authenticator as stauth
import yaml
from yaml.loader import SafeLoader
import warnings
import io
from io import BytesIO
import subprocess
import json
import os

# Tentar importar bibliotecas para exportação (opcionais)
try:
    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    from openpyxl.utils.dataframe import dataframe_to_rows
    EXCEL_AVAILABLE = True
except ImportError:
    EXCEL_AVAILABLE = False

warnings.filterwarnings('ignore')

# Configuração da página
st.set_page_config(
    page_title="Dashboard de Vendas - Santa Casa",
    page_icon="🎲",
    layout="wide",
    initial_sidebar_state="expanded"
)

# CSS customizado
st.markdown("""
<style>
    .main-header {
        font-size: 3rem;
        font-weight: bold;
        text-align: center;
        color: #1f77b4;
        margin-bottom: 2rem;
    }
    .metric-card {
        background-color: #f0f2f6;
        padding: 1rem;
        border-radius: 0.5rem;
        border-left: 4px solid #1f77b4;
    }
    .stTabs [data-baseweb="tab-list"] {
        gap: 2rem;
    }

    /* Cards profissionais para KPIs */
    .kpi-card {
        background: linear-gradient(135deg, #667eea 0%, #764ba2 100%);
        padding: 25px;
        border-radius: 15px;
        box-shadow: 0 8px 20px rgba(0, 0, 0, 0.15);
        color: white;
        text-align: center;
        margin: 10px 0;
        transition: transform 0.2s;
    }
    .kpi-card:hover {
        transform: translateY(-5px);
        box-shadow: 0 12px 25px rgba(0, 0, 0, 0.2);
    }
    .kpi-card-green {
        background: linear-gradient(135deg, #11998e 0%, #38ef7d 100%);
    }
    .kpi-card-blue {
        background: linear-gradient(135deg, #667eea 0%, #764ba2 100%);
    }
    .kpi-card-orange {
        background: linear-gradient(135deg, #f093fb 0%, #f5576c 100%);
    }
    .kpi-card-yellow {
        background: linear-gradient(135deg, #ffecd2 0%, #fcb69f 100%);
        color: #333;
    }
    .kpi-value {
        font-size: 2.8rem;
        font-weight: 800;
        margin: 15px 0;
        text-shadow: 2px 2px 4px rgba(0,0,0,0.2);
    }
    .kpi-label {
        font-size: 1rem;
        opacity: 0.95;
        font-weight: 500;
        letter-spacing: 0.5px;
    }
    .kpi-icon {
        font-size: 2.5rem;
        margin-bottom: 10px;
        opacity: 0.9;
    }
    .kpi-delta {
        font-size: 0.9rem;
        margin-top: 10px;
        font-weight: 600;
    }

    /* Cards executivos */
    .exec-card {
        background: white;
        padding: 20px;
        border-radius: 12px;
        box-shadow: 0 4px 15px rgba(0, 0, 0, 0.1);
        border-left: 5px solid #667eea;
        margin: 10px 0;
    }

    /* Sparkline container */
    .sparkline-container {
        height: 40px;
        margin: 10px 0;
    }

    /* Top game card */
    .top-game-card {
        background: linear-gradient(135deg, #ffecd2 0%, #fcb69f 100%);
        padding: 15px;
        border-radius: 10px;
        margin: 8px 0;
        box-shadow: 0 4px 10px rgba(0, 0, 0, 0.1);
    }
    .top-game-rank {
        font-size: 2rem;
        font-weight: bold;
        color: #764ba2;
    }
    .top-game-name {
        font-size: 1.2rem;
        font-weight: 600;
        color: #333;
    }
    .top-game-value {
        font-size: 1.5rem;
        font-weight: bold;
        color: #11998e;
    }
</style>
""", unsafe_allow_html=True)

# Mapeamento de nomes de jogos para exibição
NOMES_JOGOS = {
    'Subtotal (LI)': 'Lotaria Instantânea',
    'Subtotal (LP)': 'Lotaria Popular',
    'Subtotal (LC)': 'Lotaria Clássica'
}

# Percentagens de remuneração por jogo
REMUNERACAO_JOGOS = {
    'Euromilhões': 5.0,
    'EuroDreams': 5.0,
    'Totoloto': 7.0,
    'M1lhao': 5.0,
    'Totobola': 7.0,
    'Hípicas': 7.0,
    'Placard': 5.0,
    'Lotaria Instantânea': 10.0,
    'Lotaria Popular': 12.5,
    'Lotaria Clássica': 12.7
}

# Objetivos semanais por jogo (em euros)
OBJETIVOS_SEMANAIS = {
    'Raspadinha': 2627.70,
    'Euromilhões': 785.40,  # Objetivo atualizado (88% do total Euromilhões+M1lhão)
    'M1lhao': 107.10,  # Objetivo atualizado (12% do total Euromilhões+M1lhão)
    'EuroDreams': 301.65,
    'Placard': 433.55,
    'Lotaria Clássica': 72.15,
    'Lotaria Popular': 107.40,
    'Lotaria Instantânea': 2627.70,  # Mesmo que Raspadinha (jogo equivalente)
    'Totoloto': 153.00,
    'Totobola': 3.00,
    'Hípicas': 0.0  # Sem vendas/objetivo
}

# Proporção Euromilhões + M1lhão (por cada aposta de 2,5€)
PROPORCAO_EURO_M1LHAO = {
    'Euromilhões': 2.20 / 2.50,  # 88%
    'M1lhao': 0.30 / 2.50         # 12%
}


def renomear_jogo(nome):
    """Renomeia jogo se estiver no mapeamento."""
    return NOMES_JOGOS.get(nome, nome)


@st.cache_data
def carregar_dados():
    """Carrega e processa os dados do arquivo."""
    # Caminho fixo para os dados no Raspberry Pi
    data_file = Path('/home/jorge/Documentos/Santa casa/dados/dados_extracao.txt')

    if not data_file.exists():
        st.error(f"Arquivo não encontrado: {data_file}")
        return None

    # Lendo o arquivo
    try:
        text = data_file.read_text(encoding='utf-8')
    except UnicodeDecodeError:
        try:
            text = data_file.read_text(encoding='cp1252')
        except UnicodeDecodeError:
            text = data_file.read_text(encoding='latin-1')

    data = text.splitlines()
    records = []

    # Regex para extrair dados
    pattern = r'^(.+?):\s*(\d+\.\d+)\s*\(Data de Emiss[ãa]o:\s*(\d{2}-\d{2}-\d{4})\)'

    for line in data:
        line = line.strip()
        if not line:
            continue

        match = re.match(pattern, line)
        if match:
            jogo = match.group(1).strip()
            valor = float(match.group(2))
            data_emissao = match.group(3)

            records.append({
                'Jogo': jogo,
                'Valor': valor,
                'Data_Emissao': data_emissao
            })

    # Criar DataFrame
    df = pd.DataFrame(records)
    df['Data_Emissao'] = pd.to_datetime(df['Data_Emissao'], format='%d-%m-%Y')

    # Adicionar +2 dias a todas as datas
    df['Data_Emissao'] = df['Data_Emissao'] + timedelta(days=2)

    # Renomear jogos de acordo com o mapeamento
    df['Jogo'] = df['Jogo'].map(lambda x: renomear_jogo(x))

    # Adicionar colunas temporais
    df['Ano'] = df['Data_Emissao'].dt.year
    df['Mes'] = df['Data_Emissao'].dt.month
    df['Mes_Nome'] = df['Data_Emissao'].dt.strftime('%B')
    df['Semana_Ano'] = df['Data_Emissao'].dt.isocalendar().week
    df['Trimestre'] = df['Data_Emissao'].dt.quarter
    df['Dia_Semana'] = df['Data_Emissao'].dt.day_name()
    df['Ano_Semana'] = df['Ano'].astype(str) + '-S' + df['Semana_Ano'].astype(str).str.zfill(2)
    df['Ano_Mes'] = df['Data_Emissao'].dt.to_period('M').astype(str)

    return df


def reprocessar_pdfs():
    """Reprocessa todos os PDFs na pasta configurada."""

    # Tentar carregar a configuração para obter o caminho
    try:
        config_path = Path('/home/jorge/Documentos/Santa casa/config.json')
        if config_path.exists():
            with open(config_path, 'r') as f:
                config = json.load(f)
                pdf_folder = config.get('pdf_folder')
                processor_script = Path(__file__).parent.parent / 'install_dashboard' / 'process_pdfs.py'
        else:
            # Caminho padrão
            pdf_folder = '/home/jorge/Documentos/Santa casa/pdf'
            processor_script = Path('/home/jorge/Documentos/install_dashboard/process_pdfs.py')

        if not processor_script.exists():
            return False, f"Script de processamento não encontrado: {processor_script}"

        # Executar o processador de PDFs
        try:
            result = subprocess.run(
                [f"python3", str(processor_script), pdf_folder],
                capture_output=True,
                text=True,
                timeout=300  # Timeout de 5 minutos
            )

            if result.returncode == 0:
                return True, "PDFs reprocessados com sucesso!"
            else:
                return False, f"Erro ao reprocessar PDFs:\n{result.stderr}"

        except subprocess.TimeoutExpired:
            return False, "O processamento demorou muito tempo (timeout)"
        except Exception as e:
            return False, f"Erro ao executar processamento: {str(e)}"

    except Exception as e:
        return False, f"Erro ao configurar reprocessamento: {str(e)}"


def calcular_previsao(df_serie, periodos_futuros=4):
    """Calcula previsão usando regressão polinomial."""
    if len(df_serie) < 3:
        return None, None

    X = np.arange(len(df_serie)).reshape(-1, 1)
    y = df_serie.values

    # Regressão polinomial de grau 2
    poly = PolynomialFeatures(degree=2)
    X_poly = poly.fit_transform(X)

    model = LinearRegression()
    model.fit(X_poly, y)

    # Prever valores futuros
    X_futuro = np.arange(len(df_serie), len(df_serie) + periodos_futuros).reshape(-1, 1)
    X_futuro_poly = poly.transform(X_futuro)
    previsao = model.predict(X_futuro_poly)

    return previsao, model.score(X_poly, y)


def calcular_metricas_periodo(df, data_referencia=None):
    """
    Calcula métricas para diferentes períodos temporais (semanal, mensal, anual).
    Usa a última semana/mês com dados disponíveis como referência.
    """
    if data_referencia is None:
        data_referencia = datetime.now()

    # Garantir que Data_Emissao é datetime
    if not pd.api.types.is_datetime64_any_dtype(df['Data_Emissao']):
        df = df.copy()
        df['Data_Emissao'] = pd.to_datetime(df['Data_Emissao'])

    # Encontrar a última data com dados disponíveis
    ultima_data_disponivel = df['Data_Emissao'].max()

    # SEMANAS - Usar a última semana com dados
    semana_ref = ultima_data_disponivel.isocalendar()[1]
    ano_ref = ultima_data_disponivel.isocalendar()[0]

    # Calcular semana anterior e SWLY baseado na última semana disponível
    data_semana_anterior = ultima_data_disponivel - timedelta(weeks=1)
    semana_anterior = data_semana_anterior.isocalendar()[1]
    ano_semana_anterior = data_semana_anterior.isocalendar()[0]

    ano_passado = ano_ref - 1

    # MESES - Usar o último mês com dados
    mes_ref = ultima_data_disponivel.month
    ano_mes_ref = ultima_data_disponivel.year

    # Mês anterior
    primeiro_dia_mes = ultima_data_disponivel.replace(day=1)
    data_mes_anterior = primeiro_dia_mes - timedelta(days=1)
    mes_anterior = data_mes_anterior.month
    ano_mes_anterior = data_mes_anterior.year

    metricas = {}
    metricas['ultima_data'] = ultima_data_disponivel.strftime('%d/%m/%Y')

    # SEMANAS
    # Última semana com dados
    dados_semana_atual = df[
        (df['Data_Emissao'].dt.isocalendar().week == semana_ref) &
        (df['Data_Emissao'].dt.isocalendar().year == ano_ref)
    ]
    metricas['semana_atual'] = dados_semana_atual['Valor'].sum() if not dados_semana_atual.empty else 0

    # Semana anterior à última
    dados_semana_anterior = df[
        (df['Data_Emissao'].dt.isocalendar().week == semana_anterior) &
        (df['Data_Emissao'].dt.isocalendar().year == ano_semana_anterior)
    ]
    metricas['semana_anterior'] = dados_semana_anterior['Valor'].sum() if not dados_semana_anterior.empty else 0

    # SWLY (Same Week Last Year)
    dados_swly = df[
        (df['Data_Emissao'].dt.isocalendar().week == semana_ref) &
        (df['Data_Emissao'].dt.isocalendar().year == ano_passado)
    ]
    metricas['swly'] = dados_swly['Valor'].sum() if not dados_swly.empty else 0

    # Variações semanais
    if metricas['semana_anterior'] > 0:
        metricas['var_semana'] = ((metricas['semana_atual'] - metricas['semana_anterior']) / metricas['semana_anterior']) * 100
    else:
        metricas['var_semana'] = 0

    if metricas['swly'] > 0:
        metricas['var_swly'] = ((metricas['semana_atual'] - metricas['swly']) / metricas['swly']) * 100
    else:
        metricas['var_swly'] = 0

    # MESES - Baseado na última data disponível
    # Último mês com dados
    dados_mes_atual = df[
        (df['Data_Emissao'].dt.month == mes_ref) &
        (df['Data_Emissao'].dt.year == ano_mes_ref)
    ]
    metricas['mes_atual'] = dados_mes_atual['Valor'].sum() if not dados_mes_atual.empty else 0

    # Mês anterior ao último
    dados_mes_anterior = df[
        (df['Data_Emissao'].dt.month == mes_anterior) &
        (df['Data_Emissao'].dt.year == ano_mes_anterior)
    ]
    metricas['mes_anterior'] = dados_mes_anterior['Valor'].sum() if not dados_mes_anterior.empty else 0

    # SMLY (Same Month Last Year)
    dados_smly = df[
        (df['Data_Emissao'].dt.month == mes_ref) &
        (df['Data_Emissao'].dt.year == ano_mes_ref - 1)
    ]
    metricas['smly'] = dados_smly['Valor'].sum() if not dados_smly.empty else 0

    # Variações mensais
    if metricas['mes_anterior'] > 0:
        metricas['var_mes'] = ((metricas['mes_atual'] - metricas['mes_anterior']) / metricas['mes_anterior']) * 100
    else:
        metricas['var_mes'] = 0

    if metricas['smly'] > 0:
        metricas['var_smly'] = ((metricas['mes_atual'] - metricas['smly']) / metricas['smly']) * 100
    else:
        metricas['var_smly'] = 0

    # ANOS - Baseado no ano da última data disponível
    # Ano da última data (YTD - Year to Date)
    dados_ano_atual = df[df['Data_Emissao'].dt.year == ano_mes_ref]
    metricas['ano_atual'] = dados_ano_atual['Valor'].sum() if not dados_ano_atual.empty else 0

    # Ano anterior completo
    dados_ano_anterior = df[df['Data_Emissao'].dt.year == ano_mes_ref - 1]
    metricas['ano_anterior'] = dados_ano_anterior['Valor'].sum() if not dados_ano_anterior.empty else 0

    # Ano anterior até à mesma data (YTD comparison)
    # Ex: Se estamos em 8 Nov 2025, pegar dados de 2024 até 8 Nov 2024
    mes_ref_num = mes_ref
    dia_ref = ultima_data_disponivel.day

    dados_ano_anterior_ytd = df[
        (df['Data_Emissao'].dt.year == ano_mes_ref - 1) &
        (
            (df['Data_Emissao'].dt.month < mes_ref_num) |
            ((df['Data_Emissao'].dt.month == mes_ref_num) & (df['Data_Emissao'].dt.day <= dia_ref))
        )
    ]
    metricas['ano_anterior_ytd'] = dados_ano_anterior_ytd['Valor'].sum() if not dados_ano_anterior_ytd.empty else 0

    # Variação anual (comparando YTD)
    if metricas['ano_anterior_ytd'] > 0:
        metricas['var_ano'] = ((metricas['ano_atual'] - metricas['ano_anterior_ytd']) / metricas['ano_anterior_ytd']) * 100
    else:
        metricas['var_ano'] = 0

    return metricas


def pagina_visao_geral(df):
    """Página com visão geral das vendas."""
    st.markdown('<h1 class="main-header">📊 Visão Geral de Vendas</h1>', unsafe_allow_html=True)

    # Filtros na sidebar
    st.sidebar.header("🔍 Filtros")

    anos_disponiveis = sorted(df['Ano'].unique())
    ano_selecionado = st.sidebar.multiselect(
        "Selecionar Anos",
        anos_disponiveis,
        default=anos_disponiveis
    )

    if not ano_selecionado:
        st.warning("Selecione pelo menos um ano")
        return

    df_filtrado = df[df['Ano'].isin(ano_selecionado)]

    # Métricas principais com cards profissionais
    total_vendas = df_filtrado['Valor'].sum()
    media_semanal = df_filtrado.groupby('Data_Emissao')['Valor'].sum().mean()
    num_registos = len(df_filtrado)
    num_jogos = df_filtrado['Jogo'].nunique()

    # Calcular variação WoW para cards
    datas_unicas = sorted(df['Data_Emissao'].unique(), reverse=True)
    delta_vendas = ""
    if len(datas_unicas) >= 2:
        vendas_ultima = df[df['Data_Emissao'] == datas_unicas[0]]['Valor'].sum()
        vendas_anterior = df[df['Data_Emissao'] == datas_unicas[1]]['Valor'].sum()
        if vendas_anterior > 0:
            var_perc = ((vendas_ultima - vendas_anterior) / vendas_anterior) * 100
            delta_vendas = f"{var_perc:+.1f}% vs semana anterior"

    col1, col2, col3, col4 = st.columns(4)

    with col1:
        st.markdown(f"""
        <div class="kpi-card kpi-card-green">
            <div class="kpi-icon">💰</div>
            <div class="kpi-label">TOTAL DE VENDAS</div>
            <div class="kpi-value">€{total_vendas:,.0f}</div>
            <div class="kpi-delta">{delta_vendas}</div>
        </div>
        """, unsafe_allow_html=True)

    with col2:
        st.markdown(f"""
        <div class="kpi-card kpi-card-blue">
            <div class="kpi-icon">📅</div>
            <div class="kpi-label">MÉDIA SEMANAL</div>
            <div class="kpi-value">€{media_semanal:,.0f}</div>
            <div class="kpi-delta">{len(df_filtrado['Data_Emissao'].unique())} semanas</div>
        </div>
        """, unsafe_allow_html=True)

    with col3:
        st.markdown(f"""
        <div class="kpi-card kpi-card-orange">
            <div class="kpi-icon">📋</div>
            <div class="kpi-label">TOTAL REGISTOS</div>
            <div class="kpi-value">{num_registos:,}</div>
            <div class="kpi-delta">Dados processados</div>
        </div>
        """, unsafe_allow_html=True)

    with col4:
        st.markdown(f"""
        <div class="kpi-card kpi-card-yellow">
            <div class="kpi-icon">🎮</div>
            <div class="kpi-label">JOGOS ATIVOS</div>
            <div class="kpi-value">{num_jogos}</div>
            <div class="kpi-delta">Produtos diferentes</div>
        </div>
        """, unsafe_allow_html=True)

    st.markdown("---")

    # Gráficos em duas colunas
    col1, col2 = st.columns(2)

    with col1:
        # Top 10 jogos por vendas totais
        st.subheader("🏆 Top 10 Jogos por Vendas")
        top_jogos = df_filtrado.groupby('Jogo')['Valor'].sum().sort_values(ascending=False).head(10)

        fig = px.bar(
            x=top_jogos.values,
            y=top_jogos.index,
            orientation='h',
            labels={'x': 'Vendas (€)', 'y': 'Jogo'},
            color=top_jogos.values,
            color_continuous_scale='Blues'
        )
        fig.update_layout(showlegend=False, height=400)
        st.plotly_chart(fig, use_container_width=True)

    with col2:
        # Distribuição por categoria
        st.subheader("📊 Distribuição por Categoria")

        df_categoria = df_filtrado.copy()
        df_categoria['Categoria'] = df_categoria['Jogo'].apply(
            lambda x: 'Lotarias' if 'Lotaria' in x else 'Jogos'
        )

        vendas_categoria = df_categoria.groupby('Categoria')['Valor'].sum()

        fig = px.pie(
            values=vendas_categoria.values,
            names=vendas_categoria.index,
            hole=0.4,
            color_discrete_sequence=['#1f77b4', '#ff7f0e']
        )
        fig.update_layout(height=400)
        st.plotly_chart(fig, use_container_width=True)

    st.markdown("---")

    # ===== MÉTRICAS DE PERFORMANCE =====
    st.markdown("### 📊 Métricas de Performance")

    # Calcular métricas temporais (usa todo o dataset para comparações temporais)
    metricas = calcular_metricas_periodo(df)

    st.info(f"📅 Dados atualizados até: **{metricas['ultima_data']}**")

    # Performance Semanal
    st.markdown("#### 📅 Performance Semanal")
    col_s1, col_s2, col_s3 = st.columns(3)

    with col_s1:
        st.metric(
            label="Última Semana",
            value=f"€{metricas['semana_atual']:,.0f}",
            delta=f"{metricas['var_semana']:.1f}% vs Semana Anterior"
        )

    with col_s2:
        st.metric(
            label="Semana Anterior",
            value=f"€{metricas['semana_anterior']:,.0f}"
        )

    with col_s3:
        st.metric(
            label="SWLY (Mesma Semana Ano Passado)",
            value=f"€{metricas['swly']:,.0f}",
            delta=f"{metricas['var_swly']:.1f}% vs SWLY"
        )

    # Performance Mensal
    st.markdown("#### 📆 Performance Mensal")
    col_m1, col_m2, col_m3 = st.columns(3)

    with col_m1:
        st.metric(
            label="Mês em Curso",
            value=f"€{metricas['mes_atual']:,.0f}",
            delta=f"{metricas['var_mes']:.1f}% vs Mês Anterior"
        )

    with col_m2:
        st.metric(
            label="Mês Anterior",
            value=f"€{metricas['mes_anterior']:,.0f}"
        )

    with col_m3:
        st.metric(
            label="SMLY (Mesmo Mês Ano Passado)",
            value=f"€{metricas['smly']:,.0f}",
            delta=f"{metricas['var_smly']:.1f}% vs SMLY"
        )

    # Performance Anual
    st.markdown("#### 📈 Performance Anual")
    col_a1, col_a2, col_a3 = st.columns(3)

    with col_a1:
        st.metric(
            label="Ano em Curso (YTD)",
            value=f"€{metricas['ano_atual']:,.0f}",
            delta=f"{metricas['var_ano']:.1f}% vs Ano Anterior YTD"
        )

    with col_a2:
        st.metric(
            label="Ano Anterior (mesma altura)",
            value=f"€{metricas['ano_anterior_ytd']:,.0f}",
            help=f"Ano anterior até {metricas['ultima_data']}"
        )

    with col_a3:
        st.metric(
            label="Ano Anterior (Total)",
            value=f"€{metricas['ano_anterior']:,.0f}"
        )

    # Evolução temporal
    st.subheader("📈 Evolução de Vendas ao Longo do Tempo")
    st.caption("💡 Nota: Os dados são resumos semanais enviados pela Santa Casa")

    # Agrupar por tipo
    tipo_visualizacao = st.radio(
        "Agrupar por:",
        ["Resumo Semanal", "Semana do Ano", "Mês"],
        horizontal=True
    )

    if tipo_visualizacao == "Resumo Semanal":
        vendas_temp = df_filtrado.groupby('Data_Emissao')['Valor'].sum().reset_index()
        x_col = 'Data_Emissao'
    elif tipo_visualizacao == "Semana do Ano":
        vendas_temp = df_filtrado.groupby('Ano_Semana')['Valor'].sum().reset_index()
        x_col = 'Ano_Semana'
    else:
        vendas_temp = df_filtrado.groupby('Ano_Mes')['Valor'].sum().reset_index()
        x_col = 'Ano_Mes'

    fig = px.line(
        vendas_temp,
        x=x_col,
        y='Valor',
        markers=True,
        labels={'Valor': 'Vendas (€)', x_col: tipo_visualizacao}
    )
    fig.update_layout(height=500)
    st.plotly_chart(fig, use_container_width=True)

    # Exportação de dados
    st.markdown("---")
    st.subheader("📥 Exportar Dados")

    col1, col2 = st.columns(2)

    # Preparar dados para exportação
    top_jogos = df_filtrado.groupby('Jogo')['Valor'].sum().sort_values(ascending=False).reset_index()
    top_jogos.columns = ['Jogo', 'Total Vendas (€)']

    with col1:
        # Exportar CSV
        csv_data = top_jogos.to_csv(index=False).encode('utf-8')
        st.download_button(
            label="📄 Baixar Resumo em CSV",
            data=csv_data,
            file_name=f"visao_geral_{'-'.join(map(str, ano_selecionado))}.csv",
            mime="text/csv"
        )

    with col2:
        # Exportar Excel
        if EXCEL_AVAILABLE:
            try:
                dataframes = {
                    'Top Jogos': top_jogos,
                    'Evolução Temporal': vendas_temp
                }
                info = {
                    'Página': 'Visão Geral',
                    'Anos Selecionados': ', '.join(map(str, ano_selecionado)),
                    'Total de Vendas': f'€{total_vendas:,.0f}',
                    'Média Semanal': f'€{media_semanal:,.0f}'
                }
                excel_data = exportar_excel_generico(dataframes, 'visao_geral', info)
                if excel_data:
                    st.download_button(
                        label="📊 Baixar Relatório em Excel",
                        data=excel_data,
                        file_name=f"visao_geral_{'-'.join(map(str, ano_selecionado))}.xlsx",
                        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                    )
            except Exception as e:
                st.error(f"Erro ao gerar Excel: {str(e)}")
        else:
            st.info("📊 Excel: Instale openpyxl para exportar em Excel")


def pagina_analise_jogos(df):
    """Análise detalhada por jogo."""
    st.markdown('<h1 class="main-header">🎮 Análise por Jogo</h1>', unsafe_allow_html=True)
    st.info("💡 Os dados apresentados são resumos semanais fornecidos pela Santa Casa da Misericórdia")

    # Seleção de jogo
    jogos_disponiveis = sorted(df['Jogo'].unique())
    jogo_selecionado = st.sidebar.selectbox("Selecionar Jogo", jogos_disponiveis)

    df_jogo = df[df['Jogo'] == jogo_selecionado]

    # Métricas do jogo
    col1, col2, col3, col4 = st.columns(4)

    with col1:
        st.metric("💰 Total", f"€{df_jogo['Valor'].sum():,.0f}")
    with col2:
        st.metric("📊 Média", f"€{df_jogo['Valor'].mean():,.0f}")
    with col3:
        st.metric("📈 Máximo", f"€{df_jogo['Valor'].max():,.0f}")
    with col4:
        st.metric("📉 Mínimo", f"€{df_jogo['Valor'].min():,.0f}")

    st.markdown("---")

    # Tabs para diferentes análises
    tab1, tab2, tab3, tab4 = st.tabs(["📈 Tendências", "📅 Sazonalidade", "🔮 Previsões", "📊 Estatísticas"])

    with tab1:
        st.subheader(f"Evolução de Vendas - {jogo_selecionado}")

        # Vendas por ano
        vendas_ano = df_jogo.groupby(['Ano', 'Data_Emissao'])['Valor'].sum().reset_index()

        fig = px.line(
            vendas_ano,
            x='Data_Emissao',
            y='Valor',
            color='Ano',
            markers=True,
            labels={'Valor': 'Vendas (€)', 'Data_Emissao': 'Data'}
        )
        fig.update_layout(height=500)
        st.plotly_chart(fig, use_container_width=True)

        # Comparação ano a ano por semana
        st.subheader("Comparação Semanal entre Anos")
        vendas_semana_ano = df_jogo.groupby(['Ano', 'Semana_Ano'])['Valor'].sum().reset_index()

        fig = px.line(
            vendas_semana_ano,
            x='Semana_Ano',
            y='Valor',
            color='Ano',
            markers=True,
            labels={'Valor': 'Vendas (€)', 'Semana_Ano': 'Semana do Ano'}
        )
        fig.update_layout(height=400)
        st.plotly_chart(fig, use_container_width=True)

    with tab2:
        st.subheader("Análise de Sazonalidade")

        col1, col2 = st.columns(2)

        with col1:
            # Vendas por mês
            vendas_mes = df_jogo.groupby('Mes')['Valor'].mean().reset_index()
            vendas_mes['Mes_Nome'] = vendas_mes['Mes'].apply(
                lambda x: datetime(2024, x, 1).strftime('%B')
            )

            fig = px.bar(
                vendas_mes,
                x='Mes_Nome',
                y='Valor',
                labels={'Valor': 'Média de Vendas (€)', 'Mes_Nome': 'Mês'},
                color='Valor',
                color_continuous_scale='Blues'
            )
            fig.update_layout(showlegend=False, height=400)
            st.plotly_chart(fig, use_container_width=True)

        with col2:
            # Vendas por trimestre
            vendas_trim = df_jogo.groupby('Trimestre')['Valor'].mean().reset_index()

            fig = px.bar(
                vendas_trim,
                x='Trimestre',
                y='Valor',
                labels={'Valor': 'Média de Vendas (€)', 'Trimestre': 'Trimestre'},
                color='Valor',
                color_continuous_scale='Greens'
            )
            fig.update_layout(showlegend=False, height=400)
            st.plotly_chart(fig, use_container_width=True)

        # Heatmap de vendas por semana e ano
        st.subheader("Mapa de Calor: Semana vs Ano")
        pivot_data = df_jogo.groupby(['Ano', 'Semana_Ano'])['Valor'].sum().reset_index()
        pivot_table = pivot_data.pivot(index='Semana_Ano', columns='Ano', values='Valor')

        fig = px.imshow(
            pivot_table,
            labels=dict(x="Ano", y="Semana", color="Vendas (€)"),
            aspect="auto",
            color_continuous_scale='RdYlGn'
        )
        fig.update_layout(height=600)
        st.plotly_chart(fig, use_container_width=True)

    with tab3:
        st.subheader("Previsão de Vendas")
        st.caption("🔮 Previsão baseada nos resumos semanais históricos")

        # Previsão baseada em vendas semanais
        vendas_semanal = df_jogo.groupby('Data_Emissao')['Valor'].sum().sort_index()

        periodos = st.slider("Semanas a Prever", 1, 12, 4)

        previsao, r2_score = calcular_previsao(vendas_semanal, periodos)

        if previsao is not None:
            # Criar DataFrame para visualização
            datas_historicas = vendas_semanal.index
            ultimas_datas = pd.date_range(
                start=datas_historicas[-1] + timedelta(days=7),
                periods=periodos,
                freq='W'
            )

            fig = go.Figure()

            # Dados históricos
            fig.add_trace(go.Scatter(
                x=datas_historicas,
                y=vendas_semanal.values,
                mode='lines+markers',
                name='Histórico',
                line=dict(color='blue')
            ))

            # Previsão
            fig.add_trace(go.Scatter(
                x=ultimas_datas,
                y=previsao,
                mode='lines+markers',
                name='Previsão',
                line=dict(color='red', dash='dash')
            ))

            fig.update_layout(
                title=f'Previsão de Vendas (R² = {r2_score:.3f})',
                xaxis_title='Data',
                yaxis_title='Vendas (€)',
                height=500
            )
            st.plotly_chart(fig, use_container_width=True)

            # Tabela de previsões
            st.subheader("Valores Previstos")
            df_previsao = pd.DataFrame({
                'Data': ultimas_datas.strftime('%d-%m-%Y'),
                'Valor Previsto (€)': [f"{v:.0f}" for v in previsao]
            })
            st.dataframe(df_previsao, use_container_width=True)
        else:
            st.warning("Dados insuficientes para previsão")

    with tab4:
        st.subheader("Estatísticas Detalhadas")

        col1, col2 = st.columns(2)

        with col1:
            st.markdown("### 📊 Medidas de Tendência Central")
            stats_tendencia = pd.DataFrame({
                'Métrica': ['Média', 'Mediana', 'Moda'],
                'Valor (€)': [
                    f"{df_jogo['Valor'].mean():.0f}",
                    f"{df_jogo['Valor'].median():.0f}",
                    f"{df_jogo['Valor'].mode()[0]:.0f}" if len(df_jogo['Valor'].mode()) > 0 else "N/A"
                ]
            })
            st.dataframe(stats_tendencia, use_container_width=True, hide_index=True)

            st.markdown("### 📏 Medidas de Dispersão")
            stats_dispersao = pd.DataFrame({
                'Métrica': ['Desvio Padrão', 'Variância', 'Coef. Variação'],
                'Valor': [
                    f"€{df_jogo['Valor'].std():.0f}",
                    f"€{df_jogo['Valor'].var():.0f}",
                    f"{(df_jogo['Valor'].std() / df_jogo['Valor'].mean() * 100):.0f}%"
                ]
            })
            st.dataframe(stats_dispersao, use_container_width=True, hide_index=True)

        with col2:
            st.markdown("### 📈 Percentis")
            percentis = [10, 25, 50, 75, 90, 95, 99]
            stats_percentis = pd.DataFrame({
                'Percentil': [f"P{p}" for p in percentis],
                'Valor (€)': [f"{df_jogo['Valor'].quantile(p/100):.0f}" for p in percentis]
            })
            st.dataframe(stats_percentis, use_container_width=True, hide_index=True)

            st.markdown("### 🎯 Extremos")
            stats_extremos = pd.DataFrame({
                'Métrica': ['Mínimo', 'Máximo', 'Amplitude'],
                'Valor (€)': [
                    f"{df_jogo['Valor'].min():.0f}",
                    f"{df_jogo['Valor'].max():.0f}",
                    f"{df_jogo['Valor'].max() - df_jogo['Valor'].min():.0f}"
                ]
            })
            st.dataframe(stats_extremos, use_container_width=True, hide_index=True)

        # Histograma
        st.subheader("Distribuição de Vendas")
        fig = px.histogram(
            df_jogo,
            x='Valor',
            nbins=30,
            labels={'Valor': 'Vendas (€)', 'count': 'Frequência'},
            color_discrete_sequence=['#1f77b4']
        )
        fig.update_layout(height=400)
        st.plotly_chart(fig, use_container_width=True)

    # Exportação de dados
    st.markdown("---")
    st.subheader("📥 Exportar Dados do Jogo")

    col1, col2 = st.columns(2)

    # Preparar dados para exportação
    df_export = df_jogo[['Data_Emissao', 'Jogo', 'Valor', 'Ano', 'Mes', 'Semana_Ano']].copy()
    df_export['Data_Emissao'] = df_export['Data_Emissao'].dt.strftime('%d-%m-%Y')

    with col1:
        # Exportar CSV
        csv_data = df_export.to_csv(index=False).encode('utf-8')
        st.download_button(
            label="📄 Baixar Dados em CSV",
            data=csv_data,
            file_name=f"analise_{jogo_selecionado.replace(' ', '_')}.csv",
            mime="text/csv"
        )

    with col2:
        # Exportar Excel
        if EXCEL_AVAILABLE:
            try:
                # Preparar estatísticas
                stats_df = pd.DataFrame({
                    'Métrica': ['Total', 'Média', 'Mediana', 'Desvio Padrão', 'Máximo', 'Mínimo'],
                    'Valor (€)': [
                        df_jogo['Valor'].sum(),
                        df_jogo['Valor'].mean(),
                        df_jogo['Valor'].median(),
                        df_jogo['Valor'].std(),
                        df_jogo['Valor'].max(),
                        df_jogo['Valor'].min()
                    ]
                })

                dataframes = {
                    'Dados': df_export,
                    'Estatísticas': stats_df
                }
                info = {
                    'Página': 'Análise por Jogo',
                    'Jogo': jogo_selecionado,
                    'Total Vendas': f'€{df_jogo["Valor"].sum():,.0f}',
                    'Média': f'€{df_jogo["Valor"].mean():.0f}'
                }
                excel_data = exportar_excel_generico(dataframes, f'analise_{jogo_selecionado}', info)
                if excel_data:
                    st.download_button(
                        label="📊 Baixar Relatório em Excel",
                        data=excel_data,
                        file_name=f"analise_{jogo_selecionado.replace(' ', '_')}.xlsx",
                        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                    )
            except Exception as e:
                st.error(f"Erro ao gerar Excel: {str(e)}")
        else:
            st.info("📊 Excel: Instale openpyxl para exportar em Excel")


def pagina_comparacao(df):
    """Comparação entre jogos e anos."""
    st.markdown('<h1 class="main-header">⚖️ Comparação e Benchmarking</h1>', unsafe_allow_html=True)

    # Seleção de jogos para comparar
    jogos_disponiveis = sorted(df['Jogo'].unique())

    jogos_selecionados = st.sidebar.multiselect(
        "Selecionar Jogos para Comparar",
        jogos_disponiveis,
        default=jogos_disponiveis[:5] if len(jogos_disponiveis) >= 5 else jogos_disponiveis
    )

    if not jogos_selecionados:
        st.warning("Selecione pelo menos um jogo")
        return

    df_filtrado = df[df['Jogo'].isin(jogos_selecionados)]

    # Tabs
    tab1, tab2, tab3 = st.tabs(["📊 Comparação Geral", "📅 Comparação Temporal", "🔗 Correlações"])

    with tab1:
        st.subheader("Comparação de Métricas")

        # Criar tabela comparativa
        comparacao = df_filtrado.groupby('Jogo')['Valor'].agg([
            ('Total', 'sum'),
            ('Média', 'mean'),
            ('Mediana', 'median'),
            ('Desvio Padrão', 'std'),
            ('Máximo', 'max'),
            ('Mínimo', 'min')
        ]).round(2)

        st.dataframe(comparacao, use_container_width=True)

        # Gráfico de barras comparativo
        st.subheader("Comparação de Vendas Totais")
        totais = df_filtrado.groupby('Jogo')['Valor'].sum().sort_values(ascending=True)

        fig = px.bar(
            x=totais.values,
            y=totais.index,
            orientation='h',
            labels={'x': 'Total de Vendas (€)', 'y': 'Jogo'},
            color=totais.values,
            color_continuous_scale='Viridis'
        )
        fig.update_layout(showlegend=False, height=400)
        st.plotly_chart(fig, use_container_width=True)

    with tab2:
        st.subheader("Evolução Comparativa ao Longo do Tempo")

        # Vendas por semana para cada jogo
        vendas_temp = df_filtrado.groupby(['Ano_Semana', 'Jogo'])['Valor'].sum().reset_index()

        fig = px.line(
            vendas_temp,
            x='Ano_Semana',
            y='Valor',
            color='Jogo',
            markers=True,
            labels={'Valor': 'Vendas (€)', 'Ano_Semana': 'Semana'}
        )
        fig.update_layout(height=500)
        st.plotly_chart(fig, use_container_width=True)

        # Comparação ano a ano
        st.subheader("Crescimento Ano a Ano (%)")

        vendas_ano_jogo = df_filtrado.groupby(['Ano', 'Jogo'])['Valor'].sum().reset_index()
        pivot_ano = vendas_ano_jogo.pivot(index='Jogo', columns='Ano', values='Valor')

        # Calcular crescimento percentual
        anos = sorted(pivot_ano.columns)
        for i in range(1, len(anos)):
            ano_atual = anos[i]
            ano_anterior = anos[i-1]
            if ano_anterior in pivot_ano.columns and ano_atual in pivot_ano.columns:
                pivot_ano[f'Crescimento {ano_anterior}-{ano_atual} (%)'] = (
                    (pivot_ano[ano_atual] - pivot_ano[ano_anterior]) / pivot_ano[ano_anterior] * 100
                ).round(2)

        st.dataframe(pivot_ano, use_container_width=True)

    with tab3:
        st.subheader("Matriz de Correlação entre Jogos")

        # Criar pivot table para correlação
        vendas_pivot = df_filtrado.groupby(['Data_Emissao', 'Jogo'])['Valor'].sum().reset_index()
        vendas_wide = vendas_pivot.pivot(index='Data_Emissao', columns='Jogo', values='Valor')
        vendas_wide = vendas_wide.fillna(0)

        # Calcular correlação
        correlacao = vendas_wide.corr()

        fig = px.imshow(
            correlacao,
            labels=dict(color="Correlação"),
            x=correlacao.columns,
            y=correlacao.columns,
            color_continuous_scale='RdBu',
            zmin=-1,
            zmax=1
        )
        fig.update_layout(height=600)
        st.plotly_chart(fig, use_container_width=True)

        st.info("💡 Valores próximos de 1 indicam forte correlação positiva, valores próximos de -1 indicam forte correlação negativa.")

    # Exportação de dados
    st.markdown("---")
    st.subheader("📥 Exportar Dados de Comparação")

    col1, col2 = st.columns(2)

    # Preparar dados de comparação
    comparacao_export = df_filtrado.groupby('Jogo')['Valor'].agg([
        ('Total', 'sum'),
        ('Média', 'mean'),
        ('Mediana', 'median'),
        ('Desvio_Padrão', 'std'),
        ('Máximo', 'max'),
        ('Mínimo', 'min')
    ]).round(2).reset_index()

    with col1:
        # Exportar CSV
        csv_data = comparacao_export.to_csv(index=False).encode('utf-8')
        st.download_button(
            label="📄 Baixar Comparação em CSV",
            data=csv_data,
            file_name=f"comparacao_jogos.csv",
            mime="text/csv"
        )

    with col2:
        # Exportar Excel
        if EXCEL_AVAILABLE:
            try:
                # Evolução temporal para exportar
                vendas_temp_export = df_filtrado.groupby(['Ano_Semana', 'Jogo'])['Valor'].sum().reset_index()

                dataframes = {
                    'Comparação Geral': comparacao_export,
                    'Evolução Temporal': vendas_temp_export,
                    'Correlações': correlacao.reset_index()
                }
                info = {
                    'Página': 'Comparação',
                    'Jogos Selecionados': ', '.join(jogos_selecionados),
                    'Total Jogos': len(jogos_selecionados)
                }
                excel_data = exportar_excel_generico(dataframes, 'comparacao', info)
                if excel_data:
                    st.download_button(
                        label="📊 Baixar Relatório em Excel",
                        data=excel_data,
                        file_name=f"comparacao_jogos.xlsx",
                        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                    )
            except Exception as e:
                st.error(f"Erro ao gerar Excel: {str(e)}")
        else:
            st.info("📊 Excel: Instale openpyxl para exportar em Excel")


def pagina_subtotais(df):
    """Análise dos subtotais."""
    st.markdown('<h1 class="main-header">📦 Análise de Lotarias</h1>', unsafe_allow_html=True)

    # Filtrar apenas lotarias (antigos subtotais)
    df_subtotais = df[df['Jogo'].str.contains('Lotaria', na=False)]

    if df_subtotais.empty:
        st.warning("Nenhuma lotaria encontrada nos dados")
        return

    # Métricas gerais
    col1, col2, col3 = st.columns(3)

    for idx, subtotal in enumerate(['Lotaria Instantânea', 'Lotaria Popular', 'Lotaria Clássica']):
        df_sub = df_subtotais[df_subtotais['Jogo'] == subtotal]
        if not df_sub.empty:
            with [col1, col2, col3][idx]:
                st.metric(
                    subtotal,
                    f"€{df_sub['Valor'].sum():,.0f}",
                    f"Média: €{df_sub['Valor'].mean():.0f}"
                )

    st.markdown("---")

    # Tabs
    tab1, tab2, tab3 = st.tabs(["📈 Tendências", "📊 Comparação", "🔍 Detalhes"])

    with tab1:
        st.subheader("Evolução das Lotarias")

        vendas_sub = df_subtotais.groupby(['Data_Emissao', 'Jogo'])['Valor'].sum().reset_index()

        fig = px.line(
            vendas_sub,
            x='Data_Emissao',
            y='Valor',
            color='Jogo',
            markers=True,
            labels={'Valor': 'Valor (€)', 'Data_Emissao': 'Data'}
        )
        fig.update_layout(height=500)
        st.plotly_chart(fig, use_container_width=True)

    with tab2:
        st.subheader("Comparação entre Lotarias")

        col1, col2 = st.columns(2)

        with col1:
            # Total por lotaria
            totais_sub = df_subtotais.groupby('Jogo')['Valor'].sum().sort_values(ascending=True)

            fig = px.bar(
                x=totais_sub.values,
                y=totais_sub.index,
                orientation='h',
                labels={'x': 'Total (€)', 'y': 'Lotaria'},
                color=totais_sub.values,
                color_continuous_scale='Oranges'
            )
            fig.update_layout(showlegend=False, height=300)
            st.plotly_chart(fig, use_container_width=True)

        with col2:
            # Proporção
            fig = px.pie(
                values=totais_sub.values,
                names=totais_sub.index,
                hole=0.4
            )
            fig.update_layout(height=300)
            st.plotly_chart(fig, use_container_width=True)

        # Evolução ano a ano
        st.subheader("Crescimento Ano a Ano")
        vendas_ano = df_subtotais.groupby(['Ano', 'Jogo'])['Valor'].sum().reset_index()

        fig = px.bar(
            vendas_ano,
            x='Ano',
            y='Valor',
            color='Jogo',
            barmode='group',
            labels={'Valor': 'Total (€)'}
        )
        fig.update_layout(height=400)
        st.plotly_chart(fig, use_container_width=True)

    with tab3:
        st.subheader("Estatísticas Detalhadas")

        stats_subtotais = df_subtotais.groupby('Jogo')['Valor'].agg([
            ('Total', 'sum'),
            ('Média', 'mean'),
            ('Mediana', 'median'),
            ('Desvio Padrão', 'std'),
            ('Máximo', 'max'),
            ('Mínimo', 'min'),
            ('Contagem', 'count')
        ]).round(2)

        st.dataframe(stats_subtotais, use_container_width=True)


def exportar_excel_generico(dataframes_dict, nome_arquivo_base, info_adicional=None):
    """
    Exporta múltiplos DataFrames para Excel com formatação profissional.

    Args:
        dataframes_dict: Dict com {nome_sheet: dataframe}
        nome_arquivo_base: Nome base do arquivo
        info_adicional: Dict com informações adicionais (opcional)

    Returns:
        BytesIO com o arquivo Excel
    """
    if not EXCEL_AVAILABLE:
        return None

    output = BytesIO()

    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        # Escrever cada DataFrame em uma sheet
        for sheet_name, df in dataframes_dict.items():
            df.to_excel(writer, sheet_name=sheet_name[:31], index=False)  # Limite de 31 caracteres

            # Obter worksheet
            worksheet = writer.sheets[sheet_name[:31]]

            # Estilos
            header_fill = PatternFill(start_color="1F77B4", end_color="1F77B4", fill_type="solid")
            header_font = Font(bold=True, color="FFFFFF", size=12)
            center_alignment = Alignment(horizontal="center", vertical="center")

            # Formatação do cabeçalho
            for cell in worksheet[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = center_alignment

            # Ajustar largura das colunas
            for column in worksheet.columns:
                max_length = 0
                column_letter = column[0].column_letter
                for cell in column:
                    try:
                        if len(str(cell.value)) > max_length:
                            max_length = len(cell.value)
                    except:
                        pass
                adjusted_width = min(max_length + 2, 50)
                worksheet.column_dimensions[column_letter].width = adjusted_width

        # Adicionar sheet de informações se fornecida
        if info_adicional:
            info_sheet = writer.book.create_sheet('Informações', 0)
            info_sheet['A1'] = 'Relatório de Vendas - Jogos Santa Casa'
            info_sheet['A1'].font = Font(bold=True, size=14)

            linha = 3
            for chave, valor in info_adicional.items():
                info_sheet[f'A{linha}'] = f'{chave}: {valor}'
                linha += 1

            info_sheet[f'A{linha+1}'] = f'Data de Geração: {datetime.now().strftime("%d/%m/%Y %H:%M")}'

    output.seek(0)
    return output


def exportar_excel_estilizado(df_performance, periodo_str):
    """Exporta relatório para Excel com formatação profissional."""
    if not EXCEL_AVAILABLE:
        return None

    output = BytesIO()

    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        # Escrever dados de performance
        df_performance.to_excel(writer, sheet_name='Performance', index=False)

        # Obter workbook e worksheet
        workbook = writer.book
        worksheet = writer.sheets['Performance']

        # Estilos
        header_fill = PatternFill(start_color="1F77B4", end_color="1F77B4", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF", size=12)
        center_alignment = Alignment(horizontal="center", vertical="center")

        # Formatação do cabeçalho
        for cell in worksheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = center_alignment

        # Ajustar largura das colunas
        for column in worksheet.columns:
            max_length = 0
            column_letter = column[0].column_letter
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(cell.value)
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            worksheet.column_dimensions[column_letter].width = adjusted_width

        # Adicionar informações adicionais
        info_sheet = workbook.create_sheet('Informações')
        info_sheet['A1'] = 'Relatório de Performance - Jogos Santa Casa'
        info_sheet['A1'].font = Font(bold=True, size=14)
        info_sheet['A3'] = f'Período: {periodo_str}'
        info_sheet['A4'] = f'Data de Geração: {datetime.now().strftime("%d/%m/%Y %H:%M")}'
        info_sheet['A6'] = 'Legenda:'
        info_sheet['A7'] = '🟢 Verde: ≥100% do objetivo'
        info_sheet['A8'] = '🟡 Amarelo: 80-99% do objetivo'
        info_sheet['A9'] = '🔴 Vermelho: <80% do objetivo'

    output.seek(0)
    return output


def exportar_excel_com_insights_e_graficos(df, insights_list, data_referencia):
    """
    Exporta Excel avançado com múltiplas sheets, insights automáticos e gráficos.

    Args:
        df: DataFrame com todos os dados
        insights_list: Lista de insights gerados automaticamente
        data_referencia: Data da última semana

    Returns:
        BytesIO com arquivo Excel
    """
    if not EXCEL_AVAILABLE:
        return None

    try:
        from openpyxl.chart import BarChart, PieChart, LineChart, Reference
        from openpyxl.chart.label import DataLabelList

        output = BytesIO()

        with pd.ExcelWriter(output, engine='openpyxl') as writer:
            # Sheet 1: Insights Automáticos
            df_insights = pd.DataFrame({'Insights Automáticos': insights_list})
            df_insights.to_excel(writer, sheet_name='Insights', index=False)

            # Sheet 2: Ranking Dinâmico
            df_ranking = calcular_ranking_dinamico(df)
            if df_ranking is not None:
                df_ranking_export = df_ranking[['Posicao_Atual', 'Jogo', 'Vendas_Atual', 'Posicao_Anterior', 'Mudanca']].copy()
                df_ranking_export.columns = ['Posição', 'Jogo', 'Vendas (€)', 'Posição Anterior', 'Mudança']
                df_ranking_export.to_excel(writer, sheet_name='Ranking Dinâmico', index=False)

            # Sheet 3: WoW (Week over Week)
            vendas_wow = calcular_comparacao_wow(df)
            ultima_semana_wow = vendas_wow[vendas_wow['Data_Emissao'] == data_referencia].dropna(subset=['Crescimento_WoW_%'])
            if not ultima_semana_wow.empty:
                df_wow_export = ultima_semana_wow[['Jogo', 'Valor', 'Valor_Semana_Anterior', 'Crescimento_WoW_%', 'Diferenca_Absoluta']].copy()
                df_wow_export.columns = ['Jogo', 'Vendas Atual (€)', 'Vendas Anterior (€)', 'Crescimento (%)', 'Diferença (€)']
                df_wow_export.to_excel(writer, sheet_name='WoW', index=False, startrow=1)

            # Sheet 4: Contribuição Percentual
            df_contrib = calcular_percentagem_contribuicao(df, data_referencia)
            df_contrib.columns = ['Jogo', 'Vendas (€)', 'Contribuição (%)']
            df_contrib.to_excel(writer, sheet_name='Contribuição %', index=False, startrow=1)

            # Sheet 5: Velocímetro de Performance
            velocimetro, _ = calcular_velocimetro_performance(df)
            vel_data = []
            for jogo, info in velocimetro.items():
                vel_data.append({
                    'Jogo': jogo,
                    'Valor Atual (€)': info['valor_atual'],
                    'Média Histórica (€)': info['media_historica'],
                    'Performance (%)': info['percentual'],
                    'Status': info['status'],
                    'Diferença (€)': info['diferenca']
                })
            df_vel = pd.DataFrame(vel_data).sort_values('Performance (%)', ascending=False)
            df_vel.to_excel(writer, sheet_name='Velocímetro', index=False, startrow=1)

            # Obter workbook
            workbook = writer.book

            # Estilos gerais
            header_fill = PatternFill(start_color="1F77B4", end_color="1F77B4", fill_type="solid")
            header_font = Font(bold=True, color="FFFFFF", size=12)
            center_alignment = Alignment(horizontal="center", vertical="center")

            # Formatar todas as sheets
            for sheet_name in workbook.sheetnames:
                worksheet = workbook[sheet_name]

                # Formatar cabeçalhos
                for cell in worksheet[1]:
                    if cell.value:
                        cell.fill = header_fill
                        cell.font = header_font
                        cell.alignment = center_alignment

                # Ajustar largura das colunas
                for column in worksheet.columns:
                    max_length = 0
                    column_letter = column[0].column_letter
                    for cell in column:
                        try:
                            if cell.value and len(str(cell.value)) > max_length:
                                max_length = len(str(cell.value))
                        except:
                            pass
                    adjusted_width = min(max_length + 2, 50)
                    worksheet.column_dimensions[column_letter].width = adjusted_width

            # Adicionar gráficos
            try:
                # Gráfico na sheet WoW
                if 'WoW' in workbook.sheetnames:
                    ws_wow = workbook['WoW']
                    ws_wow['A1'] = f'Análise Week over Week - {data_referencia.strftime("%d/%m/%Y")}'
                    ws_wow['A1'].font = Font(bold=True, size=14)

                    # Gráfico de barras para crescimento WoW
                    if ws_wow.max_row > 2:
                        chart = BarChart()
                        chart.title = "Crescimento WoW (%)"
                        chart.x_axis.title = "Jogo"
                        chart.y_axis.title = "Crescimento (%)"

                        data = Reference(ws_wow, min_col=4, min_row=2, max_row=min(ws_wow.max_row, 12), max_col=4)
                        categories = Reference(ws_wow, min_col=1, min_row=3, max_row=min(ws_wow.max_row, 12))

                        chart.add_data(data, titles_from_data=True)
                        chart.set_categories(categories)
                        chart.height = 10
                        chart.width = 20

                        ws_wow.add_chart(chart, "G2")

                # Gráfico na sheet Contribuição
                if 'Contribuição %' in workbook.sheetnames:
                    ws_contrib = workbook['Contribuição %']
                    ws_contrib['A1'] = f'Contribuição Percentual - {data_referencia.strftime("%d/%m/%Y")}'
                    ws_contrib['A1'].font = Font(bold=True, size=14)

                    # Gráfico de pizza
                    if ws_contrib.max_row > 2:
                        chart = PieChart()
                        chart.title = "Distribuição de Vendas (%)"

                        data = Reference(ws_contrib, min_col=3, min_row=2, max_row=min(ws_contrib.max_row, 12), max_col=3)
                        categories = Reference(ws_contrib, min_col=1, min_row=3, max_row=min(ws_contrib.max_row, 12))

                        chart.add_data(data, titles_from_data=True)
                        chart.set_categories(categories)
                        chart.height = 12
                        chart.width = 18

                        # Adicionar labels
                        chart.dataLabels = DataLabelList()
                        chart.dataLabels.showPercent = True

                        ws_contrib.add_chart(chart, "F2")

                # Gráfico na sheet Velocímetro
                if 'Velocímetro' in workbook.sheetnames:
                    ws_vel = workbook['Velocímetro']
                    ws_vel['A1'] = 'Velocímetro de Performance vs Média Histórica'
                    ws_vel['A1'].font = Font(bold=True, size=14)

                    # Gráfico de barras comparativo
                    if ws_vel.max_row > 2:
                        chart = BarChart()
                        chart.title = "Atual vs Média Histórica"
                        chart.x_axis.title = "Jogo"
                        chart.y_axis.title = "Valor (€)"
                        chart.grouping = "clustered"

                        # Valor atual
                        data1 = Reference(ws_vel, min_col=2, min_row=2, max_row=min(ws_vel.max_row, 12), max_col=2)
                        # Média histórica
                        data2 = Reference(ws_vel, min_col=3, min_row=2, max_row=min(ws_vel.max_row, 12), max_col=3)
                        categories = Reference(ws_vel, min_col=1, min_row=3, max_row=min(ws_vel.max_row, 12))

                        chart.add_data(data1, titles_from_data=True)
                        chart.add_data(data2, titles_from_data=True)
                        chart.set_categories(categories)
                        chart.height = 12
                        chart.width = 20

                        ws_vel.add_chart(chart, "H2")

            except Exception as e:
                # Se houver erro nos gráficos, continuar sem eles
                print(f"Aviso: Não foi possível adicionar gráficos ao Excel: {e}")

            # Sheet de Informações (primeira sheet)
            info_sheet = workbook.create_sheet('📊 Resumo Executivo', 0)
            info_sheet['A1'] = 'RELATÓRIO DE ANÁLISE SEMANAL AVANÇADA'
            info_sheet['A1'].font = Font(bold=True, size=16, color="1F77B4")
            info_sheet.merge_cells('A1:D1')

            info_sheet['A3'] = 'Dashboard de Vendas - Jogos Santa Casa da Misericórdia'
            info_sheet['A3'].font = Font(bold=True, size=12)

            linha = 5
            info_sheet[f'A{linha}'] = 'Data do Relatório:'
            info_sheet[f'B{linha}'] = datetime.now().strftime("%d/%m/%Y %H:%M")
            linha += 1

            info_sheet[f'A{linha}'] = 'Semana Analisada:'
            info_sheet[f'B{linha}'] = data_referencia.strftime("%d/%m/%Y")
            linha += 1

            info_sheet[f'A{linha}'] = 'Versão do Dashboard:'
            info_sheet[f'B{linha}'] = '2.1'
            linha += 2

            info_sheet[f'A{linha}'] = 'CONTEÚDO DO RELATÓRIO:'
            info_sheet[f'A{linha}'].font = Font(bold=True, size=12)
            linha += 1

            conteudo = [
                '• Insights Automáticos - Análise inteligente das tendências',
                '• Ranking Dinâmico - Mudanças de posição entre jogos',
                '• WoW (Week over Week) - Crescimento semanal',
                '• Contribuição % - Peso de cada jogo no total',
                '• Velocímetro - Performance vs média histórica',
                '• Gráficos visuais incluídos nas sheets'
            ]

            for item in conteudo:
                info_sheet[f'A{linha}'] = item
                linha += 1

            linha += 2
            info_sheet[f'A{linha}'] = '💡 PRINCIPAIS INSIGHTS:'
            info_sheet[f'A{linha}'].font = Font(bold=True, size=11, color="FF6B35")
            linha += 1

            for i, insight in enumerate(insights_list[:5], 1):
                info_sheet[f'A{linha}'] = f"{i}. {insight}"
                linha += 1

        output.seek(0)
        return output

    except Exception as e:
        print(f"Erro ao criar Excel avançado: {e}")
        return None


def calcular_comparacao_wow(df):
    """Calcula comparação Week over Week."""
    df_semana = df.copy()

    # Agrupar por data de emissão e jogo
    vendas_semana = df_semana.groupby(['Data_Emissao', 'Jogo'])['Valor'].sum().reset_index()
    vendas_semana = vendas_semana.sort_values(['Jogo', 'Data_Emissao'])

    # Calcular crescimento WoW
    vendas_semana['Valor_Semana_Anterior'] = vendas_semana.groupby('Jogo')['Valor'].shift(1)

    # Proteção contra divisão por zero: só calcular quando Valor_Semana_Anterior > 0
    vendas_semana['Crescimento_WoW_%'] = vendas_semana.apply(
        lambda row: ((row['Valor'] - row['Valor_Semana_Anterior']) / row['Valor_Semana_Anterior'] * 100)
        if pd.notna(row['Valor_Semana_Anterior']) and row['Valor_Semana_Anterior'] != 0
        else np.nan,
        axis=1
    )
    vendas_semana['Diferenca_Absoluta'] = vendas_semana['Valor'] - vendas_semana['Valor_Semana_Anterior']

    return vendas_semana


def calcular_comparacao_mom(df):
    """Calcula comparação Month over Month."""
    df_mes = df.copy()
    df_mes['Mes_Ref'] = df_mes['Data_Emissao'].dt.to_period('M')

    vendas_mes = df_mes.groupby(['Mes_Ref', 'Jogo'])['Valor'].sum().reset_index()
    vendas_mes['Mes_Ref_Str'] = vendas_mes['Mes_Ref'].astype(str)

    # Calcular crescimento MoM
    vendas_mes = vendas_mes.sort_values(['Jogo', 'Mes_Ref'])
    vendas_mes['Valor_Mes_Anterior'] = vendas_mes.groupby('Jogo')['Valor'].shift(1)

    # Proteção contra divisão por zero
    vendas_mes['Crescimento_MoM_%'] = vendas_mes.apply(
        lambda row: ((row['Valor'] - row['Valor_Mes_Anterior']) / row['Valor_Mes_Anterior'] * 100)
        if pd.notna(row['Valor_Mes_Anterior']) and row['Valor_Mes_Anterior'] != 0
        else np.nan,
        axis=1
    )

    return vendas_mes


def calcular_comparacao_yoy(df):
    """Calcula comparação Year over Year."""
    df_ano = df.copy()

    vendas_ano = df_ano.groupby(['Ano', 'Jogo'])['Valor'].sum().reset_index()
    vendas_ano = vendas_ano.sort_values(['Jogo', 'Ano'])

    # Calcular crescimento YoY
    vendas_ano['Valor_Ano_Anterior'] = vendas_ano.groupby('Jogo')['Valor'].shift(1)

    # Proteção contra divisão por zero
    vendas_ano['Crescimento_YoY_%'] = vendas_ano.apply(
        lambda row: ((row['Valor'] - row['Valor_Ano_Anterior']) / row['Valor_Ano_Anterior'] * 100)
        if pd.notna(row['Valor_Ano_Anterior']) and row['Valor_Ano_Anterior'] != 0
        else np.nan,
        axis=1
    )

    return vendas_ano


def prever_atingimento_objetivo(df_jogo, objetivo_semanal, semanas_restantes=4):
    """Prevê probabilidade de atingir objetivo baseado em tendência."""
    if len(df_jogo) < 3 or objetivo_semanal == 0:
        return None

    vendas_temporal = df_jogo.groupby('Data_Emissao')['Valor'].sum().sort_index()
    media_atual = vendas_temporal.mean()

    # Tendência recente (últimas 4 semanas)
    ultimas_semanas = vendas_temporal.tail(4)
    if len(ultimas_semanas) > 1:
        tendencia = ultimas_semanas.pct_change().mean() * 100
    else:
        tendencia = 0

    # Projeção
    projecao = media_atual * (1 + tendencia/100) ** semanas_restantes

    # Probabilidade de atingir objetivo
    probabilidade = min((projecao / objetivo_semanal) * 100, 100)

    return {
        'media_atual': media_atual,
        'tendencia_%': tendencia,
        'projecao': projecao,
        'objetivo': objetivo_semanal,
        'probabilidade_%': probabilidade
    }


def calcular_ranking_dinamico(df):
    """Calcula ranking de jogos e mudanças em relação à semana anterior."""
    # Obter últimas duas semanas
    datas_unicas = sorted(df['Data_Emissao'].unique(), reverse=True)

    if len(datas_unicas) < 2:
        return None

    ultima_semana = datas_unicas[0]
    penultima_semana = datas_unicas[1]

    # Rankings
    rank_atual = df[df['Data_Emissao'] == ultima_semana].groupby('Jogo')['Valor'].sum().sort_values(ascending=False)
    rank_anterior = df[df['Data_Emissao'] == penultima_semana].groupby('Jogo')['Valor'].sum().sort_values(ascending=False)

    # Criar DataFrame com mudanças
    df_rank = pd.DataFrame({
        'Jogo': rank_atual.index,
        'Vendas_Atual': rank_atual.values,
        'Posicao_Atual': range(1, len(rank_atual) + 1)
    })

    # Adicionar posição anterior
    rank_anterior_dict = {jogo: pos + 1 for pos, jogo in enumerate(rank_anterior.index)}
    df_rank['Posicao_Anterior'] = df_rank['Jogo'].map(rank_anterior_dict)
    df_rank['Posicao_Anterior'] = df_rank['Posicao_Anterior'].fillna(len(df_rank) + 1)  # Novos jogos vão para última posição

    # Calcular mudança
    df_rank['Mudanca'] = df_rank['Posicao_Anterior'] - df_rank['Posicao_Atual']

    # Adicionar indicador visual
    def get_indicador(mudanca):
        if mudanca > 0:
            return f"⬆️ +{int(mudanca)}"
        elif mudanca < 0:
            return f"⬇️ {int(mudanca)}"
        else:
            return "➡️ 0"

    df_rank['Indicador'] = df_rank['Mudanca'].apply(get_indicador)
    df_rank['Data_Atual'] = ultima_semana
    df_rank['Data_Anterior'] = penultima_semana

    return df_rank


def calcular_percentagem_contribuicao(df, data_referencia=None):
    """Calcula percentagem de contribuição de cada jogo para o total."""
    if data_referencia is None:
        # Usar última data disponível
        data_referencia = df['Data_Emissao'].max()

    df_semana = df[df['Data_Emissao'] == data_referencia]
    vendas_por_jogo = df_semana.groupby('Jogo')['Valor'].sum()
    total = vendas_por_jogo.sum()

    percentagens = (vendas_por_jogo / total * 100).sort_values(ascending=False)

    df_contrib = pd.DataFrame({
        'Jogo': percentagens.index,
        'Valor': vendas_por_jogo.values,
        'Percentagem': percentagens.values
    })

    return df_contrib


def calcular_velocimetro_performance(df, janela_historica=8):
    """
    Calcula indicador de performance comparando última semana com média histórica.

    Returns:
        Dict com informações de performance para cada jogo
    """
    datas_unicas = sorted(df['Data_Emissao'].unique(), reverse=True)

    if len(datas_unicas) < janela_historica + 1:
        janela_historica = max(len(datas_unicas) - 1, 1)

    ultima_semana = datas_unicas[0]

    # Vendas da última semana
    vendas_ultima = df[df['Data_Emissao'] == ultima_semana].groupby('Jogo')['Valor'].sum()

    # Média histórica (excluindo última semana)
    datas_historicas = datas_unicas[1:janela_historica+1]
    df_historico = df[df['Data_Emissao'].isin(datas_historicas)]
    media_historica = df_historico.groupby('Jogo')['Valor'].mean()

    # Calcular performance
    resultados = {}
    for jogo in vendas_ultima.index:
        valor_atual = vendas_ultima[jogo]
        media_hist = media_historica.get(jogo, valor_atual)

        if media_hist > 0:
            percentual = (valor_atual / media_hist) * 100
        else:
            percentual = 100

        # Determinar status
        if percentual >= 110:
            status = "🟢 Excelente"
            cor = "green"
        elif percentual >= 95:
            status = "🟡 Normal"
            cor = "orange"
        else:
            status = "🔴 Abaixo"
            cor = "red"

        resultados[jogo] = {
            'valor_atual': valor_atual,
            'media_historica': media_hist,
            'percentual': percentual,
            'status': status,
            'cor': cor,
            'diferenca': valor_atual - media_hist
        }

    return resultados, ultima_semana


def gerar_insights_automaticos(df, janela_medias=[4, 8]):
    """Gera insights automáticos sobre as tendências de vendas."""
    insights = []

    datas_unicas = sorted(df['Data_Emissao'].unique(), reverse=True)
    if len(datas_unicas) < 2:
        return ["Dados insuficientes para gerar insights"]

    ultima_semana = datas_unicas[0]

    # Análise WoW
    vendas_wow = calcular_comparacao_wow(df)
    ultimas_wow = vendas_wow[vendas_wow['Data_Emissao'] == ultima_semana].dropna(subset=['Crescimento_WoW_%'])

    for _, row in ultimas_wow.iterrows():
        crescimento = row['Crescimento_WoW_%']
        if abs(crescimento) >= 5:  # Só reportar mudanças significativas
            direcao = "cresceu" if crescimento > 0 else "caiu"
            insights.append(f"📊 {row['Jogo']}: {direcao} {abs(crescimento):.1f}% em relação à semana anterior (€{row['Diferenca_Absoluta']:+,.0f})")

    # Análise de médias móveis
    for janela in janela_medias:
        if len(datas_unicas) >= janela:
            datas_janela = datas_unicas[:janela]
            df_janela = df[df['Data_Emissao'].isin(datas_janela)]

            for jogo in df['Jogo'].unique():
                df_jogo = df_janela[df_janela['Jogo'] == jogo]
                if len(df_jogo) >= janela:
                    media_janela = df_jogo['Valor'].mean()
                    valor_atual = df_jogo[df_jogo['Data_Emissao'] == ultima_semana]['Valor'].sum()

                    if valor_atual > 0 and media_janela > 0:
                        diff_percentual = ((valor_atual - media_janela) / media_janela) * 100

                        if abs(diff_percentual) >= 10:
                            direcao = "acima" if diff_percentual > 0 else "abaixo"
                            insights.append(f"📈 {jogo}: {abs(diff_percentual):.1f}% {direcao} da média das últimas {janela} semanas")

    # Velocímetro de performance
    velocimetro, _ = calcular_velocimetro_performance(df)
    for jogo, info in velocimetro.items():
        if info['percentual'] >= 120:
            insights.append(f"⭐ {jogo}: Performance excepcional! {info['percentual']:.0f}% da média histórica")
        elif info['percentual'] < 80:
            insights.append(f"⚠️ {jogo}: Performance abaixo do esperado ({info['percentual']:.0f}% da média histórica)")

    # Top performers
    vendas_ultima = df[df['Data_Emissao'] == ultima_semana].groupby('Jogo')['Valor'].sum().sort_values(ascending=False)
    if len(vendas_ultima) > 0:
        top_jogo = vendas_ultima.index[0]
        top_valor = vendas_ultima.values[0]
        insights.insert(0, f"🏆 Destaque da semana: {top_jogo} com €{top_valor:,.0f}")

    return insights if insights else ["Sem insights significativos para esta semana"]


def calcular_performance_objetivo(valor_real, objetivo):
    """Calcula a performance em relação ao objetivo."""
    if objetivo == 0:
        return 0, "⚪"

    percentual = (valor_real / objetivo) * 100

    # Definir semáforo
    if percentual >= 100:
        semaforo = "🟢"
    elif percentual >= 80:
        semaforo = "🟡"
    else:
        semaforo = "🔴"

    return percentual, semaforo


def pagina_dashboard_executivo(df):
    """Dashboard executivo com KPIs e análise de performance."""
    st.markdown('<h1 class="main-header">🎯 Dashboard Executivo</h1>', unsafe_allow_html=True)

    # Filtros na sidebar
    st.sidebar.header("🔍 Filtros")

    # Filtro de data personalizado
    col1, col2 = st.sidebar.columns(2)
    with col1:
        data_inicio = st.date_input(
            "Data Início",
            value=df['Data_Emissao'].min(),
            min_value=df['Data_Emissao'].min(),
            max_value=df['Data_Emissao'].max(),
            key="exec_data_inicio"
        )
    with col2:
        data_fim = st.date_input(
            "Data Fim",
            value=df['Data_Emissao'].max(),
            min_value=df['Data_Emissao'].min(),
            max_value=df['Data_Emissao'].max(),
            key="exec_data_fim"
        )

    # Filtrar dados por data
    df_filtrado = df[
        (df['Data_Emissao'] >= pd.to_datetime(data_inicio)) &
        (df['Data_Emissao'] <= pd.to_datetime(data_fim))
    ]

    if df_filtrado.empty:
        st.warning("Nenhum dado disponível para o período selecionado")
        return

    # Calcular número de semanas no período
    num_semanas = len(df_filtrado['Data_Emissao'].unique())

    # KPIs Principais com cards profissionais
    st.subheader("📊 KPIs Executivos")

    total_vendas = df_filtrado['Valor'].sum()

    # Calcular objetivo do período
    jogos_presentes = df_filtrado['Jogo'].unique()
    total_objetivo_periodo = sum([
        OBJETIVOS_SEMANAIS.get(jogo, 0) * num_semanas
        for jogo in jogos_presentes
    ])

    percentual_objetivo = (total_vendas / total_objetivo_periodo * 100) if total_objetivo_periodo > 0 else 0

    # Calcular remuneração
    df_rem = df_filtrado.copy()
    df_rem['Percentagem'] = df_rem['Jogo'].map(REMUNERACAO_JOGOS)
    df_rem['Remuneracao'] = df_rem['Valor'] * df_rem['Percentagem'] / 100
    total_remuneracao = df_rem['Remuneracao'].sum()

    # Calcular YoY e MoM comparisons
    datas_unicas = sorted(df_filtrado['Data_Emissao'].unique(), reverse=True)
    yoy_comparison = ""
    mom_comparison = ""

    if len(datas_unicas) >= 2:
        # Week over Week
        vendas_ultima = df_filtrado[df_filtrado['Data_Emissao'] == datas_unicas[0]]['Valor'].sum()
        vendas_anterior = df_filtrado[df_filtrado['Data_Emissao'] == datas_unicas[1]]['Valor'].sum()
        if vendas_anterior > 0:
            wow_perc = ((vendas_ultima - vendas_anterior) / vendas_anterior) * 100
            mom_comparison = f"WoW: {wow_perc:+.1f}%"

    # Semáforo geral
    if percentual_objetivo >= 100:
        status_text = "Excelente"
        status_icon = "🟢"
        card_class = "kpi-card-green"
    elif percentual_objetivo >= 90:
        status_text = "Bom"
        status_icon = "🟡"
        card_class = "kpi-card-yellow"
    else:
        status_text = "Atenção"
        status_icon = "🔴"
        card_class = "kpi-card-orange"

    col1, col2, col3, col4 = st.columns(4)

    with col1:
        st.markdown(f"""
        <div class="kpi-card kpi-card-blue">
            <div class="kpi-icon">💰</div>
            <div class="kpi-label">TOTAL DE VENDAS</div>
            <div class="kpi-value">€{total_vendas:,.0f}</div>
            <div class="kpi-delta">{percentual_objetivo:.1f}% do objetivo</div>
        </div>
        """, unsafe_allow_html=True)

    with col2:
        st.markdown(f"""
        <div class="kpi-card kpi-card-green">
            <div class="kpi-icon">💵</div>
            <div class="kpi-label">REMUNERAÇÃO</div>
            <div class="kpi-value">€{total_remuneracao:,.0f}</div>
            <div class="kpi-delta">€{total_remuneracao/num_semanas:,.0f} / semana</div>
        </div>
        """, unsafe_allow_html=True)

    with col3:
        st.markdown(f"""
        <div class="kpi-card kpi-card-orange">
            <div class="kpi-icon">🎯</div>
            <div class="kpi-label">OBJETIVO PERÍODO</div>
            <div class="kpi-value">€{total_objetivo_periodo:,.0f}</div>
            <div class="kpi-delta">{num_semanas} semana(s)</div>
        </div>
        """, unsafe_allow_html=True)

    with col4:
        st.markdown(f"""
        <div class="kpi-card {card_class}">
            <div class="kpi-icon">{status_icon}</div>
            <div class="kpi-label">STATUS GERAL</div>
            <div class="kpi-value" style="font-size: 2.2rem;">{status_text}</div>
            <div class="kpi-delta">{percentual_objetivo:.1f}%</div>
        </div>
        """, unsafe_allow_html=True)

    st.markdown("---")

    # Top 5 Jogos em destaque
    st.subheader("🏆 Top 5 Jogos em Destaque")

    top_5_jogos = df_filtrado.groupby('Jogo')['Valor'].sum().sort_values(ascending=False).head(5).reset_index()

    col1, col2 = st.columns([1, 2])

    with col1:
        for idx, row in top_5_jogos.iterrows():
            medal = ["🥇", "🥈", "🥉", "4️⃣", "5️⃣"][idx]
            st.markdown(f"""
            <div class="top-game-card">
                <span class="top-game-rank">{medal}</span>
                <div class="top-game-name">{row['Jogo']}</div>
                <div class="top-game-value">€{row['Valor']:,.0f}</div>
            </div>
            """, unsafe_allow_html=True)

    with col2:
        # Gráfico de barras com gradiente para Top 5
        colors = ['#11998e', '#38ef7d', '#667eea', '#764ba2', '#f093fb']

        fig = go.Figure(data=[
            go.Bar(
                x=top_5_jogos['Valor'],
                y=top_5_jogos['Jogo'],
                orientation='h',
                marker=dict(
                    color=top_5_jogos['Valor'],
                    colorscale=[[0, '#f093fb'], [0.5, '#667eea'], [1, '#11998e']],
                    showscale=False,
                    line=dict(color='rgba(0,0,0,0.2)', width=1)
                ),
                text=[f'€{v:,.0f}' for v in top_5_jogos['Valor']],
                textposition='outside',
                textfont=dict(size=14, color='#333', weight='bold')
            )
        ])

        fig.update_layout(
            title=dict(text="Vendas por Jogo", font=dict(size=16, weight='bold')),
            xaxis_title="Vendas (€)",
            yaxis_title="",
            height=350,
            showlegend=False,
            plot_bgcolor='rgba(0,0,0,0)',
            paper_bgcolor='rgba(0,0,0,0)',
            margin=dict(l=10, r=10, t=40, b=40)
        )

        st.plotly_chart(fig, use_container_width=True)

    st.markdown("---")

    # Análise por Jogo com Objetivos
    st.subheader("🎮 Performance por Jogo vs Objetivo Semanal")

    # Agrupar dados por jogo
    vendas_por_jogo = df_filtrado.groupby('Jogo')['Valor'].sum().reset_index()

    # Calcular média semanal por jogo
    vendas_por_jogo['Media_Semanal'] = vendas_por_jogo['Valor'] / num_semanas

    # Adicionar objetivos
    vendas_por_jogo['Objetivo'] = vendas_por_jogo['Jogo'].map(OBJETIVOS_SEMANAIS)
    vendas_por_jogo['Objetivo'] = vendas_por_jogo['Objetivo'].fillna(0)

    # Calcular performance
    vendas_por_jogo['Performance_%'] = vendas_por_jogo.apply(
        lambda row: calcular_performance_objetivo(row['Media_Semanal'], row['Objetivo'])[0], axis=1
    )
    vendas_por_jogo['Status'] = vendas_por_jogo.apply(
        lambda row: calcular_performance_objetivo(row['Media_Semanal'], row['Objetivo'])[1], axis=1
    )

    # Calcular diferença
    vendas_por_jogo['Diferenca'] = vendas_por_jogo['Media_Semanal'] - vendas_por_jogo['Objetivo']

    # Ordenar por performance
    vendas_por_jogo = vendas_por_jogo.sort_values('Performance_%', ascending=False)

    # Sparklines - mini-gráficos de tendência
    st.subheader("📈 Tendências Recentes (Últimas 4 Semanas)")

    ultimas_4_semanas = sorted(df_filtrado['Data_Emissao'].unique(), reverse=True)[:4]

    if len(ultimas_4_semanas) >= 2:
        top_3_jogos_sparkline = vendas_por_jogo.head(3)['Jogo'].tolist()

        cols_spark = st.columns(3)

        for idx, jogo in enumerate(top_3_jogos_sparkline):
            with cols_spark[idx]:
                df_jogo_temp = df_filtrado[df_filtrado['Jogo'] == jogo]
                vendas_temp = []

                for data in reversed(ultimas_4_semanas):
                    valor = df_jogo_temp[df_jogo_temp['Data_Emissao'] == data]['Valor'].sum()
                    vendas_temp.append(valor)

                fig_spark = go.Figure()
                fig_spark.add_trace(go.Scatter(
                    y=vendas_temp,
                    mode='lines+markers',
                    line=dict(color='#667eea', width=3),
                    marker=dict(size=8, color='#11998e'),
                    fill='tozeroy',
                    fillcolor='rgba(102, 126, 234, 0.2)'
                ))

                fig_spark.update_layout(
                    title=dict(text=jogo, font=dict(size=12)),
                    height=150,
                    showlegend=False,
                    xaxis=dict(showticklabels=False, showgrid=False),
                    yaxis=dict(showticklabels=False, showgrid=False),
                    margin=dict(l=10, r=10, t=30, b=10),
                    plot_bgcolor='rgba(0,0,0,0)',
                    paper_bgcolor='rgba(0,0,0,0)'
                )

                st.plotly_chart(fig_spark, use_container_width=True)

    st.markdown("---")

    # Visualização em tabela
    col1, col2 = st.columns([2, 1])

    with col1:
        # Gráfico de barras comparativo com gradiente
        fig = go.Figure()

        fig.add_trace(go.Bar(
            name='Média Semanal Real',
            x=vendas_por_jogo['Jogo'],
            y=vendas_por_jogo['Media_Semanal'],
            marker=dict(
                color=vendas_por_jogo['Media_Semanal'],
                colorscale='Blues',
                showscale=False
            ),
            text=[f'€{v:.0f}' for v in vendas_por_jogo['Media_Semanal']],
            textposition='outside'
        ))

        fig.add_trace(go.Bar(
            name='Objetivo Semanal',
            x=vendas_por_jogo['Jogo'],
            y=vendas_por_jogo['Objetivo'],
            marker_color='rgba(255, 127, 14, 0.6)',
            text=[f'€{v:.0f}' for v in vendas_por_jogo['Objetivo']],
            textposition='outside'
        ))

        fig.update_layout(
            title=dict(text="Comparação: Real vs Objetivo", font=dict(size=16)),
            xaxis_title="Jogo",
            yaxis_title="Valor (€)",
            barmode='group',
            height=500,
            showlegend=True,
            plot_bgcolor='rgba(0,0,0,0.02)'
        )

        st.plotly_chart(fig, use_container_width=True)

    with col2:
        # Tabela de performance
        tabela_performance = vendas_por_jogo[['Status', 'Jogo', 'Performance_%']].copy()
        tabela_performance['Performance_%'] = tabela_performance['Performance_%'].apply(lambda x: f"{x:.1f}%")
        tabela_performance.columns = ['📊', 'Jogo', 'Performance']

        st.dataframe(
            tabela_performance,
            use_container_width=True,
            hide_index=True,
            height=500
        )

    # Tabela detalhada
    st.subheader("📋 Tabela Detalhada de Performance")

    tabela_detalhada = vendas_por_jogo[['Status', 'Jogo', 'Media_Semanal', 'Objetivo', 'Diferenca', 'Performance_%']].copy()
    tabela_detalhada['Media_Semanal'] = tabela_detalhada['Media_Semanal'].apply(lambda x: f"€{x:.0f}")
    tabela_detalhada['Objetivo'] = tabela_detalhada['Objetivo'].apply(lambda x: f"€{x:.0f}")
    tabela_detalhada['Diferenca'] = tabela_detalhada['Diferenca'].apply(lambda x: f"€{x:+.0f}")
    tabela_detalhada['Performance_%'] = tabela_detalhada['Performance_%'].apply(lambda x: f"{x:.1f}%")
    tabela_detalhada.columns = ['📊', 'Jogo', 'Média Semanal', 'Objetivo', 'Diferença', 'Performance']

    st.dataframe(tabela_detalhada, use_container_width=True, hide_index=True)

    st.markdown("---")

    # Análise de Tendências com Média Móvel
    st.subheader("📈 Análise de Tendências (Média Móvel)")

    # Seletor de jogos (apenas jogos com objetivo > 0)
    jogos_disponiveis = sorted([
        j for j in df_filtrado['Jogo'].unique()
        if j in OBJETIVOS_SEMANAIS and OBJETIVOS_SEMANAIS[j] > 0
    ])

    if not jogos_disponiveis:
        st.warning("Nenhum jogo com objetivo definido encontrado no período selecionado")
        return

    jogo_analise = st.selectbox("Selecionar Jogo para Análise", jogos_disponiveis, key="exec_jogo")

    col1, col2 = st.columns(2)

    with col1:
        janela_mm = st.slider("Janela de Média Móvel (semanas)", 2, 8, 4, key="exec_mm")

    with col2:
        mostrar_objetivo = st.checkbox("Mostrar Linha de Objetivo", value=True, key="exec_obj")

    # Dados do jogo selecionado
    df_jogo = df_filtrado[df_filtrado['Jogo'] == jogo_analise].copy()
    vendas_temporal = df_jogo.groupby('Data_Emissao')['Valor'].sum().sort_index()

    # Calcular média móvel
    media_movel = vendas_temporal.rolling(window=janela_mm, min_periods=1).mean()

    # Criar gráfico
    fig = go.Figure()

    # Vendas reais
    fig.add_trace(go.Scatter(
        x=vendas_temporal.index,
        y=vendas_temporal.values,
        mode='lines+markers',
        name='Vendas Semanais',
        line=dict(color='lightblue', width=2),
        marker=dict(size=6)
    ))

    # Média móvel
    fig.add_trace(go.Scatter(
        x=media_movel.index,
        y=media_movel.values,
        mode='lines',
        name=f'Média Móvel ({janela_mm} semanas)',
        line=dict(color='blue', width=3)
    ))

    # Linha de objetivo
    if mostrar_objetivo and jogo_analise in OBJETIVOS_SEMANAIS:
        objetivo_valor = OBJETIVOS_SEMANAIS[jogo_analise]
        fig.add_trace(go.Scatter(
            x=vendas_temporal.index,
            y=[objetivo_valor] * len(vendas_temporal),
            mode='lines',
            name='Objetivo Semanal',
            line=dict(color='red', width=2, dash='dash')
        ))

    fig.update_layout(
        title=f"Tendência de Vendas - {jogo_analise}",
        xaxis_title="Data",
        yaxis_title="Vendas (€)",
        height=500,
        hovermode='x unified'
    )

    st.plotly_chart(fig, use_container_width=True)

    # Estatísticas da tendência
    col1, col2, col3, col4 = st.columns(4)

    with col1:
        st.metric("📊 Média do Período", f"€{vendas_temporal.mean():.0f}")

    with col2:
        if jogo_analise in OBJETIVOS_SEMANAIS:
            objetivo = OBJETIVOS_SEMANAIS[jogo_analise]
            perf = (vendas_temporal.mean() / objetivo * 100) if objetivo > 0 else 0
            st.metric("🎯 Performance Média", f"{perf:.1f}%")
        else:
            st.metric("🎯 Performance Média", "N/A")

    with col3:
        # Tendência (comparar primeira metade vs segunda metade)
        meio = len(vendas_temporal) // 2
        if meio > 0:
            primeira_metade = vendas_temporal.iloc[:meio].mean()
            segunda_metade = vendas_temporal.iloc[meio:].mean()
            crescimento = ((segunda_metade - primeira_metade) / primeira_metade * 100) if primeira_metade > 0 else 0
            st.metric("📈 Tendência", f"{crescimento:+.1f}%",
                     delta=f"€{segunda_metade - primeira_metade:+.0f}")
        else:
            st.metric("📈 Tendência", "N/A")

    with col4:
        volatilidade = (vendas_temporal.std() / vendas_temporal.mean() * 100) if vendas_temporal.mean() > 0 else 0
        st.metric("📊 Volatilidade", f"{volatilidade:.1f}%")

    st.markdown("---")

    # Exportação de Relatório
    st.subheader("📥 Exportar Relatório")

    col1, col2, col3 = st.columns(3)

    periodo_str = f"{data_inicio.strftime('%d/%m/%Y')} - {data_fim.strftime('%d/%m/%Y')}"

    with col1:
        # Botão para CSV
        csv_data = vendas_por_jogo.to_csv(index=False).encode('utf-8')
        st.download_button(
            label="📄 Baixar CSV",
            data=csv_data,
            file_name=f"relatorio_performance_{data_inicio}_{data_fim}.csv",
            mime="text/csv"
        )

    with col2:
        # Exportação Excel
        if EXCEL_AVAILABLE:
            try:
                excel_data = exportar_excel_estilizado(vendas_por_jogo, periodo_str)
                if excel_data:
                    st.download_button(
                        label="📊 Baixar Excel",
                        data=excel_data,
                        file_name=f"relatorio_performance_{data_inicio}_{data_fim}.xlsx",
                        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                    )
            except Exception as e:
                st.error(f"Erro ao gerar Excel: {str(e)}")
        else:
            if st.button("📊 Info Excel"):
                st.info("OpenPyXL já instalado! Se vir este botão, recarregue a página.")

    with col3:
        # Status da exportação
        if EXCEL_AVAILABLE:
            st.success("✅ Excel: Disponível")
        else:
            st.info("📑 Excel: Instalar openpyxl")


def pagina_comparacoes_avancadas(df):
    """Página com comparações MoM, YoY e previsão de objetivos."""
    st.markdown('<h1 class="main-header">📊 Comparações Avançadas</h1>', unsafe_allow_html=True)

    # Tabs para diferentes análises
    tab1, tab2, tab3 = st.tabs(["📈 MoM (Mês a Mês)", "📅 YoY (Ano a Ano)", "🔮 Previsão de Objetivos"])

    with tab1:
        st.subheader("Análise Month over Month (MoM)")
        st.caption("💡 Comparação do crescimento em relação ao mês anterior")

        # Calcular MoM
        vendas_mom = calcular_comparacao_mom(df)

        # Filtro de jogo
        jogos_disponiveis = sorted(vendas_mom['Jogo'].unique())
        jogos_selecionados = st.multiselect(
            "Selecionar Jogos",
            jogos_disponiveis,
            default=jogos_disponiveis[:5] if len(jogos_disponiveis) >= 5 else jogos_disponiveis,
            key="mom_jogos"
        )

        if jogos_selecionados:
            vendas_mom_filtrado = vendas_mom[vendas_mom['Jogo'].isin(jogos_selecionados)]

            # Gráfico de linha MoM
            col1, col2 = st.columns(2)

            with col1:
                # Vendas por mês
                fig = px.line(
                    vendas_mom_filtrado,
                    x='Mes_Ref_Str',
                    y='Valor',
                    color='Jogo',
                    markers=True,
                    title="Vendas Mensais",
                    labels={'Mes_Ref_Str': 'Mês', 'Valor': 'Vendas (€)'}
                )
                fig.update_layout(height=400)
                st.plotly_chart(fig, use_container_width=True)

            with col2:
                # Crescimento MoM
                vendas_mom_sem_na = vendas_mom_filtrado.dropna(subset=['Crescimento_MoM_%'])

                fig = px.bar(
                    vendas_mom_sem_na,
                    x='Mes_Ref_Str',
                    y='Crescimento_MoM_%',
                    color='Jogo',
                    title="Crescimento MoM (%)",
                    labels={'Mes_Ref_Str': 'Mês', 'Crescimento_MoM_%': 'Crescimento (%)'},
                    barmode='group'
                )
                fig.add_hline(y=0, line_dash="dash", line_color="red", annotation_text="Zero")
                fig.update_layout(height=400)
                st.plotly_chart(fig, use_container_width=True)

            # Tabela detalhada
            st.subheader("📋 Tabela Detalhada MoM")
            tabela_mom = vendas_mom_filtrado[['Jogo', 'Mes_Ref_Str', 'Valor', 'Valor_Mes_Anterior', 'Crescimento_MoM_%']].copy()
            tabela_mom['Valor'] = tabela_mom['Valor'].apply(lambda x: f"€{x:,.0f}")
            tabela_mom['Valor_Mes_Anterior'] = tabela_mom['Valor_Mes_Anterior'].apply(
                lambda x: f"€{x:,.0f}" if pd.notna(x) else "N/A"
            )
            tabela_mom['Crescimento_MoM_%'] = tabela_mom['Crescimento_MoM_%'].apply(
                lambda x: f"{x:+.1f}%" if pd.notna(x) else "N/A"
            )
            tabela_mom.columns = ['Jogo', 'Mês', 'Vendas', 'Mês Anterior', 'Crescimento MoM']
            st.dataframe(tabela_mom, use_container_width=True, hide_index=True)

            # Estatísticas MoM
            st.subheader("📊 Estatísticas de Crescimento MoM")
            col1, col2, col3, col4 = st.columns(4)

            crescimento_medio = vendas_mom_filtrado['Crescimento_MoM_%'].mean()
            crescimento_max = vendas_mom_filtrado['Crescimento_MoM_%'].max()
            crescimento_min = vendas_mom_filtrado['Crescimento_MoM_%'].min()
            meses_positivos = (vendas_mom_filtrado['Crescimento_MoM_%'] > 0).sum()

            with col1:
                st.metric("📈 Crescimento Médio", f"{crescimento_medio:.1f}%")
            with col2:
                st.metric("⬆️ Máximo", f"{crescimento_max:.1f}%")
            with col3:
                st.metric("⬇️ Mínimo", f"{crescimento_min:.1f}%")
            with col4:
                st.metric("✅ Meses Positivos", f"{meses_positivos}")

            # Exportação de dados MoM
            st.markdown("---")
            st.subheader("📥 Exportar Dados MoM")

            col1, col2 = st.columns(2)

            with col1:
                # Exportar CSV
                csv_data = tabela_mom.to_csv(index=False).encode('utf-8')
                st.download_button(
                    label="📄 Baixar MoM em CSV",
                    data=csv_data,
                    file_name=f"comparacao_mom.csv",
                    mime="text/csv",
                    key="csv_mom"
                )

            with col2:
                # Exportar Excel
                if EXCEL_AVAILABLE:
                    try:
                        dataframes = {
                            'Dados MoM': vendas_mom_filtrado[['Jogo', 'Mes_Ref_Str', 'Valor', 'Valor_Mes_Anterior', 'Crescimento_MoM_%']]
                        }
                        info = {
                            'Página': 'Comparações Avançadas - MoM',
                            'Jogos': ', '.join(jogos_selecionados),
                            'Crescimento Médio': f'{crescimento_medio:.1f}%'
                        }
                        excel_data = exportar_excel_generico(dataframes, 'mom', info)
                        if excel_data:
                            st.download_button(
                                label="📊 Baixar MoM em Excel",
                                data=excel_data,
                                file_name=f"comparacao_mom.xlsx",
                                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                                key="excel_mom"
                            )
                    except Exception as e:
                        st.error(f"Erro ao gerar Excel: {str(e)}")
                else:
                    st.info("📊 Excel: Instale openpyxl para exportar em Excel")

    with tab2:
        st.subheader("Análise Year over Year (YoY)")
        st.caption("💡 Comparação do crescimento em relação ao ano anterior")

        # Calcular YoY
        vendas_yoy = calcular_comparacao_yoy(df)

        # Filtro de jogo
        jogos_disponiveis_yoy = sorted(vendas_yoy['Jogo'].unique())
        jogos_selecionados_yoy = st.multiselect(
            "Selecionar Jogos",
            jogos_disponiveis_yoy,
            default=jogos_disponiveis_yoy[:5] if len(jogos_disponiveis_yoy) >= 5 else jogos_disponiveis_yoy,
            key="yoy_jogos"
        )

        if jogos_selecionados_yoy:
            vendas_yoy_filtrado = vendas_yoy[vendas_yoy['Jogo'].isin(jogos_selecionados_yoy)]

            # Gráficos
            col1, col2 = st.columns(2)

            with col1:
                # Vendas por ano
                fig = px.bar(
                    vendas_yoy_filtrado,
                    x='Ano',
                    y='Valor',
                    color='Jogo',
                    title="Vendas Anuais",
                    labels={'Ano': 'Ano', 'Valor': 'Vendas (€)'},
                    barmode='group'
                )
                fig.update_layout(height=400)
                st.plotly_chart(fig, use_container_width=True)

            with col2:
                # Crescimento YoY
                vendas_yoy_sem_na = vendas_yoy_filtrado.dropna(subset=['Crescimento_YoY_%'])

                fig = px.bar(
                    vendas_yoy_sem_na,
                    x='Ano',
                    y='Crescimento_YoY_%',
                    color='Jogo',
                    title="Crescimento YoY (%)",
                    labels={'Ano': 'Ano', 'Crescimento_YoY_%': 'Crescimento (%)'},
                    barmode='group'
                )
                fig.add_hline(y=0, line_dash="dash", line_color="red", annotation_text="Zero")
                fig.update_layout(height=400)
                st.plotly_chart(fig, use_container_width=True)

            # Tabela detalhada
            st.subheader("📋 Tabela Detalhada YoY")
            tabela_yoy = vendas_yoy_filtrado[['Jogo', 'Ano', 'Valor', 'Valor_Ano_Anterior', 'Crescimento_YoY_%']].copy()
            tabela_yoy['Valor'] = tabela_yoy['Valor'].apply(lambda x: f"€{x:,.0f}")
            tabela_yoy['Valor_Ano_Anterior'] = tabela_yoy['Valor_Ano_Anterior'].apply(
                lambda x: f"€{x:,.0f}" if pd.notna(x) else "N/A"
            )
            tabela_yoy['Crescimento_YoY_%'] = tabela_yoy['Crescimento_YoY_%'].apply(
                lambda x: f"{x:+.1f}%" if pd.notna(x) else "N/A"
            )
            tabela_yoy.columns = ['Jogo', 'Ano', 'Vendas', 'Ano Anterior', 'Crescimento YoY']
            st.dataframe(tabela_yoy, use_container_width=True, hide_index=True)

            # Estatísticas YoY
            st.subheader("📊 Estatísticas de Crescimento YoY")
            col1, col2, col3, col4 = st.columns(4)

            crescimento_medio_yoy = vendas_yoy_filtrado['Crescimento_YoY_%'].mean()
            crescimento_max_yoy = vendas_yoy_filtrado['Crescimento_YoY_%'].max()
            crescimento_min_yoy = vendas_yoy_filtrado['Crescimento_YoY_%'].min()
            anos_positivos = (vendas_yoy_filtrado['Crescimento_YoY_%'] > 0).sum()

            with col1:
                st.metric("📈 Crescimento Médio", f"{crescimento_medio_yoy:.1f}%")
            with col2:
                st.metric("⬆️ Máximo", f"{crescimento_max_yoy:.1f}%")
            with col3:
                st.metric("⬇️ Mínimo", f"{crescimento_min_yoy:.1f}%")
            with col4:
                st.metric("✅ Anos Positivos", f"{anos_positivos}")

            # Exportação de dados YoY
            st.markdown("---")
            st.subheader("📥 Exportar Dados YoY")

            col1, col2 = st.columns(2)

            with col1:
                # Exportar CSV
                csv_data = tabela_yoy.to_csv(index=False).encode('utf-8')
                st.download_button(
                    label="📄 Baixar YoY em CSV",
                    data=csv_data,
                    file_name=f"comparacao_yoy.csv",
                    mime="text/csv",
                    key="csv_yoy"
                )

            with col2:
                # Exportar Excel
                if EXCEL_AVAILABLE:
                    try:
                        dataframes = {
                            'Dados YoY': vendas_yoy_filtrado[['Jogo', 'Ano', 'Valor', 'Valor_Ano_Anterior', 'Crescimento_YoY_%']]
                        }
                        info = {
                            'Página': 'Comparações Avançadas - YoY',
                            'Jogos': ', '.join(jogos_selecionados_yoy),
                            'Crescimento Médio': f'{crescimento_medio_yoy:.1f}%'
                        }
                        excel_data = exportar_excel_generico(dataframes, 'yoy', info)
                        if excel_data:
                            st.download_button(
                                label="📊 Baixar YoY em Excel",
                                data=excel_data,
                                file_name=f"comparacao_yoy.xlsx",
                                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                                key="excel_yoy"
                            )
                    except Exception as e:
                        st.error(f"Erro ao gerar Excel: {str(e)}")
                else:
                    st.info("📊 Excel: Instale openpyxl para exportar em Excel")

    with tab3:
        st.subheader("🔮 Previsão de Atingimento de Objetivos")
        st.caption("💡 Probabilidade de atingir objetivos semanais baseado em tendências")

        # Seletor de jogo
        jogos_com_objetivo = [j for j in df['Jogo'].unique() if j in OBJETIVOS_SEMANAIS and OBJETIVOS_SEMANAIS[j] > 0]
        jogo_previsao = st.selectbox("Selecionar Jogo", sorted(jogos_com_objetivo), key="prev_jogo")

        col1, col2 = st.columns([1, 2])

        with col1:
            semanas_prever = st.slider("Semanas a Prever", 1, 12, 4, key="prev_semanas")

        # Calcular previsão
        df_jogo_prev = df[df['Jogo'] == jogo_previsao]
        objetivo = OBJETIVOS_SEMANAIS.get(jogo_previsao, 0)

        previsao = prever_atingimento_objetivo(df_jogo_prev, objetivo, semanas_prever)

        if previsao:
            st.markdown("---")

            # KPIs da previsão
            col1, col2, col3, col4 = st.columns(4)

            with col1:
                st.metric("📊 Média Atual", f"€{previsao['media_atual']:.0f}")

            with col2:
                delta_color = "normal" if previsao['tendencia_%'] >= 0 else "inverse"
                st.metric("📈 Tendência", f"{previsao['tendencia_%']:+.1f}%", delta_color=delta_color)

            with col3:
                st.metric("🎯 Objetivo", f"€{previsao['objetivo']:.0f}")

            with col4:
                prob = previsao['probabilidade_%']
                if prob >= 80:
                    emoji = "🟢"
                    status = "Alta"
                elif prob >= 50:
                    emoji = "🟡"
                    status = "Média"
                else:
                    emoji = "🔴"
                    status = "Baixa"

                st.metric(f"{emoji} Probabilidade", f"{prob:.0f}%", status)

            # Gráfico de previsão
            st.subheader("📈 Projeção vs Objetivo")

            vendas_temporal = df_jogo_prev.groupby('Data_Emissao')['Valor'].sum().sort_index()

            # Criar datas futuras
            ultima_data = vendas_temporal.index[-1]
            datas_futuras = pd.date_range(start=ultima_data + timedelta(days=7), periods=semanas_prever, freq='W')

            # Calcular projeções com base na tendência
            projecoes = []
            for i in range(semanas_prever):
                proj = previsao['media_atual'] * (1 + previsao['tendencia_%']/100) ** (i + 1)
                projecoes.append(proj)

            # Criar gráfico
            fig = go.Figure()

            # Histórico
            fig.add_trace(go.Scatter(
                x=vendas_temporal.index,
                y=vendas_temporal.values,
                mode='lines+markers',
                name='Histórico',
                line=dict(color='blue', width=2)
            ))

            # Projeção
            fig.add_trace(go.Scatter(
                x=datas_futuras,
                y=projecoes,
                mode='lines+markers',
                name='Projeção',
                line=dict(color='orange', width=2, dash='dash')
            ))

            # Linha de objetivo
            fig.add_trace(go.Scatter(
                x=list(vendas_temporal.index) + list(datas_futuras),
                y=[objetivo] * (len(vendas_temporal) + len(datas_futuras)),
                mode='lines',
                name='Objetivo Semanal',
                line=dict(color='red', width=2, dash='dot')
            ))

            fig.update_layout(
                title=f"Previsão de Vendas - {jogo_previsao}",
                xaxis_title="Data",
                yaxis_title="Vendas (€)",
                height=500,
                hovermode='x unified'
            )

            st.plotly_chart(fig, use_container_width=True)

            # Recomendações
            st.subheader("💡 Recomendações")

            if prob >= 80:
                st.success(f"""
                ✅ **Excelente!** Com a tendência atual, há {prob:.0f}% de probabilidade de atingir o objetivo.

                - Continue com as estratégias atuais
                - Mantenha o foco na qualidade do serviço
                """)
            elif prob >= 50:
                st.warning(f"""
                ⚠️ **Atenção!** Probabilidade moderada ({prob:.0f}%) de atingir o objetivo.

                - Considere ações promocionais
                - Analise possíveis melhorias no atendimento
                - Monitore a concorrência
                """)
            else:
                st.error(f"""
                🚨 **Alerta!** Baixa probabilidade ({prob:.0f}%) de atingir o objetivo.

                - Ações urgentes necessárias
                - Revisar estratégia de vendas
                - Considerar promoções agressivas
                - Analisar causas da queda
                """)

        else:
            st.warning("Dados insuficientes para previsão")


def pagina_remuneracao(df):
    """Análise de remuneração por jogo."""
    st.markdown('<h1 class="main-header">💰 Análise de Remuneração</h1>', unsafe_allow_html=True)
    st.info("💡 Remuneração calculada com base nas percentagens de comissão de cada jogo")

    # Adicionar coluna de remuneração
    df_rem = df.copy()
    df_rem['Percentagem'] = df_rem['Jogo'].map(REMUNERACAO_JOGOS)
    df_rem['Remuneracao'] = df_rem['Valor'] * df_rem['Percentagem'] / 100

    # Filtros na sidebar
    st.sidebar.header("🔍 Filtros")
    anos_disponiveis = sorted(df_rem['Ano'].unique())
    ano_selecionado = st.sidebar.multiselect(
        "Selecionar Anos",
        anos_disponiveis,
        default=anos_disponiveis,
        key="remuneracao_anos"
    )

    if not ano_selecionado:
        st.warning("Selecione pelo menos um ano")
        return

    df_filtrado = df_rem[df_rem['Ano'].isin(ano_selecionado)]

    # Métricas principais
    col1, col2, col3, col4 = st.columns(4)

    total_vendas = df_filtrado['Valor'].sum()
    total_remuneracao = df_filtrado['Remuneracao'].sum()
    percentagem_media = (total_remuneracao / total_vendas * 100) if total_vendas > 0 else 0
    media_semanal_rem = df_filtrado.groupby('Data_Emissao')['Remuneracao'].sum().mean()

    with col1:
        st.metric("💰 Remuneração Total", f"€{total_remuneracao:,.0f}")
    with col2:
        st.metric("📊 Percentagem Média", f"{percentagem_media:.0f}%")
    with col3:
        st.metric("📅 Média Semanal", f"€{media_semanal_rem:,.0f}")
    with col4:
        st.metric("📈 Total Vendas", f"€{total_vendas:,.0f}")

    st.markdown("---")

    # Tabs
    tab1, tab2, tab3 = st.tabs(["📊 Por Jogo", "📈 Evolução Temporal", "📉 Análise Detalhada"])

    with tab1:
        st.subheader("Remuneração por Jogo")

        # Agrupar por jogo
        rem_por_jogo = df_filtrado.groupby('Jogo').agg({
            'Valor': 'sum',
            'Remuneracao': 'sum',
            'Percentagem': 'first'
        }).reset_index()

        rem_por_jogo = rem_por_jogo.sort_values('Remuneracao', ascending=False)

        # Gráfico de barras
        col1, col2 = st.columns(2)

        with col1:
            fig = px.bar(
                rem_por_jogo,
                x='Remuneracao',
                y='Jogo',
                orientation='h',
                labels={'Remuneracao': 'Remuneração (€)', 'Jogo': 'Jogo'},
                color='Percentagem',
                color_continuous_scale='Greens',
                text='Percentagem'
            )
            fig.update_traces(texttemplate='%{text}%', textposition='inside')
            fig.update_layout(showlegend=False, height=500)
            st.plotly_chart(fig, use_container_width=True)

        with col2:
            # Tabela detalhada
            rem_tabela = rem_por_jogo.copy()
            rem_tabela['Valor'] = rem_tabela['Valor'].apply(lambda x: f"€{x:,.0f}")
            rem_tabela['Remuneracao'] = rem_tabela['Remuneracao'].apply(lambda x: f"€{x:,.0f}")
            rem_tabela['Percentagem'] = rem_tabela['Percentagem'].apply(lambda x: f"{x}%")
            rem_tabela = rem_tabela.rename(columns={
                'Valor': 'Vendas',
                'Remuneracao': 'Remuneração',
                'Percentagem': '%'
            })
            st.dataframe(rem_tabela, use_container_width=True, hide_index=True, height=500)

    with tab2:
        st.subheader("Evolução da Remuneração ao Longo do Tempo")

        # Seletor de jogos
        jogos_disponiveis = sorted(df_filtrado['Jogo'].unique())
        jogos_selecionados = st.multiselect(
            "Selecionar Jogos",
            jogos_disponiveis,
            default=jogos_disponiveis[:5] if len(jogos_disponiveis) >= 5 else jogos_disponiveis
        )

        if jogos_selecionados:
            df_sel = df_filtrado[df_filtrado['Jogo'].isin(jogos_selecionados)]

            # Evolução por jogo
            rem_temp = df_sel.groupby(['Data_Emissao', 'Jogo'])['Remuneracao'].sum().reset_index()

            fig = px.line(
                rem_temp,
                x='Data_Emissao',
                y='Remuneracao',
                color='Jogo',
                markers=True,
                labels={'Remuneracao': 'Remuneração (€)', 'Data_Emissao': 'Data'}
            )
            fig.update_layout(height=500)
            st.plotly_chart(fig, use_container_width=True)

            # Total por período
            st.subheader("Remuneração Total por Período")
            rem_total = df_filtrado.groupby('Data_Emissao')['Remuneracao'].sum().reset_index()

            fig = px.area(
                rem_total,
                x='Data_Emissao',
                y='Remuneracao',
                labels={'Remuneracao': 'Remuneração Total (€)', 'Data_Emissao': 'Data'}
            )
            fig.update_layout(height=400)
            st.plotly_chart(fig, use_container_width=True)

    with tab3:
        st.subheader("Análise Comparativa")

        col1, col2 = st.columns(2)

        with col1:
            # Comparação vendas vs remuneração
            st.markdown("### 💼 Comparação Vendas vs Remuneração")

            comparacao = df_filtrado.groupby('Jogo').agg({
                'Valor': 'sum',
                'Remuneracao': 'sum'
            }).reset_index()

            comparacao = comparacao.sort_values('Remuneracao', ascending=True)

            fig = go.Figure()
            fig.add_trace(go.Bar(
                name='Vendas',
                x=comparacao['Valor'],
                y=comparacao['Jogo'],
                orientation='h',
                marker=dict(color='lightblue')
            ))
            fig.add_trace(go.Bar(
                name='Remuneração',
                x=comparacao['Remuneracao'],
                y=comparacao['Jogo'],
                orientation='h',
                marker=dict(color='green')
            ))
            fig.update_layout(barmode='group', height=500)
            st.plotly_chart(fig, use_container_width=True)

        with col2:
            # Distribuição da remuneração
            st.markdown("### 📊 Distribuição da Remuneração")

            rem_dist = df_filtrado.groupby('Jogo')['Remuneracao'].sum()

            fig = px.pie(
                values=rem_dist.values,
                names=rem_dist.index,
                hole=0.4
            )
            fig.update_layout(height=500)
            st.plotly_chart(fig, use_container_width=True)

        # Estatísticas por ano
        st.subheader("📅 Remuneração Anual")
        rem_anual = df_filtrado.groupby(['Ano', 'Jogo'])['Remuneracao'].sum().reset_index()

        fig = px.bar(
            rem_anual,
            x='Ano',
            y='Remuneracao',
            color='Jogo',
            barmode='stack',
            labels={'Remuneracao': 'Remuneração (€)', 'Ano': 'Ano'}
        )
        fig.update_layout(height=400)
        st.plotly_chart(fig, use_container_width=True)

        # Tabela de percentagens
        st.subheader("📋 Tabela de Percentagens por Jogo")
        tabela_perc = pd.DataFrame(list(REMUNERACAO_JOGOS.items()), columns=['Jogo', 'Percentagem (%)'])
        tabela_perc = tabela_perc.sort_values('Percentagem (%)', ascending=False)
        st.dataframe(tabela_perc, use_container_width=True, hide_index=True)

    # Exportação de dados
    st.markdown("---")
    st.subheader("📥 Exportar Dados de Remuneração")

    col1, col2 = st.columns(2)

    # Preparar dados para exportação
    rem_export = df_filtrado.groupby('Jogo').agg({
        'Valor': 'sum',
        'Remuneracao': 'sum',
        'Percentagem': 'first'
    }).reset_index()
    rem_export.columns = ['Jogo', 'Total Vendas (€)', 'Total Remuneração (€)', 'Percentagem (%)']

    with col1:
        # Exportar CSV
        csv_data = rem_export.to_csv(index=False).encode('utf-8')
        st.download_button(
            label="📄 Baixar Remuneração em CSV",
            data=csv_data,
            file_name=f"remuneracao_{'-'.join(map(str, ano_selecionado))}.csv",
            mime="text/csv"
        )

    with col2:
        # Exportar Excel
        if EXCEL_AVAILABLE:
            try:
                # Remuneração anual
                rem_anual_export = df_filtrado.groupby(['Ano', 'Jogo'])['Remuneracao'].sum().reset_index()
                rem_anual_export.columns = ['Ano', 'Jogo', 'Remuneração (€)']

                # Evolução temporal
                rem_temp_export = df_filtrado.groupby(['Data_Emissao', 'Jogo'])['Remuneracao'].sum().reset_index()
                rem_temp_export['Data_Emissao'] = rem_temp_export['Data_Emissao'].dt.strftime('%d-%m-%Y')
                rem_temp_export.columns = ['Data', 'Jogo', 'Remuneração (€)']

                dataframes = {
                    'Resumo por Jogo': rem_export,
                    'Remuneração Anual': rem_anual_export,
                    'Evolução Temporal': rem_temp_export,
                    'Percentagens': tabela_perc
                }
                info = {
                    'Página': 'Remuneração',
                    'Anos': ', '.join(map(str, ano_selecionado)),
                    'Total Vendas': f'€{total_vendas:,.0f}',
                    'Total Remuneração': f'€{total_remuneracao:,.0f}',
                    'Percentagem Média': f'{percentagem_media:.0f}%'
                }
                excel_data = exportar_excel_generico(dataframes, 'remuneracao', info)
                if excel_data:
                    st.download_button(
                        label="📊 Baixar Relatório em Excel",
                        data=excel_data,
                        file_name=f"remuneracao_{'-'.join(map(str, ano_selecionado))}.xlsx",
                        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                    )
            except Exception as e:
                st.error(f"Erro ao gerar Excel: {str(e)}")
        else:
            st.info("📊 Excel: Instale openpyxl para exportar em Excel")


def pagina_analise_semanal_avancada(df):
    """Análise semanal avançada com WoW, ranking dinâmico, percentagens e velocímetro."""
    st.markdown('<h1 class="main-header">🔬 Análise Semanal Avançada</h1>', unsafe_allow_html=True)
    st.caption("💡 Análise detalhada das últimas semanas com comparações WoW, rankings e insights automáticos")

    # Verificar dados suficientes
    datas_unicas = sorted(df['Data_Emissao'].unique(), reverse=True)
    if len(datas_unicas) < 2:
        st.warning("Dados insuficientes para análise semanal. São necessárias pelo menos 2 semanas de dados.")
        return

    ultima_semana = datas_unicas[0]
    penultima_semana = datas_unicas[1]

    # Header com informação da semana
    st.info(f"📅 Semana Atual: **{ultima_semana.strftime('%d/%m/%Y')}** | Semana Anterior: **{penultima_semana.strftime('%d/%m/%Y')}**")

    # Insights Automáticos
    st.subheader("💡 Insights Automáticos")
    with st.spinner("Gerando insights..."):
        insights = gerar_insights_automaticos(df, janela_medias=[4, 8])

    for insight in insights[:10]:  # Mostrar top 10 insights
        st.markdown(f"- {insight}")

    st.markdown("---")

    # Tabs para diferentes análises
    tab1, tab2, tab3, tab4, tab5 = st.tabs([
        "📊 WoW (Semana a Semana)",
        "🏆 Ranking Dinâmico",
        "🥧 Contribuição %",
        "🎯 Velocímetro Performance",
        "📈 Médias Móveis"
    ])

    with tab1:
        st.subheader("Comparação Week over Week (WoW)")
        st.caption("Análise de crescimento semanal")

        # Calcular WoW
        vendas_wow = calcular_comparacao_wow(df)

        # Filtro de jogos
        jogos_disponiveis = sorted(df['Jogo'].unique())
        jogos_selecionados_wow = st.multiselect(
            "Selecionar Jogos para Análise WoW",
            jogos_disponiveis,
            default=jogos_disponiveis[:5] if len(jogos_disponiveis) >= 5 else jogos_disponiveis,
            key="wow_jogos"
        )

        if jogos_selecionados_wow:
            vendas_wow_filtrado = vendas_wow[vendas_wow['Jogo'].isin(jogos_selecionados_wow)]

            # Últimas 8 semanas
            ultimas_datas = datas_unicas[:8]
            vendas_wow_recente = vendas_wow_filtrado[vendas_wow_filtrado['Data_Emissao'].isin(ultimas_datas)]

            col1, col2 = st.columns(2)

            with col1:
                # Gráfico de vendas
                fig = px.line(
                    vendas_wow_recente,
                    x='Data_Emissao',
                    y='Valor',
                    color='Jogo',
                    markers=True,
                    title="Vendas Semanais (Últimas 8 Semanas)",
                    labels={'Data_Emissao': 'Data', 'Valor': 'Vendas (€)'}
                )
                fig.update_layout(height=400)
                st.plotly_chart(fig, use_container_width=True)

            with col2:
                # Gráfico de crescimento WoW
                vendas_wow_sem_na = vendas_wow_recente.dropna(subset=['Crescimento_WoW_%'])

                fig = px.bar(
                    vendas_wow_sem_na,
                    x='Data_Emissao',
                    y='Crescimento_WoW_%',
                    color='Jogo',
                    title="Crescimento WoW (%)",
                    labels={'Data_Emissao': 'Data', 'Crescimento_WoW_%': 'Crescimento (%)'},
                    barmode='group'
                )
                fig.add_hline(y=0, line_dash="dash", line_color="red", annotation_text="Zero")
                fig.update_layout(height=400)
                st.plotly_chart(fig, use_container_width=True)

            # Tabela detalhada da última semana
            st.subheader(f"📋 Detalhes da Semana de {ultima_semana.strftime('%d/%m/%Y')}")
            ultima_semana_data = vendas_wow_filtrado[vendas_wow_filtrado['Data_Emissao'] == ultima_semana].dropna(subset=['Crescimento_WoW_%'])

            if not ultima_semana_data.empty:
                tabela_wow = ultima_semana_data[['Jogo', 'Valor', 'Valor_Semana_Anterior', 'Crescimento_WoW_%', 'Diferenca_Absoluta']].copy()
                tabela_wow = tabela_wow.sort_values('Crescimento_WoW_%', ascending=False)
                tabela_wow['Valor'] = tabela_wow['Valor'].apply(lambda x: f"€{x:,.0f}")
                tabela_wow['Valor_Semana_Anterior'] = tabela_wow['Valor_Semana_Anterior'].apply(lambda x: f"€{x:,.0f}")
                tabela_wow['Crescimento_WoW_%'] = tabela_wow['Crescimento_WoW_%'].apply(lambda x: f"{x:+.1f}%")
                tabela_wow['Diferenca_Absoluta'] = tabela_wow['Diferenca_Absoluta'].apply(lambda x: f"€{x:+,.0f}")
                tabela_wow.columns = ['Jogo', 'Vendas Atual', 'Vendas Anterior', 'Crescimento %', 'Diferença €']

                st.dataframe(tabela_wow, use_container_width=True, hide_index=True)
            else:
                st.info("Sem dados WoW para a última semana")

            # Estatísticas WoW
            st.subheader("📊 Estatísticas WoW")
            col1, col2, col3, col4 = st.columns(4)

            crescimento_medio = vendas_wow_sem_na['Crescimento_WoW_%'].mean()
            crescimento_max = vendas_wow_sem_na['Crescimento_WoW_%'].max()
            crescimento_min = vendas_wow_sem_na['Crescimento_WoW_%'].min()
            semanas_positivas = (vendas_wow_sem_na['Crescimento_WoW_%'] > 0).sum()

            with col1:
                st.metric("📈 Crescimento Médio", f"{crescimento_medio:.1f}%")
            with col2:
                st.metric("⬆️ Máximo", f"{crescimento_max:.1f}%")
            with col3:
                st.metric("⬇️ Mínimo", f"{crescimento_min:.1f}%")
            with col4:
                st.metric("✅ Semanas Positivas", f"{semanas_positivas}")

    with tab2:
        st.subheader("🏆 Ranking Dinâmico de Jogos")
        st.caption("Mudanças de posição no ranking em relação à semana anterior")

        df_ranking = calcular_ranking_dinamico(df)

        if df_ranking is not None:
            col1, col2 = st.columns([2, 1])

            with col1:
                # Tabela de ranking
                tabela_rank = df_ranking[['Posicao_Atual', 'Jogo', 'Vendas_Atual', 'Posicao_Anterior', 'Indicador']].copy()
                tabela_rank['Vendas_Atual'] = tabela_rank['Vendas_Atual'].apply(lambda x: f"€{x:,.0f}")
                tabela_rank['Posicao_Anterior'] = tabela_rank['Posicao_Anterior'].apply(lambda x: f"#{int(x)}")
                tabela_rank.columns = ['#', 'Jogo', 'Vendas', 'Posição Anterior', 'Mudança']

                st.dataframe(tabela_rank, use_container_width=True, hide_index=True, height=600)

            with col2:
                # Estatísticas de mudanças
                st.markdown("### 📊 Estatísticas")

                maiores_subidas = df_ranking[df_ranking['Mudanca'] > 0].nlargest(3, 'Mudanca')
                maiores_descidas = df_ranking[df_ranking['Mudanca'] < 0].nsmallest(3, 'Mudanca')

                if not maiores_subidas.empty:
                    st.markdown("**🚀 Maiores Subidas:**")
                    for _, row in maiores_subidas.iterrows():
                        st.success(f"{row['Jogo']}: +{int(row['Mudanca'])} posições")

                if not maiores_descidas.empty:
                    st.markdown("**⬇️ Maiores Descidas:**")
                    for _, row in maiores_descidas.iterrows():
                        st.error(f"{row['Jogo']}: {int(row['Mudanca'])} posições")

                # Gráfico de mudanças
                fig = px.bar(
                    df_ranking,
                    x='Mudanca',
                    y='Jogo',
                    orientation='h',
                    title="Mudanças de Posição",
                    labels={'Mudanca': 'Posições', 'Jogo': ''},
                    color='Mudanca',
                    color_continuous_scale='RdYlGn',
                    color_continuous_midpoint=0
                )
                fig.update_layout(height=400, showlegend=False)
                st.plotly_chart(fig, use_container_width=True)
        else:
            st.warning("Dados insuficientes para calcular ranking dinâmico")

    with tab3:
        st.subheader("🥧 Percentagem de Contribuição")
        st.caption("Contribuição de cada jogo para o total de vendas")

        # Seletor de semana
        col1, col2 = st.columns([1, 3])

        with col1:
            data_selecionada = st.selectbox(
                "Selecionar Semana",
                datas_unicas[:10],  # Últimas 10 semanas
                format_func=lambda x: x.strftime('%d/%m/%Y'),
                key="contrib_data"
            )

        # Calcular contribuição
        df_contrib = calcular_percentagem_contribuicao(df, data_selecionada)

        col1, col2 = st.columns(2)

        with col1:
            # Gráfico de pizza
            fig = px.pie(
                df_contrib,
                values='Percentagem',
                names='Jogo',
                title=f"Distribuição de Vendas - {data_selecionada.strftime('%d/%m/%Y')}",
                hole=0.4
            )
            fig.update_traces(textposition='inside', textinfo='percent+label')
            fig.update_layout(height=500)
            st.plotly_chart(fig, use_container_width=True)

        with col2:
            # Tabela detalhada
            tabela_contrib = df_contrib.copy()
            tabela_contrib['Valor'] = tabela_contrib['Valor'].apply(lambda x: f"€{x:,.0f}")
            tabela_contrib['Percentagem'] = tabela_contrib['Percentagem'].apply(lambda x: f"{x:.0f}%")

            st.markdown("### 📋 Detalhes")
            st.dataframe(tabela_contrib, use_container_width=True, hide_index=True, height=500)

            # Total
            total = df_contrib['Valor'].apply(lambda x: float(x.replace('€', '').replace(',', '')) if isinstance(x, str) else x).sum()
            st.metric("💰 Total da Semana", f"€{df[df['Data_Emissao'] == data_selecionada]['Valor'].sum():,.0f}")

        # Comparação temporal de contribuições
        st.subheader("📈 Evolução das Contribuições (%)")

        # Top 5 jogos por contribuição atual
        top_jogos = df_contrib.head(5)['Jogo'].tolist()

        # Calcular contribuições ao longo do tempo
        contrib_temporal = []
        for data in datas_unicas[:12]:  # Últimas 12 semanas
            df_temp = calcular_percentagem_contribuicao(df, data)
            df_temp['Data'] = data
            contrib_temporal.append(df_temp)

        df_contrib_temporal = pd.concat(contrib_temporal)
        df_contrib_temporal = df_contrib_temporal[df_contrib_temporal['Jogo'].isin(top_jogos)]

        fig = px.line(
            df_contrib_temporal,
            x='Data',
            y='Percentagem',
            color='Jogo',
            markers=True,
            title=f"Evolução da Contribuição % - Top 5 Jogos",
            labels={'Percentagem': 'Contribuição (%)', 'Data': 'Data'}
        )
        fig.update_layout(height=400)
        st.plotly_chart(fig, use_container_width=True)

    with tab4:
        st.subheader("🎯 Velocímetro de Performance")
        st.caption("Comparação da última semana com a média histórica")

        # Controles
        col1, col2 = st.columns([1, 3])

        with col1:
            janela_velocimetro = st.slider(
                "Semanas Históricas",
                min_value=4,
                max_value=12,
                value=8,
                key="vel_janela"
            )

        # Calcular velocímetro
        velocimetro, data_atual = calcular_velocimetro_performance(df, janela_velocimetro)

        # Criar DataFrame para visualização
        vel_data = []
        for jogo, info in velocimetro.items():
            vel_data.append({
                'Jogo': jogo,
                'Valor Atual': info['valor_atual'],
                'Média Histórica': info['media_historica'],
                'Performance (%)': info['percentual'],
                'Status': info['status'],
                'Diferença (€)': info['diferenca']
            })

        df_vel = pd.DataFrame(vel_data).sort_values('Performance (%)', ascending=False)

        # Gráficos gauge para top jogos
        st.subheader(f"📊 Performance vs Média de {janela_velocimetro} Semanas")

        # Gráfico de barras comparativo
        col1, col2 = st.columns([2, 1])

        with col1:
            fig = go.Figure()

            fig.add_trace(go.Bar(
                name='Valor Atual',
                x=df_vel['Jogo'],
                y=df_vel['Valor Atual'],
                marker_color='lightblue'
            ))

            fig.add_trace(go.Bar(
                name='Média Histórica',
                x=df_vel['Jogo'],
                y=df_vel['Média Histórica'],
                marker_color='orange'
            ))

            fig.update_layout(
                title="Comparação: Atual vs Média Histórica",
                xaxis_title="Jogo",
                yaxis_title="Valor (€)",
                barmode='group',
                height=500
            )

            st.plotly_chart(fig, use_container_width=True)

        with col2:
            # Tabela de performance
            tabela_vel = df_vel[['Status', 'Jogo', 'Performance (%)']].copy()
            tabela_vel['Performance (%)'] = tabela_vel['Performance (%)'].apply(lambda x: f"{x:.0f}%")

            st.dataframe(tabela_vel, use_container_width=True, hide_index=True, height=500)

        # Detalhes numéricos
        st.subheader("📋 Detalhes de Performance")
        tabela_vel_det = df_vel.copy()
        tabela_vel_det['Valor Atual'] = tabela_vel_det['Valor Atual'].apply(lambda x: f"€{x:,.0f}")
        tabela_vel_det['Média Histórica'] = tabela_vel_det['Média Histórica'].apply(lambda x: f"€{x:,.0f}")
        tabela_vel_det['Performance (%)'] = tabela_vel_det['Performance (%)'].apply(lambda x: f"{x:.1f}%")
        tabela_vel_det['Diferença (€)'] = tabela_vel_det['Diferença (€)'].apply(lambda x: f"€{x:+,.0f}")

        st.dataframe(tabela_vel_det, use_container_width=True, hide_index=True)

    with tab5:
        st.subheader("📈 Análise de Médias Móveis")
        st.caption("Suavização de tendências com múltiplas janelas temporais")

        # Seletor de jogo
        jogo_mm = st.selectbox(
            "Selecionar Jogo",
            sorted(df['Jogo'].unique()),
            key="mm_jogo"
        )

        # Janelas de média móvel
        col1, col2, col3 = st.columns(3)

        with col1:
            janela_1 = st.number_input("Janela 1 (semanas)", min_value=2, max_value=12, value=4, key="mm_j1")
        with col2:
            janela_2 = st.number_input("Janela 2 (semanas)", min_value=2, max_value=12, value=8, key="mm_j2")
        with col3:
            mostrar_objetivo_mm = st.checkbox("Mostrar Objetivo", value=True, key="mm_obj")

        # Calcular médias móveis
        df_jogo_mm = df[df['Jogo'] == jogo_mm].copy()
        vendas_temporal = df_jogo_mm.groupby('Data_Emissao')['Valor'].sum().sort_index()

        mm_1 = vendas_temporal.rolling(window=janela_1, min_periods=1).mean()
        mm_2 = vendas_temporal.rolling(window=janela_2, min_periods=1).mean()

        # Gráfico
        fig = go.Figure()

        # Vendas reais
        fig.add_trace(go.Scatter(
            x=vendas_temporal.index,
            y=vendas_temporal.values,
            mode='lines+markers',
            name='Vendas Reais',
            line=dict(color='lightgray', width=1),
            marker=dict(size=4)
        ))

        # Média móvel 1
        fig.add_trace(go.Scatter(
            x=mm_1.index,
            y=mm_1.values,
            mode='lines',
            name=f'MM {janela_1} semanas',
            line=dict(color='blue', width=2)
        ))

        # Média móvel 2
        fig.add_trace(go.Scatter(
            x=mm_2.index,
            y=mm_2.values,
            mode='lines',
            name=f'MM {janela_2} semanas',
            line=dict(color='green', width=2)
        ))

        # Objetivo
        if mostrar_objetivo_mm and jogo_mm in OBJETIVOS_SEMANAIS:
            objetivo = OBJETIVOS_SEMANAIS[jogo_mm]
            fig.add_trace(go.Scatter(
                x=vendas_temporal.index,
                y=[objetivo] * len(vendas_temporal),
                mode='lines',
                name='Objetivo Semanal',
                line=dict(color='red', width=2, dash='dash')
            ))

        fig.update_layout(
            title=f"Médias Móveis - {jogo_mm}",
            xaxis_title="Data",
            yaxis_title="Vendas (€)",
            height=500,
            hovermode='x unified'
        )

        st.plotly_chart(fig, use_container_width=True)

        # Estatísticas
        st.subheader("📊 Estatísticas das Médias Móveis")
        col1, col2, col3, col4 = st.columns(4)

        with col1:
            st.metric(f"MM {janela_1} (Última)", f"€{mm_1.iloc[-1]:.0f}")
        with col2:
            st.metric(f"MM {janela_2} (Última)", f"€{mm_2.iloc[-1]:.0f}")
        with col3:
            st.metric("Valor Atual", f"€{vendas_temporal.iloc[-1]:.0f}")
        with col4:
            if jogo_mm in OBJETIVOS_SEMANAIS:
                obj = OBJETIVOS_SEMANAIS[jogo_mm]
                perf = (vendas_temporal.iloc[-1] / obj * 100) if obj > 0 else 0
                st.metric("Performance vs Objetivo", f"{perf:.1f}%")

        # Tendência
        st.info(f"""
        💡 **Interpretação:**
        - MM {janela_1} semanas captura tendências de curto prazo
        - MM {janela_2} semanas mostra a tendência de longo prazo
        - Quando MM curta cruza acima da longa: sinal de alta
        - Quando MM curta cruza abaixo da longa: sinal de baixa
        """)

    st.markdown("---")

    # ANÁLISES AVANÇADAS
    st.markdown("## 🔬 Análises Avançadas")

    tab_av1, tab_av2, tab_av3, tab_av4 = st.tabs([
        "🔥 Heatmap Performance",
        "🔗 Correlações",
        "📅 Padrões Temporais",
        "📊 Previsões"
    ])

    with tab_av1:
        st.subheader("Heatmap de Performance Semanal")
        st.caption("Visualização de performance por jogo ao longo das semanas")

        # Preparar dados para heatmap
        ultimas_8_semanas_heat = sorted(df['Data_Emissao'].unique(), reverse=True)[:8]
        df_heat = df[df['Data_Emissao'].isin(ultimas_8_semanas_heat)]

        # Pivot table para heatmap
        pivot_heat = df_heat.pivot_table(
            values='Valor',
            index='Jogo',
            columns='Data_Emissao',
            aggfunc='sum',
            fill_value=0
        )

        # Normalizar por objetivo para mostrar performance
        for jogo in pivot_heat.index:
            if jogo in OBJETIVOS_SEMANAIS and OBJETIVOS_SEMANAIS[jogo] > 0:
                pivot_heat.loc[jogo] = (pivot_heat.loc[jogo] / OBJETIVOS_SEMANAIS[jogo]) * 100

        # Criar heatmap
        fig_heat = go.Figure(data=go.Heatmap(
            z=pivot_heat.values,
            x=[d.strftime('%d/%m') for d in pivot_heat.columns],
            y=pivot_heat.index,
            colorscale=[
                [0, '#d32f2f'],
                [0.5, '#ffa726'],
                [0.9, '#ffeb3b'],
                [1, '#4caf50']
            ],
            text=np.round(pivot_heat.values, 1),
            texttemplate='%{text}%',
            textfont=dict(size=10),
            colorbar=dict(title="Performance %")
        ))

        fig_heat.update_layout(
            title="Performance por Jogo e Semana (% do Objetivo)",
            xaxis_title="Semana",
            yaxis_title="Jogo",
            height=500
        )

        st.plotly_chart(fig_heat, use_container_width=True)

        # Análise de volatilidade
        st.subheader("📊 Indicadores de Volatilidade")

        volatilidade = pivot_heat.std(axis=1).sort_values(ascending=False)

        col1, col2 = st.columns(2)

        with col1:
            st.markdown("**Jogos Mais Voláteis**")
            top_vol = volatilidade.head(5)
            for jogo, vol in top_vol.items():
                st.metric(jogo, f"{vol:.1f}%", "Alta volatilidade")

        with col2:
            st.markdown("**Jogos Mais Estáveis**")
            bottom_vol = volatilidade.tail(5)
            for jogo, vol in bottom_vol.items():
                st.metric(jogo, f"{vol:.1f}%", "Baixa volatilidade")

    with tab_av2:
        st.subheader("Matriz de Correlação entre Jogos")
        st.caption("Análise de como as vendas de diferentes jogos se correlacionam")

        # Criar matriz de vendas por jogo e semana
        ultimas_12_semanas = sorted(df['Data_Emissao'].unique(), reverse=True)[:12]
        df_corr = df[df['Data_Emissao'].isin(ultimas_12_semanas)]

        # Pivot para correlação
        pivot_corr = df_corr.pivot_table(
            values='Valor',
            index='Data_Emissao',
            columns='Jogo',
            aggfunc='sum',
            fill_value=0
        )

        # Calcular correlação
        if len(pivot_corr) >= 3:
            corr_matrix = pivot_corr.corr()

            # Heatmap de correlação
            fig_corr = go.Figure(data=go.Heatmap(
                z=corr_matrix.values,
                x=corr_matrix.columns,
                y=corr_matrix.index,
                colorscale='RdBu',
                zmid=0,
                text=np.round(corr_matrix.values, 2),
                texttemplate='%{text}',
                textfont=dict(size=8),
                colorbar=dict(title="Correlação")
            ))

            fig_corr.update_layout(
                title="Matriz de Correlação - Vendas entre Jogos",
                height=600,
                xaxis=dict(tickangle=-45)
            )

            st.plotly_chart(fig_corr, use_container_width=True)

            # Insights de correlação
            st.info("""
            💡 **Como interpretar:**
            - Valores próximos de +1: jogos com vendas que sobem/descem juntas
            - Valores próximos de -1: quando um sobe, o outro tende a descer
            - Valores próximos de 0: sem relação aparente
            """)

            # Correlações mais fortes
            st.subheader("🔗 Correlações Mais Fortes")

            # Extrair pares únicos
            corr_pairs = []
            for i in range(len(corr_matrix.columns)):
                for j in range(i+1, len(corr_matrix.columns)):
                    corr_pairs.append({
                        'Jogo 1': corr_matrix.columns[i],
                        'Jogo 2': corr_matrix.columns[j],
                        'Correlação': corr_matrix.iloc[i, j]
                    })

            df_pairs = pd.DataFrame(corr_pairs).sort_values('Correlação', ascending=False)

            col1, col2 = st.columns(2)

            with col1:
                st.markdown("**Correlações Positivas**")
                st.dataframe(
                    df_pairs.head(5)[['Jogo 1', 'Jogo 2', 'Correlação']],
                    hide_index=True
                )

            with col2:
                st.markdown("**Correlações Negativas**")
                st.dataframe(
                    df_pairs.tail(5)[['Jogo 1', 'Jogo 2', 'Correlação']],
                    hide_index=True
                )
        else:
            st.warning("Dados insuficientes para análise de correlação (mínimo 3 semanas)")

    with tab_av3:
        st.subheader("Análise de Padrões Temporais")
        st.caption("Identificação de padrões por dia da semana e períodos")

        # Adicionar dia da semana
        df_padroes = df.copy()
        df_padroes['Dia_Semana'] = df_padroes['Data_Emissao'].dt.day_name()
        df_padroes['Semana_Mes'] = df_padroes['Data_Emissao'].dt.isocalendar().week % 4 + 1

        # Média por dia da semana
        vendas_dia_semana = df_padroes.groupby('Dia_Semana')['Valor'].mean().reindex([
            'Monday', 'Tuesday', 'Wednesday', 'Thursday', 'Friday', 'Saturday', 'Sunday'
        ])

        dias_pt = ['Segunda', 'Terça', 'Quarta', 'Quinta', 'Sexta', 'Sábado', 'Domingo']

        fig_dias = go.Figure(data=[
            go.Bar(
                x=dias_pt,
                y=vendas_dia_semana.values,
                marker=dict(
                    color=vendas_dia_semana.values,
                    colorscale='Viridis',
                    showscale=False
                ),
                text=[f'€{v:,.0f}' for v in vendas_dia_semana.values],
                textposition='outside'
            )
        ])

        fig_dias.update_layout(
            title="Média de Vendas por Dia da Semana",
            xaxis_title="Dia",
            yaxis_title="Vendas Médias (€)",
            height=400
        )

        st.plotly_chart(fig_dias, use_container_width=True)

        # Matriz WoW vs MoM
        st.subheader("📈 Matriz de Crescimento (WoW vs MoM)")

        if len(datas_unicas) >= 8:
            # Calcular WoW e MoM para cada jogo
            jogos_crescimento = []

            for jogo in df['Jogo'].unique():
                df_jogo = df[df['Jogo'] == jogo]

                # WoW
                vendas_ultima = df_jogo[df_jogo['Data_Emissao'] == datas_unicas[0]]['Valor'].sum()
                vendas_anterior = df_jogo[df_jogo['Data_Emissao'] == datas_unicas[1]]['Valor'].sum()
                wow = ((vendas_ultima - vendas_anterior) / vendas_anterior * 100) if vendas_anterior > 0 else 0

                # MoM (últimas 4 semanas vs 4 anteriores)
                vendas_mes_atual = df_jogo[df_jogo['Data_Emissao'].isin(datas_unicas[:4])]['Valor'].sum()
                vendas_mes_anterior = df_jogo[df_jogo['Data_Emissao'].isin(datas_unicas[4:8])]['Valor'].sum()
                mom = ((vendas_mes_atual - vendas_mes_anterior) / vendas_mes_anterior * 100) if vendas_mes_anterior > 0 else 0

                jogos_crescimento.append({
                    'Jogo': jogo,
                    'WoW': wow,
                    'MoM': mom,
                    'Tamanho': vendas_ultima
                })

            df_cresc = pd.DataFrame(jogos_crescimento)

            # Scatter plot
            fig_cresc = px.scatter(
                df_cresc,
                x='WoW',
                y='MoM',
                size='Tamanho',
                text='Jogo',
                color='MoM',
                color_continuous_scale='RdYlGn',
                labels={'WoW': 'Crescimento WoW (%)', 'MoM': 'Crescimento MoM (%)'}
            )

            fig_cresc.add_hline(y=0, line_dash="dash", line_color="gray")
            fig_cresc.add_vline(x=0, line_dash="dash", line_color="gray")

            fig_cresc.update_traces(textposition='top center')
            fig_cresc.update_layout(
                title="Matriz de Crescimento: WoW vs MoM",
                height=500,
                showlegend=False
            )

            st.plotly_chart(fig_cresc, use_container_width=True)

            st.info("""
            💡 **Quadrantes:**
            - Superior Direito: Crescimento em ambos (ideal)
            - Superior Esquerdo: Crescimento MoM mas queda WoW (atenção)
            - Inferior Direito: Crescimento WoW mas queda MoM (possível flutuação)
            - Inferior Esquerdo: Queda em ambos (preocupante)
            """)

    with tab_av4:
        st.subheader("Previsões Estatísticas Simples")
        st.caption("Projeções baseadas em regressão linear")

        # Selecionar jogo para previsão
        jogo_prev = st.selectbox(
            "Selecionar Jogo para Previsão",
            sorted(df['Jogo'].unique()),
            key="prev_jogo"
        )

        semanas_prev = st.slider(
            "Semanas para Prever",
            min_value=1,
            max_value=8,
            value=4,
            key="prev_semanas"
        )

        # Preparar dados
        df_jogo_prev = df[df['Jogo'] == jogo_prev].sort_values('Data_Emissao')
        vendas_historico = df_jogo_prev.groupby('Data_Emissao')['Valor'].sum().reset_index()

        if len(vendas_historico) >= 4:
            # Preparar features
            vendas_historico['Semana_Num'] = range(len(vendas_historico))
            X = vendas_historico[['Semana_Num']].values
            y = vendas_historico['Valor'].values

            # Treinar modelo
            model = LinearRegression()
            model.fit(X, y)

            # Prever
            ultimas_semanas_num = vendas_historico['Semana_Num'].max()
            semanas_futuro = np.array([[ultimas_semanas_num + i] for i in range(1, semanas_prev + 1)])
            previsoes = model.predict(semanas_futuro)

            # Gráfico
            fig_prev = go.Figure()

            # Histórico
            fig_prev.add_trace(go.Scatter(
                x=vendas_historico['Data_Emissao'],
                y=vendas_historico['Valor'],
                mode='lines+markers',
                name='Histórico',
                line=dict(color='#667eea', width=2),
                marker=dict(size=8)
            ))

            # Previsões
            datas_futuro = pd.date_range(
                start=vendas_historico['Data_Emissao'].max() + pd.Timedelta(days=7),
                periods=semanas_prev,
                freq='W'
            )

            fig_prev.add_trace(go.Scatter(
                x=datas_futuro,
                y=previsoes,
                mode='lines+markers',
                name='Previsão',
                line=dict(color='#f5576c', width=2, dash='dash'),
                marker=dict(size=8, symbol='diamond')
            ))

            # Objetivo
            if jogo_prev in OBJETIVOS_SEMANAIS:
                objetivo = OBJETIVOS_SEMANAIS[jogo_prev]
                fig_prev.add_hline(
                    y=objetivo,
                    line_dash="dot",
                    line_color="green",
                    annotation_text="Objetivo"
                )

            fig_prev.update_layout(
                title=f"Previsão de Vendas - {jogo_prev}",
                xaxis_title="Data",
                yaxis_title="Vendas (€)",
                height=500,
                hovermode='x unified'
            )

            st.plotly_chart(fig_prev, use_container_width=True)

            # Métricas da previsão
            col1, col2, col3 = st.columns(3)

            with col1:
                st.metric(
                    "Previsão Próxima Semana",
                    f"€{previsoes[0]:,.0f}"
                )

            with col2:
                media_historica = vendas_historico['Valor'].mean()
                st.metric(
                    "Média Histórica",
                    f"€{media_historica:,.0f}"
                )

            with col3:
                if jogo_prev in OBJETIVOS_SEMANAIS:
                    obj = OBJETIVOS_SEMANAIS[jogo_prev]
                    perf_prev = (previsoes[0] / obj * 100) if obj > 0 else 0
                    st.metric(
                        "Performance Prevista",
                        f"{perf_prev:.1f}%"
                    )

            st.warning("⚠️ Estas previsões são baseadas em regressão linear simples e devem ser usadas apenas como referência. Fatores externos não são considerados.")

        else:
            st.warning("Dados insuficientes para previsão (mínimo 4 semanas)")

    # Exportação com insights
    st.markdown("---")
    st.subheader("📥 Exportar Análise Semanal")

    col1, col2 = st.columns(2)

    with col1:
        # CSV básico
        if df_ranking is not None:
            csv_data = df_ranking.to_csv(index=False).encode('utf-8')
            st.download_button(
                label="📄 Baixar Ranking em CSV",
                data=csv_data,
                file_name=f"analise_semanal_{ultima_semana.strftime('%Y%m%d')}.csv",
                mime="text/csv"
            )

    with col2:
        # Excel com insights e gráficos
        if EXCEL_AVAILABLE:
            try:
                with st.spinner("Gerando Excel com insights e gráficos..."):
                    excel_data = exportar_excel_com_insights_e_graficos(df, insights, ultima_semana)

                if excel_data:
                    st.download_button(
                        label="📊 Baixar Relatório Completo em Excel",
                        data=excel_data,
                        file_name=f"analise_semanal_avancada_{ultima_semana.strftime('%Y%m%d')}.xlsx",
                        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        help="Inclui insights automáticos, gráficos visuais e múltiplas análises"
                    )
                    st.success("✅ Excel com insights e gráficos disponível!")
                else:
                    st.error("Erro ao gerar Excel com insights")
            except Exception as e:
                st.error(f"Erro ao gerar Excel: {str(e)}")
        else:
            st.info("📊 Excel: Instale openpyxl para exportar em Excel")


def pagina_prestacoes_contas(df):
    """Análise de prestações de contas - quanto dever à Santa Casa."""
    st.markdown('<h1 class="main-header">📋 Prestações de Contas</h1>', unsafe_allow_html=True)
    st.info("💡 Gestão de débitos e créditos com a Santa Casa após acertos com vendas e prémios")

    # Carregar dados de prestações - usar caminho absoluto baseado na home
    home = Path.home()
    caminho_prestacao = home / "Documentos" / "Santa casa" / "dados" / "prestacao.csv"

    if not caminho_prestacao.exists():
        st.error(f"Ficheiro de prestações não encontrado em: {caminho_prestacao}")
        return

    try:
        # Carregar CSV com tratamento de encoding
        prestacao_df = pd.read_csv(
            caminho_prestacao,
            encoding='utf-8',
            dtype={'Filename': str, 'Value': str, 'Date': str}
        )

        # Converter a coluna Value (substituir vírgula por ponto)
        prestacao_df['Value'] = prestacao_df['Value'].str.replace('.', '').str.replace(',', '.').astype(float)

        # Converter a coluna Date
        prestacao_df['Date'] = pd.to_datetime(prestacao_df['Date'], format='%d-%m-%Y')

        # Aplicar delta de +2 dias (mesmo que nas vendas)
        prestacao_df['Date'] = prestacao_df['Date'] + timedelta(days=2)

        # Adicionar colunas de período
        prestacao_df['Mes_Ano'] = prestacao_df['Date'].dt.strftime('%m/%Y')
        prestacao_df['Ano'] = prestacao_df['Date'].dt.year
        prestacao_df['Mes'] = prestacao_df['Date'].dt.month

    except Exception as e:
        st.error(f"Erro ao carregar dados de prestações: {e}")
        return

    st.markdown("---")

    # Filtros na sidebar
    st.sidebar.header("🔍 Filtros - Prestações")

    # Filtro por ano
    anos_disponiveis = sorted(prestacao_df['Ano'].unique(), reverse=True)
    anos_selecionados = st.sidebar.multiselect(
        "Selecionar Anos",
        anos_disponiveis,
        default=anos_disponiveis,
        key="prest_anos"
    )

    if not anos_selecionados:
        st.warning("Selecione pelo menos um ano")
        return

    # Filtrar dados
    prest_filtrada = prestacao_df[prestacao_df['Ano'].isin(anos_selecionados)]

    # Métricas principais
    col1, col2, col3, col4 = st.columns(4)

    total_prestacao = prest_filtrada['Value'].sum()
    num_prestacoes = len(prest_filtrada)
    media_prestacao = prest_filtrada['Value'].mean()
    prestacao_maior = prest_filtrada['Value'].max()

    with col1:
        st.metric("💰 Total a Pagar", f"€{total_prestacao:,.2f}")
    with col2:
        st.metric("📄 Nº de Prestações", num_prestacoes)
    with col3:
        st.metric("📊 Média por Prestação", f"€{media_prestacao:,.2f}")
    with col4:
        st.metric("📈 Maior Prestação", f"€{prestacao_maior:,.2f}")

    st.markdown("---")

    # Tabs para diferentes análises
    tab1, tab2, tab3, tab4 = st.tabs([
        "📈 Evolução Temporal",
        "📊 Por Mês",
        "📋 Detalhes",
        "📌 Resumo Anual"
    ])

    with tab1:
        st.subheader("Evolução de Prestações ao Longo do Tempo")

        # Agrupar por data
        prest_por_data = prest_filtrada.sort_values('Date').groupby('Date')['Value'].sum().reset_index()

        fig = px.line(
            prest_por_data,
            x='Date',
            y='Value',
            markers=True,
            title="Prestações de Contas por Data",
            labels={'Date': 'Data', 'Value': 'Valor (€)'},
            color_discrete_sequence=['#d62728']
        )
        fig.update_traces(line=dict(width=2), marker=dict(size=8))
        fig.update_layout(height=500, hovermode='x unified')
        st.plotly_chart(fig, use_container_width=True)

        # Estatísticas de evolução
        col1, col2, col3 = st.columns(3)

        with col1:
            prestacao_minima = prest_filtrada['Value'].min()
            st.metric("💚 Prestação Menor", f"€{prestacao_minima:,.2f}")

        with col2:
            datas_unicas = prest_filtrada['Date'].nunique()
            st.metric("📅 Datas com Prestação", datas_unicas)

        with col3:
            periodo = f"{prest_filtrada['Date'].min().strftime('%d/%m/%Y')} a {prest_filtrada['Date'].max().strftime('%d/%m/%Y')}"
            st.metric("📌 Período", periodo.split(' a ')[1][-4:])  # Mostrar só o ano final

    with tab2:
        st.subheader("Distribuição de Prestações por Mês")

        # Agrupar por mês
        prest_por_mes = prest_filtrada.groupby('Mes_Ano')['Value'].agg(['sum', 'count']).reset_index()
        prest_por_mes.columns = ['Mês', 'Total', 'Nº_Prestações']

        col1, col2 = st.columns([2, 1])

        with col1:
            # Gráfico de barras
            fig = px.bar(
                prest_por_mes,
                x='Mês',
                y='Total',
                title="Total de Prestações por Mês",
                labels={'Mês': 'Mês/Ano', 'Total': 'Valor Total (€)'},
                color='Total',
                color_continuous_scale='Reds',
                text='Total'
            )
            fig.update_traces(texttemplate='€%{text:.0f}', textposition='outside')
            fig.update_layout(height=450, showlegend=False)
            st.plotly_chart(fig, use_container_width=True)

        with col2:
            st.markdown("### 📋 Detalhes por Mês")
            tabela_mes = prest_por_mes.copy()
            tabela_mes['Total'] = tabela_mes['Total'].apply(lambda x: f"€{x:,.2f}")
            tabela_mes = tabela_mes.rename(columns={'Mês': 'Período', 'Nº_Prestações': 'Nº'})
            st.dataframe(tabela_mes, use_container_width=True, hide_index=True, height=450)

    with tab3:
        st.subheader("Detalhes Completos das Prestações")

        # Criar tabela formatada
        prest_tabela = prest_filtrada.copy()
        prest_tabela = prest_tabela.sort_values('Date', ascending=False)
        prest_tabela['Date'] = prest_tabela['Date'].dt.strftime('%d/%m/%Y')
        prest_tabela['Value'] = prest_tabela['Value'].apply(lambda x: f"€{x:,.2f}")

        # Colunas para mostrar
        colunas_exibir = ['Date', 'Filename', 'Value', 'Mes_Ano']
        prest_tabela = prest_tabela[colunas_exibir]
        prest_tabela = prest_tabela.rename(columns={
            'Date': 'Data',
            'Filename': 'Documento',
            'Value': 'Valor',
            'Mes_Ano': 'Período'
        })

        st.dataframe(prest_tabela, use_container_width=True, hide_index=True, height=600)

        # Download de dados
        csv = prest_tabela.to_csv(index=False, encoding='utf-8-sig')
        st.download_button(
            label="📥 Descarregar como CSV",
            data=csv,
            file_name=f"prestacoes_contas_{pd.Timestamp.now().strftime('%Y%m%d')}.csv",
            mime="text/csv"
        )

    with tab4:
        st.subheader("Resumo Anual")

        # Agrupar por ano
        prest_por_ano = prest_filtrada.groupby('Ano')['Value'].agg(['sum', 'count', 'mean']).reset_index()
        prest_por_ano.columns = ['Ano', 'Total', 'Nº_Prestações', 'Média']
        prest_por_ano = prest_por_ano.sort_values('Ano', ascending=False)

        col1, col2 = st.columns([2, 1])

        with col1:
            # Gráfico comparativo de anos
            fig = px.bar(
                prest_por_ano,
                x='Ano',
                y='Total',
                title="Total de Prestações por Ano",
                labels={'Ano': 'Ano', 'Total': 'Valor Total (€)'},
                color='Total',
                color_continuous_scale='Blues',
                text='Total'
            )
            fig.update_traces(texttemplate='€%{text:.0f}', textposition='outside')
            fig.update_layout(height=450, showlegend=False)
            st.plotly_chart(fig, use_container_width=True)

        with col2:
            st.markdown("### 📊 Resumo Anual")
            tabela_ano = prest_por_ano.copy()
            tabela_ano['Total'] = tabela_ano['Total'].apply(lambda x: f"€{x:,.2f}")
            tabela_ano['Média'] = tabela_ano['Média'].apply(lambda x: f"€{x:,.2f}")
            tabela_ano = tabela_ano.rename(columns={'Nº_Prestações': 'Nº'})
            st.dataframe(tabela_ano, use_container_width=True, hide_index=True, height=450)

    # Comparação com vendas (se houver dados)
    st.markdown("---")
    st.subheader("📊 Comparação com Vendas")

    total_vendas = df[df['Ano'].isin(anos_selecionados)]['Valor'].sum()

    col1, col2, col3 = st.columns(3)

    with col1:
        st.metric("💰 Total de Vendas", f"€{total_vendas:,.2f}")
    with col2:
        st.metric("💸 Total de Prestações", f"€{total_prestacao:,.2f}")
    with col3:
        saldo = total_vendas - total_prestacao
        cor = "🟢" if saldo >= 0 else "🔴"
        st.metric(f"{cor} Saldo", f"€{saldo:,.2f}")

    # Gráfico de comparação
    col1, col2 = st.columns([2, 1])

    with col1:
        comparacao_data = pd.DataFrame({
            'Tipo': ['Vendas', 'Prestações'],
            'Valor': [total_vendas, total_prestacao]
        })

        fig = px.bar(
            comparacao_data,
            x='Tipo',
            y='Valor',
            title="Comparação: Vendas vs Prestações",
            labels={'Tipo': '', 'Valor': 'Valor (€)'},
            color='Tipo',
            color_discrete_map={'Vendas': '#1f77b4', 'Prestações': '#d62728'},
            text='Valor'
        )
        fig.update_traces(texttemplate='€%{text:.0f}', textposition='outside')
        fig.update_layout(height=400, showlegend=False)
        st.plotly_chart(fig, use_container_width=True)

    with col2:
        st.markdown("### 💡 Análise")

        percentagem_prestacao = (total_prestacao / total_vendas * 100) if total_vendas > 0 else 0

        if saldo > 0:
            st.success(f"✅ Saldo positivo de €{saldo:,.2f}")
            st.info(f"📊 Prestações representam {percentagem_prestacao:.1f}% das vendas")
        elif saldo < 0:
            st.error(f"❌ Saldo negativo de €{abs(saldo):,.2f}")
            st.warning(f"⚠️ Precisa pagar mais €{abs(saldo):,.2f} à Santa Casa")
        else:
            st.warning("⚖️ Saldo equilibrado")


def pagina_premios_pagos(df):
    """Análise de Prémios Pagos ao longo do tempo."""
    st.markdown('<h1 class="main-header">🎁 Prémios Pagos</h1>', unsafe_allow_html=True)
    st.info("💡 Análise dos prémios pagos (Vendas - Remunerações - Prestações) ao longo do tempo")

    # Carregar dados de prestações
    home = Path.home()
    caminho_prestacao = home / "Documentos" / "Santa casa" / "dados" / "prestacao.csv"

    if not caminho_prestacao.exists():
        st.error(f"Ficheiro de prestações não encontrado em: {caminho_prestacao}")
        return

    try:
        # Carregar CSV com tratamento de encoding
        prestacao_df = pd.read_csv(
            caminho_prestacao,
            encoding='utf-8',
            dtype={'Filename': str, 'Value': str, 'Date': str}
        )

        # Converter a coluna Value (substituir vírgula por ponto)
        prestacao_df['Value'] = prestacao_df['Value'].str.replace('.', '').str.replace(',', '.').astype(float)

        # Converter a coluna Date
        prestacao_df['Date'] = pd.to_datetime(prestacao_df['Date'], format='%d-%m-%Y')

        # Aplicar delta de +2 dias (mesmo que nas vendas)
        prestacao_df['Date'] = prestacao_df['Date'] + timedelta(days=2)

        # Adicionar colunas de período
        prestacao_df['Data_Emissao'] = prestacao_df['Date'].dt.date
        prestacao_df['Ano'] = prestacao_df['Date'].dt.year

    except Exception as e:
        st.error(f"Erro ao carregar dados de prestações: {e}")
        return

    st.markdown("---")

    # Filtros na sidebar
    st.sidebar.header("🔍 Filtros - Prémios")

    # Filtro por ano
    anos_df = sorted(df['Ano'].unique(), reverse=True)
    anos_selecionados = st.sidebar.multiselect(
        "Selecionar Anos",
        anos_df,
        default=anos_df,
        key="premios_anos"
    )

    if not anos_selecionados:
        st.warning("Selecione pelo menos um ano")
        return

    # Filtrar dados por ano
    df_filtrado = df[df['Ano'].isin(anos_selecionados)].copy()
    prest_filtrada = prestacao_df[prestacao_df['Ano'].isin(anos_selecionados)].copy()

    # Calcular prémios pagos por data de emissão (de vendas)
    df_vendas = df_filtrado.copy()
    df_vendas['Percentagem'] = df_vendas['Jogo'].map(REMUNERACAO_JOGOS)
    df_vendas['Remuneracao'] = df_vendas['Valor'] * df_vendas['Percentagem'] / 100

    # Agregar por data de emissão (data das vendas)
    vendas_por_data = df_vendas.groupby('Data_Emissao').agg({
        'Valor': 'sum',
        'Remuneracao': 'sum'
    }).reset_index()
    vendas_por_data.columns = ['Data', 'Total_Vendas', 'Total_Remuneracao']
    vendas_por_data['Data'] = pd.to_datetime(vendas_por_data['Data'])

    # Agregar prestações por data
    prestacoes_por_data = prest_filtrada.groupby('Data_Emissao')['Value'].sum().reset_index()
    prestacoes_por_data.columns = ['Data', 'Prestacao']
    prestacoes_por_data['Data'] = pd.to_datetime(prestacoes_por_data['Data'])

    # Mesclar dados
    analise_premios = vendas_por_data.copy()
    analise_premios = analise_premios.merge(prestacoes_por_data, on='Data', how='left')
    analise_premios['Prestacao'] = analise_premios['Prestacao'].fillna(0)
    analise_premios['Premios_Pagos'] = analise_premios['Total_Vendas'] - analise_premios['Total_Remuneracao'] - analise_premios['Prestacao']
    analise_premios = analise_premios.sort_values('Data')

    # Métricas principais
    total_vendas = analise_premios['Total_Vendas'].sum()
    total_remuneracao = analise_premios['Total_Remuneracao'].sum()
    total_prestacao = analise_premios['Prestacao'].sum()
    total_premios = analise_premios['Premios_Pagos'].sum()

    col1, col2, col3 = st.columns(3)

    with col1:
        st.metric("💰 Total de Vendas", f"€{total_vendas:,.2f}")
    with col2:
        st.metric("👔 Total Remunerações", f"€{total_remuneracao:,.2f}")
    with col3:
        cor_premio = "🎁" if total_premios > 0 else "⚠️"
        st.metric(f"{cor_premio} Total Prémios", f"€{total_premios:,.2f}")

    st.markdown("---")

    # Tabs para diferentes análises
    tab1, tab2, tab3, tab4 = st.tabs([
        "📈 Evolução Temporal",
        "📊 Decomposição",
        "📋 Detalhes",
        "🔍 Análise Comparativa"
    ])

    with tab1:
        st.subheader("Evolução dos Prémios Pagos ao Longo do Tempo")

        col1, col2 = st.columns(2)

        with col1:
            # Gráfico de linha dos prémios
            fig = px.line(
                analise_premios,
                x='Data',
                y='Premios_Pagos',
                markers=True,
                title="Prémios Pagos por Semana",
                labels={'Data': 'Data', 'Premios_Pagos': 'Prémios (€)'},
                color_discrete_sequence=['#2ca02c']
            )
            fig.update_traces(line=dict(width=2), marker=dict(size=8))
            fig.update_layout(height=400, hovermode='x unified')
            st.plotly_chart(fig, use_container_width=True)

        with col2:
            # Gráfico de área dos prémios acumulados
            analise_premios['Premios_Acumulados'] = analise_premios['Premios_Pagos'].cumsum()

            fig = px.area(
                analise_premios,
                x='Data',
                y='Premios_Acumulados',
                title="Prémios Acumulados",
                labels={'Data': 'Data', 'Premios_Acumulados': 'Prémios Acumulados (€)'},
                color_discrete_sequence=['#2ca02c']
            )
            fig.update_layout(height=400)
            st.plotly_chart(fig, use_container_width=True)

    with tab2:
        st.subheader("Decomposição das Vendas ao Longo do Tempo")

        # Gráfico empilhado mostrando os componentes
        fig = go.Figure()

        fig.add_trace(go.Bar(
            x=analise_premios['Data'],
            y=analise_premios['Total_Remuneracao'],
            name='Remunerações',
            marker_color='#1f77b4'
        ))

        fig.add_trace(go.Bar(
            x=analise_premios['Data'],
            y=analise_premios['Prestacao'],
            name='Prestações',
            marker_color='#d62728'
        ))

        fig.add_trace(go.Bar(
            x=analise_premios['Data'],
            y=analise_premios['Premios_Pagos'],
            name='Prémios Pagos',
            marker_color='#2ca02c'
        ))

        fig.update_layout(
            barmode='stack',
            title='Decomposição das Vendas (Remunerações + Prestações + Prémios)',
            xaxis_title='Data',
            yaxis_title='Valor (€)',
            height=500,
            hovermode='x unified'
        )

        st.plotly_chart(fig, use_container_width=True)

    with tab3:
        st.subheader("Detalhes Completos")

        # Criar tabela formatada
        tabela_premios = analise_premios.copy()
        tabela_premios['Data'] = tabela_premios['Data'].dt.strftime('%d/%m/%Y')
        tabela_premios['Total_Vendas'] = tabela_premios['Total_Vendas'].apply(lambda x: f"€{x:,.2f}")
        tabela_premios['Total_Remuneracao'] = tabela_premios['Total_Remuneracao'].apply(lambda x: f"€{x:,.2f}")
        tabela_premios['Prestacao'] = tabela_premios['Prestacao'].apply(lambda x: f"€{x:,.2f}")
        tabela_premios['Premios_Pagos'] = tabela_premios['Premios_Pagos'].apply(lambda x: f"€{x:,.2f}")

        # Colunas para mostrar
        colunas_exibir = ['Data', 'Total_Vendas', 'Total_Remuneracao', 'Prestacao', 'Premios_Pagos']
        tabela_premios = tabela_premios[colunas_exibir]
        tabela_premios = tabela_premios.rename(columns={
            'Data': 'Data',
            'Total_Vendas': 'Vendas (€)',
            'Total_Remuneracao': 'Remunerações (€)',
            'Prestacao': 'Prestações (€)',
            'Premios_Pagos': 'Prémios (€)'
        })

        st.dataframe(tabela_premios, use_container_width=True, hide_index=True, height=600)

        # Download de dados
        csv = tabela_premios.to_csv(index=False, encoding='utf-8-sig')
        st.download_button(
            label="📥 Descarregar como CSV",
            data=csv,
            file_name=f"premios_pagos_{pd.Timestamp.now().strftime('%Y%m%d')}.csv",
            mime="text/csv"
        )

    with tab4:
        st.subheader("Análise Comparativa")

        col1, col2 = st.columns(2)

        with col1:
            # Gráfico de linhas mostrando os três componentes
            fig = go.Figure()

            fig.add_trace(go.Scatter(
                x=analise_premios['Data'],
                y=analise_premios['Total_Vendas'],
                name='Vendas',
                mode='lines+markers',
                line=dict(color='#1f77b4', width=2)
            ))

            fig.add_trace(go.Scatter(
                x=analise_premios['Data'],
                y=analise_premios['Total_Remuneracao'],
                name='Remunerações',
                mode='lines+markers',
                line=dict(color='#ff7f0e', width=2)
            ))

            fig.add_trace(go.Scatter(
                x=analise_premios['Data'],
                y=analise_premios['Premios_Pagos'],
                name='Prémios Pagos',
                mode='lines+markers',
                line=dict(color='#2ca02c', width=2)
            ))

            fig.update_layout(
                title='Comparação: Vendas vs Remunerações vs Prémios',
                xaxis_title='Data',
                yaxis_title='Valor (€)',
                height=450,
                hovermode='x unified'
            )

            st.plotly_chart(fig, use_container_width=True)

        with col2:
            # Estatísticas
            st.markdown("### 📊 Estatísticas")

            media_premios = analise_premios['Premios_Pagos'].mean()
            max_premios = analise_premios['Premios_Pagos'].max()
            min_premios = analise_premios['Premios_Pagos'].min()
            total_semanas = len(analise_premios)

            col1_stat, col2_stat = st.columns(2)

            with col1_stat:
                st.metric("📈 Média Semanal", f"€{media_premios:,.2f}")
                st.metric("⬆️ Máximo", f"€{max_premios:,.2f}")
                st.metric("⬇️ Mínimo", f"€{min_premios:,.2f}")

            with col2_stat:
                st.metric("📊 Total de Semanas", total_semanas)

            # Tabela de análise
            st.markdown("---")
            st.markdown("### 💡 Análise")

            percentagem_premios = (total_premios / total_vendas * 100) if total_vendas > 0 else 0
            percentagem_remuneracao = (total_remuneracao / total_vendas * 100) if total_vendas > 0 else 0
            percentagem_prestacao = (total_prestacao / total_vendas * 100) if total_vendas > 0 else 0

            st.markdown(f"""
            - 💰 **Prémios representam {percentagem_premios:.1f}% das vendas**
            - 👔 **Remunerações representam {percentagem_remuneracao:.1f}% das vendas**
            - 📋 **Prestações representam {percentagem_prestacao:.1f}% das vendas**
            """)

            if total_premios > 0:
                st.success(f"✅ Prémios totais positivos: €{total_premios:,.2f}")
            else:
                st.error(f"❌ Prémios totais negativos: €{abs(total_premios):,.2f}")


def main():
    """Função principal do dashboard."""

    # Título principal
    st.sidebar.title("🎲 Dashboard de Vendas dos Jogos da Santa Casa")
    st.sidebar.markdown("### Café Martins")
    st.sidebar.markdown("---")

    # Carregar dados
    with st.spinner("Carregando dados..."):
        df = carregar_dados()

    if df is None:
        st.error("Erro ao carregar os dados. Verifique se o arquivo 'dados_extracao.txt' está na pasta correta.")
        return

    # Informações gerais na sidebar
    st.sidebar.markdown("### 📊 Informações Gerais")
    st.sidebar.info(f"""
    **Período:** {df['Data_Emissao'].min().strftime('%d/%m/%Y')} - {df['Data_Emissao'].max().strftime('%d/%m/%Y')}

    **Total de Registos:** {len(df):,}

    **Jogos:** {df['Jogo'].nunique()}

    **Anos:** {', '.join(map(str, sorted(df['Ano'].unique())))}
    """)

    st.sidebar.markdown("---")

    # Menu de navegação
    pagina = st.sidebar.radio(
        "Navegação",
        ["📈 Visão Geral", "🎯 Dashboard Executivo", "🔬 Análise Semanal", "🎮 Análise por Jogo", "⚖️ Comparação", "📊 Comparações Avançadas", "💰 Remuneração", "📋 Prestações de Contas", "🎁 Prémios Pagos"]
    )

    # Renderizar página selecionada
    if pagina == "📈 Visão Geral":
        pagina_visao_geral(df)
    elif pagina == "🎯 Dashboard Executivo":
        pagina_dashboard_executivo(df)
    elif pagina == "🔬 Análise Semanal":
        pagina_analise_semanal_avancada(df)
    elif pagina == "🎮 Análise por Jogo":
        pagina_analise_jogos(df)
    elif pagina == "⚖️ Comparação":
        pagina_comparacao(df)
    elif pagina == "📊 Comparações Avançadas":
        pagina_comparacoes_avancadas(df)
    elif pagina == "💰 Remuneração":
        pagina_remuneracao(df)
    elif pagina == "📋 Prestações de Contas":
        pagina_prestacoes_contas(df)
    elif pagina == "🎁 Prémios Pagos":
        pagina_premios_pagos(df)


if __name__ == "__main__":
    # Carregar configuração de autenticação
    config_file = Path(__file__).parent / 'config.yaml'

    if not config_file.exists():
        st.error("Ficheiro de configuração config.yaml não encontrado!")
        st.stop()

    with open(config_file) as file:
        config = yaml.load(file, Loader=SafeLoader)

    # Criar objeto de autenticação
    authenticator = stauth.Authenticate(
        config['credentials'],
        config['cookie']['name'],
        config['cookie']['key'],
        config['cookie']['expiry_days']
    )

    # Renderizar formulário de login
    try:
        authenticator.login()
    except Exception as e:
        st.error(f'Erro ao processar login: {e}')
        st.stop()

    # Verificar status de autenticação
    if st.session_state.get("authentication_status"):
        # Usuário autenticado - adicionar logout na sidebar
        authenticator.logout()
        st.sidebar.write(f'Bem-vindo, **{st.session_state["name"]}**!')

        # Executar aplicação principal
        main()

    elif st.session_state.get("authentication_status") is False:
        st.error('Utilizador ou password incorretos')
    elif st.session_state.get("authentication_status") is None:
        st.warning('Por favor, introduza as suas credenciais')
